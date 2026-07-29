#!/usr/bin/env python3
"""Build manuscript-ready outputs for the optimizer-matched CamCan comparison.

This is a presentation-only layer over the validated schema-v2 outputs from
``camcan_personalized_comparison.py``. It does not read or modify source
simulation images. The comparison contains every combination of seven
optimized subjects and four target ROIs, with generic and ROI-specific
personalized montages evaluated on the same corrected subject head.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from CamCan_Experiment.post.build_camcan_supervisor_figure_guide import (
    BLUE,
    GRAY,
    NAVY,
    PALE_BLUE,
    PALE_ORANGE,
    add_bullet,
    add_callout,
    add_heading,
    add_table,
    configure_document,
    set_cell_shading,
)


ROI_ORDER = [
    "Left_Hippocampus",
    "Left_M1",
    "Right_DLPC",
    "Right_Thalamus",
]
ROI_LABELS = {
    "Left_Hippocampus": "Left hippocampus",
    "Left_M1": "Left M1",
    "Right_DLPC": "Right DLPFC",
    "Right_Thalamus": "Right thalamus",
}
CONDITION_ORDER = ["generic", "personalized"]
ROLE_LABELS = {
    "best": "ROI-specific generic-best subject",
    "worst": "ROI-specific generic-worst subject",
    "cross_target": "Other optimized subject",
}
ROLE_LINESTYLES = {
    "best": ":",
    "worst": "--",
    "cross_target": "-",
}

GENERIC_BLUE = "#2F6B9A"
PERSONAL_ORANGE = "#D97706"
LIGHT_GRAY = "#CBD1D8"
DARK_GRAY = "#4B5563"
GRID_GRAY = "#E6E9ED"

SUBJECT_PALETTE = [
    "#0072B2",
    "#E69F00",
    "#009E73",
    "#CC79A7",
    "#D55E00",
    "#56B4E9",
    "#6F6259",
]

FIGURE_SUMMARY = {
    "designation": "Figure 1",
    "stem": "figure_personalization_all_subject_changes",
    "title": "Optimizer-target changes across all subjects and ROIs",
}
ARROW_FIGURES = [
    {
        "designation": "Figure 2",
        "stem": "figure_personalization_effectiveness_spread_ge_0p18",
        "title": "Effectiveness-spread trade-offs at 0.18 V/m",
        "threshold": 0.18,
    },
    {
        "designation": "Supplementary Figure S1",
        "stem": "figure_personalization_effectiveness_spread_ge_0p20",
        "title": "Effectiveness-spread sensitivity analysis at 0.20 V/m",
        "threshold": 0.20,
    },
    {
        "designation": "Supplementary Figure S2",
        "stem": "figure_personalization_effectiveness_spread_ge_0p15",
        "title": "Effectiveness-spread sensitivity analysis at 0.15 V/m",
        "threshold": 0.15,
    },
]
REPEAT_FIGURES = [
    {
        "designation": f"Supplementary Figure S{index + 3}",
        "stem": f"figure_personalization_technical_repeats_{roi.lower()}",
        "title": f"Technical-repeat distributions: {ROI_LABELS[roi]}",
        "roi": roi,
    }
    for index, roi in enumerate(ROI_ORDER)
]
FIGURES = [FIGURE_SUMMARY, *ARROW_FIGURES, *REPEAT_FIGURES]


def threshold_token(threshold: float) -> str:
    return str(threshold).replace(".", "p")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def short_subject(subject: str) -> str:
    return str(subject).replace("sub-", "")


def set_figure_style() -> None:
    mpl.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 8.0,
            "axes.titlesize": 9.4,
            "axes.labelsize": 8.5,
            "xtick.labelsize": 7.3,
            "ytick.labelsize": 7.3,
            "legend.fontsize": 7.2,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.linewidth": 0.7,
            "xtick.major.width": 0.7,
            "ytick.major.width": 0.7,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "savefig.facecolor": "white",
        }
    )


def save_figure(figure: plt.Figure, figures_dir: Path, stem: str) -> None:
    figure.savefig(
        figures_dir / f"{stem}.png",
        dpi=400,
        bbox_inches="tight",
        pad_inches=0.04,
        facecolor="white",
    )
    figure.savefig(
        figures_dir / f"{stem}.pdf",
        bbox_inches="tight",
        pad_inches=0.04,
        facecolor="white",
    )
    plt.close(figure)


def load_and_validate(
    input_dir: Path,
) -> tuple[dict, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    required = [
        "analysis_manifest.json",
        "paired_personalized_vs_generic.csv",
        "condition_repeat_mean_metrics.csv",
        "repeat_level_metrics.csv",
        "table_main_selected_metrics.csv",
        "table_stimulation_parameters.csv",
        "table_optimizer_roi_definitions.csv",
        "selection_allowlist.csv",
    ]
    missing = [name for name in required if not (input_dir / name).is_file()]
    if missing:
        raise FileNotFoundError("Missing comparison output(s): " + ", ".join(missing))

    manifest = json.loads((input_dir / "analysis_manifest.json").read_text())
    paired = pd.read_csv(input_dir / "paired_personalized_vs_generic.csv")
    condition = pd.read_csv(input_dir / "condition_repeat_mean_metrics.csv")
    repeats = pd.read_csv(input_dir / "repeat_level_metrics.csv")

    expected_manifest = {
        "status": "complete",
        "comparison_schema_version": 2,
        "subject_roi_configurations": 28,
        "originally_selected_extreme_configurations": 8,
        "cross_target_configurations": 20,
        "unique_subjects": 7,
        "repeat_level_records": 560,
        "condition_repeat_mean_records": 56,
    }
    for key, expected in expected_manifest.items():
        if manifest.get(key) != expected:
            raise RuntimeError(
                f"Unexpected manifest value {key}={manifest.get(key)!r}; "
                f"expected {expected!r}"
            )
    if manifest.get("thresholds_v_per_m") != [0.2, 0.18, 0.15]:
        raise RuntimeError("Expected thresholds [0.2, 0.18, 0.15] V/m")
    if len(paired) != 28 or set(paired["roi"]) != set(ROI_ORDER):
        raise RuntimeError("Expected seven subjects for each of four ROIs")
    if paired["subject"].nunique() != 7:
        raise RuntimeError("Expected seven unique optimized subjects")
    expected_roles = {"best", "worst", "cross_target"}
    if set(paired["selection_role"]) != expected_roles:
        raise RuntimeError(f"Expected selection roles {sorted(expected_roles)}")
    role_counts = paired.groupby(["roi", "selection_role"]).size()
    for roi in ROI_ORDER:
        if role_counts.get((roi, "best"), 0) != 1:
            raise RuntimeError(f"Expected exactly one best subject for {roi}")
        if role_counts.get((roi, "worst"), 0) != 1:
            raise RuntimeError(f"Expected exactly one worst subject for {roi}")
        if role_counts.get((roi, "cross_target"), 0) != 5:
            raise RuntimeError(f"Expected exactly five other subjects for {roi}")
    if len(condition) != 56 or set(condition["condition"]) != set(CONDITION_ORDER):
        raise RuntimeError("Expected 56 condition means across two conditions")
    if len(repeats) != 560:
        raise RuntimeError("Expected 560 repeat-level records")
    if repeats.duplicated(["pair_index", "condition", "repeat"]).any():
        raise RuntimeError("Duplicate pair/condition/repeat records")
    counts = repeats.groupby(["pair_index", "condition"]).size()
    if not (counts == 10).all():
        raise RuntimeError("Every subject/ROI/condition must have ten repeats")

    plotted_metrics = [
        "roi_min_v_per_m",
        "roi_median_v_per_m",
        "target_coverage_percent_ge_0p2",
        "off_target_coverage_percent_ge_0p2",
        "threshold_localization_percent_in_roi_ge_0p2",
        "target_coverage_percent_ge_0p18",
        "off_target_coverage_percent_ge_0p18",
        "threshold_localization_percent_in_roi_ge_0p18",
        "target_coverage_percent_ge_0p15",
        "off_target_coverage_percent_ge_0p15",
        "threshold_localization_percent_in_roi_ge_0p15",
    ]
    paired_columns: list[str] = []
    for metric in plotted_metrics:
        paired_columns.extend(
            [
                f"{metric}__generic_repeat_mean",
                f"{metric}__generic_repeat_sd",
                f"{metric}__personalized_repeat_mean",
                f"{metric}__personalized_repeat_sd",
                f"{metric}__absolute_change",
            ]
        )
    checks = [
        ("paired plotted metrics", paired, paired_columns),
        ("repeat plotted metrics", repeats, plotted_metrics),
    ]
    for name, frame, columns in checks:
        missing_columns = sorted(set(columns) - set(frame.columns))
        if missing_columns:
            raise RuntimeError(f"{name} missing columns: {missing_columns}")
        if not np.isfinite(frame[columns].to_numpy(dtype=float)).all():
            raise RuntimeError(f"{name} contains non-finite values")

    roi_rank = {roi: index for index, roi in enumerate(ROI_ORDER)}
    subjects = sorted(paired["subject"].unique())
    subject_rank = {subject: index for index, subject in enumerate(subjects)}
    paired = paired.assign(
        _roi_order=paired["roi"].map(roi_rank),
        _subject_order=paired["subject"].map(subject_rank),
    ).sort_values(["_roi_order", "_subject_order"])
    condition = condition.assign(
        _roi_order=condition["roi"].map(roi_rank),
        _subject_order=condition["subject"].map(subject_rank),
    ).sort_values(["_roi_order", "_subject_order", "condition"])
    repeats = repeats.assign(
        _roi_order=repeats["roi"].map(roi_rank),
        _subject_order=repeats["subject"].map(subject_rank),
    ).sort_values(["_roi_order", "_subject_order", "condition", "repeat"])
    return manifest, paired, condition, repeats


def subject_colors(frame: pd.DataFrame) -> dict[str, str]:
    subjects = sorted(frame["subject"].unique())
    if len(subjects) != len(SUBJECT_PALETTE):
        raise RuntimeError("Subject palette expects exactly seven subjects")
    return dict(zip(subjects, SUBJECT_PALETTE))


def role_suffix(role: str) -> str:
    if role == "best":
        return " (best)"
    if role == "worst":
        return " (worst)"
    return ""


def plot_all_subject_changes(paired: pd.DataFrame, figures_dir: Path) -> None:
    panels = [
        ("roi_min_v_per_m", "Minimum target field\n(V/m)"),
        ("roi_median_v_per_m", "Median target field\n(V/m)"),
        (
            "target_coverage_percent_ge_0p18",
            "Target coverage\n>=0.18 V/m (%)",
        ),
        (
            "off_target_coverage_percent_ge_0p18",
            "Off-target coverage\n>=0.18 V/m (%)",
        ),
    ]
    figure, axes = plt.subplots(
        4,
        4,
        figsize=(7.35, 8.8),
        sharey="row",
        constrained_layout=False,
    )
    figure.subplots_adjust(
        left=0.17,
        right=0.99,
        top=0.92,
        bottom=0.075,
        hspace=0.34,
        wspace=0.28,
    )
    for roi_index, roi in enumerate(ROI_ORDER):
        subset = paired.loc[paired["roi"] == roi].sort_values("_subject_order")
        y = np.arange(len(subset))[::-1]
        labels = [
            short_subject(row["subject"]) + role_suffix(str(row["selection_role"]))
            for _, row in subset.iterrows()
        ]
        for metric_index, (metric, title) in enumerate(panels):
            axis = axes[roi_index, metric_index]
            generic = subset[f"{metric}__generic_repeat_mean"].to_numpy()
            personalized = subset[
                f"{metric}__personalized_repeat_mean"
            ].to_numpy()
            generic_sd = subset[f"{metric}__generic_repeat_sd"].to_numpy()
            personalized_sd = subset[
                f"{metric}__personalized_repeat_sd"
            ].to_numpy()
            for yi, start, end in zip(y, generic, personalized):
                axis.plot(
                    [start, end],
                    [yi, yi],
                    color=LIGHT_GRAY,
                    lw=1.25,
                    zorder=1,
                )
            axis.errorbar(
                generic,
                y,
                xerr=generic_sd,
                fmt="o",
                ms=4.0,
                mfc="white",
                mec=GENERIC_BLUE,
                mew=1.05,
                ecolor=GENERIC_BLUE,
                elinewidth=0.7,
                capsize=1.5,
                zorder=3,
            )
            axis.errorbar(
                personalized,
                y,
                xerr=personalized_sd,
                fmt="o",
                ms=4.0,
                mfc=PERSONAL_ORANGE,
                mec=PERSONAL_ORANGE,
                mew=0.9,
                ecolor=PERSONAL_ORANGE,
                elinewidth=0.7,
                capsize=1.5,
                zorder=4,
            )
            if roi_index == 0:
                axis.set_title(title, weight="bold", pad=6)
            axis.set_yticks(y)
            if metric_index == 0:
                axis.set_yticklabels(labels)
            else:
                axis.tick_params(axis="y", labelleft=False)
            axis.grid(axis="x", color=GRID_GRAY, lw=0.55)
            axis.set_axisbelow(True)
            if "coverage_percent" in metric:
                axis.set_xlim(-3, 103)
            else:
                values = np.concatenate(
                    [
                        generic - generic_sd,
                        generic + generic_sd,
                        personalized - personalized_sd,
                        personalized + personalized_sd,
                    ]
                )
                span = max(float(np.ptp(values)), 0.02)
                axis.set_xlim(
                    max(0.0, float(values.min()) - span * 0.08),
                    float(values.max()) + span * 0.08,
                )
            if metric_index == 0:
                axis.set_ylabel(ROI_LABELS[roi], weight="bold", labelpad=7)
    handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            color=GENERIC_BLUE,
            markerfacecolor="white",
            markeredgecolor=GENERIC_BLUE,
            lw=0.8,
            label="Generic montage",
        ),
        Line2D(
            [0],
            [0],
            marker="o",
            color=PERSONAL_ORANGE,
            markerfacecolor=PERSONAL_ORANGE,
            markeredgecolor=PERSONAL_ORANGE,
            lw=0.8,
            label="Personalized montage",
        ),
    ]
    figure.legend(
        handles=handles,
        loc="upper center",
        bbox_to_anchor=(0.58, 0.985),
        ncol=2,
        frameon=False,
        handletextpad=0.45,
        columnspacing=1.4,
    )
    save_figure(figure, figures_dir, FIGURE_SUMMARY["stem"])


def _off_target_limits(frame: pd.DataFrame, threshold: float) -> tuple[float, float]:
    token = threshold_token(threshold)
    columns = [
        f"off_target_coverage_percent_ge_{token}__generic_repeat_mean",
        f"off_target_coverage_percent_ge_{token}__personalized_repeat_mean",
    ]
    maximum = float(frame[columns].to_numpy().max())
    padding = max(0.25, maximum * 0.10)
    return 0.0, min(100.0, maximum + padding)


def plot_tradeoff(
    paired: pd.DataFrame,
    figures_dir: Path,
    threshold: float,
) -> None:
    token = threshold_token(threshold)
    x_metric = f"target_coverage_percent_ge_{token}"
    y_metric = f"off_target_coverage_percent_ge_{token}"
    colors = subject_colors(paired)
    figure, axes = plt.subplots(2, 2, figsize=(7.35, 7.35))
    figure.subplots_adjust(
        left=0.09,
        right=0.985,
        top=0.77,
        bottom=0.19,
        hspace=0.40,
        wspace=0.31,
    )
    for panel_index, roi in enumerate(ROI_ORDER):
        axis = axes.flat[panel_index]
        subset = paired.loc[paired["roi"] == roi].sort_values("_subject_order")
        for _, row in subset.iterrows():
            color = colors[str(row["subject"])]
            linestyle = ROLE_LINESTYLES[str(row["selection_role"])]
            x0 = float(row[f"{x_metric}__generic_repeat_mean"])
            y0 = float(row[f"{y_metric}__generic_repeat_mean"])
            x1 = float(row[f"{x_metric}__personalized_repeat_mean"])
            y1 = float(row[f"{y_metric}__personalized_repeat_mean"])
            axis.annotate(
                "",
                xy=(x1, y1),
                xytext=(x0, y0),
                arrowprops={
                    "arrowstyle": "-|>",
                    "color": color,
                    "linestyle": linestyle,
                    "lw": 1.7,
                    "mutation_scale": 9.5,
                    "shrinkA": 4.2,
                    "shrinkB": 4.2,
                },
                zorder=2,
            )
            axis.scatter(
                x0,
                y0,
                s=31,
                facecolors="white",
                edgecolors=color,
                linewidths=1.15,
                zorder=3,
            )
            axis.scatter(
                x1,
                y1,
                s=33,
                facecolors=color,
                edgecolors="white",
                linewidths=0.55,
                zorder=4,
            )
        axis.set_title(ROI_LABELS[roi], weight="bold", pad=4)
        axis.set_xlim(-3, 103)
        axis.set_ylim(*_off_target_limits(subset, threshold))
        axis.set_xlabel(f"Optimizer-target coverage >= {threshold:.2f} V/m (%)")
        axis.set_ylabel(f"Off-target coverage >= {threshold:.2f} V/m (%)")
        axis.grid(color=GRID_GRAY, lw=0.55)
        axis.set_axisbelow(True)
        axis.text(
            0.01,
            1.025,
            chr(ord("A") + panel_index),
            transform=axis.transAxes,
            weight="bold",
            fontsize=10.5,
        )

    subject_handles = [
        Line2D(
            [0],
            [0],
            color=colors[subject],
            lw=2.2,
            label=short_subject(subject),
        )
        for subject in sorted(colors)
    ]
    role_handles = [
        Line2D(
            [0],
            [0],
            color=DARK_GRAY,
            lw=1.8,
            linestyle=ROLE_LINESTYLES[role],
            label=ROLE_LABELS[role],
        )
        for role in ["best", "worst", "cross_target"]
    ]
    condition_handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            color="none",
            markerfacecolor="white",
            markeredgecolor=DARK_GRAY,
            markersize=5,
            label="Generic start",
        ),
        Line2D(
            [0],
            [0],
            marker="o",
            color="none",
            markerfacecolor=DARK_GRAY,
            markeredgecolor="white",
            markersize=5,
            label="Personalized end",
        ),
    ]
    figure.legend(
        handles=subject_handles,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.985),
        ncol=4,
        frameon=False,
        title="Subject (arrow color)",
        title_fontsize=7.7,
        handlelength=1.8,
        handletextpad=0.4,
        columnspacing=1.15,
    )
    figure.legend(
        handles=[*condition_handles, *role_handles],
        loc="lower center",
        bbox_to_anchor=(0.5, 0.015),
        ncol=3,
        frameon=False,
        handlelength=2.1,
        handletextpad=0.45,
        columnspacing=1.0,
    )
    stem = next(
        item["stem"]
        for item in ARROW_FIGURES
        if float(item["threshold"]) == threshold
    )
    save_figure(figure, figures_dir, stem)


def plot_repeat_distribution_for_roi(
    repeats: pd.DataFrame,
    figures_dir: Path,
    roi: str,
) -> None:
    subset = repeats.loc[repeats["roi"] == roi].copy()
    subjects = sorted(subset["subject"].unique())
    role_map = (
        subset[["subject", "selection_role"]]
        .drop_duplicates()
        .set_index("subject")["selection_role"]
        .to_dict()
    )
    figure, axes = plt.subplots(2, 4, figsize=(7.35, 4.75), sharey=True)
    figure.subplots_adjust(
        left=0.08,
        right=0.985,
        top=0.86,
        bottom=0.12,
        hspace=0.42,
        wspace=0.20,
    )
    rng = np.random.default_rng(20260728)
    all_values = subset["roi_min_v_per_m"].to_numpy()
    padding = max(0.008, float(np.ptp(all_values)) * 0.08)
    limits = (
        max(0.0, float(all_values.min()) - padding),
        float(all_values.max()) + padding,
    )
    for index, subject in enumerate(subjects):
        axis = axes.flat[index]
        subject_data = subset.loc[subset["subject"] == subject]
        values = [
            subject_data.loc[
                subject_data["condition"] == condition, "roi_min_v_per_m"
            ].to_numpy()
            for condition in CONDITION_ORDER
        ]
        box = axis.boxplot(
            values,
            positions=[1, 2],
            widths=0.48,
            showfliers=False,
            patch_artist=True,
            medianprops={"color": "white", "lw": 1.3},
            whiskerprops={"color": DARK_GRAY, "lw": 0.75},
            capprops={"color": DARK_GRAY, "lw": 0.75},
            boxprops={"edgecolor": DARK_GRAY, "lw": 0.75},
        )
        for patch, color in zip(box["boxes"], [GENERIC_BLUE, PERSONAL_ORANGE]):
            patch.set_facecolor(color)
            patch.set_alpha(0.84)
        for position, group, color in zip(
            [1, 2],
            values,
            [GENERIC_BLUE, PERSONAL_ORANGE],
        ):
            jitter = rng.uniform(-0.09, 0.09, size=len(group))
            axis.scatter(
                position + jitter,
                group,
                s=13,
                facecolors="white",
                edgecolors=color,
                linewidths=0.7,
                alpha=0.95,
                zorder=3,
            )
            axis.scatter(
                position,
                float(np.mean(group)),
                marker="D",
                s=24,
                facecolors=color,
                edgecolors="white",
                linewidths=0.6,
                zorder=4,
            )
        role = role_map[subject]
        axis.set_title(
            short_subject(subject) + role_suffix(role),
            fontsize=8.2,
            weight="bold" if role in {"best", "worst"} else "normal",
            pad=3,
        )
        axis.set_xticks([1, 2])
        axis.set_xticklabels(["Generic", "Personalized"], fontsize=6.8)
        axis.set_ylim(*limits)
        axis.grid(axis="y", color=GRID_GRAY, lw=0.55)
        axis.set_axisbelow(True)
        if index % 4 != 0:
            axis.tick_params(axis="y", labelleft=False)
    axes.flat[-1].axis("off")
    axes.flat[-1].legend(
        handles=[
            Line2D(
                [0],
                [0],
                marker="s",
                color=GENERIC_BLUE,
                markerfacecolor=GENERIC_BLUE,
                lw=0,
                label="Generic",
            ),
            Line2D(
                [0],
                [0],
                marker="s",
                color=PERSONAL_ORANGE,
                markerfacecolor=PERSONAL_ORANGE,
                lw=0,
                label="Personalized",
            ),
            Line2D(
                [0],
                [0],
                marker="D",
                color=DARK_GRAY,
                markerfacecolor=DARK_GRAY,
                lw=0,
                label="Repeat mean",
            ),
        ],
        loc="center",
        frameon=False,
        title="Encoding",
        title_fontsize=8,
    )
    figure.suptitle(
        f"{ROI_LABELS[roi]}: technical-repeat minimum target fields",
        fontsize=10.3,
        weight="bold",
        y=0.965,
    )
    figure.supylabel(
        "Minimum optimizer-target field (V/m)",
        x=0.018,
        fontsize=8.5,
    )
    stem = next(item["stem"] for item in REPEAT_FIGURES if item["roi"] == roi)
    save_figure(figure, figures_dir, stem)


def metric_change_counts(paired: pd.DataFrame, threshold: float) -> dict[str, int]:
    token = threshold_token(threshold)
    return {
        "target_increased": int(
            (
                paired[
                    f"target_coverage_percent_ge_{token}__absolute_change"
                ]
                > 0
            ).sum()
        ),
        "off_target_decreased": int(
            (
                paired[
                    f"off_target_coverage_percent_ge_{token}__absolute_change"
                ]
                < 0
            ).sum()
        ),
        "localization_increased": int(
            (
                paired[
                    "threshold_localization_percent_in_roi_ge_"
                    f"{token}__absolute_change"
                ]
                > 0
            ).sum()
        ),
    }


def build_key_change_table(paired: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "pair_index",
        "subject",
        "roi",
        "selection_role",
    ]
    result = paired[columns].copy()
    result["arrow_line_style"] = result["selection_role"].map(
        {
            "best": "dotted",
            "worst": "dashed",
            "cross_target": "solid",
        }
    )
    for metric in [
        "roi_min_v_per_m",
        "roi_median_v_per_m",
        "target_coverage_percent_ge_0p18",
        "off_target_coverage_percent_ge_0p18",
        "threshold_localization_percent_in_roi_ge_0p18",
    ]:
        for suffix in [
            "generic_repeat_mean",
            "generic_repeat_sd",
            "personalized_repeat_mean",
            "personalized_repeat_sd",
            "absolute_change",
        ]:
            source = f"{metric}__{suffix}"
            result[source] = paired[source].to_numpy()
    return result


def build_roi_summary_table(paired: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for roi in ROI_ORDER:
        subset = paired.loc[paired["roi"] == roi]
        row: dict[str, object] = {
            "roi": roi,
            "subjects": len(subset),
            "best_subject": subset.loc[
                subset["selection_role"] == "best", "subject"
            ].iloc[0],
            "worst_subject": subset.loc[
                subset["selection_role"] == "worst", "subject"
            ].iloc[0],
        }
        for metric in [
            "roi_min_v_per_m",
            "roi_median_v_per_m",
            "target_coverage_percent_ge_0p18",
            "off_target_coverage_percent_ge_0p18",
            "threshold_localization_percent_in_roi_ge_0p18",
        ]:
            change = subset[f"{metric}__absolute_change"]
            row[f"{metric}__median_change"] = float(change.median())
            row[f"{metric}__minimum_change"] = float(change.min())
            row[f"{metric}__maximum_change"] = float(change.max())
            row[f"{metric}__positive_change_count"] = int((change > 0).sum())
            row[f"{metric}__negative_change_count"] = int((change < 0).sum())
        rows.append(row)
    return pd.DataFrame(rows)


def build_captions(paired: pd.DataFrame) -> dict[str, str]:
    common = (
        "Seven optimized CamCan participants were evaluated for each of four "
        "temporal-interference targets (left hippocampus, left primary motor "
        "cortex, right dorsolateral prefrontal cortex, and right thalamus). "
        "For every subject-target configuration, an MNI152-derived generic "
        "montage and a subject- and target-specific Pareto montage were "
        "simulated on the same corrected anatomical head. The primary target "
        "was a sphere centered on the anatomical parcel volume centroid and "
        "clipped to that parcel (100 mm3 for cortical targets; 200 mm3 for "
        "subcortical targets). Metrics were calculated independently in ten "
        "remeshing repeats and then arithmetic-mean aggregated within each "
        "condition; repeat numbers were not paired between conditions. The "
        "seven subjects were selected as generic best or worst cases for at "
        "least one target, so this is a descriptive comparison rather than "
        "population inference. "
    )
    captions = {
        FIGURE_SUMMARY["stem"]: (
            "Figure 1. Generic-versus-personalized changes across all 28 "
            "subject-target configurations. "
            + common
            + "Rows correspond to targets and columns show the minimum "
            "electric-field magnitude in the optimizer-matched target, the "
            "target median, target coverage at or above 0.18 V/m, and "
            "off-target coverage at or above 0.18 V/m. Open blue circles "
            "denote generic condition means, filled orange circles denote "
            "personalized condition means, horizontal gray segments connect "
            "the two conditions for the same subject and target, and error "
            "bars show sample SD across ten technical repeats. Labels identify "
            "the ROI-specific generic-best and generic-worst subjects; the "
            "other five subjects in each row were originally selected for "
            "another target. Coverage is the percentage of target voxels, or "
            "finite brain voxels outside the target, meeting the threshold."
        )
    }
    for item in ARROW_FIGURES:
        threshold = float(item["threshold"])
        counts = metric_change_counts(paired, threshold)
        captions[item["stem"]] = (
            f"{item['designation']}. Generic-versus-personalized "
            f"effectiveness-spread changes at {threshold:.2f} V/m. "
            + common
            + f"Horizontal position is the percentage of optimizer-target "
            f"voxels with field magnitude at or above {threshold:.2f} V/m; "
            f"vertical position is the percentage of finite brain voxels "
            f"outside that target meeting the same threshold. Each arrow "
            f"starts at the generic montage (open circle) and ends at the "
            f"personalized montage (filled circle). Arrow color identifies "
            f"the subject. Within each ROI, the generic-best subject is "
            f"dotted, the generic-worst subject is dashed, and the other "
            f"five optimized subjects are solid. Rightward movement denotes "
            f"greater target coverage and downward movement denotes less "
            f"off-target spread; neither axis alone defines overall "
            f"superiority. Across the 28 configurations, target coverage "
            f"increased in {counts['target_increased']}, off-target coverage "
            f"decreased in {counts['off_target_decreased']}, and "
            f"suprathreshold localization inside the target increased in "
            f"{counts['localization_increased']}."
        )
    for item in REPEAT_FIGURES:
        roi = str(item["roi"])
        captions[item["stem"]] = (
            f"{item['designation']}. Technical-repeat distributions of the "
            f"minimum optimizer-target electric field for "
            f"{ROI_LABELS[roi]}. "
            + common
            + "Each subject panel compares the generic and personalized "
            "conditions. Small open points are individual remeshing repeats; "
            "boxes show the interquartile range, internal white lines show "
            "the repeat median, whiskers extend to the most extreme value "
            "within 1.5 interquartile ranges, and diamonds show the arithmetic "
            "mean used in the condition comparison. All subject panels share "
            "the same y-axis scale within this figure. Best and worst labels "
            "refer only to ranking under the generic montage for this ROI."
        )
    return captions


def write_caption_files(output_dir: Path, captions: dict[str, str]) -> None:
    captions_dir = output_dir / "captions"
    captions_dir.mkdir(parents=True, exist_ok=True)
    markdown = ["# Self-contained figure captions", ""]
    csv_rows = []
    for figure in FIGURES:
        stem = figure["stem"]
        caption = captions[stem]
        (captions_dir / f"{stem}_caption.txt").write_text(
            caption + "\n",
            encoding="utf-8",
        )
        markdown.extend([f"## {figure['designation']}", "", caption, ""])
        csv_rows.append(
            {
                "designation": figure["designation"],
                "figure_file": f"figures/{stem}.png",
                "caption": caption,
            }
        )
    (output_dir / "figure_captions.md").write_text(
        "\n".join(markdown),
        encoding="utf-8",
    )
    with (output_dir / "figure_captions.csv").open(
        "w",
        newline="",
        encoding="utf-8",
    ) as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["designation", "figure_file", "caption"],
        )
        writer.writeheader()
        writer.writerows(csv_rows)


def write_results_highlights(output_dir: Path, paired: pd.DataFrame) -> None:
    lines = [
        "# Optimizer-matched personalized versus generic results",
        "",
        "## Scope",
        "",
        "- 28 configurations: seven optimized subjects x four target ROIs.",
        "- Generic and personalized montages were evaluated on the same "
        "corrected head, with ten independently remeshed simulations per "
        "condition.",
        "- The primary ROI is the MakeROIs.m-equivalent, parcel-clipped "
        "optimizer sphere; full anatomical-parcel metrics remain available "
        "under the `anatomical_` prefix in the source results.",
        "- All results are descriptive because the seven participants were "
        "selected as generic outcome extremes for one or more ROIs.",
        "",
        "## Directional summary",
        "",
    ]
    for threshold in [0.18, 0.20, 0.15]:
        counts = metric_change_counts(paired, threshold)
        lines.append(
            f"- At {threshold:.2f} V/m: target coverage increased in "
            f"{counts['target_increased']}/28 configurations, off-target "
            f"coverage decreased in {counts['off_target_decreased']}/28, and "
            f"target localization increased in "
            f"{counts['localization_increased']}/28."
        )
    lines.extend(
        [
            "- Rightward-and-downward arrows combine greater target coverage "
            "with less off-target spread. Other directions represent "
            "trade-offs and should be described rather than collapsed into "
            "a single claim of improvement.",
            "- The minimum target field is the evaluation quantity closest "
            "to the optimizer's TI_free.Emin objective. The target median is "
            "retained as a complementary distributional summary.",
            "",
            "## Arrow-plot encoding",
            "",
            "- Subject identity is encoded by color.",
            "- ROI-specific generic-best subject: dotted arrow.",
            "- ROI-specific generic-worst subject: dashed arrow.",
            "- All other optimized subjects: solid arrows.",
            "- Open circle: generic condition; filled circle: personalized "
            "condition.",
            "",
            "Exact values are provided in "
            "`tables/table_all_subject_key_changes.csv` and "
            "`tables/table_roi_directional_summary.csv`.",
            "",
        ]
    )
    (output_dir / "results_highlights.md").write_text(
        "\n".join(lines),
        encoding="utf-8",
    )


def figure_interpretation(figure: dict) -> tuple[str, str]:
    if figure["stem"] == FIGURE_SUMMARY["stem"]:
        return (
            "Read each row as one anatomical target and each labeled line as "
            "one subject. The left two columns describe field magnitude "
            "inside the optimizer-matched sphere; the right two describe "
            "target engagement and off-target spread at 0.18 V/m. A "
            "rightward condition shift is an increase in the plotted metric, "
            "which is favorable for target magnitude and target coverage but "
            "unfavorable for off-target coverage.",
            "Use this figure to show the complete seven-subject comparison, "
            "not to estimate a population-average treatment effect. The "
            "best/worst labels are ROI-specific ranks under the generic "
            "montage and do not rank the personalized results.",
        )
    if "threshold" in figure:
        threshold = float(figure["threshold"])
        return (
            f"Each panel plots target coverage against off-target coverage at "
            f"{threshold:.2f} V/m. Follow an arrow from its open generic start "
            f"to its filled personalized endpoint. Right-and-down combines "
            f"more target engagement with less off-target spread. Right-and-up "
            f"and left-and-down are trade-offs; left-and-up is unfavorable on "
            f"both plotted measures.",
            "Arrow color identifies subjects across panels. Line style is "
            "reassigned within each ROI: dotted marks that ROI's generic-best "
            "subject, dashed marks its generic-worst subject, and solid marks "
            "the five subjects selected because they were extreme for another "
            "ROI. Panel-specific y-axis limits make within-ROI movement "
            "visible, so compare numerical axes rather than apparent arrow "
            "length across panels.",
        )
    return (
        "Each small panel shows the ten minimum target-field values behind "
        "the generic and personalized condition means for one subject. The "
        "box summarizes the repeat distribution and the diamond is the "
        "arithmetic mean used in the comparison.",
        "This is a technical-repeat QC figure. It separates shifts between "
        "montages from variability introduced by independent remeshing. All "
        "subject panels in the same ROI figure share a y-axis scale.",
    )


def build_markdown_guide(
    output_dir: Path,
    captions: dict[str, str],
) -> None:
    lines = [
        "# Supervisor figure interpretation guide",
        "",
        "## Analysis definition",
        "",
        "The figures compare an MNI152-derived generic montage with a "
        "subject- and ROI-specific personalized Pareto montage on the same "
        "corrected CamCan head. All seven optimized participants are shown "
        "for all four ROIs. The primary evaluation target is a parcel-clipped "
        "sphere centered on the anatomical parcel volume centroid (100 mm3 "
        "for cortical targets and 200 mm3 for subcortical targets). Each "
        "metric was calculated per repeat and then averaged across ten "
        "independent remeshing repeats within condition.",
        "",
        "The analysis is descriptive. The seven subjects were selected as "
        "generic best or worst cases for at least one ROI and are not a "
        "random or representative sample.",
        "",
    ]
    for figure in FIGURES:
        how, guardrail = figure_interpretation(figure)
        lines.extend(
            [
                f"## {figure['designation']}: {figure['title']}",
                "",
                f"File: `figures/{figure['stem']}.png`",
                "",
                "### How to read it",
                "",
                how,
                "",
                "### Interpretation guardrail",
                "",
                guardrail,
                "",
                "### Self-contained caption",
                "",
                captions[figure["stem"]],
                "",
            ]
        )
    (output_dir / "Figure_Interpretation_Guide.md").write_text(
        "\n".join(lines),
        encoding="utf-8",
    )


def add_document_title(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(45)
    run = paragraph.add_run(
        "Optimizer-matched Personalized versus Generic\n"
        "Temporal-Interference Stimulation"
    )
    run.font.name = "Arial"
    run.font.size = Pt(25)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Supervisor figure review guide, computation notes, and "
        "self-contained captions"
    ).font.size = Pt(13)
    rule = document.add_table(rows=1, cols=1)
    rule.alignment = 1
    set_cell_shading(rule.cell(0, 0), BLUE)
    rule.cell(0, 0).text = ""
    scope = document.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_before = Pt(18)
    scope.add_run(
        "Seven optimized subjects x four ROIs x two montage conditions\n"
        "Ten independently remeshed simulations per condition\n"
        "Descriptive analysis of an outcome-selected sample"
    )
    prepared = document.add_paragraph()
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared.paragraph_format.space_before = Pt(18)
    prepared_run = prepared.add_run(
        f"Prepared {date.today().strftime('%d %B %Y')}"
    )
    prepared_run.font.size = Pt(9)
    prepared_run.font.color.rgb = RGBColor.from_string(GRAY)


def build_review_guide(
    output_dir: Path,
    captions: dict[str, str],
    paired: pd.DataFrame,
) -> Path:
    figures_dir = output_dir / "figures"
    output = output_dir / "Optimizer_Matched_Figure_Review_Guide.docx"
    document = Document()
    configure_document(document)
    add_document_title(document)

    document.add_page_break()
    add_heading(document, "1. Study question and analysis scope")
    document.add_paragraph(
        "This analysis asks how target-specific personalized Pareto montages "
        "change target engagement and off-target spread relative to the "
        "MNI152-derived generic montage when both are simulated on the same "
        "corrected anatomical head."
    )
    add_callout(
        document,
        "Primary ROI",
        "The primary target is the MakeROIs.m-equivalent sphere centered on "
        "the anatomical parcel volume centroid and clipped to that parcel. "
        "Requested volume is 100 mm3 for cortical M1 and DLPFC targets and "
        "200 mm3 for subcortical hippocampal and thalamic targets.",
        PALE_BLUE,
    )
    add_callout(
        document,
        "Inference boundary",
        "All seven participants were selected because they were the generic "
        "best or worst case for at least one target. The 28 subject-target "
        "configurations are therefore descriptive and cannot estimate a "
        "population-average personalization effect.",
        PALE_ORANGE,
    )
    for text in [
        "Every subject was evaluated for all four targets.",
        "The personalized montage used the ROI-specific electrode pairs and "
        "current amplitudes extracted from that subject's optimization MAT "
        "file.",
        "Metrics were calculated independently for each of ten remeshing "
        "repeats and then arithmetic-mean aggregated within condition.",
        "Generic and personalized repeat indices were not statistically "
        "paired.",
    ]:
        add_bullet(document, text)
    add_table(
        document,
        ["Arrow property", "Meaning"],
        [
            ["Color", "Subject identity, held constant across all ROI panels."],
            ["Dotted line", "ROI-specific generic-best subject."],
            ["Dashed line", "ROI-specific generic-worst subject."],
            ["Solid line", "Other optimized subject in that ROI."],
            ["Open circle", "Generic montage condition mean."],
            ["Filled circle", "Personalized montage condition mean."],
        ],
        widths_cm=[5.5, 16.0],
    )

    document.add_page_break()
    add_heading(document, "2. Exact metric definitions")
    add_table(
        document,
        ["Metric", "Definition"],
        [
            [
                "Minimum optimizer-target field",
                "Minimum finite TI-field magnitude among voxels in the "
                "parcel-clipped optimizer sphere; closest evaluation metric "
                "to the optimizer's TI_free.Emin objective.",
            ],
            [
                "Median optimizer-target field",
                "Median finite TI-field magnitude among voxels in the same "
                "optimizer sphere.",
            ],
            [
                "Target coverage at threshold t",
                "100 x optimizer-target voxels at or above t / all "
                "optimizer-target voxels.",
            ],
            [
                "Off-target coverage at threshold t",
                "100 x finite brain voxels outside the optimizer target at "
                "or above t / all finite brain voxels outside the target.",
            ],
            [
                "Localization in target at threshold t",
                "100 x suprathreshold target voxels / all suprathreshold "
                "finite brain voxels.",
            ],
        ],
        widths_cm=[6.4, 15.2],
    )
    document.add_paragraph(
        "The source CSVs also retain full anatomical-parcel metrics under "
        "the `anatomical_` prefix. Those metrics are secondary and are not "
        "used for the primary publication figures in this package."
    )

    counts_rows = []
    for threshold in [0.18, 0.20, 0.15]:
        counts = metric_change_counts(paired, threshold)
        counts_rows.append(
            [
                f"{threshold:.2f}",
                f"{counts['target_increased']}/28",
                f"{counts['off_target_decreased']}/28",
                f"{counts['localization_increased']}/28",
            ]
        )
    add_heading(document, "Directional overview", level=2)
    add_table(
        document,
        [
            "Threshold (V/m)",
            "Target coverage increased",
            "Off-target coverage decreased",
            "Localization increased",
        ],
        counts_rows,
        widths_cm=[4.0, 5.6, 5.9, 5.2],
    )

    section = 3
    for figure in FIGURES:
        document.add_page_break()
        add_heading(
            document,
            f"{section}. {figure['designation']}: {figure['title']}",
        )
        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.add_run().add_picture(
            str(figures_dir / f"{figure['stem']}.png"),
            height=Inches(4.85),
        )
        how, guardrail = figure_interpretation(figure)
        add_heading(document, "How to read it", level=2)
        document.add_paragraph(how)
        add_callout(
            document,
            "Interpretation guardrail",
            guardrail,
            PALE_ORANGE,
        )
        document.add_page_break()
        add_heading(
            document,
            f"{figure['designation']}: self-contained caption",
        )
        paragraph = document.add_paragraph(captions[figure["stem"]])
        paragraph.style = document.styles["Caption"]
        section += 1

    document.add_page_break()
    add_heading(document, f"{section}. Reproducibility notes")
    for text in [
        "Source: validated schema-v2 output of "
        "camcan_personalized_comparison.py.",
        "Presentation layer: "
        "build_camcan_optimizer_matched_publication_outputs.py.",
        "Source simulation images were read-only and no image-level metric "
        "was recomputed by this presentation layer.",
        "The package manifest records SHA256 hashes for every generated and "
        "copied file.",
    ]:
        add_bullet(document, text)
    document.save(output)
    return output


def convert_docx_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    with tempfile.TemporaryDirectory(prefix="optimizer_matched_guide_") as temp:
        temporary = Path(temp)
        profile = temporary / "profile"
        runtime = temporary / "runtime"
        home = temporary / "home"
        profile.mkdir()
        runtime.mkdir()
        home.mkdir()
        command = [
            "soffice",
            "--headless",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            temp,
            str(docx_path),
        ]
        environment = os.environ.copy()
        environment.update(
            {
                "HOME": str(home),
                "XDG_CACHE_HOME": str(temporary / "cache"),
                "XDG_CONFIG_HOME": str(temporary / "config"),
                "XDG_RUNTIME_DIR": str(runtime),
            }
        )
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            env=environment,
        )
        if completed.returncode != 0:
            raise RuntimeError(
                "LibreOffice conversion failed: "
                + completed.stdout
                + completed.stderr
            )
        source = temporary / f"{docx_path.stem}.pdf"
        if not source.is_file():
            raise RuntimeError("LibreOffice did not create the expected PDF")
        destination = output_dir / source.name
        shutil.copy2(source, destination)
    return destination


def copy_source_tables(input_dir: Path, tables_dir: Path) -> None:
    for filename in [
        "table_main_selected_metrics.csv",
        "table_stimulation_parameters.csv",
        "table_optimizer_roi_definitions.csv",
        "selection_allowlist.csv",
        "condition_repeat_mean_metrics.csv",
    ]:
        shutil.copy2(input_dir / filename, tables_dir / filename)


def write_manifest_and_archive(
    output_dir: Path,
    input_dir: Path,
    source_manifest: dict,
) -> None:
    excluded = {
        "publication_output_manifest.json",
        "optimizer_matched_manuscript_package.zip",
        "optimizer_matched_manuscript_package.zip.sha256",
    }
    files = []
    for path in sorted(item for item in output_dir.rglob("*") if item.is_file()):
        relative = path.relative_to(output_dir).as_posix()
        if relative in excluded:
            continue
        files.append(
            {
                "path": relative,
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }
        )
    manifest = {
        "publication_package_schema_version": 2,
        "status": "complete",
        "created": date.today().isoformat(),
        "source_directory": str(input_dir.resolve()),
        "source_comparison_schema_version": source_manifest.get(
            "comparison_schema_version"
        ),
        "subject_roi_configurations": 28,
        "unique_subjects": 7,
        "repeat_level_records": 560,
        "primary_roi": (
            "MakeROIs.m-equivalent parcel-clipped optimizer sphere centered "
            "on anatomical parcel volume centroid"
        ),
        "arrow_role_line_styles": {
            "best": "dotted",
            "worst": "dashed",
            "cross_target": "solid",
        },
        "inference": "descriptive outcome-selected seven-subject analysis",
        "files": files,
    }
    manifest_path = output_dir / "publication_output_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")

    archive = output_dir / "optimizer_matched_manuscript_package.zip"
    with zipfile.ZipFile(
        archive,
        "w",
        compression=zipfile.ZIP_DEFLATED,
        compresslevel=6,
    ) as handle:
        for path in sorted(item for item in output_dir.rglob("*") if item.is_file()):
            if path == archive or path.name == f"{archive.name}.sha256":
                continue
            handle.write(path, path.relative_to(output_dir))
    (output_dir / f"{archive.name}.sha256").write_text(
        f"{sha256_file(archive)}  {archive.name}\n"
    )


def build(input_dir: Path, output_dir: Path, *, force: bool) -> dict:
    input_dir = input_dir.resolve()
    output_dir = output_dir.resolve()
    if output_dir.exists():
        if not force:
            raise FileExistsError(
                f"Output directory exists: {output_dir}; pass --force to replace it"
            )
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True)
    figures_dir = output_dir / "figures"
    tables_dir = output_dir / "tables"
    figures_dir.mkdir()
    tables_dir.mkdir()

    manifest, paired, condition, repeats = load_and_validate(input_dir)
    set_figure_style()
    plot_all_subject_changes(paired, figures_dir)
    for threshold in [0.18, 0.20, 0.15]:
        plot_tradeoff(paired, figures_dir, threshold)
    for roi in ROI_ORDER:
        plot_repeat_distribution_for_roi(repeats, figures_dir, roi)

    key_changes = build_key_change_table(paired)
    key_changes.to_csv(
        tables_dir / "table_all_subject_key_changes.csv",
        index=False,
    )
    roi_summary = build_roi_summary_table(paired)
    roi_summary.to_csv(
        tables_dir / "table_roi_directional_summary.csv",
        index=False,
    )
    copy_source_tables(input_dir, tables_dir)

    captions = build_captions(paired)
    write_caption_files(output_dir, captions)
    write_results_highlights(output_dir, paired)
    build_markdown_guide(output_dir, captions)
    docx_path = build_review_guide(output_dir, captions, paired)
    pdf_path = convert_docx_to_pdf(docx_path, output_dir)
    write_manifest_and_archive(output_dir, input_dir, manifest)

    result = {
        "status": "complete",
        "source": str(input_dir),
        "output": str(output_dir),
        "subject_roi_configurations": len(paired),
        "figures_png": len(list(figures_dir.glob("*.png"))),
        "figures_pdf": len(list(figures_dir.glob("*.pdf"))),
        "tables": len(list(tables_dir.glob("*.csv"))),
        "captions": len(captions),
        "review_guide_docx": str(docx_path),
        "review_guide_pdf": str(pdf_path),
        "archive": str(output_dir / "optimizer_matched_manuscript_package.zip"),
    }
    print(json.dumps(result, indent=2))
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument(
        "--force",
        action="store_true",
        help="replace an existing output directory",
    )
    return parser.parse_args()


def main() -> int:
    arguments = parse_args()
    build(arguments.input_dir, arguments.out_dir, force=arguments.force)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
