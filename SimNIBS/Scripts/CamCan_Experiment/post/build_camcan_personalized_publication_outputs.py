#!/usr/bin/env python3
"""Build publication outputs for the selected personalized-vs-generic cases.

This is a presentation layer over the validated outputs produced by
``camcan_personalized_comparison.py``.  It never reads or modifies source
simulation images.  The eight subject/ROI cases were selected as the highest
and lowest generic median target-ROI fields in each of four ROIs, so all
comparisons are descriptive and must not be treated as population inference.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import os
import shutil
import subprocess
import tempfile
import zipfile
from datetime import date
from pathlib import Path

import matplotlib as mpl
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from CamCan_Experiment.post.build_camcan_supervisor_figure_guide import (
    BLACK,
    BLUE,
    GRAY,
    NAVY,
    ORANGE,
    PALE_BLUE,
    PALE_ORANGE,
    add_bullet,
    add_callout,
    add_heading,
    add_number,
    add_table,
    configure_document,
    set_cell_shading,
)


ROI_ORDER = [
    "Left_Hippocampus",
    "Left_M1",
    "Right_DLPC",
    "Right_Thalamus",
]
ROI_LABELS = {
    "Left_Hippocampus": "Left hippocampus",
    "Left_M1": "Left M1",
    "Right_DLPC": "Right DLPFC",
    "Right_Thalamus": "Right thalamus",
}
ROLE_ORDER = ["best", "worst"]
CONDITION_ORDER = ["generic", "personalized"]

GENERIC_BLUE = "#2F6B9A"
PERSONAL_ORANGE = "#D97706"
BEST_GREEN = "#168A76"
WORST_PURPLE = "#7C5BA7"
LIGHT_GRAY = "#D4D8DE"
DARK_GRAY = "#4B5563"
GRID_GRAY = "#E6E9ED"

FIGURES = [
    {
        "designation": "Figure 1",
        "stem": "figure_personalization_selected_case_changes",
        "title": "Selected-case changes with personalized stimulation",
        "one_line": (
            "Generic and personalized stimulation are compared on the same "
            "subject head; points are means and error bars are sample SDs "
            "across ten independent remeshing repeats."
        ),
    },
    {
        "designation": "Figure 2",
        "stem": "figure_personalization_effectiveness_spread_ge_0p18",
        "title": "Effectiveness-spread trade-offs at 0.18 V/m",
        "one_line": (
            "Each arrow starts at the generic montage and ends at the "
            "personalized montage for one selected subject/ROI case."
        ),
    },
    {
        "designation": "Supplementary Figure S1",
        "stem": "figure_personalization_effectiveness_spread_ge_0p20",
        "title": "Effectiveness-spread sensitivity analysis at 0.20 V/m",
        "one_line": (
            "The same selected-case trade-off analysis is repeated using "
            "the optimizer-aligned 0.20 V/m evaluation threshold."
        ),
    },
    {
        "designation": "Supplementary Figure S2",
        "stem": "figure_personalization_effectiveness_spread_ge_0p15",
        "title": "Effectiveness-spread sensitivity analysis at 0.15 V/m",
        "one_line": (
            "The same selected-case trade-off analysis is repeated using "
            "the more permissive 0.15 V/m threshold."
        ),
    },
    {
        "designation": "Supplementary Figure S3",
        "stem": "figure_personalization_technical_repeat_distributions",
        "title": "Technical-repeat distributions of median target field",
        "one_line": (
            "Each panel shows all ten independently remeshed repeat values "
            "for one selected subject/ROI case and condition."
        ),
    },
]


def threshold_token(threshold: float) -> str:
    return str(threshold).replace(".", "p")


def threshold_file_token(threshold: float) -> str:
    return f"{threshold:.2f}".replace(".", "p")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_figure_style() -> None:
    mpl.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 8.5,
            "axes.titlesize": 10,
            "axes.labelsize": 9,
            "xtick.labelsize": 8,
            "ytick.labelsize": 8,
            "legend.fontsize": 8,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "axes.linewidth": 0.7,
            "xtick.major.width": 0.7,
            "ytick.major.width": 0.7,
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
            "savefig.facecolor": "white",
        }
    )


def save_figure(figure: plt.Figure, figures_dir: Path, stem: str) -> None:
    figure.savefig(
        figures_dir / f"{stem}.png",
        dpi=300,
        bbox_inches="tight",
        pad_inches=0.035,
    )
    figure.savefig(
        figures_dir / f"{stem}.pdf",
        bbox_inches="tight",
        pad_inches=0.035,
    )
    plt.close(figure)


def load_and_validate(input_dir: Path) -> tuple[dict, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    required = [
        "analysis_manifest.json",
        "paired_personalized_vs_generic.csv",
        "condition_repeat_mean_metrics.csv",
        "repeat_level_metrics.csv",
        "table_main_selected_metrics.csv",
        "table_stimulation_parameters.csv",
        "selection_allowlist.csv",
    ]
    missing = [name for name in required if not (input_dir / name).is_file()]
    if missing:
        raise FileNotFoundError("Missing comparison output(s): " + ", ".join(missing))

    manifest = json.loads((input_dir / "analysis_manifest.json").read_text())
    paired = pd.read_csv(input_dir / "paired_personalized_vs_generic.csv")
    condition = pd.read_csv(input_dir / "condition_repeat_mean_metrics.csv")
    repeats = pd.read_csv(input_dir / "repeat_level_metrics.csv")

    expected_manifest = {
        "status": "complete",
        "selected_subject_roi_pairs": 8,
        "unique_subjects": 7,
        "repeat_level_records": 160,
        "condition_repeat_mean_records": 16,
    }
    for key, expected in expected_manifest.items():
        if manifest.get(key) != expected:
            raise RuntimeError(
                f"Unexpected manifest value {key}={manifest.get(key)!r}; "
                f"expected {expected!r}"
            )
    if manifest.get("thresholds_v_per_m") != [0.2, 0.18, 0.15]:
        raise RuntimeError("Expected thresholds [0.2, 0.18, 0.15] V/m")
    if set(paired["roi"]) != set(ROI_ORDER) or len(paired) != 8:
        raise RuntimeError("Expected two selected cases for each of four ROIs")
    if set(paired["selection_role"]) != set(ROLE_ORDER):
        raise RuntimeError("Expected best and worst selection roles")
    if len(condition) != 16 or set(condition["condition"]) != set(CONDITION_ORDER):
        raise RuntimeError("Expected 16 condition means across two conditions")
    if len(repeats) != 160:
        raise RuntimeError("Expected 160 repeat-level records")
    if repeats.duplicated(["pair_index", "condition", "repeat"]).any():
        raise RuntimeError("Duplicate pair/condition/repeat records")
    counts = repeats.groupby(["pair_index", "condition"]).size()
    if not (counts == 10).all():
        raise RuntimeError("Every selected case/condition must have ten repeats")
    plotted_metrics = [
        "roi_median_v_per_m",
        "target_coverage_percent_ge_0p2",
        "off_target_coverage_percent_ge_0p2",
        "threshold_localization_percent_in_roi_ge_0p2",
        "target_coverage_percent_ge_0p18",
        "off_target_coverage_percent_ge_0p18",
        "threshold_localization_percent_in_roi_ge_0p18",
        "target_coverage_percent_ge_0p15",
        "off_target_coverage_percent_ge_0p15",
        "threshold_localization_percent_in_roi_ge_0p15",
    ]
    paired_required = []
    for metric in plotted_metrics:
        paired_required.extend(
            [
                f"{metric}__generic_repeat_mean",
                f"{metric}__generic_repeat_sd",
                f"{metric}__personalized_repeat_mean",
                f"{metric}__personalized_repeat_sd",
                f"{metric}__absolute_change",
            ]
        )
    checks = [
        ("paired plotted metrics", paired, paired_required),
        ("condition plotted metrics", condition, plotted_metrics),
        ("repeat plotted metrics", repeats, plotted_metrics),
    ]
    for name, frame, columns in checks:
        missing_columns = sorted(set(columns) - set(frame.columns))
        if missing_columns:
            raise RuntimeError(f"{name} missing columns: {missing_columns}")
        numeric = frame[columns].to_numpy(dtype=float)
        if not np.isfinite(numeric).all():
            raise RuntimeError(f"{name} contains non-finite numeric values")

    roi_rank = {roi: index for index, roi in enumerate(ROI_ORDER)}
    role_rank = {role: index for index, role in enumerate(ROLE_ORDER)}
    paired = paired.assign(
        _roi_order=paired["roi"].map(roi_rank),
        _role_order=paired["selection_role"].map(role_rank),
    ).sort_values(["_roi_order", "_role_order"])
    condition = condition.assign(
        _roi_order=condition["roi"].map(roi_rank),
        _role_order=condition["selection_role"].map(role_rank),
    ).sort_values(["_roi_order", "_role_order", "condition"])
    repeats = repeats.assign(
        _roi_order=repeats["roi"].map(roi_rank),
        _role_order=repeats["selection_role"].map(role_rank),
    ).sort_values(["_roi_order", "_role_order", "condition", "repeat"])
    return manifest, paired, condition, repeats


def case_label(row: pd.Series, *, include_subject: bool = True) -> str:
    role = str(row["selection_role"]).capitalize()
    roi = ROI_LABELS[str(row["roi"])]
    if include_subject:
        subject = str(row["subject"]).replace("sub-", "")
        return f"{roi} - {role}\n{subject}"
    return f"{roi} - {role}"


def plot_change_summary(paired: pd.DataFrame, figures_dir: Path) -> None:
    panels = [
        ("A", "roi_median_v_per_m", "Median target-ROI field (V/m)"),
        (
            "B",
            "target_coverage_percent_ge_0p18",
            "Target coverage >=0.18 V/m (%)",
        ),
        (
            "C",
            "off_target_coverage_percent_ge_0p18",
            "Off-target coverage >=0.18 V/m (%)",
        ),
        (
            "D",
            "threshold_localization_percent_in_roi_ge_0p18",
            "Localization in target >=0.18 V/m (%)",
        ),
    ]
    figure, axes = plt.subplots(2, 2, figsize=(7.35, 7.15), constrained_layout=True)
    y = np.arange(len(paired))[::-1]
    labels = [case_label(row) for _, row in paired.iterrows()]
    for index, (panel, metric, xlabel) in enumerate(panels):
        axis = axes.flat[index]
        generic = paired[f"{metric}__generic_repeat_mean"].to_numpy()
        personalized = paired[f"{metric}__personalized_repeat_mean"].to_numpy()
        generic_sd = paired[f"{metric}__generic_repeat_sd"].to_numpy()
        personalized_sd = paired[f"{metric}__personalized_repeat_sd"].to_numpy()
        for yi, start, end in zip(y, generic, personalized):
            axis.plot([start, end], [yi, yi], color=LIGHT_GRAY, lw=1.4, zorder=1)
        axis.errorbar(
            generic,
            y,
            xerr=generic_sd,
            fmt="o",
            ms=4.8,
            mfc="white",
            mec=GENERIC_BLUE,
            mew=1.25,
            ecolor=GENERIC_BLUE,
            elinewidth=0.8,
            capsize=1.8,
            label="Generic",
            zorder=3,
        )
        axis.errorbar(
            personalized,
            y,
            xerr=personalized_sd,
            fmt="o",
            ms=4.8,
            mfc=PERSONAL_ORANGE,
            mec=PERSONAL_ORANGE,
            mew=1.0,
            ecolor=PERSONAL_ORANGE,
            elinewidth=0.8,
            capsize=1.8,
            label="Personalized",
            zorder=4,
        )
        axis.set_xlabel(xlabel)
        axis.set_yticks(y)
        axis.set_yticklabels(labels if index % 2 == 0 else [])
        axis.grid(axis="x", color=GRID_GRAY, lw=0.55)
        axis.set_axisbelow(True)
        axis.text(
            -0.14,
            1.035,
            panel,
            transform=axis.transAxes,
            weight="bold",
            fontsize=11,
        )
    handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            color=GENERIC_BLUE,
            markerfacecolor="white",
            markeredgecolor=GENERIC_BLUE,
            lw=0.8,
            label="Generic",
        ),
        Line2D(
            [0],
            [0],
            marker="o",
            color=PERSONAL_ORANGE,
            markerfacecolor=PERSONAL_ORANGE,
            markeredgecolor=PERSONAL_ORANGE,
            lw=0.8,
            label="Personalized",
        ),
    ]
    figure.legend(
        handles=handles,
        loc="outside upper center",
        ncol=2,
        frameon=False,
        handletextpad=0.4,
        columnspacing=1.4,
    )
    save_figure(figure, figures_dir, "figure_personalization_selected_case_changes")


def _tradeoff_limits(frame: pd.DataFrame, threshold: float) -> tuple[float, float]:
    token = threshold_token(threshold)
    columns = [
        f"off_target_coverage_percent_ge_{token}__generic_repeat_mean",
        f"off_target_coverage_percent_ge_{token}__personalized_repeat_mean",
    ]
    maximum = float(frame[columns].to_numpy().max())
    if maximum <= 2:
        return 0.0, max(0.5, maximum * 1.22)
    if maximum <= 10:
        return 0.0, maximum * 1.18
    return 0.0, min(100.0, maximum * 1.10)


def plot_tradeoff(
    paired: pd.DataFrame,
    figures_dir: Path,
    threshold: float,
) -> None:
    token = threshold_token(threshold)
    x_metric = f"target_coverage_percent_ge_{token}"
    y_metric = f"off_target_coverage_percent_ge_{token}"
    figure, axes = plt.subplots(2, 2, figsize=(7.35, 6.5), constrained_layout=True)
    for panel_index, roi in enumerate(ROI_ORDER):
        axis = axes.flat[panel_index]
        subset = paired.loc[paired["roi"] == roi].sort_values("_role_order")
        for _, row in subset.iterrows():
            color = BEST_GREEN if row["selection_role"] == "best" else WORST_PURPLE
            x0 = float(row[f"{x_metric}__generic_repeat_mean"])
            y0 = float(row[f"{y_metric}__generic_repeat_mean"])
            x1 = float(row[f"{x_metric}__personalized_repeat_mean"])
            y1 = float(row[f"{y_metric}__personalized_repeat_mean"])
            axis.annotate(
                "",
                xy=(x1, y1),
                xytext=(x0, y0),
                arrowprops={
                    "arrowstyle": "-|>",
                    "color": color,
                    "lw": 1.35,
                    "mutation_scale": 10,
                    "shrinkA": 4,
                    "shrinkB": 4,
                },
                zorder=2,
            )
            axis.scatter(
                x0,
                y0,
                s=38,
                facecolors="white",
                edgecolors=color,
                linewidths=1.3,
                zorder=3,
            )
            axis.scatter(
                x1,
                y1,
                s=40,
                facecolors=color,
                edgecolors=color,
                linewidths=0.8,
                zorder=4,
            )
        axis.set_title(ROI_LABELS[roi], weight="bold", pad=4)
        axis.set_xlim(-2, 102)
        axis.set_ylim(*_tradeoff_limits(subset, threshold))
        axis.set_xlabel(f"Target coverage >={threshold:.2f} V/m (%)")
        axis.set_ylabel(f"Off-target coverage >={threshold:.2f} V/m (%)")
        axis.grid(color=GRID_GRAY, lw=0.55)
        axis.set_axisbelow(True)
        axis.text(
            -0.14,
            1.035,
            chr(ord("A") + panel_index),
            transform=axis.transAxes,
            weight="bold",
            fontsize=11,
        )
    handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            color="none",
            markerfacecolor="white",
            markeredgecolor=DARK_GRAY,
            label="Generic",
        ),
        Line2D(
            [0],
            [0],
            marker="o",
            color="none",
            markerfacecolor=DARK_GRAY,
            markeredgecolor=DARK_GRAY,
            label="Personalized",
        ),
        Line2D([0], [0], color=BEST_GREEN, lw=2, label="Generic-best case"),
        Line2D([0], [0], color=WORST_PURPLE, lw=2, label="Generic-worst case"),
    ]
    figure.legend(
        handles=handles,
        loc="outside lower center",
        ncol=4,
        frameon=False,
        handletextpad=0.45,
        columnspacing=1.2,
    )
    figure.suptitle(
        f"Effectiveness-spread changes after personalization ({threshold:.2f} V/m)",
        fontsize=11,
        weight="bold",
    )
    save_figure(
        figure,
        figures_dir,
        (
            "figure_personalization_effectiveness_spread_ge_"
            f"{threshold_file_token(threshold)}"
        ),
    )


def plot_repeat_distributions(repeats: pd.DataFrame, figures_dir: Path) -> None:
    figure, axes = plt.subplots(4, 2, figsize=(7.35, 8.8), constrained_layout=True)
    rng = np.random.default_rng(20260728)
    for row_index, roi in enumerate(ROI_ORDER):
        for column_index, role in enumerate(ROLE_ORDER):
            axis = axes[row_index, column_index]
            subset = repeats.loc[
                (repeats["roi"] == roi) & (repeats["selection_role"] == role)
            ]
            subject = str(subset["subject"].iloc[0]).replace("sub-", "")
            values = [
                subset.loc[
                    subset["condition"] == condition, "roi_median_v_per_m"
                ].to_numpy()
                for condition in CONDITION_ORDER
            ]
            positions = [1, 2]
            box = axis.boxplot(
                values,
                positions=positions,
                widths=0.46,
                showfliers=False,
                patch_artist=True,
                medianprops={"color": "white", "lw": 1.4},
                whiskerprops={"color": DARK_GRAY, "lw": 0.8},
                capprops={"color": DARK_GRAY, "lw": 0.8},
                boxprops={"edgecolor": DARK_GRAY, "lw": 0.8},
            )
            for patch, color in zip(box["boxes"], [GENERIC_BLUE, PERSONAL_ORANGE]):
                patch.set_facecolor(color)
                patch.set_alpha(0.82)
            for position, group, color in zip(
                positions,
                values,
                [GENERIC_BLUE, PERSONAL_ORANGE],
            ):
                jitter = rng.uniform(-0.085, 0.085, size=len(group))
                axis.scatter(
                    position + jitter,
                    group,
                    s=14,
                    facecolors="white",
                    edgecolors=color,
                    linewidths=0.75,
                    alpha=0.9,
                    zorder=3,
                )
                axis.scatter(
                    position,
                    np.mean(group),
                    marker="D",
                    s=26,
                    facecolors=color,
                    edgecolors="white",
                    linewidths=0.65,
                    zorder=4,
                )
            axis.set_xticks(positions)
            axis.set_xticklabels(["Generic", "Personalized"])
            axis.set_title(
                f"{ROI_LABELS[roi]} - {role.capitalize()} ({subject})",
                fontsize=8.8,
                weight="bold",
                pad=3,
            )
            axis.set_ylabel("Median target-ROI field (V/m)")
            axis.grid(axis="y", color=GRID_GRAY, lw=0.55)
            axis.set_axisbelow(True)
            axis.text(
                -0.16,
                1.03,
                chr(ord("A") + row_index * 2 + column_index),
                transform=axis.transAxes,
                weight="bold",
                fontsize=10,
            )
    figure.suptitle(
        "Technical-repeat distributions for selected subject/ROI cases",
        fontsize=11,
        weight="bold",
    )
    save_figure(
        figure,
        figures_dir,
        "figure_personalization_technical_repeat_distributions",
    )


def build_key_change_table(paired: pd.DataFrame) -> pd.DataFrame:
    metrics = [
        ("roi_median_v_per_m", "Median target-ROI field", "V/m"),
        (
            "target_coverage_percent_ge_0p18",
            "Target coverage >=0.18 V/m",
            "percentage points",
        ),
        (
            "off_target_coverage_percent_ge_0p18",
            "Off-target coverage >=0.18 V/m",
            "percentage points",
        ),
        (
            "threshold_localization_percent_in_roi_ge_0p18",
            "Localization in target >=0.18 V/m",
            "percentage points",
        ),
    ]
    rows: list[dict[str, object]] = []
    for _, source in paired.iterrows():
        for metric, label, unit in metrics:
            generic = float(source[f"{metric}__generic_repeat_mean"])
            personalized = float(source[f"{metric}__personalized_repeat_mean"])
            rows.append(
                {
                    "pair_index": int(source["pair_index"]),
                    "subject": source["subject"],
                    "roi": source["roi"],
                    "selection_role": source["selection_role"],
                    "metric": metric,
                    "metric_label": label,
                    "unit": unit,
                    "generic_repeat_mean": generic,
                    "generic_repeat_sd": float(
                        source[f"{metric}__generic_repeat_sd"]
                    ),
                    "personalized_repeat_mean": personalized,
                    "personalized_repeat_sd": float(
                        source[f"{metric}__personalized_repeat_sd"]
                    ),
                    "personalized_minus_generic": personalized - generic,
                }
            )
    return pd.DataFrame(rows)


def metric_change_counts(paired: pd.DataFrame, threshold: float) -> dict[str, int]:
    token = threshold_token(threshold)
    target = paired[f"target_coverage_percent_ge_{token}__absolute_change"]
    off_target = paired[f"off_target_coverage_percent_ge_{token}__absolute_change"]
    localization = paired[
        f"threshold_localization_percent_in_roi_ge_{token}__absolute_change"
    ]
    return {
        "target_increased": int((target > 0).sum()),
        "off_target_decreased": int((off_target < 0).sum()),
        "localization_increased": int((localization > 0).sum()),
    }


def build_captions() -> dict[str, str]:
    common = (
        "Eight subject/ROI cases were selected from a 132-participant CamCan "
        "temporal-interference simulation cohort: the highest and lowest "
        "generic median target-ROI field for each of four targets (left "
        "hippocampus, left M1, right DLPFC, and right thalamus). The generic "
        "MNI152-derived montage and subject-personalized Pareto montage were "
        "then simulated on the same corrected subject head. Every metric was "
        "calculated separately in ten independently remeshed simulations and "
        "then arithmetic-mean aggregated within condition. The selected "
        "extreme cases support descriptive comparison only and do not "
        "estimate a population-average personalization effect."
    )
    return {
        "figure_personalization_selected_case_changes": (
            "Figure 1. Selected-case changes with personalized temporal-"
            "interference stimulation. "
            + common
            + " Panels show (A) median electric-field magnitude within the "
            "anatomical target ROI, (B) percentage of target voxels reaching "
            "0.18 V/m, (C) percentage of finite brain voxels outside the "
            "target reaching 0.18 V/m, and (D) percentage of all "
            "suprathreshold brain voxels located inside the target. Open blue "
            "circles denote the generic montage, filled orange circles denote "
            "the personalized montage, horizontal gray segments connect the "
            "two condition means for the same selected case, and error bars "
            "show sample SD across ten remeshing repeats. Repeat numbers are "
            "not paired between conditions. Personalization increased the "
            "median target field in three of eight cases, all originally "
            "generic-worst M1, DLPFC, or thalamus cases, while the four "
            "generic-best cases decreased. These opposing changes indicate "
            "compression of the selected extreme-case range, not uniform "
            "maximization of the target-ROI median."
        ),
        "figure_personalization_effectiveness_spread_ge_0p18": (
            "Figure 2. Effectiveness-spread changes at 0.18 V/m after "
            "personalized temporal-interference stimulation. "
            + common
            + " Horizontal position is target coverage, defined as the "
            "percentage of anatomical target voxels at or above 0.18 V/m; "
            "vertical position is off-target coverage, defined as the "
            "percentage of finite brain voxels outside the target at or "
            "above 0.18 V/m. Each arrow begins at the generic montage (open "
            "symbol) and ends at the personalized montage (filled symbol). "
            "Green arrows identify cases selected as generic-best and purple "
            "arrows identify cases selected as generic-worst. The x-axis is "
            "shared, whereas y-axis limits are target-specific because "
            "off-target ranges differ markedly between targets. Movement "
            "rightward indicates greater target coverage and downward "
            "indicates less off-target spread; neither direction alone "
            "defines overall superiority. At this threshold, target coverage "
            "increased in three of eight cases, off-target coverage decreased "
            "in four, and suprathreshold localization in the target increased "
            "in six."
        ),
        "figure_personalization_effectiveness_spread_ge_0p20": (
            "Supplementary Figure S1. Sensitivity of selected-case "
            "effectiveness-spread changes to a 0.20 V/m threshold. "
            + common
            + " Target and off-target coverage use the same definitions as "
            "Figure 2 but require fields at or above 0.20 V/m. Arrows run "
            "from generic (open) to personalized (filled); green and purple "
            "identify generic-best and generic-worst cases, respectively. "
            "Panel-specific y-axis limits are used. The optimizer's recorded "
            "E_target near 0.20 V/m is an optimization objective and is not "
            "the same quantity as the median field across every anatomical "
            "target voxel. At 0.20 V/m, target coverage increased in three "
            "of eight selected cases, off-target coverage decreased in four, "
            "and suprathreshold localization increased in seven."
        ),
        "figure_personalization_effectiveness_spread_ge_0p15": (
            "Supplementary Figure S2. Sensitivity of selected-case "
            "effectiveness-spread changes to a 0.15 V/m threshold. "
            + common
            + " Target and off-target coverage use the same definitions as "
            "Figure 2 but require fields at or above 0.15 V/m. Arrows run "
            "from generic (open) to personalized (filled); green and purple "
            "identify generic-best and generic-worst cases, respectively. "
            "Panel-specific y-axis limits are used. At 0.15 V/m, target "
            "coverage increased in three of eight selected cases, off-target "
            "coverage decreased in five, and suprathreshold localization "
            "increased in four, demonstrating that the apparent trade-off "
            "depends on the reporting threshold."
        ),
        "figure_personalization_technical_repeat_distributions": (
            "Supplementary Figure S3. Technical-repeat distributions of the "
            "median target-ROI field. "
            + common
            + " Rows correspond to target ROIs and columns to cases selected "
            "as generic-best or generic-worst. Each small open point is one "
            "independently remeshed simulation; boxes show the interquartile "
            "range, internal white lines show the repeat median, whiskers "
            "extend to the most extreme value within 1.5 interquartile "
            "ranges, and diamonds show the arithmetic repeat mean used in "
            "the main comparison. Each panel has its own y-axis range to "
            "make technical variation visible. Condition shifts are "
            "generally larger than within-condition repeat variation, except "
            "for the small change in the generic-worst left-hippocampus case."
        ),
    }


def write_caption_files(output_dir: Path, captions: dict[str, str]) -> None:
    captions_dir = output_dir / "captions"
    captions_dir.mkdir(parents=True, exist_ok=True)
    markdown = ["# Self-contained figure captions", ""]
    csv_rows = []
    for figure in FIGURES:
        stem = figure["stem"]
        caption = captions[stem]
        (captions_dir / f"{stem}_caption.txt").write_text(
            caption + "\n", encoding="utf-8"
        )
        markdown.extend([f"## {figure['designation']}", "", caption, ""])
        csv_rows.append(
            {
                "designation": figure["designation"],
                "figure_file": f"figures/{stem}.png",
                "caption": caption,
            }
        )
    (output_dir / "figure_captions.md").write_text(
        "\n".join(markdown), encoding="utf-8"
    )
    with (output_dir / "figure_captions.csv").open(
        "w", newline="", encoding="utf-8"
    ) as handle:
        writer = csv.DictWriter(
            handle, fieldnames=["designation", "figure_file", "caption"]
        )
        writer.writeheader()
        writer.writerows(csv_rows)


def write_highlights(
    output_dir: Path,
    paired: pd.DataFrame,
    captions: dict[str, str],
) -> None:
    median_change = paired["roi_median_v_per_m__absolute_change"]
    generic = paired["roi_median_v_per_m__generic_repeat_mean"]
    personalized = paired["roi_median_v_per_m__personalized_repeat_mean"]
    counts_018 = metric_change_counts(paired, 0.18)
    lines = [
        "# Personalized versus generic selected-case results",
        "",
        "## Scope",
        "",
        "- Eight subject/ROI cases were deliberately selected as the generic "
        "best and worst median target-field cases in each of four ROIs.",
        "- Generic and personalized montages were evaluated on the same "
        "corrected subject head, using ten independent remeshing repeats per "
        "condition.",
        "- Results are descriptive selected-case comparisons. They do not "
        "estimate an average effect in the full CamCan population.",
        "",
        "## Main observations",
        "",
        f"- Generic median target fields spanned {generic.min():.3f} to "
        f"{generic.max():.3f} V/m; personalized fields spanned "
        f"{personalized.min():.3f} to {personalized.max():.3f} V/m.",
        f"- Personalization increased median target field in "
        f"{int((median_change > 0).sum())}/8 cases and decreased it in "
        f"{int((median_change < 0).sum())}/8.",
        "- All four generic-best cases decreased in median target field. "
        "Among generic-worst cases, left M1, right DLPFC, and right thalamus "
        "increased; left hippocampus changed little and decreased.",
        f"- At 0.18 V/m, target coverage increased in "
        f"{counts_018['target_increased']}/8 cases, off-target coverage "
        f"decreased in {counts_018['off_target_decreased']}/8, and "
        f"localization in the target increased in "
        f"{counts_018['localization_increased']}/8.",
        "- The generic-best cases consistently traded reduced target "
        "coverage for reduced off-target coverage at 0.18 V/m.",
        "- The generic-worst right-thalamus case gained target coverage but "
        "also had a large increase in off-target coverage, illustrating why "
        "target effectiveness and spatial spread must be interpreted "
        "together.",
        "- The optimizer's E_target value near 0.20 V/m is not the median "
        "field over every anatomical target voxel. A personalized montage "
        "can satisfy the optimizer's Pareto objective while producing a "
        "lower evaluation-stage ROI median.",
        "",
        "## Recommended manuscript placement",
        "",
        "- Figure 1 is the clearest main-text summary of selected-case changes.",
        "- Figure 2 can be main text if the effectiveness-spread trade-off is "
        "central; otherwise place it in the supplement.",
        "- Supplementary Figures S1-S3 document threshold sensitivity and "
        "technical repeatability.",
        "- Avoid hypothesis tests, confidence intervals for a population "
        "effect, or language implying that these outcome-selected cases are "
        "representative.",
        "",
        "## Caption files",
        "",
        "Self-contained captions are provided in `figure_captions.md`, "
        "`figure_captions.csv`, and as one text file per figure in `captions/`.",
        "",
    ]
    (output_dir / "results_highlights.md").write_text(
        "\n".join(lines), encoding="utf-8"
    )


def _format_change(value: float, unit: str) -> str:
    if unit == "V/m":
        return f"{value:+.3f}"
    return f"{value:+.1f}"


def add_review_figure_page(
    document: Document,
    designation: str,
    title: str,
    image_path: Path,
    one_line: str,
) -> None:
    """Add a landscape-safe figure page without image spillover."""
    document.add_page_break()
    heading = document.add_heading(f"{designation}. {title}", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_after = Pt(2)
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_paragraph.add_run().add_picture(
        str(image_path),
        height=Inches(5.25),
    )
    caption = document.add_paragraph(one_line, style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True


def build_review_guide(
    output_dir: Path,
    paired: pd.DataFrame,
    stimulation: pd.DataFrame,
    captions: dict[str, str],
) -> Path:
    figures_dir = output_dir / "figures"
    output = output_dir / "Personalized_vs_Generic_Figure_Review_Guide.docx"
    document = Document()
    configure_document(document)

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(50)
    run = paragraph.add_run(
        "Personalized versus Generic\nTemporal-Interference Stimulation"
    )
    run.font.name = "Arial"
    run.font.size = Pt(27)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Supervisor review guide: computation, interpretation, captions, "
        "and reporting guardrails"
    ).font.size = Pt(13.5)
    rule = document.add_table(rows=1, cols=1)
    rule.alignment = 1
    set_cell_shading(rule.cell(0, 0), BLUE)
    rule.cell(0, 0).text = ""
    scope = document.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_before = Pt(21)
    scope.add_run(
        "Eight selected subject/ROI cases | Four anatomical targets\n"
        "Generic MNI152-derived versus subject-personalized Pareto montage\n"
        "Ten independent remeshing repeats per condition"
    )
    prepared = document.add_paragraph()
    prepared.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prepared.paragraph_format.space_before = Pt(18)
    prepared_run = prepared.add_run(
        f"Prepared {date.today().strftime('%d %B %Y')} | "
        "Descriptive selected-case analysis"
    )
    prepared_run.font.size = Pt(9)
    prepared_run.font.color.rgb = RGBColor.from_string(GRAY)
    document.add_page_break()

    add_heading(document, "1. Study question and analysis scope")
    document.add_paragraph(
        "This analysis asks how a subject-personalized Pareto montage changes "
        "target engagement and off-target spread relative to the fixed "
        "MNI152-derived generic montage when both are simulated on the same "
        "corrected subject head."
    )
    add_callout(
        document,
        "Selection rule",
        "Within each of four ROIs, one case was the highest and one the "
        "lowest generic median target-ROI field in the 132-participant "
        "cohort. The resulting eight subject/ROI pairs include seven unique "
        "participants because one participant was selected for two ROIs.",
    )
    add_callout(
        document,
        "Inference boundary",
        "These are outcome-selected extremes, not a random sample. The "
        "figures describe what happened in these cases and cannot estimate "
        "the average benefit of personalization in the population.",
        PALE_ORANGE,
    )
    for text in [
        "Generic and personalized conditions were simulated on the same "
        "corrected-v4 anatomical head.",
        "Each condition used ten independently generated meshes and "
        "simulations.",
        "Metrics were calculated separately for each repeat and then "
        "arithmetic-mean aggregated within condition.",
        "Repeat numbers are not paired across conditions; connecting lines "
        "and arrows link condition means for the same subject/ROI case.",
        "Only the 80 personalized simulations corresponding to the intended "
        "subject/ROI selections were included; 200 out-of-scope simulations "
        "from the earlier over-complete run were excluded.",
    ]:
        add_bullet(document, text)

    add_heading(document, "Why the optimized value need not equal the ROI median", level=2)
    add_callout(
        document,
        "Critical distinction",
        "The optimization files record an E_target near 0.20 V/m for the "
        "selected Pareto solution (TI_free.Emin). That objective is defined "
        "inside the optimizer and is not the evaluation-stage median across "
        "all voxels in the anatomical ROI. Therefore a successful exhaustive "
        "search does not imply that the post-processing ROI median must be "
        "0.20 V/m or higher.",
        PALE_ORANGE,
    )
    document.add_paragraph(
        "The optimizer and the manuscript analysis can also differ in target "
        "sampling, objective formulation, mesh realization, and the way the "
        "anatomical ROI mask is evaluated. The QC field overlays currently "
        "being generated are the appropriate next check for spatial "
        "alignment and field placement; they do not change the scalar "
        "comparison reported here."
    )

    document.add_page_break()
    add_heading(document, "2. How the metrics were computed")
    for text in [
        "Load the brain-only TI field image and subject-space atlas for one "
        "selected subject, ROI, condition, and remeshing repeat.",
        "Construct the anatomical target mask and finite whole-brain and "
        "off-target masks.",
        "Calculate target-field, threshold-coverage, localization, robust "
        "maximum, and rank-based metrics for that repeat only.",
        "Repeat for all ten independently remeshed simulations.",
        "Calculate the arithmetic mean and sample SD of each metric across "
        "the ten repeats within the generic and personalized conditions.",
        "Subtract the generic mean from the personalized mean for each "
        "selected subject/ROI pair. No population hypothesis test is run.",
    ]:
        add_number(document, text)
    add_table(
        document,
        ["Metric", "Exact interpretation"],
        [
            [
                "Median target-ROI field",
                "Median finite TI-field magnitude among voxels inside the anatomical target.",
            ],
            [
                "Target coverage at threshold t",
                "100 x target voxels at or above t / all anatomical target voxels.",
            ],
            [
                "Off-target coverage at threshold t",
                "100 x finite non-target brain voxels at or above t / all finite non-target brain voxels.",
            ],
            [
                "Localization in target at threshold t",
                "100 x suprathreshold target voxels / all suprathreshold finite brain voxels.",
            ],
            [
                "P99.9 robust maximum",
                "99.9th percentile of finite field magnitudes, used instead of a single-voxel maximum.",
            ],
        ],
        widths_cm=[6.3, 14.0],
    )
    add_callout(
        document,
        "Why effectiveness and localization can disagree",
        "Target coverage uses the anatomical target as its denominator, "
        "whereas localization uses all suprathreshold brain voxels. A montage "
        "can reduce target coverage yet improve localization if off-target "
        "suprathreshold tissue falls even more strongly.",
    )

    key_rows = []
    for _, row in paired.iterrows():
        key_rows.append(
            [
                ROI_LABELS[row["roi"]],
                str(row["selection_role"]).capitalize(),
                str(row["subject"]).replace("sub-", ""),
                f"{row['roi_median_v_per_m__generic_repeat_mean']:.3f}",
                f"{row['roi_median_v_per_m__personalized_repeat_mean']:.3f}",
                _format_change(
                    row["roi_median_v_per_m__absolute_change"], "V/m"
                ),
                _format_change(
                    row[
                        "target_coverage_percent_ge_0p18__absolute_change"
                    ],
                    "percentage points",
                ),
                _format_change(
                    row[
                        "off_target_coverage_percent_ge_0p18__absolute_change"
                    ],
                    "percentage points",
                ),
            ]
        )
    document.add_page_break()
    add_heading(document, "3. Numerical orientation")
    add_table(
        document,
        [
            "Target",
            "Role",
            "Subject",
            "Generic ROI median",
            "Personalized ROI median",
            "Delta ROI median",
            "Delta target coverage at 0.18",
            "Delta off-target coverage at 0.18",
        ],
        key_rows,
        widths_cm=[3.1, 2.0, 2.8, 3.0, 3.2, 2.5, 3.3, 3.5],
        font_size=7.4,
    )
    add_bullet(
        document,
        "All four generic-best cases decreased in median target field after "
        "personalization.",
    )
    add_bullet(
        document,
        "The generic-worst left M1, right DLPFC, and right-thalamus cases "
        "increased in median target field; the worst left-hippocampus case "
        "changed little and decreased.",
    )
    add_bullet(
        document,
        "The selected generic range (0.072-0.296 V/m) narrowed under "
        "personalization (0.122-0.209 V/m). Because the cases were selected "
        "from the generic extremes, this compression is descriptive and may "
        "partly reflect extreme-case selection.",
    )

    guide_sections = {
        "figure_personalization_selected_case_changes": [
            "The four panels place target amplitude, target coverage, "
            "off-target spread, and localization on the same selected-case "
            "map. The connecting segment is a within-case condition contrast, "
            "not a pairing of individual repeat numbers.",
            "The central pattern is not uniform improvement. Personalization "
            "raises weak generic cases for M1, DLPFC, and thalamus, but lowers "
            "all generic-best cases. This looks like selected-range "
            "compression rather than universal maximization.",
            "Use this as the principal summary figure. Report exact values "
            "from the accompanying table and keep the selected-case qualifier "
            "in the Results and caption.",
        ],
        "figure_personalization_effectiveness_spread_ge_0p18": [
            "Rightward movement means more target tissue reaches 0.18 V/m; "
            "downward movement means less non-target brain tissue reaches "
            "0.18 V/m. The most unambiguously favorable direction is "
            "right-and-down, but other directions encode trade-offs.",
            "The four generic-best cases move left and down: target coverage "
            "falls, but off-target spread also falls. Worst M1 and DLPFC move "
            "right with only small off-target increases. Worst thalamus gains "
            "target coverage with a large off-target increase. Worst "
            "hippocampus changes little.",
            "Y-axis scales differ between ROI panels. Compare arrow direction "
            "within a panel, not raw visual arrow length across panels.",
        ],
        "figure_personalization_effectiveness_spread_ge_0p20": [
            "This repeats the trade-off analysis at 0.20 V/m, the threshold "
            "closest to the optimizer's recorded E_target.",
            "At this stricter threshold, target coverage increases in three "
            "cases, off-target coverage decreases in four, and localization "
            "increases in seven. Localization ratios can be visually dramatic "
            "when very few voxels reach threshold, so interpret them with "
            "target and off-target coverage.",
            "Treat this as a sensitivity figure, not proof that the optimizer "
            "failed or succeeded according to a different metric.",
        ],
        "figure_personalization_effectiveness_spread_ge_0p15": [
            "This repeats the same analysis at the more permissive 0.15 V/m "
            "threshold.",
            "The pattern changes with threshold: target coverage increases "
            "in three cases, off-target coverage decreases in five, and "
            "localization increases in four. This threshold dependence is "
            "why all three thresholds should remain available.",
            "Use as a supplementary robustness check and avoid choosing a "
            "threshold solely because it produces the most favorable pattern.",
        ],
        "figure_personalization_technical_repeat_distributions": [
            "Each panel exposes the ten repeat values behind the condition "
            "means. Boxes summarize the repeat distribution; diamonds are "
            "the arithmetic means used elsewhere.",
            "Most generic-personalized shifts are substantially larger than "
            "technical repeat SD. The worst left-hippocampus change is the "
            "notable small-shift case, where the condition difference is "
            "closer to repeat variability.",
            "Panel-specific y scales deliberately magnify technical "
            "variation. Do not compare apparent vertical separation across "
            "panels without reading the axes.",
        ],
    }

    section_number = 4
    for figure in FIGURES:
        stem = figure["stem"]
        add_review_figure_page(
            document,
            figure["designation"],
            figure["title"],
            figures_dir / f"{stem}.png",
            figure["one_line"],
        )
        document.add_page_break()
        add_heading(
            document,
            f"{section_number}. {figure['designation']} interpretation",
        )
        add_heading(document, "How to read it", level=2)
        document.add_paragraph(guide_sections[stem][0])
        add_heading(document, "What the selected cases show", level=2)
        document.add_paragraph(guide_sections[stem][1])
        add_heading(document, "Reporting guardrail", level=2)
        add_callout(
            document,
            "Use in the manuscript",
            guide_sections[stem][2],
            PALE_ORANGE,
        )
        add_heading(document, "Self-contained caption", level=2)
        paragraph = document.add_paragraph(captions[stem])
        paragraph.style = document.styles["Caption"]
        section_number += 1

    document.add_page_break()
    add_heading(document, f"{section_number}. Personalized stimulation parameters")
    parameter_rows = []
    for _, row in stimulation.sort_values("pair_index").iterrows():
        parameter_rows.append(
            [
                ROI_LABELS[row["roi"]],
                str(row["selection_role"]).capitalize(),
                str(row["subject"]).replace("sub-", ""),
                f"{row['generic_pair1']} / {row['generic_pair2']}",
                (
                    f"{row['personalized_pair1']} / "
                    f"{row['personalized_pair2']}"
                ),
                (
                    f"{row['personalized_current1_ma']:.3f}, "
                    f"{row['personalized_current2_ma']:.3f}"
                ),
                f"{row['personalized_optimization_e_target_v_per_m']:.3f}",
            ]
        )
    add_table(
        document,
        [
            "Target",
            "Role",
            "Subject",
            "Generic pairs",
            "Personalized pairs",
            "Personalized pair amplitudes (mA)",
            "Optimizer E_target (V/m)",
        ],
        parameter_rows,
        widths_cm=[3.0, 1.8, 2.6, 4.3, 4.3, 3.8, 2.8],
        font_size=7.2,
    )
    document.add_paragraph(
        "The complete parameter table, including configurations, source MAT "
        "files, hashes, and recorded Pareto metadata, is included as "
        "tables/table_stimulation_parameters.csv."
    )

    document.add_page_break()
    add_heading(document, f"{section_number + 1}. Conclusions and next checks")
    add_callout(
        document,
        "Defensible conclusion",
        "In these deliberately selected generic-best and generic-worst "
        "subject/ROI cases, personalization changed both target engagement "
        "and off-target spread but did not uniformly increase the median "
        "field across the anatomical target. It improved several initially "
        "weak cases and reduced amplitude and spread in initially strong "
        "cases.",
    )
    add_heading(document, "Before manuscript lock", level=2)
    for text in [
        "Inspect the paired generic-personalized E-field QC overlays for all "
        "eight selected cases to confirm anatomical alignment and spatial "
        "field placement.",
        "Confirm with the optimizer author the precise definition of "
        "TI_free.Emin and E_target, including the target sampling mask and "
        "whether the objective represents a minimum, percentile, or other "
        "target statistic.",
        "Decide whether the selected-case comparison belongs in the main "
        "text or is best framed as a proof-of-concept/sensitivity analysis.",
        "Keep all interpretation descriptive unless a prospectively selected "
        "or random sample of participants is personalized.",
    ]:
        add_bullet(document, text)

    document.add_page_break()
    add_heading(document, f"{section_number + 2}. Data lineage and reproducibility")
    add_bullet(
        document,
        "Source: validated output of camcan_personalized_comparison.py.",
    )
    add_bullet(
        document,
        "Figure layer: build_camcan_personalized_publication_outputs.py.",
    )
    add_bullet(
        document,
        "Source simulation images were read-only; no image metric was "
        "recomputed by this publication layer.",
    )
    add_bullet(
        document,
        "The output manifest records SHA256 hashes for all packaged files.",
    )
    document.save(output)
    return output


def convert_docx_to_pdf(docx_path: Path, output_dir: Path) -> Path:
    with tempfile.TemporaryDirectory(prefix="personalized_guide_soffice_") as temp:
        profile = Path(temp) / "profile"
        runtime = Path(temp) / "runtime"
        home = Path(temp) / "home"
        profile.mkdir()
        runtime.mkdir()
        home.mkdir()
        command = [
            "soffice",
            "--headless",
            f"-env:UserInstallation={profile.as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            temp,
            str(docx_path),
        ]
        environment = os.environ.copy()
        environment.update(
            {
                "HOME": str(home),
                "XDG_CACHE_HOME": str(Path(temp) / "cache"),
                "XDG_CONFIG_HOME": str(Path(temp) / "config"),
                "XDG_RUNTIME_DIR": str(runtime),
            }
        )
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            env=environment,
        )
        if completed.returncode != 0:
            raise RuntimeError(
                "LibreOffice conversion failed: "
                + completed.stdout
                + completed.stderr
            )
        source = Path(temp) / f"{docx_path.stem}.pdf"
        if not source.is_file():
            raise RuntimeError("LibreOffice did not create the expected PDF")
        destination = output_dir / source.name
        shutil.copy2(source, destination)
    return destination


def write_manifest_and_archive(
    output_dir: Path,
    input_dir: Path,
    source_manifest: dict,
) -> None:
    excluded = {
        "publication_output_manifest.json",
        "personalized_vs_generic_manuscript_package.zip",
        "personalized_vs_generic_manuscript_package.zip.sha256",
    }
    files = []
    for path in sorted(p for p in output_dir.rglob("*") if p.is_file()):
        relative = path.relative_to(output_dir).as_posix()
        if relative in excluded:
            continue
        files.append(
            {
                "path": relative,
                "bytes": path.stat().st_size,
                "sha256": sha256_file(path),
            }
        )
    manifest = {
        "publication_package_schema_version": 1,
        "status": "complete",
        "created": date.today().isoformat(),
        "source_directory": str(input_dir.resolve()),
        "source_comparison_schema_version": source_manifest.get(
            "comparison_schema_version"
        ),
        "source_manuscript_analysis_schema_version": source_manifest.get(
            "manuscript_analysis_schema_version"
        ),
        "selected_subject_roi_pairs": 8,
        "unique_subjects": 7,
        "repeat_level_records": 160,
        "inference": "descriptive selected-case analysis only",
        "files": files,
    }
    manifest_path = output_dir / "publication_output_manifest.json"
    manifest_path.write_text(json.dumps(manifest, indent=2) + "\n")

    archive = output_dir / "personalized_vs_generic_manuscript_package.zip"
    with zipfile.ZipFile(
        archive, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6
    ) as handle:
        for path in sorted(p for p in output_dir.rglob("*") if p.is_file()):
            if path == archive or path.name == f"{archive.name}.sha256":
                continue
            handle.write(path, path.relative_to(output_dir))
    (output_dir / f"{archive.name}.sha256").write_text(
        f"{sha256_file(archive)}  {archive.name}\n"
    )


def build(input_dir: Path, output_dir: Path, *, force: bool) -> dict:
    input_dir = input_dir.resolve()
    output_dir = output_dir.resolve()
    if output_dir.exists():
        if not force:
            raise FileExistsError(
                f"Output directory exists: {output_dir}; pass --force to replace it"
            )
        shutil.rmtree(output_dir)
    output_dir.mkdir(parents=True)
    figures_dir = output_dir / "figures"
    tables_dir = output_dir / "tables"
    figures_dir.mkdir()
    tables_dir.mkdir()

    manifest, paired, condition, repeats = load_and_validate(input_dir)
    set_figure_style()
    plot_change_summary(paired, figures_dir)
    for threshold in [0.18, 0.20, 0.15]:
        plot_tradeoff(paired, figures_dir, threshold)
    plot_repeat_distributions(repeats, figures_dir)

    key_changes = build_key_change_table(paired)
    key_changes.to_csv(tables_dir / "table_selected_case_key_changes.csv", index=False)
    for filename in [
        "table_main_selected_metrics.csv",
        "table_stimulation_parameters.csv",
        "selection_allowlist.csv",
    ]:
        shutil.copy2(input_dir / filename, tables_dir / filename)
    condition.to_csv(
        tables_dir / "condition_repeat_mean_metrics.csv",
        index=False,
    )

    captions = build_captions()
    write_caption_files(output_dir, captions)
    write_highlights(output_dir, paired, captions)
    stimulation = pd.read_csv(input_dir / "table_stimulation_parameters.csv")
    docx_path = build_review_guide(output_dir, paired, stimulation, captions)
    pdf_path = convert_docx_to_pdf(docx_path, output_dir)
    write_manifest_and_archive(output_dir, input_dir, manifest)

    result = {
        "status": "complete",
        "source": str(input_dir),
        "output": str(output_dir),
        "figures_png": len(list(figures_dir.glob("*.png"))),
        "figures_pdf": len(list(figures_dir.glob("*.pdf"))),
        "tables": len(list(tables_dir.glob("*.csv"))),
        "captions": len(captions),
        "review_guide_docx": str(docx_path),
        "review_guide_pdf": str(pdf_path),
        "archive": str(
            output_dir / "personalized_vs_generic_manuscript_package.zip"
        ),
    }
    print(json.dumps(result, indent=2))
    return result


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input-dir", required=True, type=Path)
    parser.add_argument("--out-dir", required=True, type=Path)
    parser.add_argument(
        "--force",
        action="store_true",
        help="replace an existing output directory",
    )
    return parser.parse_args()


def main() -> int:
    arguments = parse_args()
    build(arguments.input_dir, arguments.out_dir, force=arguments.force)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
