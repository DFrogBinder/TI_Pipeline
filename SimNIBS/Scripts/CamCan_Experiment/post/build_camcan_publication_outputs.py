#!/usr/bin/env python3
"""Build publication-ready tables and descriptive figures for CamCan TI results.

The input is the schema-2 output of ``camcan_manuscript_analysis.py``. Metrics
have already been calculated independently for each remeshing repeat and then
arithmetic-mean aggregated across the ten repeats for each subject and ROI.

This script intentionally performs descriptive analyses only:

* no cross-ROI hypothesis tests;
* no inferential comparison against the single MNI152 reference head;
* no individualized-optimization analysis.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
from typing import Iterable

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from matplotlib.lines import Line2D
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy.stats import pearsonr, percentileofscore, spearmanr


ROI_ORDER = [
    "Left_Hippocampus",
    "Left_M1",
    "Right_DLPC",
    "Right_Thalamus",
]

ROI_LABELS = {
    "Left_Hippocampus": "Left hippocampus",
    "Left_M1": "Left M1",
    "Right_DLPC": "Right DLPFC",
    "Right_Thalamus": "Right thalamus",
}

PRIMARY_METRICS = [
    (
        "Field magnitude",
        "Median target-ROI TI field",
        "roi_median_v_per_m",
        "V/m",
    ),
    (
        "Field magnitude",
        "Robust maximum target-ROI TI field (P99.9)",
        "roi_robust_max_p99_9_v_per_m",
        "V/m",
    ),
    (
        "Threshold performance (0.18 V/m)",
        "Target coverage ≥0.18 V/m",
        "target_coverage_percent_ge_0p18",
        "%",
    ),
    (
        "Threshold performance (0.18 V/m)",
        "Off-target coverage ≥0.18 V/m",
        "off_target_coverage_percent_ge_0p18",
        "%",
    ),
    (
        "Threshold performance (0.18 V/m)",
        "Whole-brain coverage ≥0.18 V/m",
        "whole_brain_coverage_percent_ge_0p18",
        "%",
    ),
    (
        "Threshold performance (0.18 V/m)",
        "Suprathreshold localization in target ≥0.18 V/m",
        "threshold_localization_percent_in_roi_ge_0p18",
        "%",
    ),
    (
        "Rank-based sensitivity",
        "Target coverage by whole-brain top 5% field",
        "top_5_percent_target_coverage_percent",
        "%",
    ),
    (
        "Rank-based sensitivity",
        "Localization of whole-brain top 5% field in target",
        "top_5_percent_localization_percent_in_roi",
        "%",
    ),
]

PRIMARY_SHORT_LABELS = {
    "roi_median_v_per_m": "ROI median field",
    "roi_robust_max_p99_9_v_per_m": "ROI P99.9",
    "target_coverage_percent_ge_0p18": "Target coverage ≥0.18",
    "off_target_coverage_percent_ge_0p18": "Off-target coverage ≥0.18",
    "whole_brain_coverage_percent_ge_0p18": "Whole-brain coverage ≥0.18",
    "threshold_localization_percent_in_roi_ge_0p18": "Localization ≥0.18",
    "top_5_percent_target_coverage_percent": "Top-5% target coverage",
    "top_5_percent_localization_percent_in_roi": "Top-5% localization",
}

THRESHOLD_METRICS = {
    "Target coverage": "target_coverage_percent",
    "Off-target coverage": "off_target_coverage_percent",
    "Whole-brain coverage": "whole_brain_coverage_percent",
    "Suprathreshold localization in target": (
        "threshold_localization_percent_in_roi"
    ),
}

REGION_ROBUST_METRICS = {
    "Target ROI": (
        "roi_robust_max_p99_9_v_per_m",
        "roi_upper_1_percent_median_v_per_m",
    ),
    "Whole brain": (
        "whole_brain_robust_max_p99_9_v_per_m",
        "whole_brain_upper_1_percent_median_v_per_m",
    ),
    "Off target": (
        "off_target_robust_max_p99_9_v_per_m",
        "off_target_upper_1_percent_median_v_per_m",
    ),
}

QC_METRICS = [
    ("Median target-ROI TI field", "roi_median_v_per_m", "V/m"),
    (
        "Target coverage ≥0.18 V/m",
        "target_coverage_percent_ge_0p18",
        "%",
    ),
    (
        "Off-target coverage ≥0.18 V/m",
        "off_target_coverage_percent_ge_0p18",
        "%",
    ),
    (
        "Suprathreshold localization in target ≥0.18 V/m",
        "threshold_localization_percent_in_roi_ge_0p18",
        "%",
    ),
]

NAVY = "17365D"
BLUE = "5B9BD5"
PALE_BLUE = "D9EAF7"
PALE_TEAL = "DDEBF7"
PALE_GRAY = "E7E6E6"
PALE_ORANGE = "FCE4D6"
WHITE = "FFFFFF"
BLACK = "000000"
GREEN = "00876C"
ORANGE = "D97706"
RED = "C44E52"
FIGURE_BLUE = "#2F6B9A"
FIGURE_ORANGE = "#E07A1F"
FIGURE_RED = "#B94A48"
FIGURE_GRAY = "#6B7280"
FIGURE_GREEN = "#00876C"


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def ensure_finite(frame: pd.DataFrame, name: str) -> None:
    numeric = frame.select_dtypes(include=[np.number])
    bad = int((~np.isfinite(numeric.to_numpy(dtype=float))).sum())
    if bad:
        raise RuntimeError(f"{name} contains {bad} non-finite numeric cell(s)")


def validate_inputs(
    manifest: dict,
    repeat_frame: pd.DataFrame,
    subject_frame: pd.DataFrame,
    mni_frame: pd.DataFrame,
    supplementary: pd.DataFrame,
) -> None:
    if manifest.get("analysis_schema_version") != 2:
        raise RuntimeError("Expected analysis_schema_version=2")
    if manifest.get("status") != "complete":
        raise RuntimeError("Source analysis is not complete")
    if manifest.get("subjects") != 132:
        raise RuntimeError("Expected 132 subjects")
    if manifest.get("repeat_level_records") != 5280:
        raise RuntimeError("Expected 5,280 repeat-level records")
    if manifest.get("subject_level_records") != 528:
        raise RuntimeError("Expected 528 subject-level records")
    if manifest.get("mni_baselines") != 4:
        raise RuntimeError("Expected four MNI152 baseline records")
    if manifest.get("thresholds_v_per_m") != [0.18, 0.15]:
        raise RuntimeError("Expected thresholds [0.18, 0.15] V/m")
    if list(manifest.get("rois", [])) != ROI_ORDER:
        raise RuntimeError("Unexpected ROI order or membership")

    expected_repeat_keys = ["subject", "roi", "repeat"]
    expected_subject_keys = ["subject", "roi"]
    if repeat_frame.duplicated(expected_repeat_keys).any():
        raise RuntimeError("Duplicate repeat-level subject/ROI/repeat keys")
    if subject_frame.duplicated(expected_subject_keys).any():
        raise RuntimeError("Duplicate subject-level subject/ROI keys")
    if len(repeat_frame) != 5280 or len(subject_frame) != 528:
        raise RuntimeError("Unexpected input row counts")
    if set(mni_frame["roi"]) != set(ROI_ORDER) or len(mni_frame) != 4:
        raise RuntimeError("MNI152 baseline ROI coverage is incomplete")

    for roi in ROI_ORDER:
        repeat_roi = repeat_frame.loc[repeat_frame["roi"] == roi]
        subject_roi = subject_frame.loc[subject_frame["roi"] == roi]
        if len(repeat_roi) != 1320 or len(subject_roi) != 132:
            raise RuntimeError(f"Unexpected record count for {roi}")
        if set(repeat_roi["repeat"].astype(str).str.zfill(2)) != {
            f"{value:02d}" for value in range(1, 11)
        }:
            raise RuntimeError(f"Repeat coverage is incomplete for {roi}")
        if not (subject_roi["repeat_count"] == 10).all():
            raise RuntimeError(f"Subject repeat_count is not uniformly 10 for {roi}")

    if len(supplementary) != 108:
        raise RuntimeError("Expected 108 rows in the source supplementary table")
    ensure_finite(repeat_frame, "repeat_level_metrics.csv")
    ensure_finite(subject_frame, "subject_level_repeat_mean_metrics.csv")
    ensure_finite(mni_frame, "mni152_baseline_metrics.csv")


def metric_decimals(unit: str) -> int:
    if unit == "V/m":
        return 3
    if unit == "%":
        return 1
    if unit == "mm³":
        return 0
    return 2


def fmt_number(value: float, unit: str) -> str:
    decimals = metric_decimals(unit)
    if decimals == 0:
        return f"{value:,.0f}"
    return f"{value:.{decimals}f}"


def fmt_mean_sd(values: pd.Series, unit: str) -> str:
    return (
        f"{fmt_number(float(values.mean()), unit)} "
        f"({fmt_number(float(values.std(ddof=1)), unit)})"
    )


def summarize_metric(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
    roi: str,
    metric: str,
    label: str,
    unit: str,
    section: str,
) -> dict:
    values = subject_frame.loc[subject_frame["roi"] == roi, metric].astype(float)
    mni = float(mni_by_roi.loc[roi, metric])
    return {
        "section": section,
        "outcome": label,
        "metric": metric,
        "unit": unit,
        "roi": ROI_LABELS[roi],
        "roi_key": roi,
        "subject_n": int(len(values)),
        "subject_mean": float(values.mean()),
        "subject_sd": float(values.std(ddof=1)),
        "subject_median": float(values.median()),
        "subject_q1": float(values.quantile(0.25)),
        "subject_q3": float(values.quantile(0.75)),
        "subject_min": float(values.min()),
        "subject_max": float(values.max()),
        "camcan_mean_sd": fmt_mean_sd(values, unit),
        "mni152": mni,
        "mni152_formatted": fmt_number(mni, unit),
        "mni_percentile_within_camcan": float(
            percentileofscore(values, mni, kind="mean")
        ),
        "mni_standardized_difference": float(
            (mni - values.mean()) / values.std(ddof=1)
        ),
    }


def build_primary_tables(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    rows = []
    for section, label, metric, unit in PRIMARY_METRICS:
        for roi in ROI_ORDER:
            rows.append(
                summarize_metric(
                    subject_frame,
                    mni_by_roi,
                    roi,
                    metric,
                    label,
                    unit,
                    section,
                )
            )
    long_frame = pd.DataFrame(rows)

    wide_rows = []
    for section, label, metric, unit in PRIMARY_METRICS:
        row = {
            "Section": section,
            "Outcome": label,
            "Unit": unit,
        }
        selected = long_frame.loc[long_frame["metric"] == metric].set_index(
            "roi_key"
        )
        for roi in ROI_ORDER:
            display = ROI_LABELS[roi]
            row[f"{display} — CamCan mean (SD)"] = selected.loc[
                roi, "camcan_mean_sd"
            ]
            row[f"{display} — MNI152"] = selected.loc[
                roi, "mni152_formatted"
            ]
        wide_rows.append(row)
    return pd.DataFrame(wide_rows), long_frame


def build_threshold_table(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
) -> pd.DataFrame:
    rows = []
    for roi in ROI_ORDER:
        roi_subjects = subject_frame.loc[subject_frame["roi"] == roi]
        for threshold, slug in [(0.18, "0p18"), (0.15, "0p15")]:
            row = {
                "ROI": ROI_LABELS[roi],
                "Threshold (V/m)": threshold,
                "CamCan n": int(len(roi_subjects)),
            }
            for label, stem in THRESHOLD_METRICS.items():
                metric = f"{stem}_ge_{slug}"
                values = roi_subjects[metric].astype(float)
                row[f"{label} — CamCan mean (SD), %"] = fmt_mean_sd(
                    values, "%"
                )
                row[f"{label} — MNI152, %"] = fmt_number(
                    float(mni_by_roi.loc[roi, metric]), "%"
                )
            rows.append(row)
    return pd.DataFrame(rows)


def build_robustness_table(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
) -> pd.DataFrame:
    rows = []
    for roi in ROI_ORDER:
        roi_subjects = subject_frame.loc[subject_frame["roi"] == roi]
        for region, (p999_metric, upper_metric) in REGION_ROBUST_METRICS.items():
            p999 = roi_subjects[p999_metric].astype(float)
            upper = roi_subjects[upper_metric].astype(float)
            rows.append(
                {
                    "ROI": ROI_LABELS[roi],
                    "Region": region,
                    "P99.9 — CamCan mean (SD), V/m": fmt_mean_sd(
                        p999, "V/m"
                    ),
                    "P99.9 — MNI152, V/m": fmt_number(
                        float(mni_by_roi.loc[roi, p999_metric]), "V/m"
                    ),
                    "Upper-1% median — CamCan mean (SD), V/m": fmt_mean_sd(
                        upper, "V/m"
                    ),
                    "Upper-1% median — MNI152, V/m": fmt_number(
                        float(mni_by_roi.loc[roi, upper_metric]), "V/m"
                    ),
                    "CamCan Spearman ρ, P99.9 vs upper-1% median": float(
                        spearmanr(p999, upper).statistic
                    ),
                    "Mean upper-1% median / P99.9 ratio": float(
                        (upper / p999).mean()
                    ),
                }
            )
    return pd.DataFrame(rows)


def build_mni_context_table(primary_long: pd.DataFrame) -> pd.DataFrame:
    columns = [
        "roi",
        "section",
        "outcome",
        "unit",
        "subject_n",
        "subject_mean",
        "subject_sd",
        "subject_median",
        "subject_q1",
        "subject_q3",
        "mni152",
        "mni_percentile_within_camcan",
        "mni_standardized_difference",
    ]
    return (
        primary_long[columns]
        .rename(
            columns={
                "roi": "ROI",
                "section": "Section",
                "outcome": "Outcome",
                "unit": "Unit",
                "subject_n": "CamCan n",
                "subject_mean": "CamCan mean",
                "subject_sd": "CamCan SD",
                "subject_median": "CamCan median",
                "subject_q1": "CamCan Q1",
                "subject_q3": "CamCan Q3",
                "mni152": "MNI152",
                "mni_percentile_within_camcan": (
                    "MNI152 percentile within CamCan"
                ),
                "mni_standardized_difference": (
                    "MNI152 standardized difference from CamCan mean"
                ),
            }
        )
        .reset_index(drop=True)
    )


def build_association_table(subject_frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    relationships = [
        (
            "Effectiveness–spread coupling",
            "Target coverage",
            "Off-target coverage",
            "target_coverage_percent",
            "off_target_coverage_percent",
        ),
        (
            "Effectiveness–localization relationship",
            "Target coverage",
            "Suprathreshold localization in target",
            "target_coverage_percent",
            "threshold_localization_percent_in_roi",
        ),
    ]
    for roi in ROI_ORDER:
        roi_subjects = subject_frame.loc[subject_frame["roi"] == roi]
        for threshold, slug in [(0.18, "0p18"), (0.15, "0p15")]:
            for relationship, x_label, y_label, x_stem, y_stem in relationships:
                x = roi_subjects[f"{x_stem}_ge_{slug}"].astype(float)
                y = roi_subjects[f"{y_stem}_ge_{slug}"].astype(float)
                rho = spearmanr(x, y)
                pearson = pearsonr(x, y)
                rows.append(
                    {
                        "ROI": ROI_LABELS[roi],
                        "Threshold (V/m)": threshold,
                        "Relationship": relationship,
                        "X outcome": x_label,
                        "Y outcome": y_label,
                        "n": int(len(x)),
                        "Spearman ρ": float(rho.statistic),
                        "Spearman p (descriptive)": float(rho.pvalue),
                        "Pearson r": float(pearson.statistic),
                        "Pearson p (descriptive)": float(pearson.pvalue),
                    }
                )
    return pd.DataFrame(rows)


def build_threshold_shift_table(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
) -> pd.DataFrame:
    rows = []
    for roi in ROI_ORDER:
        roi_subjects = subject_frame.loc[subject_frame["roi"] == roi]
        for label, stem in THRESHOLD_METRICS.items():
            metric_018 = f"{stem}_ge_0p18"
            metric_015 = f"{stem}_ge_0p15"
            values_018 = roi_subjects[metric_018].astype(float)
            values_015 = roi_subjects[metric_015].astype(float)
            delta = values_015 - values_018
            mni_018 = float(mni_by_roi.loc[roi, metric_018])
            mni_015 = float(mni_by_roi.loc[roi, metric_015])
            rows.append(
                {
                    "ROI": ROI_LABELS[roi],
                    "Outcome": label,
                    "CamCan mean at 0.18 V/m, %": float(values_018.mean()),
                    "CamCan mean at 0.15 V/m, %": float(values_015.mean()),
                    "Paired mean change (0.15 − 0.18), percentage points": (
                        float(delta.mean())
                    ),
                    "Paired SD of change, percentage points": float(
                        delta.std(ddof=1)
                    ),
                    "Paired median change, percentage points": float(
                        delta.median()
                    ),
                    "Paired change Q1, percentage points": float(
                        delta.quantile(0.25)
                    ),
                    "Paired change Q3, percentage points": float(
                        delta.quantile(0.75)
                    ),
                    "MNI152 at 0.18 V/m, %": mni_018,
                    "MNI152 at 0.15 V/m, %": mni_015,
                    "MNI152 change (0.15 − 0.18), percentage points": (
                        mni_015 - mni_018
                    ),
                }
            )
    return pd.DataFrame(rows)


def build_attainment_table(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
) -> pd.DataFrame:
    rows = []
    for roi in ROI_ORDER:
        values = subject_frame.loc[
            subject_frame["roi"] == roi, "target_coverage_percent_ge_0p18"
        ].astype(float)
        rows.append(
            {
                "ROI": ROI_LABELS[roi],
                "n": int(len(values)),
                "Subjects with ≤1% target coverage, %": float(
                    100 * (values <= 1).mean()
                ),
                "Subjects with ≥25% target coverage, %": float(
                    100 * (values >= 25).mean()
                ),
                "Subjects with ≥50% target coverage, %": float(
                    100 * (values >= 50).mean()
                ),
                "Subjects with ≥75% target coverage, %": float(
                    100 * (values >= 75).mean()
                ),
                "Subjects with ≥90% target coverage, %": float(
                    100 * (values >= 90).mean()
                ),
                "MNI152 target coverage ≥0.18 V/m, %": float(
                    mni_by_roi.loc[
                        roi, "target_coverage_percent_ge_0p18"
                    ]
                ),
            }
        )
    return pd.DataFrame(rows)


def build_repeat_qc_table(subject_frame: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for roi in ROI_ORDER:
        roi_subjects = subject_frame.loc[subject_frame["roi"] == roi]
        for label, metric, unit in QC_METRICS:
            between_sd = float(roi_subjects[metric].std(ddof=1))
            within_sd = roi_subjects[f"{metric}__repeat_sd"].astype(float)
            rows.append(
                {
                    "ROI": ROI_LABELS[roi],
                    "Outcome": label,
                    "Unit": unit,
                    "Between-subject SD of ten-repeat means": between_sd,
                    "Median within-subject repeat SD": float(
                        within_sd.median()
                    ),
                    "Within-subject Q1 repeat SD": float(
                        within_sd.quantile(0.25)
                    ),
                    "Within-subject Q3 repeat SD": float(
                        within_sd.quantile(0.75)
                    ),
                    "Median within-subject SD / between-subject SD": float(
                        within_sd.median() / between_sd
                    ),
                }
            )
    return pd.DataFrame(rows)


def add_title(
    ws,
    title: str,
    subtitle: str,
    max_column: int,
    header_row: int = 4,
) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max_column)
    title_cell = ws.cell(1, 1, title)
    title_cell.font = Font(name="Arial", size=14, bold=True, color=WHITE)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max_column)
    subtitle_cell = ws.cell(2, 1, subtitle)
    subtitle_cell.font = Font(name="Arial", size=9, italic=True, color="404040")
    subtitle_cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 34
    ws.freeze_panes = f"A{header_row + 1}"
    ws.sheet_view.showGridLines = False


def write_frame_sheet(
    wb: Workbook,
    name: str,
    title: str,
    subtitle: str,
    frame: pd.DataFrame,
    percent_columns: Iterable[str] = (),
    number_formats: dict[str, str] | None = None,
    source_comment: str | None = None,
) -> None:
    ws = wb.create_sheet(name)
    add_title(ws, title, subtitle, len(frame.columns), header_row=4)
    header_row = 4
    thin_gray = Side(style="thin", color="BFBFBF")
    for col_idx, column in enumerate(frame.columns, start=1):
        cell = ws.cell(header_row, col_idx, column)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = Border(bottom=thin_gray)
        if source_comment and col_idx == 1:
            cell.comment = Comment(source_comment, "Codex")

    percent_columns = set(percent_columns)
    number_formats = number_formats or {}
    for row_idx, row in enumerate(frame.itertuples(index=False), start=5):
        for col_idx, (column, value) in enumerate(
            zip(frame.columns, row), start=1
        ):
            if pd.isna(value):
                value = ""
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = Font(name="Arial", size=9, color=BLACK)
            cell.alignment = Alignment(
                horizontal=(
                    "right"
                    if isinstance(value, (int, float, np.integer, np.floating))
                    else "left"
                ),
                vertical="top",
                wrap_text=True,
            )
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7F9FB")
            if column in percent_columns and isinstance(
                value, (int, float, np.integer, np.floating)
            ):
                cell.number_format = "0.0"
            if column in number_formats and isinstance(
                value, (int, float, np.integer, np.floating)
            ):
                cell.number_format = number_formats[column]

    widths = {}
    for col_idx, column in enumerate(frame.columns, start=1):
        values = [str(column)]
        values.extend(
            "" if pd.isna(value) else str(value)
            for value in frame[column].head(200)
        )
        width = min(max(max(map(len, values)) + 2, 10), 36)
        widths[col_idx] = width
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(len(frame.columns))}"
        f"{header_row + len(frame)}"
    )
    ws.print_title_rows = f"1:{header_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Page &P of &N"


def write_readme_sheet(
    wb: Workbook,
    input_dir: Path,
    manifest: dict,
    source_hashes: dict[str, str],
) -> None:
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "CamCan TI manuscript — publication tables"
    ws["A1"].font = Font(name="Arial", size=16, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(horizontal="left")
    ws.row_dimensions[1].height = 26

    notes = [
        ("Status", "Validated schema-2 results; publication tables generated reproducibly."),
        ("Cohort", "132 CamCan subjects; four target ROIs; ten remeshing repeats per subject and ROI."),
        (
            "Aggregation",
            "Every metric was calculated independently for each repeat, then arithmetic-mean aggregated across the ten repeats for each subject and ROI.",
        ),
        (
            "Primary threshold",
            "0.18 V/m. The 0.15 V/m analysis is reported as a sensitivity analysis.",
        ),
        (
            "Robust maximum",
            "99.9th percentile (P99.9). The median of the upper 1% is reported as a sensitivity definition.",
        ),
        (
            "MNI152",
            "A single descriptive reference head. MNI152 values are not treated as an inferential sample.",
        ),
        (
            "Inference",
            "No formal cross-ROI comparisons are performed. Exploratory correlations are descriptive within each ROI.",
        ),
        (
            "Primary reporting convention",
            "CamCan values are mean (SD) across subjects. MNI152 values are shown separately.",
        ),
        (
            "Missing/nonfinite policy",
            "Nonfinite anatomical ROI voxels count as unstimulated. Localization is 0% when no finite whole-brain voxel reaches the threshold.",
        ),
        ("Source directory", str(input_dir)),
        ("Analysis schema", str(manifest["analysis_schema_version"])),
    ]
    ws["A3"] = "Item"
    ws["B3"] = "Details"
    for cell in ws[3][:2]:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for row_idx, (label, value) in enumerate(notes, start=4):
        ws.cell(row_idx, 1, label).font = Font(
            name="Arial", size=10, bold=True, color=BLACK
        )
        ws.cell(row_idx, 2, value).font = Font(
            name="Arial", size=10, color=BLACK
        )
        ws.cell(row_idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
        if row_idx % 2 == 0:
            for cell in ws[row_idx][:2]:
                cell.fill = PatternFill("solid", fgColor="F7F9FB")

    hash_start = 4 + len(notes) + 2
    ws.cell(hash_start, 1, "Validated source file")
    ws.cell(hash_start, 2, "SHA-256")
    for cell in ws[hash_start][:2]:
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
    for offset, (filename, digest) in enumerate(source_hashes.items(), start=1):
        ws.cell(hash_start + offset, 1, filename)
        ws.cell(hash_start + offset, 2, digest)
        ws.cell(hash_start + offset, 2).font = Font(
            name="Courier New", size=8, color=GREEN
        )
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 105
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def write_primary_sheet(
    wb: Workbook,
    primary_wide: pd.DataFrame,
    source_comment: str,
) -> None:
    ws = wb.create_sheet("Table 1 Primary")
    total_columns = 3 + 2 * len(ROI_ORDER)
    subtitle = (
        "CamCan values are mean (SD) across 132 subjects. Each subject value "
        "is the arithmetic mean of ten repeat-level metric values. MNI152 is "
        "a single descriptive reference."
    )
    add_title(ws, "Table 1. Primary effectiveness and focality outcomes", subtitle, total_columns, header_row=5)

    ws.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
    ws.merge_cells(start_row=4, start_column=2, end_row=5, end_column=2)
    ws.merge_cells(start_row=4, start_column=3, end_row=5, end_column=3)
    for col, label in [(1, "Section"), (2, "Outcome"), (3, "Unit")]:
        ws.cell(4, col, label)
    ws["A4"].comment = Comment(source_comment, "Codex")

    col = 4
    for roi in ROI_ORDER:
        ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
        ws.cell(4, col, ROI_LABELS[roi])
        ws.cell(5, col, "CamCan mean (SD)")
        ws.cell(5, col + 1, "MNI152")
        col += 2

    for row in ws.iter_rows(min_row=4, max_row=5, min_col=1, max_col=total_columns):
        for cell in row:
            cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    current_section = None
    output_row = 6
    for record in primary_wide.to_dict(orient="records"):
        section = record["Section"]
        if section != current_section:
            ws.merge_cells(
                start_row=output_row,
                start_column=1,
                end_row=output_row,
                end_column=total_columns,
            )
            cell = ws.cell(output_row, 1, section)
            cell.font = Font(name="Arial", size=9, bold=True, color=NAVY)
            cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
            output_row += 1
            current_section = section
        ws.cell(output_row, 1, "")
        ws.cell(output_row, 2, record["Outcome"])
        ws.cell(output_row, 3, record["Unit"])
        col = 4
        for roi in ROI_ORDER:
            display = ROI_LABELS[roi]
            ws.cell(output_row, col, record[f"{display} — CamCan mean (SD)"])
            ws.cell(output_row, col + 1, record[f"{display} — MNI152"])
            col += 2
        for cell in ws[output_row][:total_columns]:
            cell.font = Font(name="Arial", size=9, color=BLACK)
            cell.alignment = Alignment(
                horizontal="left" if cell.column == 2 else "center",
                vertical="top",
                wrap_text=True,
            )
        if output_row % 2 == 0:
            for cell in ws[output_row][:total_columns]:
                cell.fill = PatternFill("solid", fgColor="F7F9FB")
        output_row += 1

    note_row = output_row + 1
    ws.merge_cells(
        start_row=note_row,
        start_column=1,
        end_row=note_row,
        end_column=total_columns,
    )
    ws.cell(
        note_row,
        1,
        (
            "Notes: ROI, region of interest; TI, temporal interference. "
            "Target coverage is the percentage of anatomical ROI voxels at or "
            "above the threshold. Off-target coverage uses finite whole-brain "
            "voxels outside the target ROI as its denominator. Localization "
            "is the percentage of suprathreshold whole-brain voxels lying "
            "inside the target ROI. P99.9 avoids single-voxel extrema."
        ),
    )
    ws.cell(note_row, 1).font = Font(name="Arial", size=8, italic=True, color="404040")
    ws.cell(note_row, 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[note_row].height = 52
    widths = [9, 44, 8] + [19, 11] * len(ROI_ORDER)
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "D6"
    ws.sheet_view.showGridLines = False
    ws.print_title_rows = "1:5"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Page &P of &N"


def latex_escape(value: str) -> str:
    replacements = {
        "\\": r"\textbackslash{}",
        "&": r"\&",
        "%": r"\%",
        "$": r"\$",
        "#": r"\#",
        "_": r"\_",
        "{": r"\{",
        "}": r"\}",
        "~": r"\textasciitilde{}",
        "^": r"\textasciicircum{}",
        "≥": r"$\geq$",
        "−": "--",
        "ρ": r"$\rho$",
        "³": r"$^3$",
    }
    return "".join(replacements.get(char, char) for char in str(value))


def write_primary_latex(primary_wide: pd.DataFrame, path: Path) -> None:
    lines = [
        r"\begin{table*}[htbp]",
        r"\centering",
        r"\caption{Primary effectiveness and focality outcomes.}",
        r"\label{tab:camcan_primary}",
        r"\small",
        r"\setlength{\tabcolsep}{3.5pt}",
        r"\begin{tabular}{ll" + "cc" * len(ROI_ORDER) + "}",
        r"\toprule",
        r"& & "
        + " & ".join(
            rf"\multicolumn{{2}}{{c}}{{{latex_escape(ROI_LABELS[roi])}}}"
            for roi in ROI_ORDER
        )
        + r" \\",
        r"\cmidrule(lr){3-4}\cmidrule(lr){5-6}\cmidrule(lr){7-8}\cmidrule(lr){9-10}",
        r"Outcome & Unit & "
        + " & ".join(["CamCan mean (SD) & MNI152"] * len(ROI_ORDER))
        + r" \\",
        r"\midrule",
    ]
    current_section = None
    for record in primary_wide.to_dict(orient="records"):
        if record["Section"] != current_section:
            if current_section is not None:
                lines.append(r"\addlinespace")
            lines.append(
                rf"\multicolumn{{10}}{{l}}{{\textit{{{latex_escape(record['Section'])}}}}} \\"
            )
            current_section = record["Section"]
        values = [latex_escape(record["Outcome"]), latex_escape(record["Unit"])]
        for roi in ROI_ORDER:
            display = ROI_LABELS[roi]
            values.extend(
                [
                    latex_escape(record[f"{display} — CamCan mean (SD)"]),
                    latex_escape(record[f"{display} — MNI152"]),
                ]
            )
        lines.append(" & ".join(values) + r" \\")
    lines.extend(
        [
            r"\bottomrule",
            r"\end{tabular}",
            r"\begin{minipage}{0.99\textwidth}",
            r"\footnotesize",
            (
                r"\textit{Notes:} CamCan values are mean (SD) across 132 "
                r"subjects. Each subject value is the arithmetic mean of ten "
                r"repeat-level metric values. MNI152 is a single descriptive "
                r"reference and is not treated as an inferential sample. "
                r"P99.9 denotes the 99.9th percentile."
            ),
            r"\end{minipage}",
            r"\end{table*}",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def write_threshold_latex(frame: pd.DataFrame, path: Path) -> None:
    lines = [
        r"\begin{table*}[htbp]",
        r"\centering",
        r"\caption{Threshold sensitivity of effectiveness and spatial spread.}",
        r"\label{tab:camcan_threshold_sensitivity}",
        r"\scriptsize",
        r"\setlength{\tabcolsep}{3pt}",
        r"\begin{tabular}{llcccccccc}",
        r"\toprule",
        r"ROI & Threshold & \multicolumn{2}{c}{Target coverage} & \multicolumn{2}{c}{Off-target coverage} & \multicolumn{2}{c}{Whole-brain coverage} & \multicolumn{2}{c}{Localization in target} \\",
        r"\cmidrule(lr){3-4}\cmidrule(lr){5-6}\cmidrule(lr){7-8}\cmidrule(lr){9-10}",
        r"& (V/m) & CamCan mean (SD) & MNI152 & CamCan mean (SD) & MNI152 & CamCan mean (SD) & MNI152 & CamCan mean (SD) & MNI152 \\",
        r"\midrule",
    ]
    columns = [
        "ROI",
        "Threshold (V/m)",
        "Target coverage — CamCan mean (SD), %",
        "Target coverage — MNI152, %",
        "Off-target coverage — CamCan mean (SD), %",
        "Off-target coverage — MNI152, %",
        "Whole-brain coverage — CamCan mean (SD), %",
        "Whole-brain coverage — MNI152, %",
        "Suprathreshold localization in target — CamCan mean (SD), %",
        "Suprathreshold localization in target — MNI152, %",
    ]
    for idx, record in enumerate(frame.to_dict(orient="records")):
        values = []
        for column in columns:
            value = record[column]
            if column == "Threshold (V/m)":
                value = f"{float(value):.2f}"
            values.append(latex_escape(value))
        lines.append(" & ".join(values) + r" \\")
        if idx % 2 == 1 and idx != len(frame) - 1:
            lines.append(r"\addlinespace")
    lines.extend(
        [
            r"\bottomrule",
            r"\end{tabular}",
            r"\begin{minipage}{0.99\textwidth}",
            r"\footnotesize",
            (
                r"\textit{Notes:} Values are percentages. CamCan values are "
                r"mean (SD) across 132 subjects after repeat-level metric "
                r"calculation and ten-repeat arithmetic-mean aggregation."
            ),
            r"\end{minipage}",
            r"\end{table*}",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def setup_matplotlib() -> None:
    plt.rcParams.update(
        {
            "font.family": "DejaVu Sans",
            "font.size": 9,
            "axes.titlesize": 10,
            "axes.labelsize": 9,
            "legend.fontsize": 8,
            "xtick.labelsize": 8,
            "ytick.labelsize": 8,
            "axes.spines.top": False,
            "axes.spines.right": False,
            "figure.dpi": 150,
            "savefig.dpi": 400,
            "savefig.bbox": "tight",
            "pdf.fonttype": 42,
            "ps.fonttype": 42,
        }
    )


def save_figure(fig, base_path: Path) -> None:
    fig.savefig(base_path.with_suffix(".png"), dpi=400, facecolor="white")
    fig.savefig(base_path.with_suffix(".pdf"), facecolor="white")
    plt.close(fig)


def plot_effectiveness_spread(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
    figures_dir: Path,
) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(10.6, 8.4))
    x_metric = "target_coverage_percent_ge_0p18"
    y_metric = "off_target_coverage_percent_ge_0p18"
    for ax, roi in zip(axes.flat, ROI_ORDER):
        data = subject_frame.loc[subject_frame["roi"] == roi]
        x = data[x_metric].astype(float)
        y = data[y_metric].astype(float)
        rho = spearmanr(x, y).statistic
        ax.scatter(
            x,
            y,
            s=20,
            alpha=0.52,
            color=FIGURE_BLUE,
            edgecolors="none",
            label="CamCan subject",
        )
        if np.ptp(x) > 0:
            coefficients = np.polyfit(x, y, 1)
            grid = np.linspace(float(x.min()), float(x.max()), 100)
            ax.plot(
                grid,
                coefficients[0] * grid + coefficients[1],
                color=FIGURE_GRAY,
                linewidth=1.2,
                linestyle="--",
            )
        ax.scatter(
            [mni_by_roi.loc[roi, x_metric]],
            [mni_by_roi.loc[roi, y_metric]],
            s=70,
            marker="D",
            color=FIGURE_ORANGE,
            edgecolors="white",
            linewidths=0.8,
            zorder=5,
            label="MNI152",
        )
        ax.text(
            0.04,
            0.95,
            rf"Spearman $\rho$ = {rho:.2f}",
            transform=ax.transAxes,
            ha="left",
            va="top",
            fontsize=9,
            bbox={"facecolor": "white", "alpha": 0.8, "edgecolor": "none"},
        )
        ax.set_title(ROI_LABELS[roi])
        ax.set_xlabel("Target coverage ≥0.18 V/m (%)")
        ax.set_ylabel("Off-target coverage ≥0.18 V/m (%)")
        ax.grid(True, color="#E5E7EB", linewidth=0.6)
        ax.set_xlim(left=min(0.0, float(x.min()) - 2))
        ax.set_ylim(bottom=min(0.0, float(y.min()) - 1))
    handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            color="none",
            markerfacecolor=FIGURE_BLUE,
            alpha=0.65,
            markersize=6,
            label="CamCan subject (ten-repeat mean)",
        ),
        Line2D(
            [0],
            [0],
            marker="D",
            color="none",
            markerfacecolor=FIGURE_ORANGE,
            markersize=7,
            label="MNI152 reference",
        ),
    ]
    fig.subplots_adjust(
        left=0.08,
        right=0.98,
        bottom=0.08,
        top=0.84,
        wspace=0.26,
        hspace=0.35,
    )
    fig.legend(
        handles=handles,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.925),
        ncol=2,
        frameon=False,
    )
    fig.suptitle(
        "Within-ROI effectiveness–spread relationship at 0.18 V/m",
        fontsize=13,
        fontweight="bold",
        y=0.985,
    )
    save_figure(
        fig,
        figures_dir / "figure_effectiveness_off_target_relationship_ge_0p18",
    )


def plot_threshold_sensitivity(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
    figures_dir: Path,
) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(10.6, 8.2))
    x = np.array([0, 1])
    for ax, roi in zip(axes.flat, ROI_ORDER):
        data = subject_frame.loc[subject_frame["roi"] == roi]
        for label, stem, color in [
            ("Target coverage", "target_coverage_percent", FIGURE_BLUE),
            ("Off-target coverage", "off_target_coverage_percent", FIGURE_RED),
        ]:
            metrics = [f"{stem}_ge_0p18", f"{stem}_ge_0p15"]
            means = np.array([data[metric].mean() for metric in metrics])
            sds = np.array([data[metric].std(ddof=1) for metric in metrics])
            mni = np.array([mni_by_roi.loc[roi, metric] for metric in metrics])
            ax.errorbar(
                x - 0.035,
                means,
                yerr=sds,
                marker="o",
                color=color,
                linewidth=1.8,
                capsize=3,
                markersize=5,
                label=f"CamCan {label}",
            )
            ax.plot(
                x + 0.035,
                mni,
                marker="D",
                color=color,
                linewidth=1.2,
                linestyle=":",
                markersize=5,
                alpha=0.9,
                label=f"MNI152 {label}",
            )
        ax.set_xticks(x, ["0.18", "0.15"])
        ax.set_xlabel("TI-field threshold (V/m)")
        ax.set_ylabel("Coverage (%)")
        ax.set_ylim(0, 105)
        ax.set_title(ROI_LABELS[roi])
        ax.grid(True, axis="y", color="#E5E7EB", linewidth=0.6)
    handles, labels = axes.flat[0].get_legend_handles_labels()
    fig.subplots_adjust(
        left=0.08,
        right=0.98,
        bottom=0.08,
        top=0.80,
        wspace=0.26,
        hspace=0.36,
    )
    fig.legend(
        handles,
        labels,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.91),
        ncol=2,
        frameon=False,
    )
    fig.suptitle(
        "Threshold sensitivity of target effectiveness and off-target spread",
        fontsize=13,
        fontweight="bold",
        y=0.985,
    )
    save_figure(fig, figures_dir / "figure_threshold_sensitivity")


def plot_mni_percentile_context(
    mni_context: pd.DataFrame,
    figures_dir: Path,
) -> None:
    metric_order = [metric for _, _, metric, _ in PRIMARY_METRICS]
    category_colors = {
        "Field magnitude": FIGURE_BLUE,
        "Threshold performance (0.18 V/m)": FIGURE_RED,
        "Rank-based sensitivity": FIGURE_GREEN,
    }
    fig, axes = plt.subplots(2, 2, figsize=(11.5, 9.2))
    for ax, roi in zip(axes.flat, ROI_ORDER):
        data = mni_context.loc[
            mni_context["ROI"] == ROI_LABELS[roi]
        ].copy()
        data["_order"] = data["Outcome"].map(
            {
                label: idx
                for idx, (_, label, _, _) in enumerate(PRIMARY_METRICS)
            }
        )
        data = data.sort_values("_order", ascending=False)
        y = np.arange(len(data))
        colors = [category_colors[section] for section in data["Section"]]
        ax.axvspan(25, 75, color="#F3F4F6", zorder=0)
        ax.axvline(50, color=FIGURE_GRAY, linewidth=1, linestyle="--")
        ax.scatter(
            data["MNI152 percentile within CamCan"],
            y,
            color=colors,
            s=42,
            edgecolor="white",
            linewidth=0.7,
            zorder=3,
        )
        ax.set_yticks(
            y,
            [
                PRIMARY_SHORT_LABELS[
                    next(
                        metric
                        for _, label, metric, _ in PRIMARY_METRICS
                        if label == outcome
                    )
                ]
                for outcome in data["Outcome"]
            ],
        )
        ax.set_xlim(0, 100)
        ax.set_xlabel("MNI152 percentile within CamCan (%)")
        ax.set_title(ROI_LABELS[roi])
        ax.grid(True, axis="x", color="#E5E7EB", linewidth=0.6)
    legend_handles = [
        Line2D(
            [0],
            [0],
            marker="o",
            color="none",
            markerfacecolor=color,
            markersize=7,
            label=label,
        )
        for label, color in category_colors.items()
    ]
    fig.subplots_adjust(
        left=0.16,
        right=0.98,
        bottom=0.08,
        top=0.82,
        wspace=0.50,
        hspace=0.32,
    )
    fig.legend(
        handles=legend_handles,
        loc="upper center",
        bbox_to_anchor=(0.5, 0.92),
        ncol=3,
        frameon=False,
    )
    fig.suptitle(
        "Position of the MNI152 reference within each CamCan distribution",
        fontsize=13,
        fontweight="bold",
        y=0.985,
    )
    save_figure(fig, figures_dir / "figure_mni152_percentile_context")


def plot_target_coverage_ecdf(
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
    figures_dir: Path,
) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(10.6, 8.2), constrained_layout=True)
    metric = "target_coverage_percent_ge_0p18"
    for ax, roi in zip(axes.flat, ROI_ORDER):
        values = np.sort(
            subject_frame.loc[subject_frame["roi"] == roi, metric].to_numpy(
                dtype=float
            )
        )
        y = np.arange(1, len(values) + 1) / len(values)
        mni = float(mni_by_roi.loc[roi, metric])
        ax.step(values, y, where="post", color=FIGURE_BLUE, linewidth=2)
        ax.axvline(
            mni,
            color=FIGURE_ORANGE,
            linestyle="--",
            linewidth=1.6,
            label=f"MNI152 ({mni:.1f}%)",
        )
        for cutoff in [25, 50, 75]:
            ax.axvline(cutoff, color="#D1D5DB", linewidth=0.7, zorder=0)
        ax.set_xlim(0, 100)
        ax.set_ylim(0, 1.02)
        ax.set_title(ROI_LABELS[roi])
        ax.set_xlabel("Target coverage ≥0.18 V/m (%)")
        ax.set_ylabel("Cumulative proportion of subjects")
        ax.legend(frameon=False, loc="lower right")
        ax.grid(True, axis="y", color="#E5E7EB", linewidth=0.6)
    fig.suptitle(
        "Inter-individual distribution of target coverage",
        fontsize=13,
        fontweight="bold",
    )
    save_figure(fig, figures_dir / "figure_target_coverage_ecdf_ge_0p18")


def write_results_highlights(
    path: Path,
    subject_frame: pd.DataFrame,
    mni_by_roi: pd.DataFrame,
    association: pd.DataFrame,
    shift: pd.DataFrame,
    attainment: pd.DataFrame,
    robustness: pd.DataFrame,
) -> None:
    def value(roi: str, metric: str) -> tuple[float, float]:
        values = subject_frame.loc[subject_frame["roi"] == roi, metric]
        return float(values.mean()), float(values.std(ddof=1))

    lines = [
        "# CamCan manuscript analysis: publication-facing descriptive findings",
        "",
        "These notes are generated from the validated schema-2 outputs. They are descriptive and do not introduce formal cross-ROI inference.",
        "",
        "## Principal patterns",
        "",
        "### 1. Effectiveness and spatial spread are tightly coupled within subjects",
        "",
    ]
    assoc_018 = association.loc[
        (association["Threshold (V/m)"] == 0.18)
        & (
            association["Relationship"]
            == "Effectiveness–spread coupling"
        )
    ]
    for roi in ROI_ORDER:
        rho = float(
            assoc_018.loc[
                assoc_018["ROI"] == ROI_LABELS[roi], "Spearman ρ"
            ].iloc[0]
        )
        lines.append(f"- {ROI_LABELS[roi]}: Spearman ρ = {rho:.3f}.")
    lines.extend(
        [
            "",
            "Subjects with greater target coverage also tended to have greater off-target coverage. This makes the effectiveness–spread plane more informative than either outcome alone. It should be described as an anatomical/field-distribution relationship, not as proof of a causal trade-off.",
            "",
            "### 2. Lowering the threshold changes deep- and cortical-target summaries differently",
            "",
        ]
    )
    for roi in ROI_ORDER:
        roi_shift = shift.loc[shift["ROI"] == ROI_LABELS[roi]].set_index(
            "Outcome"
        )
        target_delta = roi_shift.loc[
            "Target coverage",
            "Paired mean change (0.15 − 0.18), percentage points",
        ]
        off_delta = roi_shift.loc[
            "Off-target coverage",
            "Paired mean change (0.15 − 0.18), percentage points",
        ]
        lines.append(
            f"- {ROI_LABELS[roi]}: target coverage +{target_delta:.1f} percentage points; off-target coverage +{off_delta:.1f} percentage points."
        )
    lines.extend(
        [
            "",
            "At 0.15 V/m, target coverage increases for every ROI. The accompanying increase in off-target coverage is much larger for the hippocampal and thalamic montages than for M1 and DLPFC. This supports retaining 0.18 V/m as the primary threshold and presenting 0.15 V/m as sensitivity analysis rather than pooling the two.",
            "",
            "### 3. Target coverage is highly heterogeneous for the cortical targets",
            "",
        ]
    )
    for _, row in attainment.iterrows():
        lines.append(
            f"- {row['ROI']}: {row['Subjects with ≤1% target coverage, %']:.1f}% of subjects had ≤1% target coverage; {row['Subjects with ≥50% target coverage, %']:.1f}% reached ≥50% coverage."
        )
    lines.extend(
        [
            "",
            "The right-DLPFC montage is especially heterogeneous at 0.18 V/m: many subjects have little or no suprathreshold target volume, while a small minority have substantially greater coverage. The ECDF figure communicates this more faithfully than mean ± SD alone.",
            "",
            "### 4. MNI152 is not uniformly representative of the cohort",
            "",
        ]
    )
    context_metrics = [
        ("roi_median_v_per_m", "median target field"),
        ("roi_robust_max_p99_9_v_per_m", "target P99.9"),
        ("target_coverage_percent_ge_0p18", "target coverage"),
        ("off_target_coverage_percent_ge_0p18", "off-target coverage"),
        (
            "threshold_localization_percent_in_roi_ge_0p18",
            "suprathreshold localization",
        ),
    ]
    for roi in ROI_ORDER:
        roi_subjects = subject_frame.loc[subject_frame["roi"] == roi]
        bits = []
        for metric, label in context_metrics:
            percentile = percentileofscore(
                roi_subjects[metric],
                float(mni_by_roi.loc[roi, metric]),
                kind="mean",
            )
            bits.append(f"{label} percentile rank {percentile:.1f}")
        lines.append(f"- {ROI_LABELS[roi]} MNI percentiles: " + "; ".join(bits) + ".")
    lines.extend(
        [
            "",
            "For the left hippocampus and right thalamus, MNI152 generally has lower field magnitude, target coverage, and off-target spread than the CamCan mean, while its localization percentage is relatively high. Right DLPFC is much closer to the cohort center. The MNI baseline should therefore be presented as a reference configuration, not as a typical participant.",
            "",
            "### 5. The robust-maximum conclusion is insensitive to the exact robust estimator",
            "",
        ]
    )
    target_robust = robustness.loc[robustness["Region"] == "Target ROI"]
    rho_min = float(
        target_robust[
            "CamCan Spearman ρ, P99.9 vs upper-1% median"
        ].min()
    )
    rho_max = float(
        target_robust[
            "CamCan Spearman ρ, P99.9 vs upper-1% median"
        ].max()
    )
    lines.extend(
        [
            f"Across the four target ROIs, P99.9 and the upper-1% median rank subjects almost identically (Spearman ρ = {rho_min:.3f}–{rho_max:.3f}). P99.9 is therefore defensible as the primary robust maximum, with the upper-1% median as a sensitivity check.",
            "",
            "## Suggested figure placement",
            "",
            "- Main text: the effectiveness–off-target relationship at 0.18 V/m.",
            "- Main text or supplement: target-coverage ECDFs with MNI152 reference.",
            "- Supplement: threshold-sensitivity figure (0.18 versus 0.15 V/m).",
            "- Supplement: MNI152 percentile-context profile.",
            "",
            "## Interpretation guardrails",
            "",
            "- Do not claim that one ROI or montage is statistically superior to another; target anatomy, ROI volume, montage, and field distribution differ.",
            "- Localization percentages are influenced by target ROI size. Interpret them alongside target and off-target coverage.",
            "- MNI152 is one deterministic reference head and must not be assigned a standard error or used in a two-sample test.",
            "- Correlations are descriptive within ROI and do not establish causality.",
            "- The cohort simulations used SimNIBS 4.0.1, whereas the corrected MNI152 baselines were generated with SimNIBS 4.5.0; disclose this version difference when the MNI comparison is reported.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def build_figure_captions(association: pd.DataFrame) -> pd.DataFrame:
    assoc_018 = association.loc[
        (association["Threshold (V/m)"] == 0.18)
        & (
            association["Relationship"]
            == "Effectiveness–spread coupling"
        )
    ].set_index("ROI")
    rho_text = ", ".join(
        f"{ROI_LABELS[roi]} ρ={assoc_018.loc[ROI_LABELS[roi], 'Spearman ρ']:.2f}"
        for roi in ROI_ORDER
    )
    records = [
        {
            "Designation": "Figure 1",
            "Figure file stem": (
                "figure_effectiveness_off_target_relationship_ge_0p18"
            ),
            "Short title": (
                "Target effectiveness and off-target spread at 0.18 V/m"
            ),
            "Caption": (
                "Figure 1. Relationship between target effectiveness and "
                "off-target spread during temporal-interference (TI) "
                "stimulation in 132 CamCan participants. Four fixed "
                "electrode montages were evaluated separately, targeting the "
                "left hippocampus, left primary motor cortex (M1), right "
                "dorsolateral prefrontal cortex (DLPFC), and right thalamus. "
                "Target coverage (horizontal axis) is the percentage of "
                "voxels in the anatomical target region of interest (ROI) "
                "whose TI electric-field magnitude was at least 0.18 V/m; "
                "off-target coverage (vertical axis) is the percentage of "
                "finite whole-brain voxels outside that ROI meeting the same "
                "threshold. Each blue point represents one participant. "
                "Metrics were calculated separately for ten independently "
                "remeshed head models generated from the same corrected "
                "tissue segmentation and were then averaged within "
                "participant and ROI. The orange diamond is the corrected "
                "MNI152 reference head simulated with the same "
                "target-specific montage and current settings. Gray dashed "
                "lines are descriptive least-squares fits, and the "
                "within-ROI Spearman rank correlations are shown in the "
                f"panels ({rho_text}). Movement to the right indicates more "
                "of the target reaching the threshold, whereas upward "
                "movement indicates more suprathreshold tissue outside the "
                "target. Correlations are descriptive; no inferential "
                "comparison was made across ROIs. CamCan simulations used "
                "SimNIBS 4.0.1 and the corrected MNI152 reference used "
                "SimNIBS 4.5.0."
            ),
        },
        {
            "Designation": "Figure 2",
            "Figure file stem": "figure_target_coverage_ecdf_ge_0p18",
            "Short title": (
                "Distribution of target coverage at 0.18 V/m"
            ),
            "Caption": (
                "Figure 2. Inter-individual distribution of target "
                "coverage during temporal-interference (TI) stimulation in "
                "132 CamCan participants. Separate panels show fixed "
                "electrode montages targeting the left hippocampus, left "
                "primary motor cortex (M1), right dorsolateral prefrontal "
                "cortex (DLPFC), and right thalamus. Target coverage is the "
                "percentage of voxels in the anatomical target region of "
                "interest whose TI electric-field magnitude was at least "
                "0.18 V/m. The blue empirical cumulative distribution "
                "function gives, at each horizontal-axis value, the "
                "proportion of participants with target coverage less than "
                "or equal to that value; a curve shifted farther right "
                "therefore indicates generally greater target coverage. "
                "Each participant's value is the arithmetic mean of metrics "
                "calculated independently on ten remeshed head models "
                "generated from the same corrected tissue segmentation. "
                "The orange dashed vertical line marks the corrected MNI152 "
                "reference value obtained with the same target-specific "
                "montage and current settings. The light-gray vertical "
                "lines at 25%, 50%, and 75% are coverage guides, not cohort "
                "quartiles. The distributions are descriptive and were not "
                "used for formal cross-ROI comparisons. CamCan simulations "
                "used SimNIBS 4.0.1 and the corrected MNI152 reference used "
                "SimNIBS 4.5.0."
            ),
        },
        {
            "Designation": "Supplementary Figure S1",
            "Figure file stem": "figure_threshold_sensitivity",
            "Short title": (
                "Sensitivity of target and off-target coverage to threshold"
            ),
            "Caption": (
                "Supplementary Figure S1. Sensitivity of target and "
                "off-target coverage to the temporal-interference (TI) "
                "electric-field threshold in 132 CamCan participants. "
                "Separate panels show fixed electrode montages targeting "
                "the left hippocampus, left primary motor cortex (M1), "
                "right dorsolateral prefrontal cortex (DLPFC), and right "
                "thalamus. Target coverage (blue) is the percentage of "
                "voxels in the anatomical target region of interest (ROI) "
                "meeting the indicated threshold; off-target coverage "
                "(red) is the percentage of finite whole-brain voxels "
                "outside the ROI meeting that threshold. Solid lines with "
                "circles show the CamCan mean and error bars show one "
                "between-participant standard deviation. Each participant's "
                "value is the arithmetic mean of metrics calculated "
                "independently on ten remeshed head models generated from "
                "the same corrected tissue segmentation. Dotted lines with "
                "diamonds show the single corrected MNI152 reference "
                "simulated with the same target-specific montage and current "
                "settings; no error bars are assigned to this single "
                "reference head. The prespecified primary threshold was "
                "0.18 V/m, and 0.15 V/m was evaluated as a sensitivity "
                "threshold. An increase at 0.15 V/m means that more tissue "
                "is classified as suprathreshold when the criterion is "
                "lowered; it does not represent a change in the simulated "
                "electric field. CamCan simulations used SimNIBS 4.0.1 and "
                "the corrected MNI152 reference used SimNIBS 4.5.0."
            ),
        },
        {
            "Designation": "Supplementary Figure S2",
            "Figure file stem": "figure_mni152_percentile_context",
            "Short title": (
                "MNI152 position within the CamCan outcome distributions"
            ),
            "Caption": (
                "Supplementary Figure S2. Position of the corrected "
                "MNI152 reference head within outcome distributions from "
                "132 CamCan participants undergoing simulated "
                "temporal-interference (TI) stimulation. Separate panels "
                "show fixed electrode montages targeting the left "
                "hippocampus, left primary motor cortex (M1), right "
                "dorsolateral prefrontal cortex (DLPFC), and right thalamus. "
                "Each point is the percentile rank of the single MNI152 "
                "value within the corresponding distribution of 132 "
                "participant values; 0% denotes the bottom of the CamCan "
                "distribution, 50% its median, and 100% its top. Participant "
                "values are arithmetic means of metrics calculated "
                "independently on ten remeshed head models generated from "
                "the same corrected tissue segmentation. Blue points are "
                "target-ROI field-magnitude measures: the median and P99.9 "
                "(the 99.9th percentile, used as a robust maximum). Red "
                "points are 0.18-V/m threshold measures: target coverage "
                "(percentage of anatomical target voxels reaching the "
                "threshold), off-target and whole-brain coverage "
                "(percentages of their respective finite voxel sets "
                "reaching the threshold), and localization (percentage of "
                "all suprathreshold whole-brain voxels lying inside the "
                "target). Green points are rank-based measures using voxels "
                "at or above the whole-brain 95th percentile: the percentage "
                "of anatomical target voxels in this top 5% and the "
                "percentage of all top-5% voxels located inside the target. "
                "The gray band spans the 25th–75th CamCan percentiles and "
                "the dashed line marks the cohort median. Percentile "
                "direction is metric dependent—for example, a high target "
                "coverage rank and a high off-target coverage rank do not "
                "have the same interpretation. MNI152 was simulated with "
                "the same target-specific montage and current settings as "
                "the cohort. Percentiles are descriptive, not inferential. "
                "CamCan simulations used SimNIBS 4.0.1 and the corrected "
                "MNI152 reference used SimNIBS 4.5.0."
            ),
        },
    ]
    return pd.DataFrame(records)


def write_figure_captions(path: Path, captions: pd.DataFrame) -> None:
    blocks = ["# Publication-ready figure captions", ""]
    for record in captions.to_dict(orient="records"):
        blocks.extend(
            [
                f"## {record['Designation']}: {record['Short title']}",
                "",
                f"Files: `{record['Figure file stem']}.png` and "
                f"`{record['Figure file stem']}.pdf`",
                "",
                record["Caption"],
                "",
            ]
        )
    text = "\n".join(blocks)
    path.write_text(text, encoding="utf-8")


def write_individual_figure_captions(
    captions_dir: Path,
    captions: pd.DataFrame,
) -> None:
    captions_dir.mkdir(exist_ok=True)
    for record in captions.to_dict(orient="records"):
        path = captions_dir / f"{record['Figure file stem']}_caption.txt"
        path.write_text(record["Caption"] + "\n", encoding="utf-8")


def write_figure_caption_sheet(
    wb: Workbook,
    captions: pd.DataFrame,
    source_comment: str,
) -> None:
    ws = wb.create_sheet("Figure Captions")
    add_title(
        ws,
        "Publication-ready figure captions",
        (
            "Each caption is self-contained and can be copied directly into "
            "the manuscript; update figure numbering if journal layout changes."
        ),
        4,
        header_row=4,
    )
    headers = list(captions.columns)
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(4, col_idx, header)
        cell.font = Font(name="Arial", size=9, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        if col_idx == 1:
            cell.comment = Comment(source_comment, "Codex")
    for row_idx, record in enumerate(
        captions.itertuples(index=False), start=5
    ):
        for col_idx, value in enumerate(record, start=1):
            cell = ws.cell(row_idx, col_idx, value)
            cell.font = Font(name="Arial", size=9, color=BLACK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7F9FB")
        ws.row_dimensions[row_idx].height = 250
    for column, width in {"A": 24, "B": 54, "C": 46, "D": 135}.items():
        ws.column_dimensions[column].width = width
    ws.freeze_panes = "A5"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = f"A4:D{4 + len(captions)}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Page &P of &N"


def write_output_readme(path: Path, input_dir: Path, out_dir: Path) -> None:
    text = f"""# CamCan final-132 publication outputs

This directory contains publication-facing tables and descriptive figures generated from the validated schema-2 manuscript-analysis results.

## Recommended manuscript set

- **Main Table 1:** `tables/table_1_primary_publication.csv` or the journal-ready LaTeX fragment `tables/table_1_primary_publication.tex`.
- **Editable workbook:** `camcan_final132_publication_tables.xlsx`.
- **Main figure candidate:** `figures/figure_effectiveness_off_target_relationship_ge_0p18.pdf`.
- **Second main/supplement candidate:** `figures/figure_target_coverage_ecdf_ge_0p18.pdf`.
- **Supplementary threshold table:** `tables/table_s1_threshold_sensitivity.csv` or `.tex`.
- **Complete numerical supplement:** `tables/table_s3_full_descriptive_statistics.csv`.

## Additional descriptive outputs

- `table_s2_robust_maximum_sensitivity.csv`: P99.9 versus upper-1% median.
- `table_s4_mni152_context.csv`: MNI percentile and standardized context.
- `table_s5_exploratory_within_roi_associations.csv`: descriptive within-ROI correlations.
- `table_s6_threshold_paired_changes.csv`: paired 0.15-minus-0.18 changes.
- `table_s7_target_coverage_attainment.csv`: clinically readable coverage thresholds.
- `table_qc_repeat_vs_between_subject_variation.csv`: QC only; coordinate with the separate repeatability manuscript before reporting.
- `results_highlights.md`: evidence-backed interpretation and guardrails.
- `figure_captions.md`: self-contained publication-ready captions.
- `figure_captions.csv`: copy-ready caption index.
- `captions/`: one plain-text caption file for each figure.
- `publication_output_manifest.json`: source/output hashes and validation metadata.

Every CamCan subject value was produced by calculating each metric independently per repeat and then arithmetic-mean aggregating the ten repeat-level values. MNI152 is a single descriptive reference, not an inferential sample. No formal cross-ROI tests are introduced.

## Reproduce

```bash
MPLCONFIGDIR=/tmp/camcan_pub_mpl \\
  /home/boyan/anaconda3/envs/simnibs_post/bin/python \\
  CamCan_Experiment/post/build_camcan_publication_outputs.py \\
  --input-dir {json.dumps(str(input_dir))} \\
  --out-dir {json.dumps(str(out_dir))}
```
"""
    path.write_text(text, encoding="utf-8")


def write_manifest(
    path: Path,
    input_dir: Path,
    output_dir: Path,
    source_hashes: dict[str, str],
) -> None:
    outputs = []
    for item in sorted(output_dir.rglob("*")):
        if item.is_file() and item != path:
            outputs.append(
                {
                    "path": str(item.relative_to(output_dir)),
                    "bytes": item.stat().st_size,
                    "sha256": sha256_file(item),
                }
            )
    payload = {
        "status": "complete",
        "source_schema": 2,
        "source_directory": str(input_dir),
        "cohort_subjects": 132,
        "rois": ROI_ORDER,
        "repeats_per_subject_roi": 10,
        "aggregation": (
            "Metric calculated independently per repeat, then arithmetic-mean "
            "aggregated across ten repeats per subject and ROI."
        ),
        "primary_threshold_v_per_m": 0.18,
        "sensitivity_threshold_v_per_m": 0.15,
        "cross_roi_inference": False,
        "mni_inference": False,
        "source_sha256": source_hashes,
        "outputs": outputs,
    }
    path.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--input-dir",
        type=Path,
        required=True,
        help="Directory containing validated schema-2 manuscript outputs.",
    )
    parser.add_argument(
        "--out-dir",
        type=Path,
        required=True,
        help="Output directory for publication-ready tables and figures.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    input_dir = args.input_dir.expanduser().resolve()
    out_dir = args.out_dir.expanduser().resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    tables_dir = out_dir / "tables"
    figures_dir = out_dir / "figures"
    captions_dir = out_dir / "captions"
    tables_dir.mkdir(exist_ok=True)
    figures_dir.mkdir(exist_ok=True)
    captions_dir.mkdir(exist_ok=True)

    input_files = {
        "analysis_manifest.json": input_dir / "analysis_manifest.json",
        "repeat_level_metrics.csv": input_dir / "repeat_level_metrics.csv",
        "subject_level_repeat_mean_metrics.csv": (
            input_dir / "subject_level_repeat_mean_metrics.csv"
        ),
        "mni152_baseline_metrics.csv": (
            input_dir / "mni152_baseline_metrics.csv"
        ),
        "table_supplementary_descriptive_statistics.csv": (
            input_dir / "table_supplementary_descriptive_statistics.csv"
        ),
        "metric_dictionary.csv": input_dir / "metric_dictionary.csv",
    }
    missing = [str(path) for path in input_files.values() if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing input(s): " + ", ".join(missing))

    with input_files["analysis_manifest.json"].open() as handle:
        manifest = json.load(handle)
    repeat_frame = pd.read_csv(input_files["repeat_level_metrics.csv"])
    subject_frame = pd.read_csv(
        input_files["subject_level_repeat_mean_metrics.csv"]
    )
    mni_frame = pd.read_csv(input_files["mni152_baseline_metrics.csv"])
    supplementary = pd.read_csv(
        input_files["table_supplementary_descriptive_statistics.csv"]
    )
    metric_dictionary = pd.read_csv(input_files["metric_dictionary.csv"])
    validate_inputs(
        manifest,
        repeat_frame,
        subject_frame,
        mni_frame,
        supplementary,
    )
    mni_by_roi = mni_frame.set_index("roi")
    source_hashes = {
        name: sha256_file(path) for name, path in input_files.items()
    }

    primary_wide, primary_long = build_primary_tables(
        subject_frame, mni_by_roi
    )
    threshold_table = build_threshold_table(subject_frame, mni_by_roi)
    robustness = build_robustness_table(subject_frame, mni_by_roi)
    mni_context = build_mni_context_table(primary_long)
    association = build_association_table(subject_frame)
    shift = build_threshold_shift_table(subject_frame, mni_by_roi)
    attainment = build_attainment_table(subject_frame, mni_by_roi)
    repeat_qc = build_repeat_qc_table(subject_frame)
    figure_captions = build_figure_captions(association)

    table_frames = {
        "table_1_primary_publication.csv": primary_wide,
        "table_1_primary_numeric_long.csv": primary_long,
        "table_s1_threshold_sensitivity.csv": threshold_table,
        "table_s2_robust_maximum_sensitivity.csv": robustness,
        "table_s3_full_descriptive_statistics.csv": supplementary,
        "table_s4_mni152_context.csv": mni_context,
        "table_s5_exploratory_within_roi_associations.csv": association,
        "table_s6_threshold_paired_changes.csv": shift,
        "table_s7_target_coverage_attainment.csv": attainment,
        "table_qc_repeat_vs_between_subject_variation.csv": repeat_qc,
        "metric_dictionary.csv": metric_dictionary,
    }
    for filename, frame in table_frames.items():
        frame.to_csv(tables_dir / filename, index=False)

    workbook = Workbook()
    write_readme_sheet(workbook, input_dir, manifest, source_hashes)
    source_comment = (
        "Generated from validated schema-2 CamCan manuscript analysis. "
        f"Source: {input_dir}"
    )
    write_primary_sheet(workbook, primary_wide, source_comment)
    write_frame_sheet(
        workbook,
        "Table S1 Thresholds",
        "Table S1. Threshold sensitivity of effectiveness and spatial spread",
        (
            "0.18 V/m is primary; 0.15 V/m is sensitivity. CamCan values are "
            "mean (SD); MNI152 is a single descriptive reference."
        ),
        threshold_table,
        number_formats={"Threshold (V/m)": "0.00"},
        source_comment=source_comment,
    )
    write_frame_sheet(
        workbook,
        "Table S2 Robustness",
        "Table S2. Robust-maximum definition sensitivity",
        (
            "P99.9 is primary. The upper-1% median is a sensitivity estimator. "
            "Correlations are descriptive within ROI."
        ),
        robustness,
        number_formats={
            "CamCan Spearman ρ, P99.9 vs upper-1% median": "0.000",
            "Mean upper-1% median / P99.9 ratio": "0.000",
        },
        source_comment=source_comment,
    )
    write_frame_sheet(
        workbook,
        "Table S3 Full Stats",
        "Table S3. Full descriptive statistics for all prespecified outcomes",
        (
            "CamCan values summarize 132 subject-level ten-repeat means. "
            "MNI152 is a single descriptive reference."
        ),
        supplementary,
        source_comment=source_comment,
    )
    write_frame_sheet(
        workbook,
        "Table S4 MNI Context",
        "Table S4. Position of MNI152 within CamCan subject distributions",
        (
            "Percentiles and standardized differences are descriptive; no "
            "inferential comparison is made against the single MNI152 head."
        ),
        mni_context,
        percent_columns={"MNI152 percentile within CamCan"},
        number_formats={
            "CamCan mean": "0.000",
            "CamCan SD": "0.000",
            "CamCan median": "0.000",
            "CamCan Q1": "0.000",
            "CamCan Q3": "0.000",
            "MNI152": "0.000",
            "MNI152 standardized difference from CamCan mean": "0.00",
        },
        source_comment=source_comment,
    )
    write_frame_sheet(
        workbook,
        "Table S5 Associations",
        "Table S5. Exploratory within-ROI associations",
        (
            "Correlations are descriptive and were not specified as cross-ROI "
            "hypothesis tests. Exact p values are retained for transparency."
        ),
        association,
        number_formats={
            "Threshold (V/m)": "0.00",
            "Spearman ρ": "0.000",
            "Spearman p (descriptive)": "0.000E+00",
            "Pearson r": "0.000",
            "Pearson p (descriptive)": "0.000E+00",
        },
        source_comment=source_comment,
    )
    write_frame_sheet(
        workbook,
        "Table S6 Threshold Shifts",
        "Table S6. Paired subject-level changes from 0.18 to 0.15 V/m",
        (
            "Positive values denote a higher percentage at the lower field "
            "threshold. Changes are calculated within subject and ROI."
        ),
        shift,
        source_comment=source_comment,
    )
    write_frame_sheet(
        workbook,
        "Table S7 Attainment",
        "Table S7. Distributional target-coverage attainment at 0.18 V/m",
        (
            "Percentages indicate the proportion of 132 CamCan subjects whose "
            "ten-repeat mean target coverage meets each criterion."
        ),
        attainment,
        source_comment=source_comment,
    )
    write_frame_sheet(
        workbook,
        "QC Repeat Stability",
        "Exploratory QC. Within-subject repeat variation versus between-subject variation",
        (
            "This sheet supports quality control and should not be placed in "
            "the main manuscript without coordination with the separate "
            "repeatability analysis."
        ),
        repeat_qc,
        source_comment=source_comment,
    )
    write_frame_sheet(
        workbook,
        "Metric Dictionary",
        "Metric dictionary",
        "Definitions are inherited from the validated schema-2 analysis.",
        metric_dictionary,
        source_comment=source_comment,
    )
    write_figure_caption_sheet(
        workbook,
        figure_captions,
        source_comment,
    )
    workbook_path = out_dir / "camcan_final132_publication_tables.xlsx"
    workbook.save(workbook_path)

    write_primary_latex(
        primary_wide, tables_dir / "table_1_primary_publication.tex"
    )
    write_threshold_latex(
        threshold_table,
        tables_dir / "table_s1_threshold_sensitivity.tex",
    )

    setup_matplotlib()
    plot_effectiveness_spread(subject_frame, mni_by_roi, figures_dir)
    plot_threshold_sensitivity(subject_frame, mni_by_roi, figures_dir)
    plot_mni_percentile_context(mni_context, figures_dir)
    plot_target_coverage_ecdf(subject_frame, mni_by_roi, figures_dir)
    write_results_highlights(
        out_dir / "results_highlights.md",
        subject_frame,
        mni_by_roi,
        association,
        shift,
        attainment,
        robustness,
    )
    write_figure_captions(
        out_dir / "figure_captions.md",
        figure_captions,
    )
    figure_captions.to_csv(out_dir / "figure_captions.csv", index=False)
    write_individual_figure_captions(captions_dir, figure_captions)
    write_output_readme(
        out_dir / "README.md",
        input_dir,
        out_dir,
    )
    write_manifest(
        out_dir / "publication_output_manifest.json",
        input_dir,
        out_dir,
        source_hashes,
    )

    print(
        json.dumps(
            {
                "status": "complete",
                "input_dir": str(input_dir),
                "out_dir": str(out_dir),
                "workbook": str(workbook_path),
                "table_files": len(list(tables_dir.glob("*"))),
                "figure_files": len(list(figures_dir.glob("*"))),
                "caption_files": len(list(captions_dir.glob("*"))),
                "subjects": 132,
                "rois": ROI_ORDER,
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
