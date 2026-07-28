#!/usr/bin/env python3
"""Build an illustrated supervisor guide for the CamCan TI figures.

The guide is intentionally explanatory rather than manuscript-like. It embeds
the four publication figures and documents their data lineage, exact metric
definitions, repeat aggregation, visual encodings, descriptive findings, and
interpretation guardrails.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


NAVY = "17365D"
BLUE = "2F6B9A"
PALE_BLUE = "EAF2F8"
PALE_ORANGE = "FCE4D6"
PALE_GRAY = "F3F4F6"
WHITE = "FFFFFF"
BLACK = "1F2937"
GRAY = "4B5563"
ORANGE = "D97706"

ROI_ORDER = [
    "Left hippocampus",
    "Left M1",
    "Right DLPFC",
    "Right thalamus",
]

FIGURES = [
    (
        "Figure 1",
        "figure_effectiveness_off_target_relationship_ge_0p18.png",
        "Target effectiveness and off-target spread at 0.18 V/m",
    ),
    (
        "Figure 2",
        "figure_target_coverage_ecdf_ge_0p18.png",
        "Inter-individual target-coverage distributions",
    ),
    (
        "Supplementary Figure S1",
        "figure_threshold_sensitivity.png",
        "Sensitivity to the 0.18 and 0.15 V/m thresholds",
    ),
    (
        "Supplementary Figure S2",
        "figure_mni152_percentile_context.png",
        "MNI152 context within the CamCan distributions",
    ),
]


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def set_cell_shading(cell, color: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), color)


def set_cell_margins(
    cell,
    *,
    top: int = 80,
    start: int = 100,
    bottom: int = 80,
    end: int = 100,
) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.name = "Arial"
    run.font.size = Pt(8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.45)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.header_distance = Cm(0.55)
    section.footer_distance = Cm(0.55)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in [
        ("Title", 25, NAVY),
        ("Subtitle", 13, GRAY),
        ("Heading 1", 18, NAVY),
        ("Heading 2", 13, BLUE),
        ("Heading 3", 10.5, ORANGE),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)

    caption = styles["Caption"]
    caption.font.name = "Arial"
    caption.font.size = Pt(8)
    caption.font.color.rgb = RGBColor.from_string(GRAY)
    caption.font.italic = False
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(5)

    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
    else:
        callout = styles["Callout"]
    callout.font.name = "Arial"
    callout.font.size = Pt(10)
    callout.font.color.rgb = RGBColor.from_string(NAVY)
    callout.font.bold = True
    callout.paragraph_format.space_after = Pt(0)

    for doc_section in document.sections:
        header = doc_section.header
        paragraph = header.paragraphs[0]
        paragraph.text = "CamCan TI simulation figures - supervisor review guide"
        paragraph.style = styles["Normal"]
        paragraph.runs[0].font.size = Pt(8)
        paragraph.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
        footer = doc_section.footer
        add_page_number(footer.paragraphs[0])


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_bullet(document: Document, text: str, level: int = 0) -> None:
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run(text)


def add_number(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Number")
    paragraph.add_run(text)


def add_callout(
    document: Document,
    title: str,
    text: str,
    color: str = PALE_BLUE,
) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, color)
    set_cell_margins(cell, top=130, start=160, bottom=130, end=160)
    paragraph = cell.paragraphs[0]
    paragraph.style = document.styles["Callout"]
    paragraph.add_run(f"{title}: ")
    detail = paragraph.add_run(text)
    detail.bold = False
    detail.font.color.rgb = RGBColor.from_string(BLACK)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_cm: list[float] | None = None,
    font_size: float = 8.2,
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    table.autofit = True
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    prevent_row_split(header_row)
    for index, value in enumerate(headers):
        cell = header_row.cells[index]
        cell.text = value
        set_cell_shading(cell, NAVY)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        for run in cell.paragraphs[0].runs:
            run.font.name = "Arial"
            run.font.size = Pt(font_size)
            run.font.bold = True
            run.font.color.rgb = RGBColor.from_string(WHITE)
        if widths_cm:
            cell.width = Cm(widths_cm[index])
    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.text = str(value)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if row_index % 2 == 0:
                set_cell_shading(cell, PALE_GRAY)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(font_size)
                    run.font.color.rgb = RGBColor.from_string(BLACK)
            if widths_cm:
                cell.width = Cm(widths_cm[column_index])
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_figure_page(
    document: Document,
    designation: str,
    title: str,
    image_path: Path,
    one_line: str,
) -> None:
    document.add_page_break()
    heading = document.add_heading(f"{designation}. {title}", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph = document.add_paragraph()
    picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture_paragraph.paragraph_format.space_after = Pt(3)
    picture_paragraph.paragraph_format.keep_with_next = True
    picture_paragraph.add_run().add_picture(
        str(image_path),
        width=Inches(7.95),
    )
    caption = document.add_paragraph(one_line, style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.keep_together = True


def f1(value: str | float) -> str:
    return f"{float(value):.1f}"


def f2(value: str | float) -> str:
    return f"{float(value):.2f}"


def mapping(
    rows: list[dict[str, str]],
    keys: tuple[str, ...],
) -> dict[tuple[str, ...], dict[str, str]]:
    return {tuple(row[key] for key in keys): row for row in rows}


def build_document(package: Path, targets_csv: Path, output: Path) -> None:
    tables = package / "tables"
    figures = package / "figures"
    required = [
        tables / "table_1_primary_numeric_long.csv",
        tables / "table_s1_threshold_sensitivity.csv",
        tables / "table_s4_mni152_context.csv",
        tables / "table_s5_exploratory_within_roi_associations.csv",
        tables / "table_s6_threshold_paired_changes.csv",
        tables / "table_s7_target_coverage_attainment.csv",
        targets_csv,
    ]
    required.extend(figures / filename for _, filename, _ in FIGURES)
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing input(s): " + ", ".join(missing))

    primary = read_csv(tables / "table_1_primary_numeric_long.csv")
    thresholds = read_csv(tables / "table_s1_threshold_sensitivity.csv")
    mni_context = read_csv(tables / "table_s4_mni152_context.csv")
    associations = read_csv(
        tables / "table_s5_exploratory_within_roi_associations.csv"
    )
    shifts = read_csv(tables / "table_s6_threshold_paired_changes.csv")
    attainment = read_csv(tables / "table_s7_target_coverage_attainment.csv")
    targets = read_csv(targets_csv)

    primary_by = mapping(primary, ("roi", "metric"))
    threshold_by = mapping(thresholds, ("ROI", "Threshold (V/m)"))
    context_by = mapping(mni_context, ("ROI", "Outcome"))
    association_by = mapping(
        associations,
        ("ROI", "Threshold (V/m)", "Relationship"),
    )
    shift_by = mapping(shifts, ("ROI", "Outcome"))
    attainment_by = mapping(attainment, ("ROI",))
    targets_by = mapping(targets, ("roi",))

    document = Document()
    configure_document(document)

    # Cover
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(55)
    run = paragraph.add_run("CamCan Temporal-Interference\nSimulation Figures")
    run.font.name = "Arial"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run(
        "Supervisor review guide: what each figure shows, how it was "
        "computed, and how it should be interpreted"
    ).font.size = Pt(14)
    rule = document.add_table(rows=1, cols=1)
    rule.alignment = WD_TABLE_ALIGNMENT.CENTER
    rule.cell(0, 0).width = Inches(7.7)
    set_cell_shading(rule.cell(0, 0), BLUE)
    rule.cell(0, 0).text = ""
    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(22)
    cover.add_run(
        "Final cohort: 132 approved CamCan participants\n"
        "Targets: left hippocampus, left M1, right DLPFC, right thalamus\n"
        "Ten independent remeshing repeats per participant and target\n"
        "Primary threshold: 0.18 V/m; sensitivity threshold: 0.15 V/m"
    )
    status = document.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.paragraph_format.space_before = Pt(20)
    status_run = status.add_run(
        f"Prepared {date.today().strftime('%d %B %Y')} | Descriptive "
        "analysis for supervisor review"
    )
    status_run.font.size = Pt(9)
    status_run.font.color.rgb = RGBColor.from_string(GRAY)
    document.add_page_break()

    # Overview
    add_heading(document, "1. What was analysed")
    document.add_paragraph(
        "This document accompanies four descriptive figures generated from "
        "the final corrected CamCan temporal-interference (TI) simulation "
        "campaign. It is written so the figures can be reviewed without "
        "consulting the analysis code or earlier presentations."
    )
    add_callout(
        document,
        "Central reporting rule",
        "The ten TI field images were not averaged voxel-by-voxel. Every "
        "metric was calculated independently on each repeat, and the ten "
        "resulting scalar metric values were then averaged for each "
        "participant and target. All plotted participant points and cohort "
        "summaries use these ten-repeat metric means.",
    )
    add_heading(document, "Analysis scope", level=2)
    add_bullet(
        document,
        "132 participants passed the final mesh review and were included.",
    )
    add_bullet(
        document,
        "Four target-specific fixed montages were applied to every "
        "participant: left hippocampus, left primary motor cortex (M1), "
        "right dorsolateral prefrontal cortex (DLPFC), and right thalamus.",
    )
    add_bullet(
        document,
        "Each participant-target combination was remeshed and simulated ten "
        "times, giving 5,280 validated simulations.",
    )
    add_bullet(
        document,
        "The corrected MNI152 head was simulated once per target with the "
        "same target-specific montage and current settings and is shown as "
        "a descriptive reference, not as a second statistical sample.",
    )
    add_bullet(
        document,
        "The figures do not include individualized optimization and do not "
        "perform inferential comparisons between targets.",
    )
    add_heading(document, "Fixed stimulation parameters", level=2)
    target_rows = [
        (
            "Left hippocampus",
            targets_by[("Left_Hippocampus",)],
        ),
        (
            "Left M1",
            targets_by[("ctx_lh_G_precentral",)],
        ),
        (
            "Right DLPFC",
            targets_by[("ctx_rh_G_front_middle",)],
        ),
        (
            "Right thalamus",
            targets_by[("Right_Thalamus",)],
        ),
    ]
    add_table(
        document,
        ["Target", "Electrode pair 1", "Amplitude 1", "Electrode pair 2", "Amplitude 2"],
        [
            [
                label,
                row["pair1"],
                f"{float(row['current1']):.3f} mA",
                row["pair2"],
                f"{float(row['current2']):.3f} mA",
            ]
            for label, row in target_rows
        ],
        widths_cm=[4.0, 4.2, 3.0, 4.2, 3.0],
    )
    document.add_paragraph(
        "Amplitudes are the pair amplitudes recorded in utils/targets.csv. "
        "Within each pair, equal and opposite currents were applied to the "
        "two named electrodes."
    )

    # Computation
    document.add_page_break()
    add_heading(document, "2. How the figure metrics were computed")
    add_heading(document, "Per-repeat processing", level=2)
    for text in [
        "Load the brain-only TI electric-field magnitude image for one "
        "participant, one target, and one remeshing repeat.",
        "Load the participant-space anatomical atlas and construct the "
        "target region-of-interest (ROI) mask.",
        "Calculate all field-magnitude, threshold, and rank-based metrics "
        "from that repeat alone.",
        "Repeat the calculation for repeats 01 through 10.",
        "Take the arithmetic mean of each scalar metric across the ten "
        "repeats. This gives one participant-level value per target.",
        "Across the 132 participant-level values, calculate descriptive "
        "statistics, empirical distributions, and within-target "
        "correlations. Calculate the same scalar metrics once for MNI152.",
    ]:
        add_number(document, text)

    add_heading(document, "Exact percentage definitions", level=2)
    document.add_paragraph(
        "Let T be all anatomical target voxels, B all finite brain voxels, "
        "O the finite brain voxels outside T, and S(t) the finite brain "
        "voxels whose TI field is at least threshold t."
    )
    add_table(
        document,
        ["Metric", "Calculation", "Question answered"],
        [
            [
                "Target coverage",
                "100 x |T intersect S(t)| / |T|",
                "What percentage of the anatomical target reaches the threshold?",
            ],
            [
                "Off-target coverage",
                "100 x |O intersect S(t)| / |O|",
                "What percentage of tissue outside the target reaches the threshold?",
            ],
            [
                "Whole-brain coverage",
                "100 x |S(t)| / |B|",
                "What percentage of all finite brain tissue reaches the threshold?",
            ],
            [
                "Localization in target",
                "100 x |T intersect S(t)| / |S(t)|",
                "Of all suprathreshold brain voxels, what percentage lies inside the target?",
            ],
        ],
        widths_cm=[4.0, 6.0, 10.0],
    )
    add_callout(
        document,
        "Why coverage and localization can appear contradictory",
        "Target coverage uses target size as its denominator; localization "
        "uses all suprathreshold brain voxels as its denominator. A "
        "simulation can therefore stimulate only a modest fraction of the "
        "target but still have high localization if there is very little "
        "suprathreshold field elsewhere.",
        PALE_ORANGE,
    )
    add_heading(document, "Field-magnitude and rank-based metrics", level=2)
    add_bullet(
        document,
        "ROI median field is the median finite TI-field magnitude inside the target.",
    )
    add_bullet(
        document,
        "ROI P99.9 is the 99.9th percentile inside the target. It is used as "
        "a robust maximum so a single extreme voxel cannot determine the result.",
    )
    add_bullet(
        document,
        "Top-5% target coverage is the percentage of target voxels that "
        "belong to the highest 5% of finite whole-brain TI-field values.",
    )
    add_bullet(
        document,
        "Top-5% localization is the percentage of all whole-brain top-5% "
        "voxels that lie inside the target.",
    )
    document.add_paragraph(
        "Non-finite target voxels were retained in the anatomical target "
        "denominator and counted as unstimulated. If no finite whole-brain "
        "voxel reached a threshold, threshold localization was defined as 0%."
    )

    # Figure 1
    add_figure_page(
        document,
        "Figure 1",
        "Target effectiveness and off-target spread at 0.18 V/m",
        figures / FIGURES[0][1],
        "Each blue point is one participant's ten-repeat mean; the orange "
        "diamond is the corrected MNI152 reference.",
    )
    document.add_page_break()
    add_heading(document, "3. Figure 1 review guide")
    add_callout(
        document,
        "Question answered",
        "When a participant has a larger fraction of the target above "
        "0.18 V/m, does that participant also tend to have a larger fraction "
        "of non-target brain tissue above 0.18 V/m?",
    )
    add_heading(document, "How to read the panels", level=2)
    add_bullet(
        document,
        "Horizontal position is target coverage. Farther right means more "
        "of the intended anatomical target reaches 0.18 V/m.",
    )
    add_bullet(
        document,
        "Vertical position is off-target coverage. Farther upward means more "
        "brain tissue outside the target reaches 0.18 V/m.",
    )
    add_bullet(
        document,
        "The gray dashed line is a descriptive least-squares fit. Spearman "
        "rho quantifies the monotonic association without assuming a linear relationship.",
    )
    figure1_rows = []
    for roi in ROI_ORDER:
        summary = threshold_by[(roi, "0.18")]
        association = association_by[
            (roi, "0.18", "Effectiveness–spread coupling")
        ]
        figure1_rows.append(
            [
                roi,
                summary["Target coverage — CamCan mean (SD), %"],
                summary["Off-target coverage — CamCan mean (SD), %"],
                f"{float(association['Spearman ρ']):.3f}",
                f"{f1(summary['Target coverage — MNI152, %'])} / "
                f"{f1(summary['Off-target coverage — MNI152, %'])}",
            ]
        )
    add_table(
        document,
        [
            "Target",
            "CamCan target coverage, mean (SD)",
            "CamCan off-target coverage, mean (SD)",
            "Spearman rho",
            "MNI152 target / off-target",
        ],
        figure1_rows,
        widths_cm=[3.4, 4.4, 4.7, 2.8, 4.2],
    )
    add_heading(document, "Interpretation", level=2)
    add_bullet(
        document,
        "Target and off-target coverage are strongly positively associated "
        "within every target (rho = 0.681 to 0.955). Participants with more "
        "target engagement generally also show more spatial spread.",
    )
    add_bullet(
        document,
        "The association is strongest for right DLPFC and left hippocampus "
        "and weakest, although still positive, for right thalamus.",
    )
    add_bullet(
        document,
        "This is best described as effectiveness-spread coupling. It does "
        "not demonstrate that increasing target coverage causes off-target "
        "coverage, and it does not establish an optimal point.",
    )
    add_bullet(
        document,
        "The panels should not be used to claim that one target-specific "
        "montage is superior to another: ROI anatomy, volume, depth, and "
        "montage differ.",
    )
    add_heading(document, "Suggested supervisor checks", level=2)
    add_bullet(
        document,
        "Is the effectiveness-versus-spread framing appropriate for the main text?",
    )
    add_bullet(
        document,
        "Should the dashed fits remain, or should only rho be reported?",
    )

    # Figure 2
    add_figure_page(
        document,
        "Figure 2",
        "Inter-individual distribution of target coverage at 0.18 V/m",
        figures / FIGURES[1][1],
        "The blue line is the empirical cumulative distribution across 132 "
        "participant-level ten-repeat means.",
    )
    document.add_page_break()
    add_heading(document, "4. Figure 2 review guide")
    add_callout(
        document,
        "Question answered",
        "How heterogeneous is target coverage across participants, and where "
        "does the single MNI152 reference fall within that distribution?",
    )
    add_heading(document, "How an empirical cumulative distribution works", level=2)
    document.add_paragraph(
        "For any x-axis value, the y-axis gives the proportion of "
        "participants whose target coverage is less than or equal to that "
        "value. A curve shifted to the right indicates generally greater "
        "target coverage. A steep section indicates that many participants "
        "have similar coverage values; a gradual curve indicates a broad distribution."
    )
    add_bullet(
        document,
        "The orange dashed line is the corrected MNI152 target coverage.",
    )
    add_bullet(
        document,
        "The light-gray lines at 25%, 50%, and 75% are coverage-value guides, "
        "not cohort quartiles.",
    )
    coverage_rows = []
    for roi in ROI_ORDER:
        row = attainment_by[(roi,)]
        context = context_by[(roi, "Target coverage ≥0.18 V/m")]
        coverage_rows.append(
            [
                roi,
                f1(row["Subjects with ≤1% target coverage, %"]),
                f1(row["Subjects with ≥25% target coverage, %"]),
                f1(row["Subjects with ≥50% target coverage, %"]),
                f1(row["Subjects with ≥75% target coverage, %"]),
                f"{f1(row['MNI152 target coverage ≥0.18 V/m, %'])} "
                f"({f1(context['MNI152 percentile within CamCan'])}th percentile)",
            ]
        )
    add_table(
        document,
        [
            "Target",
            "Participants <=1%",
            "Participants >=25%",
            "Participants >=50%",
            "Participants >=75%",
            "MNI152 coverage (cohort percentile)",
        ],
        coverage_rows,
        widths_cm=[3.4, 3.0, 3.0, 3.0, 3.0, 5.4],
    )
    add_heading(document, "Interpretation", level=2)
    add_bullet(
        document,
        "Left hippocampus and right thalamus show consistently high target "
        "coverage: 91.7% and 97.7% of participants, respectively, reach at "
        "least 50% target coverage.",
    )
    add_bullet(
        document,
        "Left M1 is more variable and usually lower: only 9.8% of "
        "participants reach at least 50% coverage.",
    )
    add_bullet(
        document,
        "Right DLPFC is the most heterogeneous and frequently has negligible "
        "absolute coverage: 39.4% of participants have <=1% coverage, while "
        "3.0% reach >=50%. Mean and standard deviation alone would hide this shape.",
    )
    add_bullet(
        document,
        "MNI152 is below most participants for hippocampus and thalamus, "
        "near the cohort middle for M1, and around the 60th percentile for DLPFC.",
    )
    add_heading(document, "Suggested supervisor checks", level=2)
    add_bullet(
        document,
        "Does the ECDF communicate participant heterogeneity more clearly "
        "than a boxplot or violin plot would?",
    )
    add_bullet(
        document,
        "Should this remain a second main-text figure or move to the supplement?",
    )

    # Threshold sensitivity
    add_figure_page(
        document,
        "Supplementary Figure S1",
        "Sensitivity to the TI-field threshold",
        figures / FIGURES[2][1],
        "Solid circles and error bars are CamCan mean +/- one between-participant "
        "SD; dotted diamonds are the single MNI152 values.",
    )
    document.add_page_break()
    add_heading(document, "5. Supplementary Figure S1 review guide")
    add_callout(
        document,
        "Question answered",
        "Would the interpretation materially change if the suprathreshold "
        "criterion were lowered from the primary 0.18 V/m value to 0.15 V/m?",
    )
    add_heading(document, "How the plot was computed", level=2)
    add_bullet(
        document,
        "For each participant, target and off-target coverage were calculated "
        "independently at both thresholds on every repeat and then averaged across repeats.",
    )
    add_bullet(
        document,
        "The solid blue and red points are cohort means; error bars are one "
        "between-participant standard deviation, not uncertainty in MNI152.",
    )
    add_bullet(
        document,
        "The dotted diamond series is the single corrected MNI152 reference "
        "and therefore has no error bar.",
    )
    sensitivity_rows = []
    for roi in ROI_ORDER:
        target = shift_by[(roi, "Target coverage")]
        off_target = shift_by[(roi, "Off-target coverage")]
        localization = shift_by[(roi, "Suprathreshold localization in target")]
        sensitivity_rows.append(
            [
                roi,
                f"{f1(target['CamCan mean at 0.18 V/m, %'])} -> "
                f"{f1(target['CamCan mean at 0.15 V/m, %'])}",
                f"+{f1(target['Paired mean change (0.15 − 0.18), percentage points'])}",
                f"{f1(off_target['CamCan mean at 0.18 V/m, %'])} -> "
                f"{f1(off_target['CamCan mean at 0.15 V/m, %'])}",
                f"+{f1(off_target['Paired mean change (0.15 − 0.18), percentage points'])}",
                f"{float(localization['Paired mean change (0.15 − 0.18), percentage points']):+.1f}",
            ]
        )
    add_table(
        document,
        [
            "Target",
            "Target coverage 0.18 -> 0.15",
            "Target change (pp)",
            "Off-target coverage 0.18 -> 0.15",
            "Off-target change (pp)",
            "Localization change (pp)",
        ],
        sensitivity_rows,
        widths_cm=[3.3, 4.0, 3.0, 4.3, 3.0, 3.4],
    )
    add_heading(document, "Interpretation", level=2)
    add_bullet(
        document,
        "Lowering the criterion necessarily increases the amount of tissue "
        "classified as suprathreshold; it does not change the underlying simulated field.",
    )
    add_bullet(
        document,
        "All four targets gain target coverage at 0.15 V/m. The mean gains "
        "range from 10.5 to 17.8 percentage points.",
    )
    add_bullet(
        document,
        "The accompanying off-target increase is small for M1 and DLPFC "
        "(approximately 0.9 percentage points) but much larger for "
        "hippocampus (14.8 points) and thalamus (24.5 points).",
    )
    add_bullet(
        document,
        "Localization decreases at the lower threshold because newly "
        "included off-target voxels enlarge the denominator faster than "
        "new target voxels. This is especially clear for M1.",
    )
    add_bullet(
        document,
        "The figure supports retaining 0.18 V/m as the primary analysis and "
        "reporting 0.15 V/m as a sensitivity analysis rather than pooling them.",
    )
    add_heading(document, "Suggested supervisor checks", level=2)
    add_bullet(
        document,
        "Is 0.15 V/m the appropriate sensitivity threshold to report?",
    )
    add_bullet(
        document,
        "Should localization sensitivity also be plotted, or is it sufficient in the table?",
    )

    # MNI context
    add_figure_page(
        document,
        "Supplementary Figure S2",
        "MNI152 percentile context",
        figures / FIGURES[3][1],
        "Every dot is the percentile rank of one corrected MNI152 value "
        "within the corresponding 132-participant CamCan distribution.",
    )
    document.add_page_break()
    add_heading(document, "6. Supplementary Figure S2 review guide")
    add_callout(
        document,
        "Question answered",
        "For each outcome and target, is the standard MNI152 reference near "
        "the center of the participant distribution, or is it unusually low or high?",
    )
    add_heading(document, "What a percentile point means", level=2)
    document.add_paragraph(
        "A percentile converts an outcome with its original units into a "
        "relative position. A value at the 10th percentile is higher than "
        "about 10% and lower than about 90% of participant values. The gray "
        "band marks the 25th to 75th percentiles, and the dashed line marks "
        "the cohort median. The figure does not show the raw electric-field values."
    )
    add_callout(
        document,
        "Critical reading rule",
        "Farther right means a higher numerical value, not universally "
        "better performance. High target coverage or localization is usually "
        "desirable; high off-target coverage means greater unwanted spread.",
        PALE_ORANGE,
    )
    context_outcomes = [
        ("Median target-ROI TI field", "ROI median"),
        ("Robust maximum target-ROI TI field (P99.9)", "ROI P99.9"),
        ("Target coverage ≥0.18 V/m", "Target coverage"),
        ("Off-target coverage ≥0.18 V/m", "Off-target coverage"),
        ("Whole-brain coverage ≥0.18 V/m", "Whole-brain coverage"),
        (
            "Suprathreshold localization in target ≥0.18 V/m",
            "Localization",
        ),
        (
            "Target coverage by whole-brain top 5% field",
            "Top-5% target coverage",
        ),
        (
            "Localization of whole-brain top 5% field in target",
            "Top-5% localization",
        ),
    ]
    percentile_rows = []
    for outcome, short in context_outcomes:
        percentile_rows.append(
            [short]
            + [
                f1(context_by[(roi, outcome)]["MNI152 percentile within CamCan"])
                for roi in ROI_ORDER
            ]
        )
    add_table(
        document,
        ["Outcome percentile", *ROI_ORDER],
        percentile_rows,
        widths_cm=[5.2, 3.5, 3.5, 3.5, 3.5],
    )
    add_heading(document, "Panel-by-panel interpretation", level=2)
    add_bullet(
        document,
        "Left hippocampus: MNI152 has low field magnitude, target coverage, "
        "and spatial spread (approximately 2nd to 15th percentiles), but high "
        "threshold localization (85th percentile). The field is weaker and "
        "less extensive than in most participants, yet the limited "
        "suprathreshold field is comparatively concentrated in the target.",
    )
    add_bullet(
        document,
        "Left M1: target and spread measures are broadly mid-distribution, "
        "but localization is low (17th percentile) and top-5% localization "
        "is exceptionally low (2nd percentile). The strongest fields are "
        "less concentrated in M1 than in almost all participants.",
    )
    add_bullet(
        document,
        "Right DLPFC: most MNI152 outcomes fall near the cohort center "
        "(approximately 50th to 72nd percentiles). It is the most "
        "representative of the four MNI152 target simulations, although "
        "off-target spread is somewhat above the cohort median.",
    )
    add_bullet(
        document,
        "Right thalamus: MNI152 has lower absolute field, target coverage, "
        "and spread than most participants, but high threshold localization "
        "(89th percentile). Its highest-ranked field values are more typical "
        "than its absolute-threshold coverage.",
    )
    add_heading(document, "Interpretation", level=2)
    add_bullet(
        document,
        "MNI152 is not a uniformly representative or average head. Its "
        "relative position depends strongly on target and outcome.",
    )
    add_bullet(
        document,
        "MNI152 should therefore be described as a standardized reference "
        "configuration, not as the expected participant result.",
    )
    add_bullet(
        document,
        "These percentile ranks are descriptive. One MNI152 head cannot be "
        "given a standard error or compared inferentially with the cohort as "
        "if it were a second group.",
    )
    add_heading(document, "Suggested supervisor checks", level=2)
    add_bullet(
        document,
        "Is the percentile-context figure useful enough to retain, or is the "
        "mixed metric direction too cognitively demanding?",
    )
    add_bullet(
        document,
        "If retained, should every point be directly labelled with its percentile?",
    )

    # Summary
    document.add_page_break()
    add_heading(document, "7. Overall interpretation and review decisions")
    add_heading(document, "What the four figures collectively show", level=2)
    add_bullet(
        document,
        "Participant anatomy materially changes the balance between target "
        "engagement and off-target spread even when the montage and current "
        "settings are fixed.",
    )
    add_bullet(
        document,
        "Higher target coverage generally accompanies higher off-target "
        "coverage within each target-specific montage.",
    )
    add_bullet(
        document,
        "The participant distributions differ substantially in shape. Right "
        "DLPFC is especially heterogeneous at the primary threshold.",
    )
    add_bullet(
        document,
        "Lowering the threshold increases target coverage everywhere, but "
        "the accompanying increase in spatial spread differs markedly by target.",
    )
    add_bullet(
        document,
        "The standardized MNI152 reference is not consistently typical of "
        "the CamCan cohort and should not substitute for a population analysis.",
    )
    add_heading(document, "Interpretation guardrails", level=2)
    add_bullet(
        document,
        "Do not interpret visible differences between target panels as "
        "formal evidence that one target or montage is superior.",
    )
    add_bullet(
        document,
        "Do not describe the effectiveness-spread correlation as causal.",
    )
    add_bullet(
        document,
        "Interpret localization alongside target and off-target coverage; "
        "its denominator and the anatomical ROI size can make it appear high "
        "even when absolute target coverage is modest.",
    )
    add_bullet(
        document,
        "MNI152 is a single deterministic reference and is not an "
        "inferential comparison group.",
    )
    add_bullet(
        document,
        "The CamCan simulations used SimNIBS 4.0.1, whereas the corrected "
        "MNI152 baselines used SimNIBS 4.5.0. This version difference should "
        "be disclosed wherever MNI152 is reported.",
    )
    add_heading(document, "Recommended placement for review", level=2)
    add_table(
        document,
        ["Figure", "Current recommendation", "Reason"],
        [
            [
                "Figure 1",
                "Main text",
                "Directly communicates the central effectiveness-spread relationship.",
            ],
            [
                "Figure 2",
                "Main text or supplement",
                "Shows heterogeneity that mean and SD do not reveal.",
            ],
            [
                "Supplementary Figure S1",
                "Supplement",
                "Documents threshold sensitivity without displacing the primary result.",
            ],
            [
                "Supplementary Figure S2",
                "Supplement",
                "Provides MNI context but requires careful explanation.",
            ],
        ],
        widths_cm=[4.0, 4.8, 11.8],
    )

    # Data lineage appendix
    document.add_page_break()
    add_heading(document, "Appendix. Data lineage and reproducibility")
    source_rows = [
        [
            "Primary descriptive values",
            "tables/table_1_primary_numeric_long.csv",
        ],
        [
            "Threshold summaries",
            "tables/table_s1_threshold_sensitivity.csv",
        ],
        [
            "MNI152 percentile context",
            "tables/table_s4_mni152_context.csv",
        ],
        [
            "Within-target correlations",
            "tables/table_s5_exploratory_within_roi_associations.csv",
        ],
        [
            "Paired threshold changes",
            "tables/table_s6_threshold_paired_changes.csv",
        ],
        [
            "Coverage attainment",
            "tables/table_s7_target_coverage_attainment.csv",
        ],
        ["Stimulation parameters", str(targets_csv)],
    ]
    add_table(
        document,
        ["Purpose", "Source"],
        source_rows,
        widths_cm=[6.2, 14.8],
    )
    document.add_paragraph(
        "The publication package is generated from the validated schema-2 "
        "analysis. It contains 5,280 repeat-level metric records, 528 "
        "participant-target aggregate records, and four corrected MNI152 "
        "baseline records."
    )
    add_heading(document, "Figure files and checksums", level=2)
    checksum_rows = []
    for designation, filename, _ in FIGURES:
        path = figures / filename
        checksum_rows.append(
            [designation, filename, f"{path.stat().st_size:,}", sha256(path)]
        )
    add_table(
        document,
        ["Figure", "File", "Bytes", "SHA-256"],
        checksum_rows,
        widths_cm=[4.0, 7.0, 2.2, 8.0],
        font_size=7.4,
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--package",
        type=Path,
        required=True,
        help="Publication package containing figures and tables.",
    )
    parser.add_argument(
        "--targets-csv",
        type=Path,
        required=True,
        help="Validated targets.csv used by the final simulations.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        required=True,
        help="Destination .docx path.",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    build_document(
        args.package.expanduser().resolve(),
        args.targets_csv.expanduser().resolve(),
        args.output.expanduser().resolve(),
    )
    print(args.output.expanduser().resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
