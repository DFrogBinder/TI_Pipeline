#!/usr/bin/env python3
"""Build a two-head comparison deck for the defacing experiment."""
from __future__ import annotations

import argparse
import csv
import json
import math
import re
import statistics as st
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt
from scipy import stats


SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

WHITE = RGBColor(255, 255, 255)
INK = RGBColor(30, 37, 45)
MUTED = RGBColor(92, 103, 116)
LIGHT = RGBColor(246, 248, 250)
NAVY = RGBColor(11, 31, 53)
BLUE = RGBColor(30, 96, 145)
CYAN = RGBColor(26, 150, 156)
ORANGE = RGBColor(214, 126, 54)
GOLD = RGBColor(232, 174, 70)
RED = RGBColor(185, 69, 58)
GREEN = RGBColor(71, 145, 93)

ARM_NAMES = (
    "Left_Hippocampus_Intact",
    "Left_Hippocampus_Defaced",
    "Left_M1_Intact",
    "Left_M1_Defaced",
)

HEADS = {
    "sub-CCMe": {
        "label": "sub-CCMe",
        "description": "left skull trauma",
        "source_dir": "post_export_40repeats",
        "color": ORANGE,
    },
    "sub-IXI025": {
        "label": "sub-IXI025",
        "description": "healthy head",
        "source_dir": "post_export_sub-IXI025",
        "color": BLUE,
    },
}

TARGET_LABELS = {
    "left-hippocampus": "Left hippocampus",
    "left-m1": "Left M1",
}

CONDITION_LABELS = {
    "intact": "Face intact",
    "defaced": "Face removed",
}


@dataclass(frozen=True)
class MetricSpec:
    key: str
    label: str
    units: str
    decimals: int
    stored_as_fraction: bool = False


METRICS = [
    MetricSpec("roi_mean", "ROI mean field", "V/m", 3),
    MetricSpec("roi_percentile_value", "ROI P95 field", "V/m", 3),
    MetricSpec("roi_peak", "ROI peak field", "V/m", 3),
    MetricSpec("overlap_fraction", "ROI in top 5% field", "%", 1, stored_as_fraction=True),
    MetricSpec("focality_in_roi_volume_mm3_gt_threshold", "ROI volume >=0.2 V/m", "mm3", 0),
    MetricSpec("focality_volume_mm3_gt_threshold", "Whole-brain volume >=0.2 V/m", "mm3", 0),
    MetricSpec("roi_volume_mm3", "ROI volume", "mm3", 0),
    MetricSpec("whole_brain_volume_mm3", "Whole-brain volume", "mm3", 0),
    MetricSpec("csf_distance_mm", "Target-to-CSF distance", "mm", 2),
]

HEADLINE_KEYS = [
    "roi_mean",
    "roi_percentile_value",
    "overlap_fraction",
    "focality_in_roi_volume_mm3_gt_threshold",
    "focality_volume_mm3_gt_threshold",
]


def mpl(color: RGBColor) -> str:
    return "#{:02x}{:02x}{:02x}".format(color[0], color[1], color[2])


def mean(values: Iterable[float]) -> float:
    vals = list(values)
    return st.mean(vals) if vals else math.nan


def sd(values: Iterable[float]) -> float:
    vals = list(values)
    return st.stdev(vals) if len(vals) > 1 else 0.0


def sem(values: Iterable[float]) -> float:
    vals = list(values)
    return sd(vals) / math.sqrt(len(vals)) if len(vals) > 1 else math.nan


def metric_spec(key: str) -> MetricSpec:
    for spec in METRICS:
        if spec.key == key:
            return spec
    raise KeyError(key)


def value_for_display(value: float, key: str) -> float:
    spec = metric_spec(key)
    return value * 100 if spec.stored_as_fraction else value


def fmt_value(value: float, key: str) -> str:
    if value is None or not isinstance(value, (int, float)) or math.isnan(value):
        return "n/a"
    spec = metric_spec(key)
    shown = value_for_display(value, key)
    if spec.decimals == 0:
        return f"{shown:,.0f}"
    return f"{shown:.{spec.decimals}f}"


def fmt_delta(value: float, key: str) -> str:
    if value is None or not isinstance(value, (int, float)) or math.isnan(value):
        return "n/a"
    shown = value_for_display(value, key)
    spec = metric_spec(key)
    if spec.decimals == 0:
        return f"{shown:+,.0f}"
    return f"{shown:+.{spec.decimals}f}"


def pct_text(value: float, decimals: int = 1) -> str:
    if value is None or not isinstance(value, (int, float)) or math.isnan(value):
        return "n/a"
    return f"{value:+.{decimals}f}%"


def p_text(value: float) -> str:
    if value is None or not isinstance(value, (int, float)) or math.isnan(value):
        return "n/a"
    return f"{value:.3g}"


def infer_arm(path: Path, source_root: Path) -> str:
    for part in path.relative_to(source_root).parts:
        if part in ARM_NAMES:
            return part
    raise ValueError(f"Cannot infer arm from {path}")


def infer_target(arm: str) -> str:
    if "Hippocampus" in arm:
        return "left-hippocampus"
    if "M1" in arm:
        return "left-m1"
    raise ValueError(f"Cannot infer target from {arm}")


def infer_condition(arm: str) -> str:
    if "Defaced" in arm:
        return "defaced"
    if "Intact" in arm:
        return "intact"
    raise ValueError(f"Cannot infer condition from {arm}")


def infer_repeat(path: Path) -> int:
    match = re.search(r"Data_(\d+)", path.as_posix())
    if not match:
        raise ValueError(f"Cannot infer repeat from {path}")
    return int(match.group(1))


def numeric(value) -> float:
    if value is None:
        return math.nan
    try:
        return float(value)
    except (TypeError, ValueError):
        return math.nan


def load_records(data_root: Path) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for head_id, meta in HEADS.items():
        source_root = data_root / str(meta["source_dir"])
        if not source_root.is_dir():
            raise FileNotFoundError(f"Missing expected export root: {source_root}")
        for arm in ARM_NAMES:
            if not (source_root / arm).is_dir():
                raise FileNotFoundError(f"Missing expected arm folder: {source_root / arm}")
        for path in sorted(source_root.rglob("subject_metrics.json")):
            data = json.loads(path.read_text(encoding="utf-8"))
            arm = infer_arm(path, source_root)
            target = infer_target(arm)
            condition = infer_condition(arm)
            target_roi = data["target_roi"]
            roi = data["rois"][target_roi]
            ext = data["extended_metrics"]
            record: dict[str, object] = {
                "head": head_id,
                "head_label": meta["label"],
                "head_description": meta["description"],
                "source_root": str(source_root),
                "path": str(path),
                "arm": arm,
                "target": target,
                "condition": condition,
                "repeat": infer_repeat(path),
                "subject": data.get("subject", head_id),
                "target_roi": target_roi,
                "subject_status": data.get("subject_metrics_meta", {}).get("status"),
                "extended_status": data.get("extended_metrics_meta", {}).get("status"),
                "qc_status": data.get("qc_meta", {}).get("status"),
            }
            record["percentile_value"] = numeric(data.get("percentile_value"))
            record["whole_brain_volume_mm3"] = numeric(data.get("whole_brain_volume_mm3"))
            for key in (
                "roi_voxels",
                "roi_volume_mm3",
                "overlap_top_voxels",
                "overlap_volume_mm3",
                "overlap_fraction",
                "roi_percentile_value",
                "focality_in_roi_volume_mm3_gt_threshold",
            ):
                record[key] = numeric(roi.get(key))
            for key in (
                "roi_mean",
                "roi_peak",
                "focality_volume_mm3_gt_threshold",
                "csf_distance_mm",
            ):
                record[key] = numeric(ext.get(key))
            records.append(record)
    return records


def records_by(
    records: list[dict[str, object]],
    *,
    head: str | None = None,
    target: str | None = None,
    condition: str | None = None,
) -> list[dict[str, object]]:
    out = records
    if head is not None:
        out = [r for r in out if r["head"] == head]
    if target is not None:
        out = [r for r in out if r["target"] == target]
    if condition is not None:
        out = [r for r in out if r["condition"] == condition]
    return out


def paired_delta_rows(records: list[dict[str, object]]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for head in HEADS:
        for target in TARGET_LABELS:
            intact = {int(r["repeat"]): r for r in records_by(records, head=head, target=target, condition="intact")}
            defaced = {int(r["repeat"]): r for r in records_by(records, head=head, target=target, condition="defaced")}
            repeats = sorted(set(intact) & set(defaced))
            for spec in METRICS:
                xs: list[float] = []
                ys: list[float] = []
                diffs: list[float] = []
                repeat_diffs: list[dict[str, object]] = []
                for repeat in repeats:
                    x = intact[repeat].get(spec.key)
                    y = defaced[repeat].get(spec.key)
                    if isinstance(x, (int, float)) and isinstance(y, (int, float)) and not math.isnan(x) and not math.isnan(y):
                        xs.append(float(x))
                        ys.append(float(y))
                        diff = float(y) - float(x)
                        diffs.append(diff)
                        repeat_diffs.append(
                            {
                                "head": head,
                                "target": target,
                                "repeat": repeat,
                                "metric": spec.key,
                                "intact": float(x),
                                "defaced": float(y),
                                "delta": diff,
                                "percent_change": 100 * diff / float(x) if float(x) else math.nan,
                            }
                        )
                if not diffs:
                    continue
                intact_mean = mean(xs)
                defaced_mean = mean(ys)
                delta_mean = mean(diffs)
                delta_sd = sd(diffs)
                delta_sem = sem(diffs)
                if len(diffs) > 1 and delta_sem and not math.isnan(delta_sem):
                    tcrit = stats.t.ppf(0.975, len(diffs) - 1)
                    ci_low = delta_mean - tcrit * delta_sem
                    ci_high = delta_mean + tcrit * delta_sem
                    p_value = float(stats.ttest_rel(ys, xs).pvalue)
                else:
                    ci_low = ci_high = p_value = math.nan
                rows.append(
                    {
                        "head": head,
                        "head_label": HEADS[head]["label"],
                        "head_description": HEADS[head]["description"],
                        "target": target,
                        "metric": spec.key,
                        "metric_label": spec.label,
                        "units": spec.units,
                        "n_pairs": len(diffs),
                        "intact_mean": intact_mean,
                        "defaced_mean": defaced_mean,
                        "delta_mean": delta_mean,
                        "delta_sd": delta_sd,
                        "delta_ci_low": ci_low,
                        "delta_ci_high": ci_high,
                        "percent_delta": 100 * delta_mean / intact_mean if intact_mean else math.nan,
                        "paired_t_p": p_value,
                        "repeat_deltas_json": json.dumps(repeat_diffs, sort_keys=True),
                    }
                )
    return rows


def arm_summary_rows(records: list[dict[str, object]]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for head in HEADS:
        for target in TARGET_LABELS:
            for condition in ("intact", "defaced"):
                subset = records_by(records, head=head, target=target, condition=condition)
                for spec in METRICS:
                    vals = [float(r[spec.key]) for r in subset if isinstance(r.get(spec.key), (int, float)) and not math.isnan(float(r[spec.key]))]
                    rows.append(
                        {
                            "head": head,
                            "target": target,
                            "condition": condition,
                            "metric": spec.key,
                            "metric_label": spec.label,
                            "units": spec.units,
                            "n": len(vals),
                            "mean": mean(vals),
                            "sd": sd(vals),
                            "min": min(vals) if vals else math.nan,
                            "max": max(vals) if vals else math.nan,
                        }
                    )
    return rows


def cross_head_rows(delta_rows: list[dict[str, object]]) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for target in TARGET_LABELS:
        for spec in METRICS:
            ccme = delta_row(delta_rows, "sub-CCMe", target, spec.key)
            ixi = delta_row(delta_rows, "sub-IXI025", target, spec.key)
            if ccme is None or ixi is None:
                continue
            rows.append(
                {
                    "target": target,
                    "metric": spec.key,
                    "metric_label": spec.label,
                    "units": spec.units,
                    "sub_CCMe_percent_delta": ccme["percent_delta"],
                    "sub_IXI025_percent_delta": ixi["percent_delta"],
                    "percent_delta_difference_CCMe_minus_IXI025": float(ccme["percent_delta"]) - float(ixi["percent_delta"]),
                    "sub_CCMe_delta_mean": ccme["delta_mean"],
                    "sub_IXI025_delta_mean": ixi["delta_mean"],
                }
            )
    return rows


def delta_row(delta_rows: list[dict[str, object]], head: str, target: str, metric: str) -> dict[str, object] | None:
    for row in delta_rows:
        if row["head"] == head and row["target"] == target and row["metric"] == metric:
            return row
    return None


def write_csv(path: Path, rows: list[dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        path.write_text("", encoding="utf-8")
        return
    fields = list(rows[0].keys())
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def write_records_csv(path: Path, records: list[dict[str, object]]) -> None:
    serializable = []
    for record in records:
        row = dict(record)
        serializable.append(row)
    write_csv(path, serializable)


def completeness(records: list[dict[str, object]]) -> dict[str, object]:
    out: dict[str, object] = {
        "records": len(records),
        "expected_records": len(HEADS) * len(TARGET_LABELS) * len(CONDITION_LABELS) * 40,
        "heads": len({r["head"] for r in records}),
        "targets": len({r["target"] for r in records}),
        "conditions": len({r["condition"] for r in records}),
    }
    for head in HEADS:
        repeats = {
            (r["target"], r["condition"], int(r["repeat"]))
            for r in records
            if r["head"] == head
        }
        out[f"{head}_records"] = len([r for r in records if r["head"] == head])
        out[f"{head}_unique_repeat_slots"] = len(repeats)
        out[f"{head}_complete_status_count"] = sum(
            1
            for r in records
            if r["head"] == head
            and r.get("subject_status") == "complete"
            and r.get("extended_status") == "complete"
        )
    return out


def plot_effect_overview(delta_rows: list[dict[str, object]], path: Path) -> None:
    endpoints = [
        ("roi_mean", "ROI mean"),
        ("roi_percentile_value", "ROI P95"),
        ("overlap_fraction", "Top-5% overlap"),
        ("focality_in_roi_volume_mm3_gt_threshold", "ROI >=0.2 V/m"),
        ("focality_volume_mm3_gt_threshold", "Brain >=0.2 V/m"),
    ]
    rows: list[tuple[str, str, str]] = []
    for target in ("left-hippocampus", "left-m1"):
        for key, label in endpoints:
            rows.append((target, key, f"{TARGET_LABELS[target]} | {label}"))
    y_positions = list(range(len(rows)))[::-1]
    fig, ax = plt.subplots(figsize=(11.3, 5.05))
    for head, offset in (("sub-CCMe", 0.15), ("sub-IXI025", -0.15)):
        xs = []
        ys = []
        labels = []
        for y, (target, key, _) in zip(y_positions, rows):
            row = delta_row(delta_rows, head, target, key)
            value = float(row["percent_delta"]) if row else math.nan
            xs.append(value)
            ys.append(y + offset)
            labels.append(pct_text(value))
        color = mpl(HEADS[head]["color"])
        ax.scatter(xs, ys, s=82, color=color, label=f"{HEADS[head]['label']} ({HEADS[head]['description']})", zorder=3)
        for x, y, label in zip(xs, ys, labels):
            if math.isnan(x):
                continue
            ha = "left" if x >= 0 else "right"
            dx = 0.18 if x >= 0 else -0.18
            ax.text(x + dx, y, label, va="center", ha=ha, fontsize=8, color=color)
    for y in y_positions:
        ax.axhline(y, color="#e5e9ee", linewidth=0.8, zorder=0)
    ax.axvline(0, color="#5b6571", linewidth=1.1)
    ax.set_yticks(y_positions, [label for _, _, label in rows])
    ax.set_xlabel("Defaced minus intact mean change (%)")
    ax.set_xlim(-12.5, 13.5)
    ax.grid(axis="x", alpha=0.16)
    ax.legend(frameon=False, loc="lower right")
    ax.set_title("Defacing effects differ by head and target", fontsize=14, loc="left")
    fig.tight_layout(pad=1.0)
    fig.savefig(path, dpi=220)
    plt.close(fig)


def plot_baseline_intact(records: list[dict[str, object]], path: Path) -> None:
    panels = [
        ("roi_percentile_value", "ROI P95 field", "V/m"),
        ("overlap_fraction", "ROI in top 5% field", "%"),
        ("focality_in_roi_volume_mm3_gt_threshold", "ROI volume >=0.2 V/m", "mm3"),
    ]
    fig, axes = plt.subplots(1, 3, figsize=(12.1, 4.6))
    for ax, (key, title, units) in zip(axes, panels):
        x = []
        heights = []
        colors = []
        labels = []
        pos = 0
        for target in ("left-hippocampus", "left-m1"):
            for head in ("sub-CCMe", "sub-IXI025"):
                vals = [
                    float(r[key])
                    for r in records_by(records, head=head, target=target, condition="intact")
                    if isinstance(r.get(key), (int, float)) and not math.isnan(float(r[key]))
                ]
                x.append(pos)
                heights.append(value_for_display(mean(vals), key))
                colors.append(mpl(HEADS[head]["color"]))
                labels.append(f"{TARGET_LABELS[target]}\n{HEADS[head]['label']}")
                pos += 1
            pos += 0.6
        bars = ax.bar(x, heights, color=colors, alpha=0.86)
        for bar, height in zip(bars, heights):
            ax.text(bar.get_x() + bar.get_width() / 2, height, f"{height:.2f}" if abs(height) < 10 else f"{height:,.0f}", ha="center", va="bottom", fontsize=8)
        ax.set_xticks(x, labels, fontsize=8)
        ax.set_title(title, fontsize=11)
        ax.set_ylabel(units)
        ax.grid(axis="y", alpha=0.16)
    fig.tight_layout()
    fig.savefig(path, dpi=220)
    plt.close(fig)


def repeat_delta_percent(row: dict[str, object]) -> list[float]:
    payload = json.loads(str(row["repeat_deltas_json"]))
    return [float(item["percent_change"]) for item in payload if isinstance(item.get("percent_change"), (int, float))]


def plot_repeat_distributions(delta_rows: list[dict[str, object]], path: Path) -> None:
    panels = [
        ("roi_percentile_value", "ROI P95"),
        ("overlap_fraction", "Top-5% overlap"),
        ("focality_in_roi_volume_mm3_gt_threshold", "ROI volume >=0.2 V/m"),
        ("focality_volume_mm3_gt_threshold", "Whole-brain volume >=0.2 V/m"),
    ]
    fig, axes = plt.subplots(2, 2, figsize=(12.0, 5.35))
    axes = axes.flatten()
    for ax, (key, title) in zip(axes, panels):
        data = []
        labels = []
        colors = []
        for target in ("left-hippocampus", "left-m1"):
            for head in ("sub-CCMe", "sub-IXI025"):
                row = delta_row(delta_rows, head, target, key)
                data.append(repeat_delta_percent(row) if row else [])
                labels.append(f"{TARGET_LABELS[target].replace('Left ', '')}\n{HEADS[head]['label']}")
                colors.append(mpl(HEADS[head]["color"]))
        bp = ax.boxplot(data, patch_artist=True, widths=0.55, showfliers=False)
        for patch, color in zip(bp["boxes"], colors):
            patch.set(facecolor=color, alpha=0.26, edgecolor=color, linewidth=1.5)
        for median in bp["medians"]:
            median.set(color="#1e252d", linewidth=1.5)
        for idx, (vals, color) in enumerate(zip(data, colors), start=1):
            ax.scatter([idx] * len(vals), vals, color=color, alpha=0.55, s=14, zorder=3)
        ax.axhline(0, color="#59636f", linewidth=0.9)
        ax.set_title(title, fontsize=11)
        ax.set_ylabel("Repeat percent change")
        ax.set_xticks(range(1, len(labels) + 1), labels, fontsize=8)
        ax.grid(axis="y", alpha=0.16)
    fig.tight_layout(pad=1.0)
    fig.savefig(path, dpi=220)
    plt.close(fig)


def plot_controls(delta_rows: list[dict[str, object]], path: Path) -> None:
    panels = [
        ("roi_volume_mm3", "ROI volume change"),
        ("whole_brain_volume_mm3", "Whole-brain volume change"),
    ]
    fig, axes = plt.subplots(1, 2, figsize=(10.8, 4.4))
    for ax, (key, title) in zip(axes, panels):
        labels = []
        values = []
        colors = []
        for target in ("left-hippocampus", "left-m1"):
            for head in ("sub-CCMe", "sub-IXI025"):
                row = delta_row(delta_rows, head, target, key)
                labels.append(f"{TARGET_LABELS[target].replace('Left ', '')}\n{HEADS[head]['label']}")
                values.append(float(row["percent_delta"]) if row else math.nan)
                colors.append(mpl(HEADS[head]["color"]))
        bars = ax.bar(range(len(values)), values, color=colors, alpha=0.86)
        for bar, value in zip(bars, values):
            ax.text(bar.get_x() + bar.get_width() / 2, value + (0.01 if value >= 0 else -0.01), pct_text(value, 2), ha="center", va="bottom" if value >= 0 else "top", fontsize=8)
        ax.axhline(0, color="#59636f", linewidth=0.9)
        ax.set_xticks(range(len(labels)), labels, fontsize=8)
        ax.set_title(title, fontsize=11)
        ax.set_ylabel("Defaced minus intact (%)")
        ax.grid(axis="y", alpha=0.16)
    fig.tight_layout()
    fig.savefig(path, dpi=220)
    plt.close(fig)


def add_bg(slide, title: str | None = None) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    top = slide.shapes.add_shape(1, 0, 0, SLIDE_W, Inches(0.16))
    top.fill.solid()
    top.fill.fore_color.rgb = GOLD
    top.line.fill.background()
    if title:
        box = slide.shapes.add_textbox(Inches(0.55), Inches(0.33), Inches(12.1), Inches(0.46))
        p = box.text_frame.paragraphs[0]
        p.text = title
        p.font.bold = True
        p.font.size = Pt(23)
        p.font.color.rgb = NAVY


def add_footer(slide, text: str = "Defacing sensitivity experiment | two-subject comparison | technical repeats") -> None:
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.12), Inches(12.2), Inches(0.25))
    p = box.text_frame.paragraphs[0]
    p.text = text
    p.font.size = Pt(8)
    p.font.color.rgb = MUTED


def textbox(slide, x, y, w, h, text, *, size=15, bold=False, color=INK, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if align:
        p.alignment = align
    return box


def bullet_list(slide, x, y, w, h, items: list[str], *, size=15, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"- {item}"
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(7)
    return box


def metric_card(slide, x, y, w, h, label: str, value: str, subtitle: str, color: RGBColor):
    shape = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = LIGHT
    shape.line.color.rgb = color
    shape.line.width = Pt(1.15)
    textbox(slide, x + 0.12, y + 0.10, w - 0.24, 0.25, label, size=9, bold=True, color=MUTED)
    textbox(slide, x + 0.12, y + 0.40, w - 0.24, 0.42, value, size=20, bold=True, color=color)
    textbox(slide, x + 0.12, y + 0.87, w - 0.24, h - 0.91, subtitle, size=8.5, color=INK)


def add_table(slide, x, y, w, h, headers: list[str], rows: list[list[str]], font_size=8):
    table = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(x), Inches(y), Inches(w), Inches(h)).table
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
        for p in cell.text_frame.paragraphs:
            p.font.color.rgb = WHITE
            p.font.bold = True
            p.font.size = Pt(font_size)
    for r_idx, row in enumerate(rows, start=1):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor(248, 249, 251) if r_idx % 2 == 0 else WHITE
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(font_size)
                p.font.color.rgb = INK
    return table


def direction(row: dict[str, object], decimals: int = 1) -> str:
    value = float(row["percent_delta"])
    if value > 0:
        return f"increased {pct_text(value, decimals)}"
    if value < 0:
        return f"decreased {pct_text(value, decimals)}"
    return "was unchanged"


def build_deck(
    records: list[dict[str, object]],
    delta_rows: list[dict[str, object]],
    figures: dict[str, Path],
    out_dir: Path,
) -> Path:
    counts = completeness(records)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    ccme_hip_p95 = delta_row(delta_rows, "sub-CCMe", "left-hippocampus", "roi_percentile_value")
    ixi_hip_p95 = delta_row(delta_rows, "sub-IXI025", "left-hippocampus", "roi_percentile_value")
    ccme_m1_roi_vol = delta_row(delta_rows, "sub-CCMe", "left-m1", "focality_in_roi_volume_mm3_gt_threshold")
    ixi_m1_roi_vol = delta_row(delta_rows, "sub-IXI025", "left-m1", "focality_in_roi_volume_mm3_gt_threshold")

    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    stripe = slide.shapes.add_shape(1, 0, 0, Inches(0.22), SLIDE_H)
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = GOLD
    stripe.line.fill.background()
    textbox(slide, 0.75, 0.72, 11.5, 0.8, "Defacing sensitivity is anatomy-dependent", size=34, bold=True, color=WHITE)
    textbox(slide, 0.78, 1.55, 11.1, 0.45, "Comparison of a left-skull-trauma head and a healthy head across 320 SimNIBS TI simulations", size=18, color=RGBColor(212, 225, 238))
    textbox(
        slide,
        0.78,
        2.55,
        11.35,
        1.15,
        "Face removal did not impose one universal bias: the healthy head showed consistent reductions after defacing, while sub-CCMe showed smaller hippocampal amplitude changes and mixed M1 shifts.",
        size=22,
        color=WHITE,
    )
    textbox(slide, 0.80, 6.78, 11.9, 0.28, "Inputs: T1/T2 intact versus pydeface face-removed images; targets: left hippocampus and left M1; 40 technical repeats per arm.", size=9, color=RGBColor(190, 204, 218))

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Main result: the healthy head was more uniformly affected")
    add_footer(slide)
    metric_card(slide, 0.72, 1.12, 2.55, 1.23, "sub-CCMe hippocampus P95", pct_text(float(ccme_hip_p95["percent_delta"])), "Amplitude was nearly stable after defacing.", ORANGE)
    metric_card(slide, 3.55, 1.12, 2.55, 1.23, "sub-IXI025 hippocampus P95", pct_text(float(ixi_hip_p95["percent_delta"])), "Healthy head showed a clearer decrease.", BLUE)
    metric_card(slide, 6.38, 1.12, 2.55, 1.23, "sub-CCMe M1 ROI >=0.2", pct_text(float(ccme_m1_roi_vol["percent_delta"])), "M1 high-field ROI extent increased.", ORANGE)
    metric_card(slide, 9.21, 1.12, 2.55, 1.23, "sub-IXI025 M1 ROI >=0.2", pct_text(float(ixi_m1_roi_vol["percent_delta"])), "M1 high-field ROI extent decreased slightly.", BLUE)
    bullet_list(
        slide,
        0.85,
        3.05,
        11.6,
        2.3,
        [
            "Both heads used the same experiment structure, targets, condition contrast, and repeat count.",
            "The healthy head shifted in the expected direction for most endpoints: defaced runs generally had lower amplitude, overlap, and above-threshold volume.",
            "sub-CCMe was not simply a larger version of the same effect: hippocampal amplitude was stable and M1 high-field ROI volume increased.",
            "This supports treating defacing sensitivity as subject- and target-specific until more heads are tested.",
        ],
        size=15,
    )
    textbox(slide, 0.85, 6.15, 11.4, 0.48, "Statistical labels in this deck are descriptive repeat-level tests over technical repeats, not inference over a subject population.", size=11, color=MUTED)

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Experiment design: same pipeline, two anatomies")
    add_footer(slide)
    add_table(
        slide,
        0.68,
        1.02,
        6.5,
        3.25,
        ["Factor", "Values"],
        [
            ["Heads", "sub-CCMe: left skull trauma; sub-IXI025: healthy IXI head"],
            ["Input conditions", "Face intact T1/T2 versus pydeface face-removed T1/T2"],
            ["Targets", "Left hippocampus and left M1 / left precentral gyrus"],
            ["Repeats", "40 paired technical repeats per head, target, and condition"],
            ["Total runs analyzed", f"{counts['records']} / {counts['expected_records']} expected post-processed outputs"],
            ["Atlas", "Subject-space FreeSurfer/Destrieux atlas for each head"],
        ],
        font_size=8.5,
    )
    metric_card(slide, 7.65, 1.08, 2.0, 1.2, "Heads", "2", "Trauma and healthy anatomy.", ORANGE)
    metric_card(slide, 9.95, 1.08, 2.0, 1.2, "Targets", "2", "Deep and cortical.", CYAN)
    metric_card(slide, 7.65, 2.70, 2.0, 1.2, "Arms", "8", "Head x target x condition.", BLUE)
    metric_card(slide, 9.95, 2.70, 2.0, 1.2, "Repeats", "40", "Paired within target.", GOLD)
    textbox(slide, 0.75, 5.10, 11.7, 0.72, "The critical comparison is within-head and within-target: repeat N in the defaced condition is paired with repeat N in the intact condition before summarizing effects.", size=15, color=INK)
    textbox(slide, 0.75, 6.05, 11.7, 0.42, "The between-head comparison then asks whether the defacing effect size has the same direction and magnitude in the two anatomies.", size=13, color=MUTED)

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Analysis workflow: repeat-paired endpoint deltas")
    add_footer(slide)
    metric_card(slide, 0.70, 1.10, 2.55, 1.22, "1. Read outputs", "320 JSONs", "subject_metrics.json from every post-processed run.", BLUE)
    metric_card(slide, 3.48, 1.10, 2.55, 1.22, "2. Extract endpoints", f"{len(METRICS)} metrics", "ROI amplitude, overlap, high-field extent, and controls.", CYAN)
    metric_card(slide, 6.26, 1.10, 2.55, 1.22, "3. Pair repeats", "40 pairs", "Intact and defaced runs matched by repeat ID.", GOLD)
    metric_card(slide, 9.04, 1.10, 2.55, 1.22, "4. Compare heads", "Effect sizes", "Contrast defaced-minus-intact deltas across anatomies.", ORANGE)
    add_table(
        slide,
        0.68,
        3.02,
        12.0,
        2.76,
        ["Output", "Computation", "Interpretation"],
        [
            ["Arm summary", "Mean, SD, min, max for each head-target-condition arm", "Shows baseline field scale and repeat variability."],
            ["Paired delta", "Defaced value minus intact value for each repeat", "Isolates face-removal sensitivity within a head."],
            ["Percent delta", "100 x mean(delta) / intact mean", "Normalizes endpoints with different units."],
            ["Paired t-test", "Descriptive repeat-level p-value", "Summarizes repeat separation only; not population inference."],
            ["Controls", "ROI volume and whole-brain volume deltas", "Checks whether gross atlas/brain volume explains endpoint shifts."],
        ],
        font_size=8,
    )
    textbox(slide, 0.75, 6.17, 11.6, 0.44, "No MNI-space atlas comparison is used in the conclusions; endpoints are subject-space outputs from each head's own atlas.", size=11, color=MUTED)

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Endpoint definitions: amplitude, overlap, and extent")
    add_footer(slide)
    add_table(
        slide,
        0.58,
        1.00,
        7.2,
        4.75,
        ["Endpoint", "Definition"],
        [
            ["ROI mean field", "Average TI magnitude within the target ROI."],
            ["ROI P95 field", "95th percentile TI magnitude inside the target ROI; robust high-end amplitude."],
            ["ROI peak field", "Maximum target value; inspected but not used as the main claim because it is outlier-sensitive."],
            ["ROI in top 5% field", "Fraction of target voxels inside the whole-brain top 5% field."],
            ["ROI volume >=0.2 V/m", "Amount of target ROI above the field threshold."],
            ["Whole-brain volume >=0.2 V/m", "Total above-threshold volume, interpreted as global high-field extent."],
        ],
        font_size=8.4,
    )
    shape = slide.shapes.add_shape(1, Inches(8.25), Inches(1.10), Inches(4.25), Inches(2.7))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(250, 245, 233)
    shape.line.color.rgb = GOLD
    textbox(slide, 8.48, 1.34, 3.8, 1.95, "Interpretation guardrail\n\nThe unit of replication is a computational repeat. Treat p-values as repeat-level separation, not evidence that the same effect generalizes to all heads.", size=13, color=INK)
    textbox(slide, 8.25, 4.32, 4.2, 0.82, "Primary claims use direction, magnitude, and consistency of paired deltas, with controls for ROI and whole-brain volume.", size=13, color=MUTED)

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Both exports are complete for the planned comparison")
    add_footer(slide)
    rows = []
    for head in HEADS:
        head_records = [r for r in records if r["head"] == head]
        complete = sum(
            1
            for r in head_records
            if r.get("subject_status") == "complete"
            and r.get("extended_status") == "complete"
        )
        rows.append(
            [
                f"{HEADS[head]['label']} ({HEADS[head]['description']})",
                str(len(head_records)),
                "4",
                "40",
                f"{complete}/{len(head_records)}",
            ]
        )
    add_table(slide, 0.85, 1.35, 11.3, 1.45, ["Head", "Runs", "Arms", "Repeats per arm", "Complete metric files"], rows, font_size=10)
    bullet_list(
        slide,
        0.90,
        3.45,
        11.3,
        2.25,
        [
            "All four arms were present for each head: hippocampus intact, hippocampus defaced, M1 intact, and M1 defaced.",
            "Each arm contributed 40 repeat-level post-processing outputs.",
            "The comparison therefore uses 160 simulations for sub-CCMe and 160 simulations for sub-IXI025.",
        ],
        size=16,
    )

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Baseline field scale differs before defacing")
    add_footer(slide)
    slide.shapes.add_picture(str(figures["baseline"]), Inches(0.45), Inches(1.08), width=Inches(12.35))
    textbox(slide, 0.78, 6.62, 11.8, 0.3, "Bars show intact-condition means across 40 repeats. Different baseline amplitude and overlap means the same absolute delta can have different practical importance by head.", size=9, color=MUTED)

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Defacing effects do not have the same direction in both heads")
    add_footer(slide)
    slide.shapes.add_picture(str(figures["effect_overview"]), Inches(0.55), Inches(0.96), width=Inches(12.0))
    textbox(slide, 0.65, 6.76, 12.0, 0.22, "Values are defaced-minus-intact mean changes normalized to the intact mean. Negative means defaced runs were lower.", size=8.2, color=MUTED)

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Repeat-level distributions confirm the head-specific pattern")
    add_footer(slide)
    slide.shapes.add_picture(str(figures["repeat_distributions"]), Inches(0.62), Inches(1.00), width=Inches(12.0))
    textbox(slide, 0.72, 6.76, 11.9, 0.22, "Each point is one paired repeat percent change. Boxplots summarize repeat-level technical variability, not biological variability.", size=8.2, color=MUTED)

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Hippocampus: healthy head shows stronger amplitude reduction")
    add_footer(slide)
    rows = []
    for head in ("sub-CCMe", "sub-IXI025"):
        for key in ("roi_mean", "roi_percentile_value", "overlap_fraction", "focality_in_roi_volume_mm3_gt_threshold"):
            row = delta_row(delta_rows, head, "left-hippocampus", key)
            rows.append(
                [
                    HEADS[head]["label"],
                    metric_spec(key).label,
                    fmt_value(float(row["intact_mean"]), key),
                    fmt_value(float(row["defaced_mean"]), key),
                    pct_text(float(row["percent_delta"])),
                    p_text(float(row["paired_t_p"])),
                ]
            )
    add_table(slide, 0.45, 1.00, 12.45, 3.70, ["Head", "Endpoint", "Intact mean", "Defaced mean", "Change", "paired p"], rows, font_size=7.8)
    bullet_list(
        slide,
        0.70,
        5.05,
        12.0,
        1.1,
        [
            f"sub-CCMe hippocampal ROI P95 {direction(ccme_hip_p95)}; sub-IXI025 hippocampal ROI P95 {direction(ixi_hip_p95)}.",
            "Both heads showed lower top-5% overlap after defacing, but the healthy head also showed clearer amplitude and above-threshold extent reductions.",
        ],
        size=13.5,
    )

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "M1: sub-CCMe and healthy head move in opposite directions")
    add_footer(slide)
    rows = []
    for head in ("sub-CCMe", "sub-IXI025"):
        for key in ("roi_mean", "roi_percentile_value", "overlap_fraction", "focality_in_roi_volume_mm3_gt_threshold"):
            row = delta_row(delta_rows, head, "left-m1", key)
            rows.append(
                [
                    HEADS[head]["label"],
                    metric_spec(key).label,
                    fmt_value(float(row["intact_mean"]), key),
                    fmt_value(float(row["defaced_mean"]), key),
                    pct_text(float(row["percent_delta"])),
                    p_text(float(row["paired_t_p"])),
                ]
            )
    add_table(slide, 0.45, 1.00, 12.45, 3.70, ["Head", "Endpoint", "Intact mean", "Defaced mean", "Change", "paired p"], rows, font_size=7.8)
    bullet_list(
        slide,
        0.70,
        5.05,
        12.0,
        1.1,
        [
            f"sub-CCMe M1 ROI volume above threshold {direction(ccme_m1_roi_vol)}, while sub-IXI025 {direction(ixi_m1_roi_vol)}.",
            "The M1 result argues against a single global defacing correction factor; the effect changes sign across heads.",
        ],
        size=13.5,
    )

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Controls: atlas volume is fixed; brain-volume shifts are small")
    add_footer(slide)
    slide.shapes.add_picture(str(figures["controls"]), Inches(1.15), Inches(1.05), width=Inches(10.7))
    bullet_list(
        slide,
        0.92,
        5.62,
        11.5,
        0.9,
        [
            "ROI volume is unchanged because each head uses a fixed subject-space atlas for intact and defaced runs.",
            "Whole-brain volume changes are below 0.2%, so the main endpoint shifts are not explained by gross atlas or brain-mask size changes.",
        ],
        size=12.5,
    )

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Interpretation: defacing sensitivity depends on anatomy and target")
    add_footer(slide)
    bullet_list(
        slide,
        0.75,
        1.12,
        11.7,
        4.8,
        [
            "Healthy head: defacing was associated with a coherent reduction in hippocampal and M1 amplitude/extent endpoints.",
            "sub-CCMe: hippocampal amplitude was comparatively stable, while M1 showed increased high-field ROI extent after defacing.",
            "The two-head comparison is consistent with anatomy-dependent interaction between defacing, face reconstruction, meshing, and target geometry.",
            "Because there are only two heads, the correct claim is sensitivity heterogeneity, not a population estimate of trauma versus healthy anatomy.",
            "Next step: add more healthy and non-healthy heads, then model head-level variability separately from repeat-level technical variability.",
        ],
        size=16,
    )

    slide = prs.slides.add_slide(blank)
    add_bg(slide, "Deliverables and reproducibility")
    add_footer(slide)
    textbox(slide, 0.75, 1.02, 11.8, 0.36, f"Output folder: {out_dir}", size=10, color=MUTED)
    add_table(
        slide,
        0.75,
        1.72,
        11.8,
        3.05,
        ["Artifact", "Filename", "Purpose"],
        [
            ["Editable deck", "defacing_two_head_comparison.pptx", "Presentation source."],
            ["Rendered deck", "defacing_two_head_comparison.pdf", "Stable version for review/sharing."],
            ["Run metrics", "head_comparison_run_metrics.csv", "One row per run with extracted endpoints."],
            ["Paired deltas", "head_comparison_paired_deltas.csv", "Defaced-minus-intact effects by head, target, and metric."],
            ["Cross-head summary", "head_comparison_cross_head_summary.csv", "Effect-size contrasts between heads."],
            ["Narrative summary", "head_comparison_summary.md", "Methods, headline values, and interpretation boundaries."],
        ],
        font_size=8.2,
    )
    textbox(slide, 0.78, 5.35, 11.6, 0.72, "Regeneration path: run this builder against the downloaded `post_export_40repeats` and `post_export_sub-IXI025` folders. The source simulation exports are read-only.", size=14, color=INK)

    path = out_dir / "defacing_two_head_comparison.pptx"
    prs.save(path)
    return path


def write_summary_md(path: Path, records: list[dict[str, object]], delta_rows: list[dict[str, object]]) -> None:
    counts = completeness(records)
    lines = [
        "# Two-head defacing experiment comparison",
        "",
        "## Scope",
        "",
        "- Heads: `sub-CCMe` (left skull trauma) and `sub-IXI025` (healthy IXI head).",
        "- Conditions: face-intact T1/T2 versus pydeface face-removed T1/T2.",
        "- Targets: left hippocampus and left M1.",
        "- Repeats: 40 paired technical repeats per head, target, and condition.",
        f"- Runs analyzed: {counts['records']} / {counts['expected_records']}.",
        "",
        "## Analysis",
        "",
        "For every run, `subject_metrics.json` was read from the post-processing export. Endpoints were extracted for ROI amplitude, ROI overlap with the whole-brain top 5% field, target and whole-brain above-threshold volume, ROI volume, whole-brain volume, and target-to-CSF distance. For each head and target, repeat `N` in the defaced arm was paired with repeat `N` in the intact arm. The reported effect is defaced minus intact, normalized by the intact mean for percent change.",
        "",
        "Paired p-values are descriptive repeat-level summaries and should not be interpreted as population inference.",
        "",
        "## Headline paired deltas",
        "",
        "| Head | Target | Endpoint | Intact mean | Defaced mean | Change | paired p |",
        "|---|---|---|---:|---:|---:|---:|",
    ]
    for head in ("sub-CCMe", "sub-IXI025"):
        for target in ("left-hippocampus", "left-m1"):
            for key in HEADLINE_KEYS:
                row = delta_row(delta_rows, head, target, key)
                lines.append(
                    "| "
                    + " | ".join(
                        [
                            HEADS[head]["label"],
                            TARGET_LABELS[target],
                            metric_spec(key).label,
                            fmt_value(float(row["intact_mean"]), key),
                            fmt_value(float(row["defaced_mean"]), key),
                            pct_text(float(row["percent_delta"])),
                            p_text(float(row["paired_t_p"])),
                        ]
                    )
                    + " |"
                )
    lines.extend(
        [
            "",
            "## Interpretation",
            "",
            "- The healthy head showed a more uniform negative defacing effect across amplitude, overlap, and high-field extent endpoints.",
            "- sub-CCMe did not show the same pattern: hippocampal amplitude was comparatively stable and M1 above-threshold ROI extent increased after defacing.",
            "- ROI volumes were unchanged across intact and defaced conditions, confirming that fixed subject-space atlases were used for each head.",
            "- Whole-brain volume changes were small, so gross brain-mask volume does not explain the main endpoint shifts.",
            "- With only two heads, the defensible conclusion is anatomy- and target-dependent sensitivity rather than a population estimate.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def build(data_root: Path, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    figures = {
        "effect_overview": out_dir / "figures" / "effect_overview.png",
        "baseline": out_dir / "figures" / "baseline_intact.png",
        "repeat_distributions": out_dir / "figures" / "repeat_distributions.png",
        "controls": out_dir / "figures" / "controls.png",
    }
    figures["effect_overview"].parent.mkdir(parents=True, exist_ok=True)

    records = load_records(data_root)
    delta_rows = paired_delta_rows(records)
    arms = arm_summary_rows(records)
    cross = cross_head_rows(delta_rows)

    write_records_csv(out_dir / "head_comparison_run_metrics.csv", records)
    write_csv(out_dir / "head_comparison_arm_summary.csv", arms)
    write_csv(out_dir / "head_comparison_paired_deltas.csv", delta_rows)
    write_csv(out_dir / "head_comparison_cross_head_summary.csv", cross)
    write_summary_md(out_dir / "head_comparison_summary.md", records, delta_rows)

    plot_effect_overview(delta_rows, figures["effect_overview"])
    plot_baseline_intact(records, figures["baseline"])
    plot_repeat_distributions(delta_rows, figures["repeat_distributions"])
    plot_controls(delta_rows, figures["controls"])

    return build_deck(records, delta_rows, figures, out_dir)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--data-root",
        type=Path,
        default=Path("/home/boyan/sandbox/Jake_Data/defacing_experiment"),
        help="Directory containing post_export_40repeats and post_export_sub-IXI025.",
    )
    parser.add_argument(
        "--out-dir",
        type=Path,
        default=Path("/tmp/defacing_head_comparison"),
        help="Directory for generated figures, CSVs, Markdown, and PPTX.",
    )
    args = parser.parse_args()
    deck = build(args.data_root.expanduser().resolve(), args.out_dir.expanduser().resolve())
    print(f"Wrote {deck}")


if __name__ == "__main__":
    main()
