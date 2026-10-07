#!/usr/bin/env python3
"""Build the initial journal-club deck for the corrected repeatability study."""
from __future__ import annotations

import csv
import json
import math
import statistics
from dataclasses import dataclass
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


HERE = Path(__file__).resolve().parent
OUT_DIR = HERE / "outputs"

ANALYSIS_ROOT = Path("/home/boyan/sandbox/repeatability_paper_analysis_v1")
FIXED_ROOT = ANALYSIS_ROOT / "corrected_fixed_analysis"
NESTED_ROOT = ANALYSIS_ROOT / "nested_analysis"
OVERLEAF_ROOT = Path(
    "/home/boyan/sandbox/Jake_Data/Overleaf_upload_packages/"
    "TIS_repeatability_Overleaf_upload"
)

LH_FIELD = FIXED_ROOT / "figures/left_hippocampus/01_primary_median_roi_repeat_distributions.png"
M1_FIELD = FIXED_ROOT / "figures/right_m1/01_primary_median_roi_repeat_distributions.png"
LH_RANK = FIXED_ROOT / "figures/left_hippocampus/02_single_repeat_subject_ranking_uncertainty.png"
M1_RANK = FIXED_ROOT / "figures/right_m1/02_single_repeat_subject_ranking_uncertainty.png"
NESTED_FIGURE = NESTED_ROOT / "nested_mesh_by_solver_repeatability.png"
LH_TISSUE = OVERLEAF_ROOT / "figures/left_hippocampus/06_example_subject_tissue_composition.png"
M1_TISSUE = OVERLEAF_ROOT / "figures/right_m1/06_example_subject_tissue_composition.png"
MNI_VERSION_ROOT = HERE / "assets/mni_version_analysis"
MNI_VERSION_FIGURE = MNI_VERSION_ROOT / "figure_01_mean_roi_field_by_version.png"
MNI_VERSION_CONTINUOUS = MNI_VERSION_ROOT / "continuous_roi_metric_comparison.csv"
MNI_VERSION_VOXELWISE = MNI_VERSION_ROOT / "voxelwise_agreement.csv"

PPTX_PATH = OUT_DIR / "repeatability_journal_club_initial_draft.pptx"
NOTES_PATH = OUT_DIR / "repeatability_journal_club_speaker_notes.md"
CONTENT_PATH = OUT_DIR / "repeatability_journal_club_content_review.md"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Noto Sans"

NAVY = RGBColor(12, 30, 48)
DEEP_BLUE = RGBColor(25, 72, 116)
BLUE = RGBColor(37, 111, 166)
PALE_BLUE = RGBColor(224, 239, 249)
TEAL = RGBColor(27, 139, 126)
PALE_TEAL = RGBColor(224, 244, 241)
ORANGE = RGBColor(194, 91, 35)
PALE_ORANGE = RGBColor(252, 236, 226)
GOLD = RGBColor(221, 168, 63)
PALE_GOLD = RGBColor(251, 245, 224)
RED = RGBColor(177, 58, 58)
INK = RGBColor(29, 39, 48)
MUTED = RGBColor(86, 101, 114)
RULE = RGBColor(207, 216, 224)
LIGHT = RGBColor(246, 248, 250)
WHITE = RGBColor(255, 255, 255)


@dataclass(frozen=True)
class SlideRecord:
    number: int
    title: str
    purpose: str
    status: str
    notes: str


def rgb_hex(color: RGBColor) -> str:
    return "#{:02x}{:02x}{:02x}".format(*tuple(color))


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: RGBColor = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.04,
    font: str = FONT,
    line_spacing: float | None = None,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.word_wrap = True
    frame.vertical_anchor = valign
    para = frame.paragraphs[0]
    para.alignment = align
    para.text = text
    para.font.name = font
    para.font.size = Pt(size)
    para.font.bold = bold
    para.font.color.rgb = color
    if line_spacing is not None:
        para.line_spacing = line_spacing
    return box


def add_runs(
    slide,
    runs: list[tuple[str, float, RGBColor, bool]],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.04,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for idx, (text, size, color, bold) in enumerate(runs):
        run = p.add_run() if idx else p.runs[0]
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
    return box


def add_bullet_list(
    slide,
    items: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 17,
    color: RGBColor = INK,
    bullet_color: RGBColor = TEAL,
    spacing: float = 8,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(0.03)
    tf.margin_right = Inches(0.03)
    tf.margin_top = 0
    tf.margin_bottom = 0
    for idx, item in enumerate(items):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = ""
        p.space_after = Pt(spacing)
        marker = p.add_run()
        marker.text = "●  "
        marker.font.name = FONT
        marker.font.size = Pt(max(8, size - 5))
        marker.font.color.rgb = bullet_color
        run = p.add_run()
        run.text = item
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return box


def add_box(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    fill: RGBColor,
    *,
    line: RGBColor | None = None,
    radius: bool = False,
    transparency: int = 0,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.fill.transparency = transparency
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.8)
    return shape


def add_line(
    slide,
    x1: float,
    y1: float,
    x2: float,
    y2: float,
    *,
    color: RGBColor = RULE,
    width: float = 1.2,
    dash: bool = False,
):
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line.line.color.rgb = color
    line.line.width = Pt(width)
    if dash:
        line.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return line


def add_circle(slide, x: float, y: float, d: float, fill: RGBColor, *, line: RGBColor | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.0)
    return shape


def set_background(slide, color: RGBColor = LIGHT):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_header(slide, title: str, section: str, number: int, *, dark: bool = False):
    title_color = WHITE if dark else NAVY
    muted = RGBColor(186, 202, 215) if dark else MUTED
    add_text(slide, section.upper(), 0.55, 0.22, 2.35, 0.2, size=8.5, color=TEAL, bold=True)
    add_text(slide, title, 0.55, 0.48, 12.1, 0.72, size=25, color=title_color, bold=True)
    add_line(slide, 0.55, 1.25, 12.78, 1.25, color=RGBColor(67, 93, 114) if dark else RULE, width=0.8)
    add_text(slide, "TI repeatability • journal club", 0.55, 7.18, 3.7, 0.15, size=7.3, color=muted)
    add_text(slide, f"{number:02d}", 12.18, 7.15, 0.55, 0.18, size=8, color=muted, bold=True, align=PP_ALIGN.RIGHT)


def add_source(slide, text: str, *, dark: bool = False):
    add_text(
        slide,
        text,
        5.0,
        7.17,
        6.85,
        0.17,
        size=6.4,
        color=RGBColor(174, 194, 207) if dark else MUTED,
        align=PP_ALIGN.RIGHT,
    )


def add_notes(slide, notes: str):
    slide.notes_slide.notes_text_frame.text = notes.strip()


def add_picture_contain(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        iw, ih = image.size
    image_ratio = iw / ih
    box_ratio = w / h
    if image_ratio > box_ratio:
        fw = w
        fh = w / image_ratio
    else:
        fh = h
        fw = h * image_ratio
    return slide.shapes.add_picture(
        str(path),
        Inches(x + (w - fw) / 2),
        Inches(y + (h - fh) / 2),
        width=Inches(fw),
        height=Inches(fh),
    )


def add_metric_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    value: str,
    label: str,
    *,
    accent: RGBColor = TEAL,
    fill: RGBColor = WHITE,
    value_size: float = 27,
):
    add_box(slide, x, y, w, h, fill, line=RULE, radius=True)
    add_box(slide, x, y, 0.07, h, accent)
    add_text(slide, value, x + 0.22, y + 0.16, w - 0.38, 0.47, size=value_size, color=accent, bold=True)
    add_text(slide, label, x + 0.22, y + 0.72, w - 0.38, h - 0.82, size=10.5, color=MUTED)


def add_chip(slide, text: str, x: float, y: float, w: float, *, fill: RGBColor, color: RGBColor):
    add_box(slide, x, y, w, 0.34, fill, radius=True)
    add_text(slide, text, x + 0.06, y + 0.09, w - 0.12, 0.16, size=8.4, color=color, bold=True, align=PP_ALIGN.CENTER)


def add_arrow_between(slide, x1: float, y: float, x2: float, *, color: RGBColor = MUTED):
    line = add_line(slide, x1, y, x2, y, color=color, width=1.6)
    line.line.end_arrowhead = True
    return line


def require_files():
    required = [
        FIXED_ROOT / "subject_condition_summary.csv",
        FIXED_ROOT / "analysis_manifest.json",
        FIXED_ROOT / "rank_uncertainty_summary.json",
        FIXED_ROOT / "pairwise_rank_reversal_probabilities.csv",
        NESTED_ROOT / "nested_figure_values.json",
        LH_FIELD,
        M1_FIELD,
        LH_RANK,
        M1_RANK,
        NESTED_FIGURE,
        LH_TISSUE,
        M1_TISSUE,
        MNI_VERSION_FIGURE,
        MNI_VERSION_CONTINUOUS,
        MNI_VERSION_VOXELWISE,
    ]
    missing = [str(path) for path in required if not path.is_file()]
    if missing:
        raise FileNotFoundError("Missing presentation inputs:\n" + "\n".join(missing))


def load_data():
    with (FIXED_ROOT / "subject_condition_summary.csv").open(newline="", encoding="utf-8") as handle:
        summary = list(csv.DictReader(handle))
    manifest = json.loads((FIXED_ROOT / "analysis_manifest.json").read_text(encoding="utf-8"))
    rank = json.loads((FIXED_ROOT / "rank_uncertainty_summary.json").read_text(encoding="utf-8"))
    nested = json.loads((NESTED_ROOT / "nested_figure_values.json").read_text(encoding="utf-8"))
    with (FIXED_ROOT / "pairwise_rank_reversal_probabilities.csv").open(newline="", encoding="utf-8") as handle:
        pairwise = list(csv.DictReader(handle))
    return summary, manifest, rank, nested, pairwise


def load_mni_version_data() -> dict[str, float | int]:
    with MNI_VERSION_CONTINUOUS.open(newline="", encoding="utf-8") as handle:
        continuous = list(csv.DictReader(handle))
    mean_rows = [row for row in continuous if row["metric"] == "roi_mean_v_per_m"]
    shifts = [abs(float(row["percent_delta_relative_to_4p5p0"])) for row in mean_rows]

    with MNI_VERSION_VOXELWISE.open(newline="", encoding="utf-8") as handle:
        voxelwise = list(csv.DictReader(handle))
    correlations = [float(row["pearson_r"]) for row in voxelwise]
    relative_l2 = [float(row["relative_l2_difference_percent"]) for row in voxelwise]

    return {
        "roi_count": len(mean_rows),
        "lower_count": sum(float(row["delta_4p0p1_minus_4p5p0"]) < 0 for row in mean_rows),
        "shift_min_percent": min(shifts),
        "shift_max_percent": max(shifts),
        "pearson_min": min(correlations),
        "relative_l2_min_percent": min(relative_l2),
        "relative_l2_max_percent": max(relative_l2),
    }


def summarize_target(summary: list[dict[str, str]], target: str) -> dict[str, float]:
    remesh = [row for row in summary if row["target"] == target and row["condition"] == "remesh"]
    fixed = [row for row in summary if row["target"] == target and row["condition"] == "fixed_mesh"]
    cvs = [float(row["cv_percent"]) for row in remesh]
    reductions = [float(row["sd_reduction_percent"]) for row in fixed]
    fixed_cvs = [float(row["cv_percent"]) for row in fixed]
    return {
        "cv_min": min(cvs),
        "cv_median": statistics.median(cvs),
        "cv_max": max(cvs),
        "min_reduction": min(reductions),
        "fixed_cv_max": max(fixed_cvs),
    }


def top_pair(pairwise: list[dict[str, str]], target: str) -> dict[str, str]:
    rows = [row for row in pairwise if row["target"] == target]
    return max(rows, key=lambda row: float(row["single_run_reversal_probability"]))


def new_slide(prs: Presentation, *, background: RGBColor = LIGHT):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, background)
    return slide


def build_deck() -> tuple[Presentation, list[SlideRecord]]:
    summary, manifest, rank, nested, pairwise = load_data()
    mni_version = load_mni_version_data()
    lh = summarize_target(summary, "left_hippocampus")
    m1 = summarize_target(summary, "right_m1")
    lh_pair = top_pair(pairwise, "left_hippocampus")
    m1_pair = top_pair(pairwise, "right_m1")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    records: list[SlideRecord] = []

    def register(slide, title: str, purpose: str, status: str, notes: str):
        add_notes(slide, notes)
        records.append(SlideRecord(len(records) + 1, title, purpose, status, notes.strip()))

    # 1. Title
    slide = new_slide(prs, background=NAVY)
    add_box(slide, 0, 0, 13.333, 0.13, TEAL)
    add_chip(slide, "INITIAL DRAFT • CONTENT REVIEW", 0.68, 0.54, 2.55, fill=RGBColor(30, 57, 78), color=RGBColor(190, 225, 220))
    add_text(slide, "Mesh realization is\npart of the result", 0.68, 1.25, 7.3, 1.65, size=38, color=WHITE, bold=True)
    add_text(
        slide,
        "Repeatability of individualized temporal-interference E-field simulations",
        0.72,
        3.18,
        6.85,
        0.72,
        size=20,
        color=RGBColor(203, 217, 228),
    )
    add_text(slide, "Boyan Ivanov  •  Lab Journal Club  •  25 September 2026", 0.72, 5.92, 6.9, 0.3, size=11, color=RGBColor(173, 192, 207))
    add_text(slide, "Fresh remeshing produces percent-level spread; repeated solving on one mesh is nearly deterministic.", 0.72, 6.42, 6.95, 0.48, size=11.5, color=WHITE, bold=True)
    # Stylized nested experiment matrix.
    add_text(slide, "FULLY NESTED CHECK", 8.47, 0.83, 3.7, 0.25, size=9, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 8.26, 1.34, 4.15, 4.6, RGBColor(17, 42, 63), line=RGBColor(54, 84, 106), radius=True)
    for row in range(6):
        for col in range(6):
            fill = PALE_TEAL if (row + col) % 3 else RGBColor(157, 216, 207)
            add_circle(slide, 8.68 + col * 0.54, 1.78 + row * 0.54, 0.23, fill)
    add_text(slide, "40", 8.73, 5.0, 1.25, 0.55, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "independent meshes", 8.56, 5.56, 1.65, 0.34, size=9, color=RGBColor(183, 202, 215), align=PP_ALIGN.CENTER)
    add_text(slide, "×", 10.18, 5.05, 0.45, 0.45, size=24, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "40", 10.72, 5.0, 1.25, 0.55, size=30, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "solves per mesh", 10.62, 5.56, 1.45, 0.34, size=9, color=RGBColor(183, 202, 215), align=PP_ALIGN.CENTER)
    add_text(slide, "3,200 simulation records across cohort + nested experiments", 8.31, 6.55, 4.02, 0.38, size=10, color=RGBColor(203, 217, 228), bold=True, align=PP_ALIGN.CENTER)
    register(
        slide,
        "Mesh realization is part of the result",
        "Open with the scientific claim rather than a generic project title.",
        "Complete",
        """
        Opening (about 40 seconds): Individualized TI simulations are usually spoken about as though one anatomy and one montage produce one field map. This experiment asks whether that output is actually unique when the tetrahedral mesh is regenerated. The answer is no: in this workflow, the mesh realization contributes percent-level uncertainty.

        Preview the two experiments: a ten-participant remesh-versus-fixed comparison and one fully nested 40-by-40 case that separates between-mesh from within-mesh variation.
        """,
    )

    # 2. Executive summary
    slide = new_slide(prs)
    add_header(slide, "The field estimate is mesh-dependent", "Take-home first", 2)
    add_metric_card(slide, 0.62, 1.47, 2.86, 1.72, f"{lh['cv_min']:.2f}–{lh['cv_max']:.2f}%", "Left-hippocampus CV across remesh repeats", accent=BLUE)
    add_metric_card(slide, 3.62, 1.47, 2.86, 1.72, f"{m1['cv_min']:.2f}–{m1['cv_max']:.2f}%", "Right-M1 CV across remesh repeats", accent=TEAL)
    add_metric_card(slide, 6.62, 1.47, 2.86, 1.72, f"{100*rank['left_hippocampus']['probability_any_reversal']:.1f}%", "Hippocampal draws with ≥1 rank reversal", accent=ORANGE)
    add_metric_card(slide, 9.62, 1.47, 3.08, 1.72, f"{100*nested['mesh_variance_fraction']:.4f}%", "Variance share attributed to mesh in one nested subject", accent=RED, value_size=25)
    add_box(slide, 0.62, 3.54, 12.08, 1.38, NAVY, radius=True)
    add_text(slide, "Central conclusion", 0.93, 3.82, 2.0, 0.28, size=10, color=RGBColor(174, 202, 220), bold=True)
    add_text(slide, "Within the evaluated CHARM–SimNIBS workflow, mesh generation—not repeated solution on a fixed mesh—dominates repeated-run variability.", 0.93, 4.18, 11.4, 0.48, size=21, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
    add_box(slide, 0.62, 5.27, 12.08, 1.14, PALE_GOLD, line=GOLD, radius=True)
    add_text(slide, "Interpretation boundary", 0.91, 5.53, 2.05, 0.24, size=9.5, color=ORANGE, bold=True)
    add_text(slide, "The 40×40 variance decomposition is a single-participant sensitivity analysis. It strengthens source attribution, but it is not a population variance estimate.", 2.62, 5.47, 9.67, 0.56, size=14, color=INK, bold=True)
    add_source(slide, "Corrected spherical-ROI analysis freeze: 11 Sep 2026")
    register(
        slide,
        "The field estimate is mesh-dependent",
        "State the entire argument and the principal numerical results up front.",
        "Complete",
        f"""
        The cohort analysis shows remesh CVs of {lh['cv_min']:.2f}–{lh['cv_max']:.2f}% in left hippocampus and {m1['cv_min']:.2f}–{m1['cv_max']:.2f}% in right M1. A single remesh repeat often changes at least one pairwise subject ordering. In the fully nested subject, {100*nested['mesh_variance_fraction']:.4f}% of observed variance is assigned to differences between meshes.

        Emphasize the boundary: the nested percentage is conditional on one subject, one target, and this workflow. It is not a claim that all simulation uncertainty everywhere is mesh-driven.
        """,
    )

    # 3. Pipeline motivation
    slide = new_slide(prs)
    add_header(slide, "Identical inputs still pass through a variable numerical representation", "Why this matters", 3)
    stages = [
        ("MRI + labels", "fixed anatomy", DEEP_BLUE, PALE_BLUE),
        ("CHARM mesh", "regenerate or reuse", ORANGE, PALE_ORANGE),
        ("FEM fields", "two carrier fields", TEAL, PALE_TEAL),
        ("TI envelope", "max-TI combination", GOLD, PALE_GOLD),
        ("Spherical ROI", "median field", DEEP_BLUE, PALE_BLUE),
    ]
    xs = [0.58, 3.08, 5.58, 8.08, 10.58]
    for idx, ((name, sub, accent, fill), x) in enumerate(zip(stages, xs)):
        add_box(slide, x, 2.18, 2.12, 1.55, fill, line=accent, radius=True)
        add_circle(slide, x + 0.84, 1.73, 0.42, accent)
        add_text(slide, str(idx + 1), x + 0.84, 1.82, 0.42, 0.16, size=8.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, name, x + 0.14, 2.53, 1.84, 0.33, size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, sub, x + 0.13, 3.03, 1.86, 0.25, size=9.5, color=MUTED, align=PP_ALIGN.CENTER)
        if idx < len(stages) - 1:
            add_arrow_between(slide, x + 2.16, 2.95, x + 2.46)
    add_box(slide, 3.08, 4.25, 2.12, 0.08, ORANGE)
    add_text(slide, "experimental factor", 3.08, 4.46, 2.12, 0.24, size=9, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_box(slide, 0.72, 5.05, 11.92, 1.05, WHITE, line=RULE, radius=True)
    add_text(slide, "Held constant", 0.98, 5.35, 1.55, 0.23, size=10, color=DEEP_BLUE, bold=True)
    add_text(slide, "anatomy • tissue labels • electrodes • currents • conductivities • software settings • ROI definition", 2.38, 5.29, 9.85, 0.36, size=15, color=INK)
    add_text(slide, "Question: can a new valid tetrahedralization change a conclusion even when every scientific input is unchanged?", 0.82, 6.38, 11.7, 0.42, size=17, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Context: Windhoff et al., 2013; Saturnino et al., 2019; Puonti et al., 2020")
    register(
        slide,
        "Identical inputs still pass through a variable numerical representation",
        "Explain where remeshing sits in the end-to-end simulation and what was held fixed.",
        "Complete",
        """
        Walk left to right. The anatomy, labels, montage, currents, conductivities, software, and ROI stay fixed. The experimental factor is whether CHARM creates a fresh tetrahedral mesh or the same mesh workspace is reused.

        The study is not comparing two biological states. It is measuring numerical repeatability of the modelling pipeline.
        """,
    )

    # 4. Dual design
    slide = new_slide(prs)
    add_header(slide, "Two complementary experiments separate pattern from source", "Study design", 4)
    add_box(slide, 0.63, 1.44, 6.0, 4.95, WHITE, line=RULE, radius=True)
    add_chip(slide, "COHORT CONTRAST", 0.9, 1.73, 1.65, fill=PALE_BLUE, color=DEEP_BLUE)
    add_text(slide, "Does the effect recur across people and targets?", 0.92, 2.25, 5.35, 0.56, size=20, color=NAVY, bold=True)
    add_metric_card(slide, 0.94, 3.0, 1.58, 1.26, "10", "participants", accent=DEEP_BLUE, value_size=25)
    add_metric_card(slide, 2.68, 3.0, 1.58, 1.26, "2", "targets", accent=DEEP_BLUE, value_size=25)
    add_metric_card(slide, 4.42, 3.0, 1.58, 1.26, "40+40", "runs / subject–target", accent=DEEP_BLUE, value_size=22)
    add_text(slide, "Left hippocampus + right M1", 0.98, 4.68, 5.15, 0.3, size=14, color=INK, bold=True)
    add_text(slide, "5 female / 5 male • ages 26.00–79.17 years", 0.98, 5.18, 5.15, 0.3, size=13, color=MUTED)
    add_text(slide, "1,600 simulations", 0.98, 5.72, 5.15, 0.38, size=21, color=DEEP_BLUE, bold=True)

    add_box(slide, 6.83, 1.44, 5.87, 4.95, NAVY, line=RGBColor(50, 77, 98), radius=True)
    add_chip(slide, "FULLY NESTED", 7.12, 1.73, 1.55, fill=RGBColor(36, 63, 82), color=GOLD)
    add_text(slide, "Which stage contributes the variance?", 7.12, 2.25, 5.05, 0.56, size=20, color=WHITE, bold=True)
    add_metric_card(slide, 7.15, 3.0, 1.55, 1.26, "1", "participant", accent=GOLD, fill=RGBColor(20, 45, 65), value_size=25)
    add_metric_card(slide, 8.87, 3.0, 1.55, 1.26, "40", "independent meshes", accent=GOLD, fill=RGBColor(20, 45, 65), value_size=25)
    add_metric_card(slide, 10.59, 3.0, 1.55, 1.26, "40", "solves / mesh", accent=GOLD, fill=RGBColor(20, 45, 65), value_size=25)
    add_text(slide, "CC320616 • left hippocampus", 7.16, 4.68, 4.95, 0.3, size=14, color=WHITE, bold=True)
    add_text(slide, "Participant selected once with a fixed seed", 7.16, 5.18, 4.95, 0.3, size=13, color=RGBColor(182, 202, 216))
    add_text(slide, "1,600 simulations", 7.16, 5.72, 4.95, 0.38, size=21, color=GOLD, bold=True)
    add_text(slide, "3,200 completed simulation records in the presentation evidence base", 0.8, 6.69, 11.72, 0.3, size=15, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "CamCAN cohort; corrected spherical-fixed + nested analysis")
    register(
        slide,
        "Two complementary experiments separate pattern from source",
        "Show how the 10-person cohort and fully nested subject answer different questions.",
        "Complete",
        """
        The cohort contrast asks whether remesh-versus-fixed behavior is consistent across ten participants and both a deep and superficial target. The fully nested experiment asks where variation enters by crossing 40 independently generated meshes with 40 repeated solutions per mesh.

        Do not combine the two sample sizes as though 3,200 were independent biological observations. These are technical simulation records.
        """,
    )

    # 5. What the design identifies
    slide = new_slide(prs)
    add_header(slide, "The contrast changes one thing: the mesh realization", "Methods that affect interpretation", 5)
    add_box(slide, 0.62, 1.48, 3.65, 3.4, PALE_BLUE, line=BLUE, radius=True)
    add_text(slide, "REMESH", 0.94, 1.8, 1.35, 0.28, size=11, color=BLUE, bold=True)
    add_text(slide, "Fresh mesh generated\nfor every repeat", 0.94, 2.28, 2.92, 1.08, size=22, color=NAVY, bold=True)
    add_text(slide, "Variation represented", 0.94, 3.63, 2.5, 0.24, size=9.5, color=MUTED, bold=True)
    add_text(slide, "mesh generation + FEM solution + export / ROI extraction", 0.94, 3.98, 2.95, 0.62, size=13.5, color=INK)

    add_box(slide, 4.49, 1.48, 3.65, 3.4, PALE_ORANGE, line=ORANGE, radius=True)
    add_text(slide, "FIXED MESH", 4.81, 1.8, 1.45, 0.28, size=11, color=ORANGE, bold=True)
    add_text(slide, "One mesh reused\nfor every repeat", 4.81, 2.28, 2.95, 1.08, size=22, color=NAVY, bold=True)
    add_text(slide, "Variation represented", 4.81, 3.63, 2.5, 0.24, size=9.5, color=MUTED, bold=True)
    add_text(slide, "residual FEM solution + export / ROI extraction conditional on one mesh", 4.81, 3.98, 2.95, 0.62, size=13.5, color=INK)

    add_box(slide, 8.36, 1.48, 4.34, 3.4, WHITE, line=RULE, radius=True)
    add_text(slide, "PRIMARY OUTCOME", 8.68, 1.8, 1.85, 0.28, size=11, color=TEAL, bold=True)
    add_text(slide, "Median maximum-TI envelope", 8.68, 2.33, 3.58, 0.78, size=23, color=NAVY, bold=True)
    add_text(slide, "within the optimizer-matched, parcel-clipped spherical ROI", 8.68, 3.24, 3.55, 0.68, size=15, color=INK)
    add_chip(slide, "200 mm³ hippocampus", 8.7, 4.26, 1.62, fill=PALE_TEAL, color=TEAL)
    add_chip(slide, "100 mm³ M1", 10.47, 4.26, 1.34, fill=PALE_TEAL, color=TEAL)

    add_box(slide, 0.62, 5.2, 12.08, 1.34, WHITE, line=RULE, radius=True)
    add_text(slide, "Analysis", 0.94, 5.55, 0.9, 0.24, size=10, color=DEEP_BLUE, bold=True)
    add_text(slide, "Per-participant SD/CV", 1.82, 5.52, 2.18, 0.32, size=14, color=INK, bold=True)
    add_text(slide, "•", 4.03, 5.48, 0.3, 0.3, size=17, color=RULE, align=PP_ALIGN.CENTER)
    add_text(slide, "20,000 single-repeat rank draws", 4.36, 5.52, 2.92, 0.32, size=14, color=INK, bold=True)
    add_text(slide, "•", 7.31, 5.48, 0.3, 0.3, size=17, color=RULE, align=PP_ALIGN.CENTER)
    add_text(slide, "balanced one-way random-effects decomposition", 7.64, 5.52, 4.6, 0.32, size=14, color=INK, bold=True)
    add_text(slide, "Participant is the population unit; repeats quantify computational variation.", 0.96, 6.03, 11.0, 0.28, size=12.5, color=RED, bold=True)
    add_source(slide, "Corrected selector aligned with the spherical analysis ROI")
    register(
        slide,
        "The contrast changes one thing: the mesh realization",
        "Define exactly what each arm measures, the primary endpoint, and the inferential unit.",
        "Complete",
        """
        The remesh arm includes everything that can vary after the fixed inputs, including tetrahedralization. The fixed arm measures only residual solution and post-processing variation conditional on one selected mesh.

        The fixed mesh is the remesh geometry nearest the median spherical-ROI result for that participant and target. This deliberately makes it representative of the remesh distribution; it does not make it the most anatomically accurate mesh.

        Mention the correction only if useful here: the current fixed arm was rebuilt after aligning mesh selection with the optimizer-matched spherical ROI.
        """,
    )

    # 6. Left hippocampus result
    slide = new_slide(prs)
    add_header(slide, f"Hippocampal fields vary by {lh['cv_min']:.2f}–{lh['cv_max']:.2f}% under remeshing", "Result 1 • left hippocampus", 6)
    add_box(slide, 0.5, 1.3, 9.82, 5.55, WHITE, line=RULE, radius=True)
    add_picture_contain(slide, LH_FIELD, 0.62, 1.45, 9.58, 5.22)
    add_metric_card(slide, 10.54, 1.5, 2.19, 1.42, f"{lh['cv_median']:.2f}%", "median remesh CV", accent=BLUE, value_size=25)
    add_metric_card(slide, 10.54, 3.13, 2.19, 1.42, f"≥{lh['min_reduction']:.2f}%", "minimum reduction in SD with fixed mesh", accent=ORANGE, value_size=21)
    add_metric_card(slide, 10.54, 4.76, 2.19, 1.42, f"≤{lh['fixed_cv_max']:.3f}%", "maximum fixed-mesh CV", accent=TEAL, value_size=22)
    add_text(slide, "Blue spread collapses to an almost single orange value for every participant.", 10.57, 6.42, 2.1, 0.42, size=10.2, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "n=10; 40 remesh + 40 fixed-mesh repeats per participant")
    register(
        slide,
        f"Hippocampal fields vary by {lh['cv_min']:.2f}–{lh['cv_max']:.2f}% under remeshing",
        "Show the participant-level spread in the deep target and the collapse under mesh reuse.",
        "Complete",
        f"""
        Read the plot as within-participant technical variability, not between-person uncertainty. Each pale blue cloud is 40 independently remeshed simulations. The orange fixed-mesh repeats are visually collapsed.

        The corrected spherical-ROI remesh CV range is {lh['cv_min']:.2f}–{lh['cv_max']:.2f}%, with a median of {lh['cv_median']:.2f}%. The smallest participant-level SD reduction after mesh reuse is {lh['min_reduction']:.2f}%.
        """,
    )

    # 7. Right M1 result
    slide = new_slide(prs)
    add_header(slide, f"M1 shows the same pattern: {m1['cv_min']:.2f}–{m1['cv_max']:.2f}% remesh CV", "Result 2 • right M1", 7)
    add_box(slide, 0.5, 1.3, 9.82, 5.55, WHITE, line=RULE, radius=True)
    add_picture_contain(slide, M1_FIELD, 0.62, 1.45, 9.58, 5.22)
    add_metric_card(slide, 10.54, 1.5, 2.19, 1.42, f"{m1['cv_median']:.2f}%", "median remesh CV", accent=BLUE, value_size=25)
    add_metric_card(slide, 10.54, 3.13, 2.19, 1.42, f"≥{m1['min_reduction']:.2f}%", "minimum reduction in SD with fixed mesh", accent=ORANGE, value_size=21)
    add_metric_card(slide, 10.54, 4.76, 2.19, 1.42, f"≤{m1['fixed_cv_max']:.4f}%", "maximum fixed-mesh CV", accent=TEAL, value_size=20)
    add_text(slide, "Replication in a superficial target argues against a hippocampus-only artifact.", 10.57, 6.42, 2.1, 0.42, size=10.2, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "n=10; 40 remesh + 40 fixed-mesh repeats per participant")
    register(
        slide,
        f"M1 shows the same pattern: {m1['cv_min']:.2f}–{m1['cv_max']:.2f}% remesh CV",
        "Demonstrate that the field variability pattern replicates in a superficial target.",
        "Complete",
        f"""
        The right-M1 result is qualitatively the same. Corrected spherical-ROI remesh CVs range from {m1['cv_min']:.2f}% to {m1['cv_max']:.2f}%, median {m1['cv_median']:.2f}%. The fixed-mesh SD reduction is at least {m1['min_reduction']:.2f}%.

        This replication supports robustness across the two evaluated targets, but not universal generalization to all montages, meshing tools, or ROIs.
        """,
    )

    # 8. Ranking uncertainty
    slide = new_slide(prs)
    add_header(slide, "Percent-level spread matters when subjects are closely matched", "Result 3 • ranking uncertainty", 8)
    add_box(slide, 0.64, 1.52, 5.96, 4.78, WHITE, line=RULE, radius=True)
    add_chip(slide, "LEFT HIPPOCAMPUS", 0.94, 1.82, 1.63, fill=PALE_BLUE, color=DEEP_BLUE)
    add_text(slide, f"{100*rank['left_hippocampus']['probability_any_reversal']:.1f}%", 0.94, 2.48, 2.8, 0.73, size=40, color=ORANGE, bold=True)
    add_text(slide, "of 20,000 draws contained ≥1 reversed pair", 0.96, 3.3, 4.86, 0.44, size=15, color=INK, bold=True)
    add_metric_card(slide, 0.94, 4.06, 2.34, 1.47, f"{rank['left_hippocampus']['median_kendall_rank_agreement']:.3f}", "median Kendall agreement", accent=DEEP_BLUE, value_size=24)
    add_metric_card(slide, 3.53, 4.06, 2.34, 1.47, f"{100*float(lh_pair['single_run_reversal_probability']):.1f}%", "largest exact pair reversal", accent=RED, value_size=24)
    add_text(slide, f"Closest high-risk pair: {lh_pair['higher_mean_subject'].removeprefix('sub-')} vs {lh_pair['lower_mean_subject'].removeprefix('sub-')}", 0.98, 5.75, 4.9, 0.28, size=10, color=MUTED)

    add_box(slide, 6.82, 1.52, 5.88, 4.78, WHITE, line=RULE, radius=True)
    add_chip(slide, "RIGHT M1", 7.12, 1.82, 1.18, fill=PALE_TEAL, color=TEAL)
    add_text(slide, f"{100*rank['right_m1']['probability_any_reversal']:.1f}%", 7.12, 2.48, 2.8, 0.73, size=40, color=ORANGE, bold=True)
    add_text(slide, "of 20,000 draws contained ≥1 reversed pair", 7.14, 3.3, 4.84, 0.44, size=15, color=INK, bold=True)
    add_metric_card(slide, 7.12, 4.06, 2.32, 1.47, f"{rank['right_m1']['median_kendall_rank_agreement']:.3f}", "median Kendall agreement", accent=TEAL, value_size=24)
    add_metric_card(slide, 9.69, 4.06, 2.32, 1.47, f"{100*float(m1_pair['single_run_reversal_probability']):.1f}%", "largest exact pair reversal", accent=RED, value_size=24)
    add_text(slide, f"Closest high-risk pair: {m1_pair['higher_mean_subject'].removeprefix('sub-')} vs {m1_pair['lower_mean_subject'].removeprefix('sub-')}", 7.16, 5.75, 4.8, 0.28, size=10, color=MUTED)
    add_box(slide, 0.64, 6.53, 12.06, 0.45, NAVY, radius=True)
    add_text(slide, "Global ordering remains fairly concordant, but a one-mesh ranking is fragile near ties.", 0.87, 6.65, 11.6, 0.2, size=12, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "One remesh repeat selected independently for each participant per draw")
    register(
        slide,
        "Percent-level spread matters when subjects are closely matched",
        "Translate CV into a decision-relevant consequence: uncertainty in subject ordering.",
        "Complete",
        f"""
        The reference ordering uses each participant's 40-repeat mean. In each of 20,000 draws, one repeat is selected independently per participant.

        At least one pair reverses in {100*rank['left_hippocampus']['probability_any_reversal']:.1f}% of hippocampal draws and {100*rank['right_m1']['probability_any_reversal']:.1f}% of M1 draws. Median agreement remains high, so the message is not that all rankings are random. It is that close pairs are sensitive to which valid mesh was selected.

        These are uncertainty measures, not p-values.
        """,
    )

    # 9. Mesh representation
    slide = new_slide(prs)
    add_header(slide, "Remeshing changes the numerical head model—not only the final scalar", "Mechanistic evidence", 9)
    add_box(slide, 0.58, 1.4, 8.55, 5.42, WHITE, line=RULE, radius=True)
    add_picture_contain(slide, LH_TISSUE, 0.73, 1.59, 8.25, 4.92)
    add_text(slide, "Example: CC320616, left-hippocampus remesh arm", 0.86, 6.42, 7.98, 0.22, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_metric_card(slide, 9.39, 1.58, 3.3, 1.53, "7,369", "mean within-participant SD of total tetrahedral element count • hippocampus", accent=BLUE, value_size=26)
    add_metric_card(slide, 9.39, 3.35, 3.3, 1.53, "7,704", "mean within-participant SD of total tetrahedral element count • M1", accent=TEAL, value_size=26)
    add_box(slide, 9.39, 5.13, 3.3, 1.45, PALE_GOLD, line=GOLD, radius=True)
    add_text(slide, "Interpret carefully", 9.67, 5.4, 2.73, 0.26, size=11, color=ORANGE, bold=True)
    add_text(slide, "Element count and tissue fractions confirm different realizations; they do not identify the causal boundary change.", 9.67, 5.8, 2.73, 0.58, size=12, color=INK)
    add_source(slide, "Tissue deviations are percentage points from each tissue's 40-repeat mean")
    register(
        slide,
        "Remeshing changes the numerical head model—not only the final scalar",
        "Provide structural evidence that each remesh run is a genuinely different discretization.",
        "Complete",
        """
        The heatmap shows small but structured shifts in the proportion of volume assigned to tissue classes across remesh realizations. Total tetrahedral element count also varies substantially within participant, while it is constant when the mesh is reused.

        Do not claim that a particular tissue transition caused the field changes. Element count is coarse, and local boundaries, element quality, electrode–skin geometry, and interpolation can also matter.
        """,
    )

    # 10. Nested design
    slide = new_slide(prs)
    add_header(slide, "The fully nested case isolates between-mesh from within-mesh variation", "Variance attribution", 10)
    add_box(slide, 0.62, 1.48, 7.05, 4.95, WHITE, line=RULE, radius=True)
    add_text(slide, "40 independent meshes", 0.95, 1.82, 2.8, 0.32, size=18, color=DEEP_BLUE, bold=True)
    add_text(slide, "each reused for 40 complete solutions", 3.54, 1.85, 3.55, 0.28, size=13, color=MUTED, align=PP_ALIGN.RIGHT)
    for row in range(5):
        y = 2.43 + row * 0.69
        add_text(slide, f"mesh {row + 1:02d}", 0.96, y + 0.1, 0.86, 0.18, size=8.5, color=MUTED, align=PP_ALIGN.RIGHT)
        for col in range(10):
            color = PALE_BLUE if col % 2 == 0 else RGBColor(207, 229, 243)
            add_circle(slide, 2.0 + col * 0.45, y, 0.25, color, line=RGBColor(158, 198, 222))
        add_text(slide, "…", 6.63, y + 0.01, 0.3, 0.2, size=15, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "⋮", 0.96, 5.98, 0.86, 0.24, size=16, color=MUTED, align=PP_ALIGN.RIGHT)
    add_text(slide, "40 repeats conditional on the same geometry", 1.99, 5.93, 4.9, 0.3, size=11, color=DEEP_BLUE, bold=True, align=PP_ALIGN.CENTER)

    add_box(slide, 7.9, 1.48, 4.8, 2.28, NAVY, radius=True)
    add_text(slide, "Balanced random-effects model", 8.24, 1.85, 4.08, 0.3, size=14, color=GOLD, bold=True)
    add_text(slide, "yₘᵣ = μ + uₘ + εₘᵣ", 8.24, 2.4, 4.08, 0.55, size=28, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "uₘ: between meshes     εₘᵣ: within a mesh", 8.23, 3.18, 4.1, 0.25, size=10, color=RGBColor(185, 203, 216), align=PP_ALIGN.CENTER)

    add_box(slide, 7.9, 4.02, 4.8, 2.41, PALE_GOLD, line=GOLD, radius=True)
    add_text(slide, "Prespecified scope", 8.24, 4.35, 2.2, 0.28, size=11, color=ORANGE, bold=True)
    add_bullet_list(
        slide,
        [
            "CC320616 selected once from the 10-person pool",
            "left hippocampus fixed before preparation",
            "1,600 technical observations",
        ],
        8.22,
        4.83,
        4.0,
        1.25,
        size=12.5,
        bullet_color=ORANGE,
        spacing=5,
    )
    add_text(slide, "Sensitivity demonstration—not a population variance estimate", 7.99, 6.69, 4.63, 0.25, size=11, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Random-selection seed 20260831; selection persisted before execution")
    register(
        slide,
        "The fully nested case isolates between-mesh from within-mesh variation",
        "Explain the 40-by-40 crossing and the variance components before showing the result.",
        "Complete",
        """
        Each outer level is a freshly generated mesh. Each mesh is then copied and solved 40 times. The balanced one-way random-effects model separates variance among mesh means from residual variation among repeated solutions conditional on the same geometry.

        The subject was selected once from the ten eligible participants with a fixed seed and persisted. The target was prespecified as left hippocampus.
        """,
    )

    # 11. Nested result
    slide = new_slide(prs)
    add_header(slide, "The nested case assigns virtually all observed variance to mesh generation", "Result 4 • fully nested subject", 11)
    add_box(slide, 0.5, 1.3, 12.33, 5.58, WHITE, line=RULE, radius=True)
    add_picture_contain(slide, NESTED_FIGURE, 0.62, 1.42, 12.08, 5.0)
    add_box(slide, 0.82, 6.49, 11.68, 0.33, PALE_GOLD, line=GOLD, radius=True)
    add_text(slide, "Boundary: one participant • one target • one workflow. Source attribution, not population prevalence.", 1.0, 6.58, 11.34, 0.16, size=9.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Between-mesh CV 2.018% • within-mesh CV 0.0040% • SD ratio 501×")
    register(
        slide,
        "The nested case assigns virtually all observed variance to mesh generation",
        "Deliver the strongest source-attribution evidence and immediately qualify its scope.",
        "Complete",
        f"""
        Each point is the CV across 40 repeated solutions on one mesh. Most are exactly or effectively zero. The dashed line is the pooled within-mesh CV.

        The between-mesh CV is {nested['between_mesh_cv_percent']:.3f}%; pooled within-mesh CV is {nested['within_mesh_cv_percent']:.4f}%. The component SD ratio is {nested['sd_ratio_between_over_within']:.0f}-fold, and {100*nested['mesh_variance_fraction']:.4f}% of total observed variance is attributed to mesh generation.

        Keep saying 'in this participant' and 'in this workflow.'
        """,
    )

    # 12. SimNIBS-version sensitivity on a fixed MNI152 model
    slide = new_slide(prs)
    add_header(slide, "Version sensitivity is systematic in amplitude—not spatial pattern", "Sensitivity • fixed MNI152 model", 12)
    add_box(slide, 0.5, 1.4, 8.55, 5.57, WHITE, line=RULE, radius=True)
    add_picture_contain(slide, MNI_VERSION_FIGURE, 0.58, 1.47, 8.39, 5.4)

    add_metric_card(
        slide,
        9.28,
        1.47,
        3.48,
        1.32,
        f"{mni_version['lower_count']:.0f} / {mni_version['roi_count']:.0f}",
        "ROIs lower under SimNIBS 4.0.1",
        accent=ORANGE,
        fill=PALE_ORANGE,
        value_size=25,
    )
    add_metric_card(
        slide,
        9.28,
        3.02,
        3.48,
        1.32,
        f"−{mni_version['shift_min_percent']:.2f}% to −{mni_version['shift_max_percent']:.2f}%",
        "change in mean field, 4.0.1 versus 4.5.0",
        accent=DEEP_BLUE,
        fill=PALE_BLUE,
        value_size=20,
    )
    add_metric_card(
        slide,
        9.28,
        4.57,
        3.48,
        1.32,
        f"r ≥ {mni_version['pearson_min']:.5f}",
        "whole-brain voxelwise agreement",
        accent=TEAL,
        fill=PALE_TEAL,
        value_size=22,
    )
    add_box(slide, 9.28, 6.13, 3.48, 0.66, PALE_GOLD, line=GOLD, radius=True)
    add_text(slide, "One fixed model: software sensitivity, not participant-level repeatability.", 9.47, 6.28, 3.1, 0.35, size=10.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Fixed mesh, reference image, montage, ROI definitions, conductivities, and currents; runtime version changed")
    register(
        slide,
        "Version sensitivity is systematic in amplitude—not spatial pattern",
        "Show that software version can shift amplitude even when geometry and spatial pattern are held constant.",
        "Complete",
        f"""
        This is a separate sensitivity analysis on one fixed MNI152 model. The mesh, reference image, ROI definitions, montages, conductivities, currents, and electrode geometry were held constant; only the SimNIBS runtime version changed.

        SimNIBS 4.0.1 produced a lower mean ROI field than 4.5.0 in all {mni_version['roi_count']:.0f} ROIs. The relative shift ranged from −{mni_version['shift_min_percent']:.2f}% to −{mni_version['shift_max_percent']:.2f}%. Whole-brain spatial agreement remained extremely high, with Pearson r at least {mni_version['pearson_min']:.5f}; relative L2 differences ranged from {mni_version['relative_l2_min_percent']:.2f}% to {mni_version['relative_l2_max_percent']:.2f}%.

        The interpretation is a small, systematic amplitude shift with the spatial pattern largely preserved. Do not pool this fixed-model version comparison with the participant or remeshing distributions, and do not attach population inference to the four ROIs.
        """,
    )

    # 13. Evidence synthesis
    slide = new_slide(prs)
    add_header(slide, "Within-version evidence converges on mesh realization", "Interpretation", 13)
    cards = [
        ("1", "Cohort contrast", "Remesh repeats spread; fixed-mesh repeats collapse across 10 participants and two targets.", BLUE, PALE_BLUE),
        ("2", "Nested decomposition", "Between-mesh SD is 501× the residual within-mesh SD in the fully nested subject.", ORANGE, PALE_ORANGE),
        ("3", "Structural correlates", "Element count and tissue-volume composition change across remesh realizations.", TEAL, PALE_TEAL),
    ]
    for idx, (num, heading, body, accent, fill) in enumerate(cards):
        x = 0.65 + idx * 4.12
        add_box(slide, x, 1.62, 3.8, 3.45, fill, line=accent, radius=True)
        add_circle(slide, x + 0.3, 1.94, 0.52, accent)
        add_text(slide, num, x + 0.3, 2.07, 0.52, 0.18, size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, heading, x + 0.32, 2.65, 3.15, 0.36, size=18, color=NAVY, bold=True)
        add_text(slide, body, x + 0.32, 3.29, 3.12, 1.3, size=15, color=INK)
    add_arrow_between(slide, 2.56, 5.35, 6.1, color=RULE)
    add_arrow_between(slide, 6.94, 5.35, 10.52, color=RULE)
    add_box(slide, 1.45, 5.66, 10.43, 0.9, NAVY, radius=True)
    add_text(slide, "Inference: mesh realization is a numerical experimental factor and belongs in the uncertainty budget.", 1.79, 5.93, 9.75, 0.34, size=19, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Dominant source of repeated-run variation ≠ proof that the selected mesh is the most accurate anatomy.", 1.45, 6.77, 10.43, 0.26, size=11, color=RED, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Evidence triangulation across field, variance, and mesh outcomes")
    register(
        slide,
        "Within-version evidence converges on mesh realization",
        "Synthesize the argument without overclaiming causality or accuracy.",
        "Complete",
        """
        The conclusion does not rest on a single plot. The main cohort shows the condition contrast, the nested experiment attributes the variance source, and structural metrics show that remeshing creates different numerical head models.

        This is evidence that mesh realization is part of the computational specification. It is not evidence that any one realization is anatomically truer.
        """,
    )

    # 14. Practical recommendations
    slide = new_slide(prs)
    add_header(slide, "The practical response depends on the scientific comparison", "Implications", 14)
    recs = [
        ("Controlled contrasts", "Freeze and checksum the mesh when comparing conditions within the same anatomy and meshing is a nuisance factor.", DEEP_BLUE, PALE_BLUE),
        ("Percent-level effects", "Repeat remeshing—or otherwise quantify discretization uncertainty—when the expected effect is similar to the observed CV.", ORANGE, PALE_ORANGE),
        ("Subject ranking", "Avoid fine-grained selection or responder stratification from one mesh when participant means are close.", RED, RGBColor(251, 232, 232)),
        ("Reproducibility", "Report meshing policy, selected-mesh provenance, software version, seed/selection rule, and remesh variability.", TEAL, PALE_TEAL),
    ]
    positions = [(0.68, 1.55), (6.78, 1.55), (0.68, 4.08), (6.78, 4.08)]
    for (title, body, accent, fill), (x, y) in zip(recs, positions):
        add_box(slide, x, y, 5.87, 2.06, fill, line=accent, radius=True)
        add_box(slide, x, y, 0.09, 2.06, accent)
        add_text(slide, title, x + 0.34, y + 0.3, 5.1, 0.33, size=17, color=NAVY, bold=True)
        add_text(slide, body, x + 0.34, y + 0.88, 5.08, 0.88, size=14, color=INK)
    add_text(slide, "Rule of thumb for interpretation", 0.86, 6.47, 2.3, 0.22, size=9.5, color=ORANGE, bold=True)
    add_text(slide, "If the scientific contrast is smaller than the mesh-induced spread, the mesh must enter the design or uncertainty statement.", 3.08, 6.4, 9.14, 0.4, size=15, color=NAVY, bold=True)
    add_source(slide, "Recommendations are scoped to the evaluated workflow")
    register(
        slide,
        "The practical response depends on the scientific comparison",
        "Convert the findings into decisions for simulation design, ranking, and reporting.",
        "Complete",
        """
        Mesh reuse is useful for controlled comparisons because it removes a nuisance source. It is not automatically the best strategy for absolute dosimetry or accuracy claims.

        If the effect of interest is a few percent, either average or model multiple mesh realizations, or demonstrate empirically that mesh uncertainty is negligible relative to that contrast. Preserve and checksum the chosen mesh so the model can be reproduced.
        """,
    )

    # 15. Limitations
    slide = new_slide(prs)
    add_header(slide, "The result is strong about repeatability—and deliberately narrow about validity", "Limitations and claim boundaries", 15)
    add_box(slide, 0.67, 1.52, 5.86, 4.95, PALE_TEAL, line=TEAL, radius=True)
    add_text(slide, "SUPPORTED", 0.98, 1.88, 1.3, 0.27, size=11, color=TEAL, bold=True)
    add_bullet_list(
        slide,
        [
            "Fresh remeshing introduces percent-level field variation in both evaluated targets.",
            "Repeated solving on one copied mesh is effectively deterministic.",
            "A single mesh can alter close between-subject rankings.",
            "The nested case isolates mesh generation as the dominant source for one participant.",
        ],
        0.98,
        2.4,
        5.08,
        3.5,
        size=14,
        bullet_color=TEAL,
        spacing=8,
    )
    add_box(slide, 6.8, 1.52, 5.86, 4.95, RGBColor(251, 235, 235), line=RED, radius=True)
    add_text(slide, "NOT ESTABLISHED", 7.11, 1.88, 1.75, 0.27, size=11, color=RED, bold=True)
    add_bullet_list(
        slide,
        [
            "Which mesh is anatomically or physically most accurate.",
            "Generalization beyond 10 datasets, two targets, and this CHARM–SimNIBS pipeline.",
            "A specific tissue boundary or mesh metric as the causal mechanism.",
            "Population inference from the 1,600 nested technical observations.",
        ],
        7.11,
        2.4,
        5.06,
        3.5,
        size=14,
        bullet_color=RED,
        spacing=8,
    )
    add_box(slide, 1.05, 6.63, 11.2, 0.37, PALE_GOLD, line=GOLD, radius=True)
    add_text(slide, "Design caveat: the fixed mesh was selected using the primary outcome, so central agreement is partly imposed by design.", 1.23, 6.72, 10.84, 0.2, size=10.5, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Technical repeats are not independent participants")
    register(
        slide,
        "The result is strong about repeatability—and deliberately narrow about validity",
        "Make the precise claim boundary visible rather than relegating it to a footnote.",
        "Complete",
        """
        The experiment is designed to quantify numerical repeatability, not validation against measured fields. Ten anatomical datasets and two targets provide replication, but not broad coverage of every montage or software stack.

        The representative fixed mesh was selected using the primary field outcome, so its central value is expected to sit near the remesh center. The meaningful fixed-arm result is the near-zero spread conditional on that mesh—not independent accuracy.
        """,
    )

    # 16. Discussion
    slide = new_slide(prs, background=NAVY)
    add_header(slide, "If the claim is smaller than the mesh noise, the mesh belongs in the claim", "Discussion", 16, dark=True)
    questions = [
        "Should future primary analyses average across mesh realizations—or standardize one mesh and treat it as part of the model definition?",
        "What numerical-uncertainty threshold is acceptable before a simulated field difference is interpreted biologically?",
        "Which extension adds the most value: more targets, another mesher, local boundary metrics, or convergence testing?",
    ]
    for idx, question in enumerate(questions, start=1):
        y = 1.52 + (idx - 1) * 1.52
        add_circle(slide, 0.78, y, 0.62, TEAL if idx < 3 else GOLD)
        add_text(slide, str(idx), 0.78, y + 0.18, 0.62, 0.2, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, question, 1.72, y - 0.02, 10.42, 0.92, size=20, color=WHITE, bold=True)
        if idx < 3:
            add_line(slide, 1.72, y + 1.12, 12.18, y + 1.12, color=RGBColor(62, 88, 107), width=0.7)
    add_box(slide, 0.78, 6.22, 11.72, 0.65, RGBColor(22, 50, 70), line=RGBColor(50, 82, 103), radius=True)
    add_text(slide, "Take-home: a reproducible individualized field estimate requires specifying the mesh—not only the anatomy and montage.", 1.03, 6.42, 11.22, 0.27, size=15, color=GOLD, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Questions / critique / next experiment", dark=True)
    register(
        slide,
        "If the claim is smaller than the mesh noise, the mesh belongs in the claim",
        "Close with an actionable thesis and journal-club discussion prompts.",
        "Complete",
        """
        Pause on the take-home sentence, then invite discussion. These questions are intentionally methodological: whether to average or standardize, how to set a numerical tolerance, and what experiment would most efficiently test generalization.
        """,
    )

    # 17. Appendix rank LH
    slide = new_slide(prs)
    add_header(slide, "Appendix • full ranking analysis: left hippocampus", "Backup", 17)
    add_box(slide, 0.5, 1.3, 12.33, 5.64, WHITE, line=RULE, radius=True)
    add_picture_contain(slide, LH_RANK, 0.65, 1.45, 12.04, 5.34)
    add_source(slide, "Exact pairwise probabilities + 20,000 one-repeat-per-participant draws")
    register(slide, "Appendix • full ranking analysis: left hippocampus", "Provide the heatmap and Kendall distribution for questions.", "Complete", "Use only if the audience wants to inspect which pairs drive ranking uncertainty.")

    # 18. Appendix rank M1
    slide = new_slide(prs)
    add_header(slide, "Appendix • full ranking analysis: right M1", "Backup", 18)
    add_box(slide, 0.5, 1.3, 12.33, 5.64, WHITE, line=RULE, radius=True)
    add_picture_contain(slide, M1_RANK, 0.65, 1.45, 12.04, 5.34)
    add_source(slide, "Exact pairwise probabilities + 20,000 one-repeat-per-participant draws")
    register(slide, "Appendix • full ranking analysis: right M1", "Provide the heatmap and Kendall distribution for questions.", "Complete", "Use only if the audience wants to inspect which pairs drive ranking uncertainty.")

    # 19. Appendix cohort and montage
    slide = new_slide(prs)
    add_header(slide, "Appendix • cohort and stimulation parameters", "Backup", 19)
    add_box(slide, 0.63, 1.43, 5.95, 5.35, WHITE, line=RULE, radius=True)
    add_text(slide, "CamCAN cohort", 0.94, 1.75, 2.4, 0.32, size=17, color=NAVY, bold=True)
    cohort = [
        ("CC110174", "Female", "26.00"), ("CC121144", "Male", "26.00"),
        ("CC310407", "Female", "39.00"), ("CC320616", "Male", "39.08"),
        ("CC420071", "Female", "52.58"), ("CC410432", "Male", "52.58"),
        ("CC520083", "Female", "65.08"), ("CC520127", "Male", "66.00"),
        ("CC610631", "Female", "77.58"), ("CC720941", "Male", "79.17"),
    ]
    headers = ["ID", "Recorded sex", "Age"]
    col_x = [0.97, 3.05, 5.15]
    col_w = [1.85, 1.8, 0.9]
    for i, header in enumerate(headers):
        add_text(slide, header, col_x[i], 2.26, col_w[i], 0.22, size=9.5, color=MUTED, bold=True, align=PP_ALIGN.RIGHT if i == 2 else PP_ALIGN.LEFT)
    add_line(slide, 0.95, 2.58, 6.21, 2.58, color=RULE, width=0.8)
    for row_idx, row in enumerate(cohort):
        y = 2.71 + row_idx * 0.35
        if row_idx % 2 == 0:
            add_box(slide, 0.93, y - 0.03, 5.3, 0.31, LIGHT)
        for i, value in enumerate(row):
            add_text(slide, value, col_x[i], y, col_w[i], 0.2, size=9.2, color=INK, align=PP_ALIGN.RIGHT if i == 2 else PP_ALIGN.LEFT)

    add_box(slide, 6.82, 1.43, 5.88, 5.35, WHITE, line=RULE, radius=True)
    add_text(slide, "Target-specific TI montages", 7.14, 1.75, 3.4, 0.32, size=17, color=NAVY, bold=True)
    montage = [
        ("Left hippocampus", "F8–P8", "2.000 mA"),
        ("", "T7–P7", "1.588656 mA"),
        ("Right M1", "Fp2–F6", "2.000 mA"),
        ("", "C4–CP2", "0.632456 mA"),
    ]
    add_text(slide, "Target", 7.17, 2.28, 1.75, 0.22, size=9.5, color=MUTED, bold=True)
    add_text(slide, "Electrode pair", 9.1, 2.28, 1.65, 0.22, size=9.5, color=MUTED, bold=True)
    add_text(slide, "Current", 11.1, 2.28, 1.2, 0.22, size=9.5, color=MUTED, bold=True, align=PP_ALIGN.RIGHT)
    add_line(slide, 7.15, 2.58, 12.33, 2.58, color=RULE, width=0.8)
    for row_idx, row in enumerate(montage):
        y = 2.85 + row_idx * 0.55
        if row_idx % 2 == 0:
            add_box(slide, 7.13, y - 0.05, 5.22, 0.47, LIGHT)
        add_text(slide, row[0], 7.18, y, 1.75, 0.22, size=10.2, color=INK, bold=bool(row[0]))
        add_text(slide, row[1], 9.1, y, 1.65, 0.22, size=10.2, color=INK)
        add_text(slide, row[2], 11.03, y, 1.28, 0.22, size=10.2, color=INK, align=PP_ALIGN.RIGHT)
    add_box(slide, 7.14, 5.45, 5.2, 0.92, PALE_BLUE, line=BLUE, radius=True)
    add_text(slide, "Electrodes", 7.43, 5.7, 1.2, 0.22, size=10, color=DEEP_BLUE, bold=True)
    add_text(slide, "20 × 20 × 2 mm circular in-plane • conductivity 1.4 S/m", 8.55, 5.64, 3.46, 0.43, size=11.5, color=INK)
    add_source(slide, "Same participants evaluated for both target ROIs")
    register(slide, "Appendix • cohort and stimulation parameters", "Provide exact sample and montage details for methods questions.", "Complete", "The cohort is balanced by recorded sex and covers five age bands. Both target experiments use the same ten participants.")

    # 20. Appendix provenance and audits
    slide = new_slide(prs)
    add_header(slide, "Appendix • the current results use the corrected spherical-ROI workflow", "Provenance and sensitivity checks", 20)
    add_box(slide, 0.65, 1.47, 5.88, 4.94, PALE_ORANGE, line=ORANGE, radius=True)
    add_text(slide, "Fixed-mesh selection correction", 0.98, 1.83, 4.9, 0.35, size=18, color=NAVY, bold=True)
    add_text(slide, "Historical selector", 0.98, 2.5, 1.8, 0.22, size=10, color=MUTED, bold=True)
    add_text(slide, "whole anatomical parcel median", 2.54, 2.47, 3.3, 0.3, size=14, color=INK)
    add_arrow_between(slide, 1.5, 3.22, 5.65, color=ORANGE)
    add_text(slide, "Corrected selector", 0.98, 3.56, 1.8, 0.22, size=10, color=ORANGE, bold=True)
    add_text(slide, "optimizer-matched spherical ROI median", 2.54, 3.52, 3.35, 0.34, size=14, color=INK, bold=True)
    add_text(slide, "18 / 20", 0.98, 4.36, 1.82, 0.55, size=30, color=ORANGE, bold=True)
    add_text(slide, "participant–target representative meshes changed", 2.57, 4.42, 3.3, 0.52, size=14, color=INK, bold=True)
    add_text(slide, "Current deck uses the corrected fixed runs and unchanged historical remesh rows.", 0.98, 5.38, 5.1, 0.58, size=13, color=NAVY, bold=True)

    add_box(slide, 6.8, 1.47, 5.88, 4.94, PALE_TEAL, line=TEAL, radius=True)
    add_text(slide, "Finite-support sensitivity audit", 7.13, 1.83, 4.9, 0.35, size=18, color=NAVY, bold=True)
    add_metric_card(slide, 7.13, 2.38, 2.38, 1.4, "97%", "minimum finite ROI fraction • hippocampus", accent=DEEP_BLUE, fill=WHITE, value_size=25)
    add_metric_card(slide, 9.82, 2.38, 2.38, 1.4, "99%", "minimum finite ROI fraction • M1", accent=TEAL, fill=WHITE, value_size=25)
    add_text(slide, "Within participant × condition correlations between finite support and field were weak:", 7.14, 4.13, 5.04, 0.5, size=13, color=INK, bold=True)
    add_text(slide, "hippocampus r = −0.074, ρ = −0.030", 7.14, 4.82, 5.04, 0.28, size=13, color=DEEP_BLUE)
    add_text(slide, "M1 r = 0.097, ρ = 0.063", 7.14, 5.27, 5.04, 0.28, size=13, color=TEAL)
    add_text(slide, "Descriptive sensitivity check; not evidence of causality.", 7.14, 5.78, 5.04, 0.3, size=11, color=RED, bold=True)
    add_box(slide, 0.92, 6.67, 11.52, 0.31, NAVY, radius=True)
    add_text(slide, "Analysis status: complete • 1,600 corrected cohort rows + 1,600 nested rows • source outputs read only", 1.1, 6.75, 11.15, 0.16, size=9.5, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Analysis manifests dated 11 Sep 2026")
    register(
        slide,
        "Appendix • the current results use the corrected spherical-ROI workflow",
        "Document the resolved selector mismatch and the finite-support sensitivity audit.",
        "Complete",
        """
        The historical fixed mesh was selected using the anatomical parcel median, while the paper endpoint was refreshed to the optimizer-matched spherical ROI. Recomputing selection changed 18 of 20 participant–target representatives. The current presentation uses the regenerated spherical-fixed runs.

        Some spherical ROI voxels were non-finite in a subset of runs, but support remained at least 97% in hippocampus and 99% in M1. Within-group correlations with the field metric were weak and are descriptive only.
        """,
    )

    # 21. References
    slide = new_slide(prs)
    add_header(slide, "References and evidence sources", "Backup", 21)
    refs = [
        "Grossman N, et al. Noninvasive deep brain stimulation via temporally interfering electric fields. Cell. 2017;169:1029–1041.e16.",
        "Rampersad S, et al. Prospects for transcranial temporal interference stimulation in humans: A computational study. NeuroImage. 2019;202:116124.",
        "Windhoff M, Opitz A, Thielscher A. Electric field calculations in brain stimulation based on finite elements. Human Brain Mapping. 2013;34:923–935.",
        "Saturnino GB, et al. SimNIBS 2.1: A comprehensive pipeline for individualized electric-field modelling. In: Brain and Human Body Modeling. 2019.",
        "Puonti O, et al. Accurate and robust whole-head segmentation from MRI for individualized head modelling. NeuroImage. 2020;219:117044.",
        "Indahlastari A, Chauhan M, Sadleir RJ. Benchmarking transcranial electrical-stimulation finite-element models. J Neural Eng. 2019;16:026019.",
        "Shafto MA, et al. The Cambridge Centre for Ageing and Neuroscience study protocol. BMC Neurology. 2014;14:204.",
    ]
    add_box(slide, 0.64, 1.45, 12.06, 4.97, WHITE, line=RULE, radius=True)
    y = 1.78
    for idx, ref in enumerate(refs, start=1):
        add_text(slide, f"{idx}", 0.96, y, 0.32, 0.26, size=9.5, color=TEAL, bold=True, align=PP_ALIGN.RIGHT)
        add_text(slide, ref, 1.47, y - 0.02, 10.75, 0.52, size=12.5, color=INK)
        y += 0.63
    add_box(slide, 0.64, 6.62, 12.06, 0.4, PALE_BLUE, line=BLUE, radius=True)
    add_text(slide, "Primary presentation evidence: corrected spherical-fixed and nested analysis manifests generated 11 September 2026.", 0.88, 6.72, 11.56, 0.2, size=10, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_source(slide, "Internal initial draft • source paths documented in build script")
    register(slide, "References and evidence sources", "Provide literature context and identify the internal corrected data freeze.", "Complete", "These references support the TI, SimNIBS, finite-element, and CamCAN context. The numerical results in the deck come from the corrected local analysis artifacts, not the older manuscript tables.")

    return prs, records


def write_notes(records: list[SlideRecord]):
    lines = [
        "# Repeatability journal-club speaker notes",
        "",
        "Core talk: slides 1–16 (about 16–21 minutes). Slides 17–21 are backup.",
        "",
    ]
    for record in records:
        lines.extend([f"## {record.number}. {record.title}", "", record.notes, ""])
    NOTES_PATH.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_content_review(records: list[SlideRecord]):
    lines = [
        "# Repeatability journal-club deck: content review",
        "",
        "## Working thesis",
        "",
        "Within the evaluated CHARM–SimNIBS workflow, tetrahedral mesh realization is the dominant source of repeated-run variability; percent-level field differences and close subject rankings should therefore be interpreted with mesh-induced uncertainty in mind.",
        "",
        "## Assumed format",
        "",
        "- Internal lab journal club",
        "- 16–21 minute core talk plus discussion",
        "- Core slides 1–16; backup slides 17–21",
        "- Repeatability data freeze: corrected spherical-ROI and fully nested analysis generated 11 September 2026",
        "- Software-version sensitivity: fixed MNI152 comparison supplied 25 September 2026",
        "",
        "## Slide map",
        "",
        "| # | Slide | Purpose | Figure/data status |",
        "|---:|---|---|---|",
    ]
    for record in records:
        lines.append(f"| {record.number} | {record.title} | {record.purpose} | {record.status} |")
    lines.extend(
        [
            "",
            "## Placeholders / optional additions after content approval",
            "",
            "No critical scientific figure or analysis is missing from this draft, and no placeholders are required. Optional additions are:",
            "",
            "- laboratory or university branding on the title slide",
            "- an anatomical rendering of the two spherical ROIs if a preferred house figure exists",
            "- a mesh close-up that visually localizes boundary differences between two remesh realizations",
            "- final author/affiliation wording if this will be reused outside the lab",
            "",
            "## Version guardrail",
            "",
            "The deck deliberately uses the corrected spherical-ROI values rather than the older manuscript tables. The provenance slide documents the 18/20 fixed-mesh selection changes.",
        ]
    )
    CONTENT_PATH.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def validate_layout(prs: Presentation):
    errors: list[str] = []
    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            x = shape.left / Inches(1)
            y = shape.top / Inches(1)
            w = shape.width / Inches(1)
            h = shape.height / Inches(1)
            if x < -0.01 or y < -0.01 or x + w > 13.343 or y + h > 7.51:
                errors.append(
                    f"slide {slide_index}: out-of-bounds shape {getattr(shape, 'name', '?')} "
                    f"at ({x:.2f},{y:.2f},{w:.2f},{h:.2f})"
                )
    if errors:
        raise ValueError("\n".join(errors))


def main():
    require_files()
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    prs, records = build_deck()
    validate_layout(prs)
    prs.save(PPTX_PATH)
    write_notes(records)
    write_content_review(records)
    print(json.dumps({
        "status": "complete",
        "slides": len(prs.slides),
        "core_slides": 16,
        "backup_slides": 5,
        "pptx": str(PPTX_PATH),
        "speaker_notes": str(NOTES_PATH),
        "content_review": str(CONTENT_PATH),
    }, indent=2))


if __name__ == "__main__":
    main()
