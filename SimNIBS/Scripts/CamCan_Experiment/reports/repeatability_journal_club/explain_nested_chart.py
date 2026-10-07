#!/usr/bin/env python3
"""Make the fully nested result chart self-explanatory in a copied deck."""
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

import build_repeatability_journal_club as theme


ORIGINAL_TITLE = "The nested case assigns virtually all observed variance to mesh generation"


def _inches(value: int) -> float:
    return value / Inches(1)


def _delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def _replace_text_preserving_style(shape, text: str) -> None:
    frame = shape.text_frame
    paragraph = frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text
    for extra in frame.paragraphs[1:]:
        extra.text = ""


def _find_text_shape(slide, exact_text: str):
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip() == exact_text:
            return shape
    raise ValueError(f"Could not find text on slide: {exact_text!r}")


def _find_slide(prs: Presentation):
    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text.strip() == ORIGINAL_TITLE:
                return slide
    raise ValueError(f"Could not find slide titled {ORIGINAL_TITLE!r}")


def _add_explainer_card(slide, y: float, heading: str, body: str, *, accent, fill) -> None:
    theme.add_box(slide, 9.08, y, 3.7, 1.35, fill, line=accent, radius=True)
    theme.add_box(slide, 9.08, y, 0.08, 1.35, accent)
    theme.add_text(slide, heading, 9.38, y + 0.2, 3.05, 0.3, size=14.5, color=accent, bold=True)
    theme.add_text(slide, body, 9.38, y + 0.62, 3.05, 0.55, size=10.5, color=theme.INK)


def revise_nested_result_slide(prs: Presentation) -> None:
    slide = _find_slide(prs)
    section = _find_text_shape(slide, "RESULT 4 • FULLY NESTED SUBJECT")
    title = _find_text_shape(slide, ORIGINAL_TITLE)
    source = _find_text_shape(slide, "Between-mesh CV 2.018% • within-mesh CV 0.0040% • SD ratio 501×")
    _replace_text_preserving_style(section, "RESULT 4 • HOW TO READ THE NESTED CHART")
    _replace_text_preserving_style(title, "Within each fixed mesh, 40 repeated solutions are nearly identical")
    _replace_text_preserving_style(source, "Blue points: within-mesh CVs • dashed line: pooled within-mesh CV • between-mesh CV uses the 40 mesh means")

    for shape in list(slide.shapes):
        y = _inches(shape.top)
        if 1.28 <= y < 7.0:
            _delete_shape(shape)

    theme.add_box(slide, 0.5, 1.4, 8.35, 4.55, theme.WHITE, line=theme.RULE, radius=True)
    theme.add_picture_contain(slide, theme.NESTED_FIGURE, 0.61, 1.5, 8.13, 4.34)

    _add_explainer_card(
        slide,
        1.4,
        "● Blue point = one mesh",
        "Height = CV across that mesh's 40 repeated solutions. Median 0%; maximum 0.023%.",
        accent=theme.DEEP_BLUE,
        fill=theme.PALE_BLUE,
    )
    _add_explainer_card(
        slide,
        2.92,
        "– – Within meshes: CV 0.0040%",
        "The dashed line is the pooled residual. Model residual SD = 9.72 × 10⁻⁶ V/m.",
        accent=theme.ORANGE,
        fill=theme.PALE_ORANGE,
    )
    _add_explainer_card(
        slide,
        4.44,
        "Between meshes: CV 2.018%",
        "This uses the 40 mesh means. Model mesh SD = 4.87 × 10⁻³ V/m; it is not plotted here.",
        accent=theme.TEAL,
        fill=theme.PALE_TEAL,
    )

    theme.add_box(slide, 0.68, 6.08, 11.98, 0.65, theme.NAVY, radius=True)
    theme.add_text(
        slide,
        "RANDOM-EFFECTS MODEL OUTPUT",
        1.0,
        6.19,
        2.45,
        0.16,
        size=8.4,
        color=theme.GOLD,
        bold=True,
    )
    theme.add_text(
        slide,
        "between / within SD = 501×   •   mesh variance share = 99.9996%",
        3.25,
        6.24,
        9.03,
        0.24,
        size=14,
        color=theme.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    theme.add_text(
        slide,
        "Scope: one participant • one target • one workflow. This identifies a technical source; it is not population inference.",
        0.84,
        6.82,
        11.64,
        0.2,
        size=9.5,
        color=theme.RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    slide.notes_slide.notes_text_frame.text = (
        "Read this chart from the blue points outward. Each blue point is not a mesh mean: it is the CV "
        "across 40 repeated solutions performed on one fixed mesh. The orange dashed line pools that "
        "within-mesh residual variability across all 40 meshes. The between-mesh CV of 2.018% is a "
        "different summary calculated from the spread of the 40 mesh means, so it is not drawn on this "
        "small within-mesh y-axis. Even the largest within-mesh CV is only 0.023%. In this participant and "
        "workflow, the between-mesh component SD is 501-fold larger and accounts for 99.9996% of the "
        "observed variance."
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    args = parser.parse_args()

    prs = Presentation(args.pptx)
    if len(prs.slides) != 22:
        raise ValueError(f"Expected the 22-slide model-explained deck, found {len(prs.slides)} slides")
    revise_nested_result_slide(prs)
    theme.validate_layout(prs)
    prs.save(args.pptx)
    print(f"Saved {len(prs.slides)} slides to {args.pptx}")


if __name__ == "__main__":
    main()
