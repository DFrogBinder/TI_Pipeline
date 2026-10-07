#!/usr/bin/env python3
"""Revise the copied journal-club deck with a self-contained model explanation."""
from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

import build_repeatability_journal_club as theme


def _inches(value: int) -> float:
    return value / Inches(1)


def _delete_shape(shape) -> None:
    element = shape._element
    element.getparent().remove(element)


def _replace_text_preserving_style(shape, text: str) -> None:
    frame = shape.text_frame
    paragraph = frame.paragraphs[0]
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text
    for extra in frame.paragraphs[1:]:
        extra.text = ""


def _find_text_shape(slide, exact_text: str):
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.text.strip() == exact_text:
            return shape
    raise ValueError(f"Could not find text on slide: {exact_text!r}")


def _add_label_card(slide, y: float, label: str, key: str, explanation: str, *, accent, fill) -> None:
    theme.add_box(slide, 7.9, y, 4.8, 1.43, fill, line=accent, radius=True)
    theme.add_box(slide, 7.9, y, 0.08, 1.43, accent)
    theme.add_text(slide, label, 8.2, y + 0.18, 2.7, 0.2, size=9, color=accent, bold=True)
    theme.add_text(slide, key, 8.2, y + 0.47, 4.0, 0.38, size=15.2, color=theme.NAVY, bold=True)
    theme.add_text(slide, explanation, 8.2, y + 0.96, 4.0, 0.34, size=9.5, color=theme.MUTED)


def revise_design_slide(slide) -> None:
    section = _find_text_shape(slide, "VARIANCE ATTRIBUTION")
    title = _find_text_shape(slide, "The fully nested case isolates between-mesh from within-mesh variation")
    source = _find_text_shape(slide, "Random-selection seed 20260831; selection persisted before execution")
    _replace_text_preserving_style(section, "NESTED DESIGN")
    _replace_text_preserving_style(title, "The nested design creates two separable levels of technical variation")
    _replace_text_preserving_style(source, "One fixed participant and target • 40 independently generated meshes • 40 complete solutions per mesh")

    for shape in list(slide.shapes):
        x = _inches(shape.left)
        y = _inches(shape.top)
        if x >= 7.75 and y < 7.0:
            _delete_shape(shape)

    _add_label_card(
        slide,
        1.48,
        "BALANCED",
        "40 meshes × 40 solutions",
        "Every mesh contributes the same number of observations and therefore has equal weight.",
        accent=theme.DEEP_BLUE,
        fill=theme.PALE_BLUE,
    )
    _add_label_card(
        slide,
        3.08,
        "NESTED",
        "Each solution belongs to one mesh",
        "Within-mesh comparisons hold the geometry fixed; mesh means are compared across geometries.",
        accent=theme.TEAL,
        fill=theme.PALE_TEAL,
    )
    _add_label_card(
        slide,
        4.68,
        "RANDOM MESH EFFECT",
        "Meshes are sampled realizations",
        "The goal is to estimate pipeline variability—not to compare named meshes as conditions.",
        accent=theme.ORANGE,
        fill=theme.PALE_ORANGE,
    )
    theme.add_box(slide, 7.9, 6.29, 4.8, 0.51, theme.PALE_GOLD, line=theme.GOLD, radius=True)
    theme.add_text(
        slide,
        "Two questions: how far apart are mesh means, and how tightly do repeats cluster within a mesh?",
        8.12,
        6.4,
        4.35,
        0.26,
        size=10.2,
        color=theme.INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    slide.notes_slide.notes_text_frame.text = (
        "The hierarchy matters before the equation. The design is balanced because all 40 meshes have "
        "40 solutions. It is nested because every solution belongs to exactly one mesh. Mesh is a random "
        "effect because the numbered meshes are sampled realizations of the pipeline, not scientific "
        "conditions of interest. The subject and target are fixed."
    )


def add_model_slide(prs: Presentation):
    slide = theme.new_slide(prs)
    theme.add_header(
        slide,
        "The model separates between-mesh from within-mesh variation",
        "Random-effects model",
        11,
    )

    theme.add_box(slide, 0.67, 1.46, 11.99, 1.45, theme.WHITE, line=theme.RULE, radius=True)
    theme.add_text(slide, "ONE OBSERVED ROI FIELD", 0.94, 1.66, 2.1, 0.18, size=8.8, color=theme.MUTED, bold=True)

    equation_items = [
        (0.94, 1.94, 1.78, "observed field", theme.LIGHT, theme.NAVY),
        (2.82, 1.94, 0.34, "=", theme.WHITE, theme.MUTED),
        (3.25, 1.94, 2.02, "overall mean", theme.PALE_BLUE, theme.DEEP_BLUE),
        (5.38, 1.94, 0.34, "+", theme.WHITE, theme.MUTED),
        (5.82, 1.94, 2.52, "mesh-specific shift", theme.PALE_ORANGE, theme.ORANGE),
        (8.45, 1.94, 0.34, "+", theme.WHITE, theme.MUTED),
        (8.89, 1.94, 2.98, "repeat-to-repeat residual", theme.PALE_TEAL, theme.TEAL),
    ]
    for x, y, w, label, fill, color in equation_items:
        if label in {"=", "+"}:
            theme.add_text(slide, label, x, y + 0.12, w, 0.28, size=20, color=color, bold=True, align=PP_ALIGN.CENTER)
        else:
            theme.add_box(slide, x, y, w, 0.56, fill, line=color, radius=True)
            theme.add_text(slide, label, x + 0.08, y + 0.17, w - 0.16, 0.22, size=12.2, color=color, bold=True, align=PP_ALIGN.CENTER)
    theme.add_text(slide, "yₘᵣ = μ + uₘ + εₘᵣ", 0.94, 2.61, 11.35, 0.21, size=11.5, color=theme.MUTED, bold=True, align=PP_ALIGN.CENTER)

    theme.add_box(slide, 0.67, 3.15, 5.77, 2.49, theme.PALE_ORANGE, line=theme.ORANGE, radius=True)
    theme.add_text(slide, "BETWEEN MESHES  •  MESH VARIANCE", 0.98, 3.48, 4.4, 0.26, size=14, color=theme.ORANGE, bold=True)
    theme.add_text(slide, "Do independently generated geometries give different average fields?", 0.98, 3.91, 4.98, 0.55, size=17, color=theme.NAVY, bold=True)
    theme.add_text(slide, "1  Average the 40 solutions within each mesh", 0.98, 4.68, 4.95, 0.25, size=12.5, color=theme.INK)
    theme.add_text(slide, "2  Compare the 40 mesh means", 0.98, 5.05, 4.95, 0.25, size=12.5, color=theme.INK)
    theme.add_text(slide, "Large value → remeshing changes the estimated field", 0.98, 5.4, 4.98, 0.2, size=10.5, color=theme.ORANGE, bold=True)

    theme.add_box(slide, 6.89, 3.15, 5.77, 2.49, theme.PALE_TEAL, line=theme.TEAL, radius=True)
    theme.add_text(slide, "WITHIN A MESH  •  RESIDUAL VARIANCE", 7.2, 3.48, 4.5, 0.26, size=14, color=theme.TEAL, bold=True)
    theme.add_text(slide, "If geometry is fixed, how much do repeated solutions still differ?", 7.2, 3.91, 4.98, 0.55, size=17, color=theme.NAVY, bold=True)
    theme.add_text(slide, "1  Compare each solution with its own mesh mean", 7.2, 4.68, 4.95, 0.25, size=12.5, color=theme.INK)
    theme.add_text(slide, "2  Pool the remaining within-mesh scatter", 7.2, 5.05, 4.95, 0.25, size=12.5, color=theme.INK)
    theme.add_text(slide, "Small value → the same mesh solves reproducibly", 7.2, 5.4, 4.98, 0.2, size=10.5, color=theme.TEAL, bold=True)

    theme.add_box(slide, 0.9, 5.96, 11.52, 0.67, theme.NAVY, radius=True)
    theme.add_text(
        slide,
        "Total technical variance = mesh variance + within-mesh variance     •     mesh share = mesh variance / total",
        1.18,
        6.18,
        10.96,
        0.24,
        size=14,
        color=theme.WHITE,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    theme.add_text(
        slide,
        "Scope: one fixed participant and target. The model attributes technical variation; it does not estimate population variance.",
        0.92,
        6.76,
        11.48,
        0.24,
        size=10.3,
        color=theme.RED,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    theme.add_source(slide, "m = mesh • r = repeated solution • uₘ = mesh effect • εₘᵣ = residual within that mesh")
    slide.notes_slide.notes_text_frame.text = (
        "Read the equation in plain language: an observed field is the overall mean plus the shift associated "
        "with that mesh plus residual repeat-to-repeat variation. The left component is estimated from the "
        "spread of mesh means. The right component is estimated from scatter around each mesh's own mean. "
        "Because there is one fixed participant and target, this is source attribution within this workflow, "
        "not population inference."
    )
    return slide


def insert_slide_after(prs: Presentation, slide, after_index: int) -> None:
    slide_id = prs.slides._sldIdLst[-1]
    prs.slides._sldIdLst.remove(slide_id)
    prs.slides._sldIdLst.insert(after_index + 1, slide_id)


def renumber_slide_footers(prs: Presentation) -> None:
    for number, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if _inches(shape.left) < 12.0 or _inches(shape.top) < 7.0:
                continue
            if shape.text.strip().isdigit():
                _replace_text_preserving_style(shape, f"{number:02d}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("pptx", type=Path)
    args = parser.parse_args()

    prs = Presentation(args.pptx)
    if len(prs.slides) != 21:
        raise ValueError(f"Expected the copied 21-slide deck, found {len(prs.slides)} slides")

    revise_design_slide(prs.slides[9])
    model_slide = add_model_slide(prs)
    insert_slide_after(prs, model_slide, after_index=9)
    renumber_slide_footers(prs)
    theme.validate_layout(prs)
    prs.save(args.pptx)
    print(f"Saved {len(prs.slides)} slides to {args.pptx}")


if __name__ == "__main__":
    main()
