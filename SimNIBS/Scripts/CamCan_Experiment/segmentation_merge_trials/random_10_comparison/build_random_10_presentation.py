#!/usr/bin/env python3
"""Build a presentation-ready comparison of the 10-subject merge trial."""
from __future__ import annotations

import argparse
import csv
import json
import math
import statistics
from pathlib import Path
from typing import Any

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


DEFAULT_ROOT = Path("/home/boyan/sandbox/Jake_Data/segmentation-merge-random-10")
OUT_NAME = "random_10_segmentation_merge_A_vs_B_comparison.pptx"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

WHITE = RGBColor(255, 255, 255)
INK = RGBColor(28, 34, 42)
MUTED = RGBColor(91, 101, 113)
LIGHT = RGBColor(247, 248, 250)
RULE = RGBColor(210, 218, 227)
BOX = RGBColor(238, 241, 245)
BLUE = RGBColor(35, 96, 150)
ORANGE = RGBColor(182, 79, 57)
GREEN = RGBColor(37, 128, 91)
GOLD = RGBColor(221, 157, 49)

TISSUE_NAMES = {
    0: "Background / enclosed gap",
    1: "White matter",
    2: "Grey matter",
    3: "CSF",
    4: "Bone",
    5: "Skin/scalp",
    6: "Eyes",
    7: "Compact bone",
    8: "Spongy bone",
    9: "Blood",
    10: "Muscle",
}

TISSUE_HEX = {
    0: "#cbd3dc",
    1: "#e8e8e8",
    2: "#818181",
    3: "#68a3ff",
    4: "#ffefb3",
    5: "#e28064",
    6: "#ffe020",
    7: "#e6cf8f",
    8: "#ff8a39",
    9: "#205592",
    10: "#238543",
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument("--output", type=Path, help="Output PPTX path; defaults inside --root.")
    return parser.parse_args()


def load_rows(root: Path) -> list[dict[str, Any]]:
    path = root / "summary.csv"
    with path.open(newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))
    numeric_fields = set(rows[0]) - {"subject"}
    for row in rows:
        for field in numeric_fields:
            row[field] = float(row[field])
    return rows


def load_qc(root: Path, subject: str) -> dict[str, Any]:
    return json.loads((root / "subjects" / subject / "merge_qc.json").read_text(encoding="utf-8"))


def median(values: list[float]) -> float:
    return float(statistics.median(values))


def fmt_range(values: list[float], decimals: int = 1) -> str:
    return f"{min(values):.{decimals}f}-{max(values):.{decimals}f}"


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: int = 18,
    color: RGBColor = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.03,
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = text
    paragraph.font.name = "Aptos"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color


def add_box(slide, x: float, y: float, w: float, h: float, fill: RGBColor, line: RGBColor | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(0.6)
    return shape


def set_background(slide, color: RGBColor = LIGHT) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def add_title(slide, title: str, subtitle: str | None = None) -> None:
    add_text(slide, title, 0.42, 0.18, 12.5, 0.47, size=22, bold=True)
    if subtitle:
        add_text(slide, subtitle, 0.43, 0.66, 12.35, 0.28, size=10, color=MUTED)
    add_box(slide, 0.42, 0.98, 12.5, 0.015, RULE)


def add_footer(slide, page: int, text: str = "10-subject segmentation merge trial | No meshing or e-field simulation") -> None:
    add_text(slide, text, 0.42, 7.18, 11.8, 0.16, size=7, color=MUTED)
    add_text(slide, str(page), 12.32, 7.18, 0.55, 0.16, size=7, color=MUTED, align=PP_ALIGN.RIGHT)


def add_label_bar(slide, x: float, y: float, w: float, text: str, color: RGBColor) -> None:
    add_box(slide, x, y, w, 0.28, color)
    add_text(slide, text, x + 0.05, y + 0.055, w - 0.1, 0.16, size=8, color=WHITE, bold=True, align=PP_ALIGN.CENTER)


def add_picture_fit(slide, path: Path, x: float, y: float, w: float, h: float) -> None:
    with Image.open(path) as image:
        image_w, image_h = image.size
    image_ratio = image_w / image_h
    box_ratio = w / h
    if image_ratio > box_ratio:
        final_w = w
        final_h = w / image_ratio
    else:
        final_h = h
        final_w = h * image_ratio
    slide.shapes.add_picture(
        str(path),
        Inches(x + (w - final_w) / 2),
        Inches(y + (h - final_h) / 2),
        width=Inches(final_w),
        height=Inches(final_h),
    )


def make_summary_chart(root: Path, rows: list[dict[str, Any]]) -> Path:
    ordered = sorted(rows, key=lambda row: row["additional_skin_volume_cm3_b_minus_a"])
    subjects = [str(row["subject"]).removeprefix("sub-") for row in ordered]
    volume = [float(row["additional_skin_volume_cm3_b_minus_a"]) for row in ordered]
    percent = [float(row["changed_percent_of_candidate_a_head"]) for row in ordered]

    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.0), facecolor="#f7f8fa")
    fig.subplots_adjust(left=0.12, right=0.98, bottom=0.13, top=0.86, wspace=0.42)
    y = np.arange(len(subjects))
    axes[0].barh(y, volume, color="#b64f39", height=0.62)
    axes[0].set_yticks(y, subjects, fontsize=9)
    axes[0].set_xlabel("Additional skin volume in B (cm3)", fontsize=10)
    axes[0].set_title("B adds skin in every sampled subject", fontsize=12, fontweight="bold", loc="left")
    axes[1].barh(y, percent, color="#236096", height=0.62)
    axes[1].set_yticks(y, subjects, fontsize=9)
    axes[1].set_xlabel("Changed voxels as % of Approach A head", fontsize=10)
    axes[1].set_title("The reassignment is a substantial tissue change", fontsize=12, fontweight="bold", loc="left")
    for axis in axes:
        axis.grid(axis="x", color="#d2dae3", linewidth=0.7)
        axis.set_axisbelow(True)
        axis.spines[["top", "right", "left"]].set_visible(False)
        axis.tick_params(axis="y", length=0)
        axis.set_facecolor("#f7f8fa")
    path = root / "presentation_assets" / "cohort_summary.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=190, facecolor=fig.get_facecolor())
    plt.close(fig)
    return path


def make_relabel_chart(root: Path, rows: list[dict[str, Any]]) -> tuple[Path, dict[int, float]]:
    totals: dict[int, int] = {}
    voxel_volume_by_subject = {str(row["subject"]): float(row["voxel_volume_mm3"]) for row in rows}
    volumes: dict[int, float] = {}
    for row in rows:
        subject = str(row["subject"])
        qc = load_qc(root, subject)
        for label_text, count in qc["candidate_a_labels_changed_by_b"].items():
            label = int(label_text)
            totals[label] = totals.get(label, 0) + int(count)
            volumes[label] = volumes.get(label, 0.0) + int(count) * voxel_volume_by_subject[subject] / 1000.0

    labels = sorted(volumes, key=volumes.get, reverse=True)
    values = [volumes[label] for label in labels]
    names = [TISSUE_NAMES.get(label, f"Label {label}") for label in labels]
    colors = [TISSUE_HEX.get(label, "#9da8b4") for label in labels]
    fig, axis = plt.subplots(figsize=(11.8, 4.8), facecolor="#f7f8fa")
    fig.subplots_adjust(left=0.22, right=0.97, bottom=0.15, top=0.88)
    y = np.arange(len(labels))
    axis.barh(y, values, color=colors, edgecolor="#8d98a5", linewidth=0.4, height=0.62)
    axis.set_yticks(y, names, fontsize=9.5)
    axis.invert_yaxis()
    axis.set_xlabel("Aggregate volume relabelled as skin across 10 subjects (cm3)", fontsize=10)
    axis.set_title("Approach B replaces multiple CHARM tissue classes with skin", fontsize=13, fontweight="bold", loc="left")
    axis.grid(axis="x", color="#d2dae3", linewidth=0.7)
    axis.set_axisbelow(True)
    axis.spines[["top", "right", "left"]].set_visible(False)
    axis.tick_params(axis="y", length=0)
    axis.set_facecolor("#f7f8fa")
    for index, value in enumerate(values):
        axis.text(value, index, f"  {value:,.0f}", va="center", fontsize=8.5, color="#434d59")
    path = root / "presentation_assets" / "relabelled_tissues.png"
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=190, facecolor=fig.get_facecolor())
    plt.close(fig)
    return path, volumes


def add_title_slide(prs: Presentation, rows: list[dict[str, Any]], manifest: dict[str, Any]) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_box(slide, 0, 0, 13.333, 0.16, BLUE)
    add_box(slide, 0, 0.16, 13.333, 0.08, ORANGE)
    add_text(slide, "Segmentation Merge Trial", 0.68, 0.9, 11.9, 0.55, size=31, bold=True)
    add_text(slide, "Approach A vs Approach B across 10 randomly sampled CamCAN subjects", 0.7, 1.48, 11.9, 0.34, size=16, color=MUTED)

    volumes = [float(row["additional_skin_volume_cm3_b_minus_a"]) for row in rows]
    percents = [float(row["changed_percent_of_candidate_a_head"]) for row in rows]
    add_box(slide, 0.72, 2.25, 7.08, 2.55, WHITE, RULE)
    add_text(slide, "Main result", 1.0, 2.55, 6.45, 0.25, size=12, color=MUTED, bold=True)
    add_text(
        slide,
        "Both candidates preserve every manual non-skin voxel.\n\nApproach B then relabels a median "
        f"{median(volumes):.1f} cm3 ({median(percents):.1f}% of the A head map) as skin.",
        1.0,
        2.92,
        6.45,
        1.38,
        size=18,
        bold=True,
    )

    add_box(slide, 8.17, 2.25, 4.45, 2.55, WHITE, RULE)
    add_text(slide, "Design", 8.45, 2.55, 3.88, 0.25, size=12, color=MUTED, bold=True)
    add_text(
        slide,
        f"Eligible maps: {manifest['eligible_subject_count']}\nSample: n={len(rows)}\nRandom seed: {manifest['seed']}\nInputs: precomputed CHARM + manual labels\nEndpoints: merged maps and map-level QC",
        8.45,
        2.92,
        3.88,
        1.35,
        size=13,
    )
    add_text(slide, "Descriptive comparison only; no meshes or electric-field simulations were generated.", 0.72, 6.83, 11.6, 0.28, size=9, color=MUTED)


def add_method_slide(prs: Presentation, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(
        slide,
        "The manual overlay is identical; only the fallback base changes",
        "Manual labels are resampled to the 0.5 mm CHARM grid using nearest-neighbour interpolation; manual label 5 is excluded from both overlays.",
    )

    add_box(slide, 0.55, 1.35, 3.05, 4.95, WHITE, RULE)
    add_text(slide, "Shared inputs", 0.84, 1.68, 2.5, 0.28, size=15, bold=True)
    add_text(slide, "Precomputed CHARM\ntissue map", 0.92, 2.25, 2.25, 0.62, size=15, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "+", 1.75, 3.03, 0.6, 0.4, size=22, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "T1-registered manual\nsegmentation", 0.92, 3.58, 2.25, 0.62, size=15, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Keep manual labels >0 except skin label 5", 0.88, 4.65, 2.36, 0.62, size=11, color=MUTED, align=PP_ALIGN.CENTER)

    add_box(slide, 3.98, 1.35, 4.1, 4.95, WHITE, RULE)
    add_label_bar(slide, 4.32, 1.72, 3.42, "Approach A", BLUE)
    add_text(slide, "Full CHARM fallback", 4.42, 2.28, 3.2, 0.34, size=18, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "1. Start with every CHARM tissue label\n\n2. Paste manual non-skin labels on top\n\n3. Leave unclaimed voxels as CHARM labelled them", 4.42, 3.0, 3.2, 1.75, size=13)
    add_text(slide, "Fallback remains tissue-specific", 4.42, 5.25, 3.2, 0.35, size=12, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

    add_box(slide, 8.45, 1.35, 4.32, 4.95, WHITE, RULE)
    add_label_bar(slide, 8.8, 1.72, 3.62, "Approach B", ORANGE)
    add_text(slide, "Solid head as skin", 8.95, 2.28, 3.3, 0.34, size=18, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "1. Fill the CHARM head envelope\n\n2. Label the entire solid envelope as skin\n\n3. Paste manual non-skin labels on top", 8.95, 3.0, 3.3, 1.75, size=13)
    add_text(slide, "Every unclaimed head voxel becomes skin", 8.9, 5.25, 3.45, 0.45, size=12, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, page)


def add_sampling_slide(prs: Presentation, rows: list[dict[str, Any]], manifest: dict[str, Any], page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "The 10-subject cohort is reproducible and every source pair passed input checks", "The random sample is fixed by seed and drawn from all 175 subjects with both required maps.")

    add_box(slide, 0.62, 1.38, 4.0, 4.95, WHITE, RULE)
    add_text(slide, "Sampling", 0.92, 1.72, 3.3, 0.28, size=15, bold=True)
    add_text(slide, f"Eligible paired maps\n{manifest['eligible_subject_count']}", 0.95, 2.28, 3.25, 0.66, size=17, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"Random sample\nn = {len(rows)}", 0.95, 3.23, 3.25, 0.66, size=17, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"Seed\n{manifest['seed']}", 0.95, 4.18, 3.25, 0.66, size=17, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "The original three-subject trial is not reused in this sample.", 0.95, 5.48, 3.25, 0.42, size=10, color=MUTED, align=PP_ALIGN.CENTER)

    add_box(slide, 5.02, 1.38, 3.25, 4.95, WHITE, RULE)
    add_text(slide, "Selected IDs", 5.3, 1.72, 2.7, 0.28, size=15, bold=True)
    subjects = [str(row["subject"]) for row in rows]
    left = "\n".join(subjects[:5])
    right = "\n".join(subjects[5:])
    add_text(slide, left, 5.32, 2.28, 1.35, 2.8, size=11)
    add_text(slide, right, 6.68, 2.28, 1.35, 2.8, size=11)

    add_box(slide, 8.67, 1.38, 4.08, 4.95, WHITE, RULE)
    add_text(slide, "QC invariants", 8.97, 1.72, 3.45, 0.28, size=15, bold=True)
    mismatch_a = sum(int(row["candidate_a_manual_overlay_mismatches"]) for row in rows)
    mismatch_b = sum(int(row["candidate_b_manual_overlay_mismatches"]) for row in rows)
    outside = [float(row["manual_overlay_voxels_outside_charm_head"]) for row in rows]
    add_text(slide, f"{mismatch_a}", 9.05, 2.30, 1.2, 0.48, size=25, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "A overlay mismatches", 10.22, 2.44, 2.05, 0.28, size=11)
    add_text(slide, f"{mismatch_b}", 9.05, 3.22, 1.2, 0.48, size=25, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "B overlay mismatches", 10.22, 3.36, 2.05, 0.28, size=11)
    add_text(slide, "All output maps use the corresponding CHARM shape and affine.", 9.05, 4.25, 3.18, 0.62, size=11)
    add_text(slide, f"Manual non-skin voxels outside the CHARM head: {int(sum(outside)):,} total; retained by both candidates.", 9.05, 5.12, 3.18, 0.68, size=10, color=MUTED)
    add_footer(slide, page)


def add_summary_slide(prs: Presentation, rows: list[dict[str, Any]], chart: Path, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    volumes = [float(row["additional_skin_volume_cm3_b_minus_a"]) for row in rows]
    percents = [float(row["changed_percent_of_candidate_a_head"]) for row in rows]
    add_title(slide, "Approach B increases skin assignment in all 10 subjects", "Bars are descriptive values for the fixed random sample; no population-level inference is made.")
    add_picture_fit(slide, chart, 0.5, 1.18, 8.65, 5.65)
    add_box(slide, 9.42, 1.42, 3.3, 4.95, WHITE, RULE)
    add_text(slide, "Median added skin", 9.78, 1.8, 2.58, 0.25, size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"{median(volumes):.1f} cm3", 9.65, 2.18, 2.85, 0.5, size=25, color=ORANGE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"range {fmt_range(volumes)} cm3", 9.75, 2.77, 2.65, 0.25, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "Median head voxels changed", 9.65, 3.55, 2.85, 0.25, size=12, color=MUTED, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"{median(percents):.1f}%", 9.65, 3.94, 2.85, 0.5, size=25, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, f"range {fmt_range(percents)}%", 9.75, 4.53, 2.65, 0.25, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, "This is a tissue-class reassignment, not merely a cosmetic cleanup of the outer scalp boundary.", 9.78, 5.24, 2.58, 0.65, size=11, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, page)


def add_relabel_slide(prs: Presentation, chart: Path, volumes: dict[int, float], page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "Approach B converts tissue-specific CHARM fallback into skin", "The chart identifies the Approach A label at every voxel that receives a different label under Approach B, aggregated across the sample.")
    add_picture_fit(slide, chart, 0.55, 1.2, 8.5, 5.65)
    add_box(slide, 9.38, 1.38, 3.36, 5.05, WHITE, RULE)
    add_text(slide, "What the difference means", 9.72, 1.73, 2.7, 0.35, size=14, bold=True, align=PP_ALIGN.CENTER)
    ordered = sorted(volumes, key=volumes.get, reverse=True)
    top = ordered[:3]
    summary = "\n\n".join(f"{TISSUE_NAMES.get(label, f'Label {label}')}\n{volumes[label]:,.0f} cm3" for label in top)
    add_text(slide, summary, 9.75, 2.35, 2.62, 2.25, size=12, align=PP_ALIGN.CENTER)
    add_text(slide, "B does not copy only the CHARM scalp label. It labels the entire solid fallback envelope as skin before the manual overlay is restored.", 9.75, 4.95, 2.62, 0.9, size=11, bold=True, align=PP_ALIGN.CENTER)
    add_footer(slide, page)


def preview_path(root: Path, subject: str, approach: str) -> Path:
    if approach == "A":
        directory = "candidate_A_full_charm_fallback"
        suffix = "candidate_a_preview.png"
    elif approach == "B":
        directory = "candidate_B_solid_charm_head_skin_base"
        suffix = "candidate_b_preview.png"
    else:
        raise ValueError(f"Unknown approach: {approach}")
    return root / "subjects" / subject / directory / f"{subject}_{suffix}"


def add_segmentation_slide(prs: Presentation, root: Path, row: dict[str, Any], index: int, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    subject = str(row["subject"])
    volume = float(row["additional_skin_volume_cm3_b_minus_a"])
    percent = float(row["changed_percent_of_candidate_a_head"])
    add_title(
        slide,
        f"{subject} segmentation maps: B adds {volume:.1f} cm3 of skin",
        "Each large panel matches the previous trial: manual resampled, CHARM original, and merged final in three planes.",
    )
    add_label_bar(slide, 0.52, 1.08, 5.95, "Approach A: full CHARM fallback", BLUE)
    add_label_bar(slide, 6.87, 1.08, 5.95, "Approach B: solid head initialized as skin", ORANGE)
    add_picture_fit(slide, preview_path(root, subject, "A"), 0.52, 1.39, 5.95, 5.95)
    add_picture_fit(slide, preview_path(root, subject, "B"), 6.87, 1.39, 5.95, 5.95)
    add_footer(
        slide,
        page,
        f"{subject} | Subject {index} of 10 | B-A skin: +{volume:.1f} cm3 ({percent:.1f}% of the A head map)",
    )


def add_voxel_change_slide(prs: Presentation, root: Path, row: dict[str, Any], index: int, page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    subject = str(row["subject"])
    volume = float(row["changed_volume_cm3_a_vs_b"])
    percent = float(row["changed_percent_of_candidate_a_head"])
    add_title(
        slide,
        f"{subject}: B changes {volume:.1f} cm3 ({percent:.1f}% of the A head map)",
        "Views are selected where A and B differ most; the three difference maps are enlarged for close visual inspection.",
    )
    figure = root / "subjects" / subject / f"{subject}_voxel_change_large.png"
    add_picture_fit(slide, figure, 0.45, 1.12, 12.43, 5.78)
    add_footer(slide, page, f"{subject} | Subject {index} of 10 | Red voxels receive a different tissue label under Approach B")


def add_interpretation_slide(prs: Presentation, rows: list[dict[str, Any]], page: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "Approach A is the safer default; B is a gap-repair stress test", "The objective is to retain all manually corrected non-skin tissues while replacing the unreliable manual skin using CHARM-derived anatomy.")

    add_box(slide, 0.7, 1.42, 5.8, 4.95, WHITE, RULE)
    add_label_bar(slide, 1.04, 1.8, 5.12, "Approach A: tissue-conservative fallback", BLUE)
    add_text(slide, "What it preserves", 1.05, 2.42, 5.02, 0.25, size=12, color=MUTED, bold=True)
    add_text(slide, "Every positive manual non-skin label, plus CHARM's own tissue identity wherever the manual map contributes nothing or contributes removed skin.", 1.05, 2.78, 5.02, 0.95, size=13)
    add_text(slide, "Implication", 1.05, 4.05, 5.02, 0.25, size=12, color=MUTED, bold=True)
    add_text(slide, "It repairs coverage using anatomically specific CHARM labels and avoids assigning skin conductivity to fallback skull, CSF, muscle, blood, or bone.", 1.05, 4.42, 5.02, 0.95, size=13, bold=True)

    add_box(slide, 6.84, 1.42, 5.8, 4.95, WHITE, RULE)
    add_label_bar(slide, 7.18, 1.8, 5.12, "Approach B: aggressive solid-head repair", ORANGE)
    add_text(slide, "What it preserves", 7.2, 2.42, 5.0, 0.25, size=12, color=MUTED, bold=True)
    add_text(slide, "Every positive manual non-skin label, but all remaining volume inside the filled CHARM head envelope is assigned to skin.", 7.2, 2.78, 5.0, 0.86, size=13)
    add_text(slide, "Implication", 7.2, 4.05, 5.0, 0.25, size=12, color=MUTED, bold=True)
    add_text(slide, "It can close segmentation gaps, but this sample shows that the operation is broad. It should be chosen only if deliberate over-assignment to skin is acceptable.", 7.2, 4.42, 5.0, 0.95, size=13, bold=True)

    add_text(slide, "Decision boundary: map-level evidence favours A for anatomical fidelity. A final production choice should still be confirmed by remeshing representative failures and visually checking scalp continuity and tissue interfaces.", 0.82, 6.68, 11.7, 0.32, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    add_footer(slide, page)


def build(root: Path, output: Path) -> None:
    rows = load_rows(root)
    manifest = json.loads((root / "sample_manifest.json").read_text(encoding="utf-8"))
    if len(rows) != 10:
        raise ValueError(f"Expected 10 completed subjects, found {len(rows)}")

    summary_chart = make_summary_chart(root, rows)
    relabel_chart, relabel_volumes = make_relabel_chart(root, rows)

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    add_title_slide(prs, rows, manifest)
    page = 2
    add_method_slide(prs, page)
    page += 1
    add_sampling_slide(prs, rows, manifest, page)
    page += 1
    add_summary_slide(prs, rows, summary_chart, page)
    page += 1
    add_relabel_slide(prs, relabel_chart, relabel_volumes, page)
    page += 1
    for index, row in enumerate(rows, start=1):
        add_segmentation_slide(prs, root, row, index, page)
        page += 1
        add_voxel_change_slide(prs, root, row, index, page)
        page += 1
    add_interpretation_slide(prs, rows, page)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    print(f"[INFO] Wrote {output} ({len(prs.slides)} slides)")


def main() -> int:
    args = parse_args()
    output = args.output or args.root / OUT_NAME
    build(args.root, output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
