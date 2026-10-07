#!/usr/bin/env python3
"""Analyze final-132 TI dosimetry by recorded sex and age.

The participant is the inferential unit. Ten independent remeshing repeats are
averaged within participant and target, while their dispersion is retained.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import os
from pathlib import Path
from typing import Iterable

os.environ.setdefault("MPLCONFIGDIR", "/tmp/final132-ti-demographic-mpl")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import seaborn as sns
import statsmodels.formula.api as smf
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from scipy import stats
from statsmodels.stats.multitest import multipletests


SCRIPT_DIR = Path(__file__).resolve().parent
SCRIPTS_ROOT = SCRIPT_DIR.parent
DEFAULT_METRICS_ROOT = Path(
    "/home/boyan/sandbox/Jake_Data/CamCan-PostData/Anatomical-Parcelation/"
    "final_132_post_processing_analysis_results/runs"
)
DEFAULT_DEMOGRAPHICS = (
    SCRIPTS_ROOT
    / "output/spreadsheet/final_132_repeatability_balanced_10/"
    "eligible_132_demographics.csv"
)
DEFAULT_SUBJECTS = (
    SCRIPTS_ROOT
    / "CamCan_Experiment/cohort_pipeline/cohorts/final_132/subjects.txt"
)
DEFAULT_OUTPUT = SCRIPT_DIR / "outputs"

EXPECTED_SUBJECTS = 132
EXPECTED_TARGETS = 4
EXPECTED_REPEATS = 10
EXPECTED_RUNS = EXPECTED_SUBJECTS * EXPECTED_TARGETS * EXPECTED_REPEATS

ROI_ORDER = ["Left_Hippocampus", "Left_M1", "Right_DLPC", "Right_Thalamus"]
ROI_LABELS = {
    "Left_Hippocampus": "Left hippocampus",
    "Left_M1": "Left M1",
    "Right_DLPC": "Right DLPFC",
    "Right_Thalamus": "Right thalamus",
}
SEX_ORDER = ["FEMALE", "MALE"]
SEX_LABELS = {"FEMALE": "Female", "MALE": "Male"}
SEX_COLORS = {"FEMALE": "#C44E52", "MALE": "#4C72B0"}

METRIC_INFO = {
    "roi_mean_v_per_m": {
        "label": "Target ROI mean (V/m)",
        "unit": "V/m",
        "favorable": "higher",
        "primary": True,
    },
    "roi_p95_v_per_m": {
        "label": "Target ROI P95 (V/m)",
        "unit": "V/m",
        "favorable": "higher",
        "primary": False,
    },
    "target_fraction_ge_0_2": {
        "label": "Target ROI fraction >=0.2 V/m",
        "unit": "fraction",
        "favorable": "higher",
        "primary": False,
    },
    "target_to_neighbor_mean_ratio": {
        "label": "Target-to-neighbour mean ratio",
        "unit": "ratio",
        "favorable": "higher",
        "primary": False,
    },
    "offtarget_fraction_ge_0_2": {
        "label": "Off-target fraction >=0.2 V/m",
        "unit": "fraction",
        "favorable": "lower",
        "primary": False,
    },
}
METRICS = list(METRIC_INFO)
PRIMARY_METRIC = "roi_mean_v_per_m"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--metrics-root", type=Path, default=DEFAULT_METRICS_ROOT)
    parser.add_argument("--demographics", type=Path, default=DEFAULT_DEMOGRAPHICS)
    parser.add_argument("--subjects", type=Path, default=DEFAULT_SUBJECTS)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_subjects(path: Path) -> list[str]:
    subjects = [line.strip() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if len(subjects) != EXPECTED_SUBJECTS or len(set(subjects)) != EXPECTED_SUBJECTS:
        raise ValueError(
            f"Expected {EXPECTED_SUBJECTS} unique subjects in {path}; found {len(subjects)} rows "
            f"and {len(set(subjects))} unique IDs."
        )
    return subjects


def read_demographics(path: Path, subjects: list[str]) -> pd.DataFrame:
    frame = pd.read_csv(path)
    required = {
        "subject",
        "age_years_precise",
        "age_years_workbook",
        "sex",
        "age_band_index",
        "age_band",
    }
    missing_columns = required - set(frame.columns)
    if missing_columns:
        raise ValueError(f"Demographics are missing columns: {sorted(missing_columns)}")
    frame = frame.loc[frame["subject"].isin(subjects), list(required)].copy()
    if len(frame) != EXPECTED_SUBJECTS or frame["subject"].nunique() != EXPECTED_SUBJECTS:
        raise ValueError("Demographic join does not contain exactly 132 unique final-cohort subjects.")
    if set(frame["subject"]) != set(subjects):
        raise ValueError("Demographic IDs differ from the authoritative final_132 subject list.")
    if frame[["age_years_precise", "sex", "age_band_index", "age_band"]].isna().any().any():
        raise ValueError("Demographic data contain missing analysis fields.")
    unexpected_sex = sorted(set(frame["sex"]) - set(SEX_ORDER))
    if unexpected_sex:
        raise ValueError(f"Unexpected recorded-sex values: {unexpected_sex}")
    frame["age_band_index"] = frame["age_band_index"].astype(int)
    frame = frame.sort_values(["age_years_precise", "subject"]).reset_index(drop=True)
    return frame


def safe_ratio(numerator: float, denominator: float) -> float:
    if not np.isfinite(numerator) or not np.isfinite(denominator) or denominator <= 0:
        return float("nan")
    return float(numerator / denominator)


def extract_run_row(path: Path) -> dict[str, object]:
    data = json.loads(path.read_text(encoding="utf-8"))
    dataset_name = path.parents[3].name
    if "_Data_" not in dataset_name:
        raise ValueError(f"Cannot parse target/repeat from {path}")
    target, repeat = dataset_name.rsplit("_Data_", 1)
    roi_keys = list(data.get("rois", {}))
    if len(roi_keys) != 1:
        raise ValueError(f"Expected one target ROI in {path}; found {roi_keys}")
    roi_values = data["rois"][roi_keys[0]]
    extended = data.get("extended_metrics", {})
    meta = data.get("subject_metrics_meta", {})
    extended_meta = data.get("extended_metrics_meta", {})

    roi_voxels = float(roi_values["roi_voxels"])
    target_above = float(roi_values["focality_in_roi_voxels_gt_threshold"])
    whole_voxels = float(data["whole_brain_voxels"])
    whole_above = float(extended["whole_brain_coverage_voxels_ge_threshold"])
    off_voxels = whole_voxels - roi_voxels
    off_above = whole_above - target_above
    if off_above < -1e-9 or off_above > off_voxels + 1e-9:
        raise ValueError(f"Invalid off-target threshold counts in {path}")

    return {
        "subject": data["subject"],
        "target": target,
        "target_label": ROI_LABELS.get(target, target),
        "repeat": int(repeat),
        "source_path": str(path),
        "status": meta.get("status"),
        "extended_status": extended_meta.get("status"),
        "target_roi_key": roi_keys[0],
        "roi_mean_v_per_m": float(extended["roi_mean"]),
        "roi_p95_v_per_m": float(roi_values["roi_percentile_value"]),
        "target_fraction_ge_0_2": safe_ratio(target_above, roi_voxels),
        "target_to_neighbor_mean_ratio": safe_ratio(
            float(extended["roi_mean"]), float(extended["neighbor_mean_of_means"])
        ),
        "offtarget_fraction_ge_0_2": safe_ratio(off_above, off_voxels),
        "roi_voxels": int(roi_voxels),
        "whole_brain_voxels": int(whole_voxels),
        "target_voxels_ge_0_2": int(target_above),
        "whole_brain_voxels_ge_0_2": int(whole_above),
        "neighbor_mean_v_per_m": float(extended["neighbor_mean_of_means"]),
    }


def collect_run_metrics(metrics_root: Path, subjects: list[str]) -> pd.DataFrame:
    paths = sorted(metrics_root.rglob("subject_metrics.json"))
    if len(paths) != EXPECTED_RUNS:
        raise ValueError(f"Expected {EXPECTED_RUNS} subject metric files; found {len(paths)}.")
    frame = pd.DataFrame(extract_run_row(path) for path in paths)
    if set(frame["subject"]) != set(subjects):
        raise ValueError("Metric subject IDs differ from the final_132 cohort.")
    if set(frame["target"]) != set(ROI_ORDER):
        raise ValueError(f"Unexpected target set: {sorted(set(frame['target']))}")
    if set(frame["repeat"]) != set(range(1, EXPECTED_REPEATS + 1)):
        raise ValueError(f"Unexpected repeat set: {sorted(set(frame['repeat']))}")
    expected_cell_count = EXPECTED_SUBJECTS * EXPECTED_TARGETS
    cell_counts = frame.groupby(["subject", "target"], observed=True).size()
    if len(cell_counts) != expected_cell_count or not (cell_counts == EXPECTED_REPEATS).all():
        raise ValueError("Each participant-target cell must contain exactly ten repeats.")
    if not (frame["status"] == "complete").all():
        raise ValueError("At least one subject metric file is not marked complete.")
    if not (frame["extended_status"] == "complete").all():
        raise ValueError("At least one extended-metric record is not marked complete.")
    if frame[METRICS].isna().any().any() or not np.isfinite(frame[METRICS].to_numpy()).all():
        raise ValueError("Primary/supporting metrics contain missing or non-finite values.")
    for fraction_metric in ["target_fraction_ge_0_2", "offtarget_fraction_ge_0_2"]:
        if not frame[fraction_metric].between(0, 1).all():
            raise ValueError(f"{fraction_metric} contains values outside [0, 1].")
    return frame.sort_values(["target", "subject", "repeat"]).reset_index(drop=True)


def aggregate_subject_targets(run_frame: pd.DataFrame, demographics: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for (subject, target), group in run_frame.groupby(["subject", "target"], observed=True):
        row: dict[str, object] = {
            "subject": subject,
            "target": target,
            "target_label": ROI_LABELS[target],
            "n_repeats": len(group),
        }
        for metric in METRICS:
            values = group[metric].to_numpy(dtype=float)
            mean_value = float(np.mean(values))
            sd_value = float(np.std(values, ddof=1))
            row[metric] = mean_value
            row[f"{metric}_repeat_sd"] = sd_value
            row[f"{metric}_repeat_cv_percent"] = (
                float(100.0 * sd_value / mean_value) if mean_value != 0 else float("nan")
            )
        rows.append(row)
    subject_frame = pd.DataFrame(rows).merge(demographics, on="subject", how="left", validate="many_to_one")
    if subject_frame["sex"].isna().any():
        raise ValueError("Subject-target metrics failed to join to demographics.")
    subject_frame["male"] = (subject_frame["sex"] == "MALE").astype(int)
    cohort_mean_age = float(demographics["age_years_precise"].mean())
    subject_frame["age_per_10"] = (subject_frame["age_years_precise"] - cohort_mean_age) / 10.0
    subject_frame["sex_label"] = subject_frame["sex"].map(SEX_LABELS)
    band_order = (
        demographics[["age_band_index", "age_band"]]
        .drop_duplicates()
        .sort_values("age_band_index")["age_band"]
        .tolist()
    )
    subject_frame["age_band"] = pd.Categorical(
        subject_frame["age_band"], categories=band_order, ordered=True
    )
    expected_rows = EXPECTED_SUBJECTS * EXPECTED_TARGETS
    if len(subject_frame) != expected_rows or not (subject_frame["n_repeats"] == EXPECTED_REPEATS).all():
        raise ValueError("Subject aggregation did not yield 132 x 4 complete rows.")
    return subject_frame.sort_values(["target", "subject"]).reset_index(drop=True)


def mean_ci(values: Iterable[float]) -> tuple[int, float, float, float, float]:
    array = np.asarray(list(values), dtype=float)
    array = array[np.isfinite(array)]
    n = len(array)
    if n == 0:
        return 0, float("nan"), float("nan"), float("nan"), float("nan")
    mean_value = float(np.mean(array))
    sd_value = float(np.std(array, ddof=1)) if n > 1 else float("nan")
    if n > 1:
        sem = sd_value / math.sqrt(n)
        margin = float(stats.t.ppf(0.975, n - 1) * sem)
        return n, mean_value, sd_value, mean_value - margin, mean_value + margin
    return n, mean_value, sd_value, float("nan"), float("nan")


def stratified_summary(
    frame: pd.DataFrame, group_columns: list[str], stratification: str
) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    grouping = ["target", *group_columns]
    for keys, group in frame.groupby(grouping, observed=True, sort=False):
        if not isinstance(keys, tuple):
            keys = (keys,)
        identifiers = dict(zip(grouping, keys))
        for metric in METRICS:
            n, mean_value, sd_value, ci_low, ci_high = mean_ci(group[metric])
            rows.append(
                {
                    "stratification": stratification,
                    **identifiers,
                    "target_label": ROI_LABELS[str(identifiers["target"])],
                    "metric": metric,
                    "metric_label": METRIC_INFO[metric]["label"],
                    "favorable_direction": METRIC_INFO[metric]["favorable"],
                    "n_participants": n,
                    "mean": mean_value,
                    "sd": sd_value,
                    "ci95_low": ci_low,
                    "ci95_high": ci_high,
                    "median": float(group[metric].median()),
                    "q1": float(group[metric].quantile(0.25)),
                    "q3": float(group[metric].quantile(0.75)),
                }
            )
    return pd.DataFrame(rows)


def extract_model_effect(
    fit,
    *,
    weights: dict[str, float],
) -> tuple[float, float, float, float, float]:
    names = list(fit.params.index)
    vector = np.asarray([weights.get(name, 0.0) for name in names], dtype=float)
    estimate = float(vector @ fit.params.to_numpy())
    covariance = fit.cov_params().to_numpy()
    variance = float(vector @ covariance @ vector)
    standard_error = math.sqrt(max(variance, 0.0))
    if standard_error == 0:
        p_value = float("nan")
        return estimate, standard_error, float("nan"), float("nan"), p_value
    z_value = estimate / standard_error
    p_value = float(2.0 * stats.norm.sf(abs(z_value)))
    return (
        estimate,
        standard_error,
        estimate - 1.959963984540054 * standard_error,
        estimate + 1.959963984540054 * standard_error,
        p_value,
    )


def fit_models(frame: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    model_specs = [
        ("sex_only", "{metric} ~ male", {"male": 1.0}, "Male - female"),
        ("age_only", "{metric} ~ age_per_10", {"age_per_10": 1.0}, "Per 10 years"),
        (
            "sex_age_combined",
            "{metric} ~ age_per_10 * male",
            {"male": 1.0},
            "Male - female at cohort mean age",
        ),
        (
            "sex_age_combined",
            "{metric} ~ age_per_10 * male",
            {"age_per_10": 1.0},
            "Age slope per 10 years in female group",
        ),
        (
            "sex_age_combined",
            "{metric} ~ age_per_10 * male",
            {"age_per_10": 1.0, "age_per_10:male": 1.0},
            "Age slope per 10 years in male group",
        ),
        (
            "sex_age_combined",
            "{metric} ~ age_per_10 * male",
            {"age_per_10:male": 1.0},
            "Age-by-sex interaction per 10 years",
        ),
    ]
    for target in ROI_ORDER:
        target_frame = frame.loc[frame["target"] == target].copy()
        for metric in METRICS:
            for model_name, formula_template, weights, effect in model_specs:
                formula = formula_template.format(metric=metric)
                fit = smf.ols(formula, data=target_frame).fit(cov_type="HC3")
                estimate, standard_error, ci_low, ci_high, p_value = extract_model_effect(
                    fit, weights=weights
                )
                rows.append(
                    {
                        "target": target,
                        "target_label": ROI_LABELS[target],
                        "metric": metric,
                        "metric_label": METRIC_INFO[metric]["label"],
                        "favorable_direction": METRIC_INFO[metric]["favorable"],
                        "primary_endpoint": METRIC_INFO[metric]["primary"],
                        "model": model_name,
                        "effect": effect,
                        "estimate": estimate,
                        "robust_se": standard_error,
                        "ci95_low": ci_low,
                        "ci95_high": ci_high,
                        "p_value": p_value,
                        "n_participants": int(fit.nobs),
                        "r_squared": float(fit.rsquared),
                    }
                )

            age_groups = [
                group[metric].to_numpy(dtype=float)
                for _, group in target_frame.groupby("age_band", observed=True, sort=True)
            ]
            statistic, p_value = stats.kruskal(*age_groups)
            rows.append(
                {
                    "target": target,
                    "target_label": ROI_LABELS[target],
                    "metric": metric,
                    "metric_label": METRIC_INFO[metric]["label"],
                    "favorable_direction": METRIC_INFO[metric]["favorable"],
                    "primary_endpoint": METRIC_INFO[metric]["primary"],
                    "model": "age_band_omnibus",
                    "effect": "Five-band Kruskal-Wallis omnibus",
                    "estimate": float(statistic),
                    "robust_se": float("nan"),
                    "ci95_low": float("nan"),
                    "ci95_high": float("nan"),
                    "p_value": float(p_value),
                    "n_participants": len(target_frame),
                    "r_squared": float("nan"),
                }
            )

    results = pd.DataFrame(rows)
    results["p_fdr"] = np.nan
    family_columns = ["metric", "model", "effect"]
    for _, indices in results.groupby(family_columns, dropna=False).groups.items():
        index_list = list(indices)
        p_values = results.loc[index_list, "p_value"].to_numpy(dtype=float)
        finite = np.isfinite(p_values)
        adjusted = np.full(len(p_values), np.nan)
        if finite.any():
            adjusted[finite] = multipletests(p_values[finite], method="fdr_bh")[1]
        results.loc[index_list, "p_fdr"] = adjusted
    return results.sort_values(["metric", "model", "effect", "target"]).reset_index(drop=True)


def build_rankings(
    summaries: list[tuple[str, pd.DataFrame, list[str]]]
) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for label, summary, group_columns in summaries:
        ranked = summary.copy()
        ranked["group"] = ranked[group_columns].astype(str).agg(" | ".join, axis=1)
        ascending = ranked["favorable_direction"].eq("lower")
        ranked["favorable_rank"] = np.nan
        for _, indices in ranked.groupby(["target", "metric"], observed=True).groups.items():
            index_list = list(indices)
            is_ascending = bool(ascending.loc[index_list].iloc[0])
            ranked.loc[index_list, "favorable_rank"] = ranked.loc[index_list, "mean"].rank(
                method="min", ascending=is_ascending
            )
        ranked["ranking_type"] = label
        frames.append(
            ranked[
                [
                    "ranking_type",
                    "target",
                    "target_label",
                    "metric",
                    "metric_label",
                    "favorable_direction",
                    "group",
                    "n_participants",
                    "mean",
                    "ci95_low",
                    "ci95_high",
                    "favorable_rank",
                ]
            ]
        )
    return pd.concat(frames, ignore_index=True)


def standardized_combined_groups(frame: pd.DataFrame) -> pd.DataFrame:
    working = frame.copy()
    working["roi_mean_within_target_z"] = working.groupby("target", observed=True)[PRIMARY_METRIC].transform(
        lambda series: (series - series.mean()) / series.std(ddof=1)
    )
    participant_scores = (
        working.groupby(
            ["subject", "sex", "sex_label", "age_band_index", "age_band"], observed=True
        )["roi_mean_within_target_z"]
        .mean()
        .reset_index(name="mean_target_standardized_roi_mean")
    )
    rows = []
    for keys, group in participant_scores.groupby(
        ["age_band_index", "age_band", "sex", "sex_label"], observed=True, sort=True
    ):
        n, mean_value, sd_value, ci_low, ci_high = mean_ci(
            group["mean_target_standardized_roi_mean"]
        )
        rows.append(
            {
                "age_band_index": keys[0],
                "age_band": keys[1],
                "sex": keys[2],
                "sex_label": keys[3],
                "n_participants": n,
                "mean_standardized_roi_mean": mean_value,
                "sd": sd_value,
                "ci95_low": ci_low,
                "ci95_high": ci_high,
            }
        )
    result = pd.DataFrame(rows)
    result["descriptive_rank"] = result["mean_standardized_roi_mean"].rank(
        ascending=False, method="min"
    )
    return result.sort_values("descriptive_rank").reset_index(drop=True)


def add_panel_label(ax, label: str) -> None:
    ax.text(
        -0.12,
        1.08,
        label,
        transform=ax.transAxes,
        fontsize=12,
        fontweight="bold",
        va="top",
    )


def save_figure(fig, output_base: Path) -> None:
    fig.savefig(output_base.with_suffix(".png"), dpi=300, bbox_inches="tight")
    fig.savefig(output_base.with_suffix(".pdf"), bbox_inches="tight")
    plt.close(fig)


def plot_sex(frame: pd.DataFrame, figures_dir: Path) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(11.5, 8.5), constrained_layout=True)
    rng = np.random.default_rng(132)
    for index, (target, ax) in enumerate(zip(ROI_ORDER, axes.flat)):
        data = frame.loc[frame["target"] == target].copy()
        sns.boxplot(
            data=data,
            x="sex_label",
            y=PRIMARY_METRIC,
            order=["Female", "Male"],
            hue="sex_label",
            palette={"Female": SEX_COLORS["FEMALE"], "Male": SEX_COLORS["MALE"]},
            showfliers=False,
            width=0.5,
            legend=False,
            ax=ax,
        )
        for x_value, sex_label in enumerate(["Female", "Male"]):
            values = data.loc[data["sex_label"] == sex_label, PRIMARY_METRIC].to_numpy()
            jitter = rng.uniform(-0.16, 0.16, len(values))
            ax.scatter(
                np.full(len(values), x_value) + jitter,
                values,
                s=14,
                alpha=0.55,
                color="#222222",
                linewidth=0,
                zorder=3,
            )
        ax.set_title(ROI_LABELS[target])
        ax.set_xlabel("Recorded sex")
        ax.set_ylabel("Participant-mean target ROI field (V/m)")
        add_panel_label(ax, chr(ord("A") + index))
    fig.suptitle(
        "Final-132 simulated TI target exposure by recorded sex\n"
        "Each point is one participant averaged across 10 independent remeshings",
        fontsize=14,
        fontweight="bold",
    )
    save_figure(fig, figures_dir / "sex_stratified_roi_mean")


def plot_age(frame: pd.DataFrame, figures_dir: Path) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(13.5, 8.5), constrained_layout=True)
    band_order = list(frame["age_band"].cat.categories)
    palette = sns.color_palette("viridis", n_colors=len(band_order))
    rng = np.random.default_rng(133)
    for index, (target, ax) in enumerate(zip(ROI_ORDER, axes.flat)):
        data = frame.loc[frame["target"] == target].copy()
        sns.boxplot(
            data=data,
            x="age_band",
            y=PRIMARY_METRIC,
            order=band_order,
            hue="age_band",
            palette=palette,
            showfliers=False,
            width=0.58,
            legend=False,
            ax=ax,
        )
        for x_value, band in enumerate(band_order):
            values = data.loc[data["age_band"] == band, PRIMARY_METRIC].to_numpy()
            jitter = rng.uniform(-0.18, 0.18, len(values))
            ax.scatter(
                np.full(len(values), x_value) + jitter,
                values,
                s=11,
                alpha=0.48,
                color="#222222",
                linewidth=0,
                zorder=3,
            )
        ax.set_title(ROI_LABELS[target])
        ax.set_xlabel("Age band (years)")
        ax.set_ylabel("Participant-mean target ROI field (V/m)")
        ax.tick_params(axis="x", rotation=22)
        add_panel_label(ax, chr(ord("A") + index))
    fig.suptitle(
        "Final-132 simulated TI target exposure by prespecified age band\n"
        "Each point is one participant averaged across 10 independent remeshings",
        fontsize=14,
        fontweight="bold",
    )
    save_figure(fig, figures_dir / "age_stratified_roi_mean")


def plot_combined_strata(
    combined_summary: pd.DataFrame, figures_dir: Path
) -> None:
    data = combined_summary.loc[combined_summary["metric"] == PRIMARY_METRIC].copy()
    band_order = list(data.sort_values("age_band_index")["age_band"].drop_duplicates())
    fig, axes = plt.subplots(2, 2, figsize=(13.5, 8.5), constrained_layout=True)
    for index, (target, ax) in enumerate(zip(ROI_ORDER, axes.flat)):
        target_data = data.loc[data["target"] == target]
        for sex in SEX_ORDER:
            group = target_data.loc[target_data["sex"] == sex].sort_values("age_band_index")
            x_values = np.arange(len(group))
            lower = group["mean"].to_numpy() - group["ci95_low"].to_numpy()
            upper = group["ci95_high"].to_numpy() - group["mean"].to_numpy()
            ax.errorbar(
                x_values,
                group["mean"],
                yerr=np.vstack([lower, upper]),
                marker="o",
                linewidth=2,
                capsize=3,
                color=SEX_COLORS[sex],
                label=SEX_LABELS[sex],
            )
        ax.set_xticks(np.arange(len(band_order)), band_order, rotation=22)
        ax.set_title(ROI_LABELS[target])
        ax.set_xlabel("Age band (years)")
        ax.set_ylabel("Mean target ROI field (V/m), 95% CI")
        add_panel_label(ax, chr(ord("A") + index))
        if index == 0:
            ax.legend(title="Recorded sex", frameon=False)
    fig.suptitle(
        "Combined age-by-recorded-sex stratification of simulated TI target exposure\n"
        "Intervals describe participant means within each demographic cell",
        fontsize=14,
        fontweight="bold",
    )
    save_figure(fig, figures_dir / "sex_age_combined_stratified_roi_mean")


def plot_combined_continuous(frame: pd.DataFrame, figures_dir: Path) -> None:
    fig, axes = plt.subplots(2, 2, figsize=(12.5, 8.5), constrained_layout=True)
    for index, (target, ax) in enumerate(zip(ROI_ORDER, axes.flat)):
        data = frame.loc[frame["target"] == target].copy()
        for sex in SEX_ORDER:
            group = data.loc[data["sex"] == sex]
            ax.scatter(
                group["age_years_precise"],
                group[PRIMARY_METRIC],
                s=18,
                alpha=0.55,
                color=SEX_COLORS[sex],
                label=SEX_LABELS[sex],
            )
        fit = smf.ols(f"{PRIMARY_METRIC} ~ age_per_10 * male", data=data).fit(cov_type="HC3")
        age_grid = np.linspace(data["age_years_precise"].min(), data["age_years_precise"].max(), 160)
        mean_age = float(data["age_years_precise"].mean())
        for sex in SEX_ORDER:
            prediction_data = pd.DataFrame(
                {
                    "age_per_10": (age_grid - mean_age) / 10.0,
                    "male": int(sex == "MALE"),
                }
            )
            prediction = fit.get_prediction(prediction_data).summary_frame(alpha=0.05)
            ax.plot(age_grid, prediction["mean"], color=SEX_COLORS[sex], linewidth=2)
            ax.fill_between(
                age_grid,
                prediction["mean_ci_lower"].to_numpy(),
                prediction["mean_ci_upper"].to_numpy(),
                color=SEX_COLORS[sex],
                alpha=0.16,
                linewidth=0,
            )
        ax.set_title(ROI_LABELS[target])
        ax.set_xlabel("Age (years)")
        ax.set_ylabel("Participant-mean target ROI field (V/m)")
        add_panel_label(ax, chr(ord("A") + index))
        if index == 0:
            ax.legend(title="Recorded sex", frameon=False)
    fig.suptitle(
        "Continuous age-by-recorded-sex models of simulated TI target exposure\n"
        "Lines are robust linear fits; shading is the 95% confidence interval",
        fontsize=14,
        fontweight="bold",
    )
    save_figure(fig, figures_dir / "sex_age_combined_continuous_roi_mean")


def write_excel(
    path: Path,
    *,
    inventory: pd.DataFrame,
    demographics: pd.DataFrame,
    subject_frame: pd.DataFrame,
    sex_summary: pd.DataFrame,
    age_summary: pd.DataFrame,
    combined_summary: pd.DataFrame,
    models: pd.DataFrame,
    rankings: pd.DataFrame,
    standardized_groups: pd.DataFrame,
) -> None:
    sheets = {
        "Inventory": inventory,
        "Demographics": demographics,
        "Subject target metrics": subject_frame,
        "Sex summary": sex_summary,
        "Age summary": age_summary,
        "Sex age summary": combined_summary,
        "Models": models,
        "Rankings": rankings,
        "Combined standardized": standardized_groups,
    }
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)

    workbook = load_workbook(path)
    navy = "17324D"
    teal = "177E89"
    white = "FFFFFF"
    light_teal = "DDEFF1"
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = PatternFill("solid", fgColor=teal)
            cell.font = Font(color=white, bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 32
        for column_cells in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:250]]
            width = min(max(max((len(value) for value in values), default=0) + 2, 10), 38)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in range(2, worksheet.max_row + 1):
            if row % 2 == 0:
                for cell in worksheet[row]:
                    cell.fill = PatternFill("solid", fgColor=light_teal)
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000"
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
    workbook["Inventory"]["A1"].fill = PatternFill("solid", fgColor=navy)
    workbook.save(path)


def format_p(value: float) -> str:
    if not np.isfinite(value):
        return "NA"
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def model_line(row: pd.Series) -> str:
    return (
        f"{row['target_label']}: estimate {row['estimate']:.4f} V/m "
        f"(95% CI {row['ci95_low']:.4f} to {row['ci95_high']:.4f}; "
        f"FDR p={format_p(float(row['p_fdr']))})"
    )


def write_report(
    path: Path,
    demographics: pd.DataFrame,
    sex_summary: pd.DataFrame,
    age_summary: pd.DataFrame,
    combined_summary: pd.DataFrame,
    models: pd.DataFrame,
    standardized_groups: pd.DataFrame,
) -> None:
    primary = models.loc[models["metric"] == PRIMARY_METRIC].copy()
    sex_rows = primary.loc[
        (primary["model"] == "sex_age_combined")
        & (primary["effect"] == "Male - female at cohort mean age")
    ].sort_values("target")
    age_rows = primary.loc[
        (primary["model"] == "age_only") & (primary["effect"] == "Per 10 years")
    ].sort_values("target")
    interaction_rows = primary.loc[
        (primary["model"] == "sex_age_combined")
        & (primary["effect"] == "Age-by-sex interaction per 10 years")
    ].sort_values("target")
    best = standardized_groups.iloc[0]

    sex_primary = sex_summary.loc[sex_summary["metric"] == PRIMARY_METRIC]
    age_primary = age_summary.loc[age_summary["metric"] == PRIMARY_METRIC]
    sex_descriptive_lines = []
    age_descriptive_lines = []
    for target in ROI_ORDER:
        target_sex = sex_primary.loc[sex_primary["target"] == target].set_index("sex")
        female_mean = float(target_sex.loc["FEMALE", "mean"])
        male_mean = float(target_sex.loc["MALE", "mean"])
        female_vs_male_percent = 100.0 * (female_mean - male_mean) / male_mean
        sex_descriptive_lines.append(
            f"- {ROI_LABELS[target]}: female {female_mean:.3f} V/m, male {male_mean:.3f} V/m "
            f"({female_vs_male_percent:+.1f}% female relative to male)."
        )

        target_age = age_primary.loc[age_primary["target"] == target].sort_values("age_band_index")
        youngest = target_age.iloc[0]
        oldest = target_age.iloc[-1]
        youngest_vs_oldest_percent = 100.0 * (youngest["mean"] - oldest["mean"]) / oldest["mean"]
        age_descriptive_lines.append(
            f"- {ROI_LABELS[target]}: {youngest['mean']:.3f} V/m in the youngest band versus "
            f"{oldest['mean']:.3f} V/m in the oldest ({youngest_vs_oldest_percent:+.1f}% youngest "
            "relative to oldest)."
        )

    lines = [
        "# Analysis report: final-132 TI dosimetry by recorded sex and age",
        "",
        "## Bottom line",
        "",
        (
            "This analysis identifies demographic differences in **simulated target exposure**, not "
            "clinical benefit. The primary endpoint is target-ROI mean field after averaging ten "
            "independent remeshings for each participant and target."
        ),
        "",
        (
            f"The cohort contains {len(demographics)} participants: "
            f"{int((demographics['sex'] == 'FEMALE').sum())} female and "
            f"{int((demographics['sex'] == 'MALE').sum())} male, aged "
            f"{demographics['age_years_precise'].min():.2f}-{demographics['age_years_precise'].max():.2f} years."
        ),
        "",
        (
            "The most defensible answer is target-specific: **younger participants had higher simulated "
            "on-target field across all four targets**. The female group also had higher on-target field "
            "for the left hippocampus and right thalamus after age adjustment, but there was no detected "
            "recorded-sex difference for left M1 or right DLPFC. No target showed an age-by-recorded-sex "
            "interaction after FDR correction, so the data do not support a claim that the age trend differs "
            "between female and male participants."
        ),
        "",
        (
            "The descriptively highest combined age-by-sex cell for the cross-target standardized "
            f"ROI-mean score was {best['sex_label'].lower()}, age {best['age_band']} years "
            f"(n={int(best['n_participants'])}, mean z={best['mean_standardized_roi_mean']:.2f}). "
            "This is a descriptive ranking, not a clinical responder classification or an adjusted "
            "causal estimate."
        ),
        "",
        "## Adjusted primary-endpoint models",
        "",
        "Recorded-sex contrasts at the cohort mean age (male minus female):",
        "",
    ]
    lines.extend(f"- {model_line(row)}" for _, row in sex_rows.iterrows())
    lines.extend(
        [
            "",
            "Pooled age slopes (change per 10 years before adding an interaction):",
            "",
        ]
    )
    lines.extend(f"- {model_line(row)}" for _, row in age_rows.iterrows())
    lines.extend(
        [
            "",
            "Age-by-recorded-sex interaction terms:",
            "",
        ]
    )
    lines.extend(f"- {model_line(row)}" for _, row in interaction_rows.iterrows())
    lines.extend(
        [
            "",
            "## Descriptive group contrasts",
            "",
            "Recorded-sex means (unadjusted):",
            "",
        ]
    )
    lines.extend(sex_descriptive_lines)
    lines.extend(
        [
            "",
            "Youngest versus oldest prespecified age bands (unadjusted):",
            "",
        ]
    )
    lines.extend(age_descriptive_lines)
    lines.extend(
        [
            "",
            "## Dose-focality trade-off",
            "",
            (
                "Higher target field was not an unqualified focality advantage. For the left hippocampus "
                "and right thalamus, the female group had higher target mean, target P95, and target fraction "
                "at or above 0.2 V/m, but also higher off-target coverage. Target-to-neighbour mean-field "
                "ratio did not differ by recorded sex at any target after adjustment. This pattern is most "
                "consistent with an absolute-exposure difference, not superior spatial selectivity."
            ),
            "",
            (
                "Increasing age was associated with both lower on-target exposure and lower off-target "
                "coverage across all targets. Target-to-neighbour ratio declined with age for left M1 and "
                "right DLPFC, but not for the left hippocampus or right thalamus after FDR correction. Thus, "
                "younger participants received more simulated target field, whereas older participants also "
                "had less field above 0.2 V/m outside the target."
            ),
            "",
            (
                "The 0.2 V/m cutoff is an analysis threshold used by the existing post-processing pipeline; "
                "it is not treated here as a validated therapeutic or clinical-response threshold."
            ),
            "",
            "## Evidence rule",
            "",
            (
                "Primary conclusions require an FDR-adjusted model result plus a coherent effect size and "
                "confidence interval. Statistical results are interpreted target by target. A non-significant "
                "contrast is not proof of equivalence, and a visually highest demographic cell is not called "
                "a responder group."
            ),
            "",
            "## Supporting outcomes",
            "",
            "The workbook and CSV tables report target P95, target fraction at or above 0.2 V/m, "
            "target-to-neighbour mean ratio, and off-target fraction at or above 0.2 V/m. These are "
            "kept separate rather than collapsed into a single clinical-benefit score.",
            "",
            "## Design and limitations",
            "",
            "- Participant means are the inferential observations; 10 remeshings are computational repeats.",
            "- Recorded sex is available, not gender identity. The analysis therefore does not make gender claims.",
            "- Age bands are prespecified equal-width bands, but continuous age models are primary for trend inference.",
            "- The final-132 cohort was selected through simulation-model QC and is not a population-random sample.",
            "- Combined age-by-sex cells are small and uneven (5-24 participants), so visual rank order is uncertain.",
            "- Simulated electric-field magnitude/focality cannot establish symptoms, efficacy, or clinical response.",
            "- No HPC computation was required because all 5,280 completed post-processing records were local.",
            "",
            "## Figures",
            "",
            "- `figures/sex_stratified_roi_mean.png`: recorded-sex stratification.",
            "- `figures/age_stratified_roi_mean.png`: five-band age stratification.",
            "- `figures/sex_age_combined_stratified_roi_mean.png`: combined stratified plot.",
            "- `figures/sex_age_combined_continuous_roi_mean.png`: continuous age-by-sex model plot.",
            "",
            "## Reproducibility",
            "",
            "Run `analyze_final132_demographics.py` from the repository root using the `simnibs_post` "
            "environment. Input paths, hashes, counts, metric definitions, and output tables are recorded "
            "in `analysis_inventory.json` and `analysis_tables.xlsx`.",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    args = parse_args()
    output_dir = args.output_dir.resolve()
    figures_dir = output_dir / "figures"
    output_dir.mkdir(parents=True, exist_ok=True)
    figures_dir.mkdir(parents=True, exist_ok=True)

    subjects = read_subjects(args.subjects.resolve())
    demographics = read_demographics(args.demographics.resolve(), subjects)
    run_frame = collect_run_metrics(args.metrics_root.resolve(), subjects)
    subject_frame = aggregate_subject_targets(run_frame, demographics)

    sex_summary = stratified_summary(subject_frame, ["sex", "sex_label"], "recorded_sex")
    age_summary = stratified_summary(
        subject_frame, ["age_band_index", "age_band"], "age_band"
    )
    combined_summary = stratified_summary(
        subject_frame,
        ["age_band_index", "age_band", "sex", "sex_label"],
        "age_band_by_recorded_sex",
    )
    models = fit_models(subject_frame)
    rankings = build_rankings(
        [
            ("recorded_sex", sex_summary, ["sex_label"]),
            ("age_band", age_summary, ["age_band"]),
            ("age_band_by_recorded_sex", combined_summary, ["age_band", "sex_label"]),
        ]
    )
    standardized_groups = standardized_combined_groups(subject_frame)

    demographic_counts = (
        demographics.groupby(["age_band_index", "age_band", "sex"], observed=True)
        .size()
        .reset_index(name="n_participants")
    )
    inventory_rows = [
        ("cohort", "final_132"),
        ("participants", len(demographics)),
        ("female_participants", int((demographics["sex"] == "FEMALE").sum())),
        ("male_participants", int((demographics["sex"] == "MALE").sum())),
        ("minimum_age_years", float(demographics["age_years_precise"].min())),
        ("maximum_age_years", float(demographics["age_years_precise"].max())),
        ("targets", subject_frame["target"].nunique()),
        ("repeats_per_participant_target", EXPECTED_REPEATS),
        ("run_level_records", len(run_frame)),
        ("subject_target_records", len(subject_frame)),
        ("complete_run_records", int((run_frame["status"] == "complete").sum())),
        ("primary_endpoint", PRIMARY_METRIC),
        ("inferential_unit", "participant"),
        ("demographics_path", str(args.demographics.resolve())),
        ("demographics_sha256", sha256_file(args.demographics.resolve())),
        ("subjects_path", str(args.subjects.resolve())),
        ("subjects_sha256", sha256_file(args.subjects.resolve())),
        ("metrics_root", str(args.metrics_root.resolve())),
    ]
    inventory = pd.DataFrame(inventory_rows, columns=["item", "value"])

    run_path = output_dir / "run_level_metrics.csv"
    subject_path = output_dir / "subject_target_metrics.csv"
    run_frame.to_csv(run_path, index=False)
    subject_frame.to_csv(subject_path, index=False)
    demographics.to_csv(output_dir / "cohort_demographics.csv", index=False)
    demographic_counts.to_csv(output_dir / "demographic_cell_counts.csv", index=False)
    sex_summary.to_csv(output_dir / "sex_stratified_summary.csv", index=False)
    age_summary.to_csv(output_dir / "age_stratified_summary.csv", index=False)
    combined_summary.to_csv(output_dir / "sex_age_stratified_summary.csv", index=False)
    models.to_csv(output_dir / "model_results.csv", index=False)
    rankings.to_csv(output_dir / "group_rankings.csv", index=False)
    standardized_groups.to_csv(output_dir / "combined_standardized_roi_mean_rankings.csv", index=False)

    plot_sex(subject_frame, figures_dir)
    plot_age(subject_frame, figures_dir)
    plot_combined_strata(combined_summary, figures_dir)
    plot_combined_continuous(subject_frame, figures_dir)

    write_excel(
        output_dir / "analysis_tables.xlsx",
        inventory=inventory,
        demographics=demographics,
        subject_frame=subject_frame,
        sex_summary=sex_summary,
        age_summary=age_summary,
        combined_summary=combined_summary,
        models=models,
        rankings=rankings,
        standardized_groups=standardized_groups,
    )
    write_report(
        output_dir / "ANALYSIS_REPORT.md",
        demographics,
        sex_summary,
        age_summary,
        combined_summary,
        models,
        standardized_groups,
    )

    inventory_json = {
        "schema_version": 1,
        "analysis": "final_132_ti_demographic_analysis",
        "interpretation": "simulated_dosimetry_not_clinical_benefit",
        "inputs": {
            "metrics_root": str(args.metrics_root.resolve()),
            "demographics": str(args.demographics.resolve()),
            "demographics_sha256": sha256_file(args.demographics.resolve()),
            "subjects": str(args.subjects.resolve()),
            "subjects_sha256": sha256_file(args.subjects.resolve()),
        },
        "validated_counts": {
            "participants": len(demographics),
            "targets": subject_frame["target"].nunique(),
            "repeats": run_frame["repeat"].nunique(),
            "run_level_records": len(run_frame),
            "complete_run_level_records": int((run_frame["status"] == "complete").sum()),
            "subject_target_records": len(subject_frame),
        },
        "analysis_choices": {
            "primary_endpoint": PRIMARY_METRIC,
            "supporting_endpoints": [metric for metric in METRICS if metric != PRIMARY_METRIC],
            "inferential_unit": "participant",
            "repeat_handling": "mean within participant and target; retain repeat SD and CV",
            "age_bands": list(subject_frame["age_band"].cat.categories),
            "sex_field": "recorded sex (FEMALE/MALE), not gender identity",
            "model_covariance": "HC3",
            "multiplicity": "Benjamini-Hochberg FDR across four targets within outcome/effect family",
        },
        "derived_data_sha256": {
            "run_level_metrics.csv": sha256_file(run_path),
            "subject_target_metrics.csv": sha256_file(subject_path),
        },
    }
    (output_dir / "analysis_inventory.json").write_text(
        json.dumps(inventory_json, indent=2) + "\n", encoding="utf-8"
    )

    print(
        json.dumps(
            {
                "status": "complete",
                "participants": len(demographics),
                "run_level_records": len(run_frame),
                "subject_target_records": len(subject_frame),
                "output_dir": str(output_dir),
            },
            indent=2,
        )
    )


if __name__ == "__main__":
    main()
