#!/usr/bin/env python3
"""Build an auditable age/sex-balanced repeatability cohort package."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import shutil
from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from statistics import mean, median, stdev
from typing import Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


COHORT_ID = "final_132_repeatability_balanced_10"
EXPECTED_APPROVED_SUBJECTS = 132
TARGET_AGE_BANDS = 5
SEX_ORDER = ("FEMALE", "MALE")
TARGET_SUBJECTS = TARGET_AGE_BANDS * len(SEX_ORDER)
WORKBOOK_NAME = "CamCAN (1).xlsx"
WORKBOOK_SHEET = "Raw data"
STANDARD_DATA_NAME = "standard_data.csv"
OLD_POOL_NAME = "subjects.txt"

NAVY = "17324D"
TEAL = "177E89"
LIGHT_TEAL = "DDEFF1"
LIGHT_GREEN = "E2F0D9"
LIGHT_PURPLE = "E4DFEC"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"
GREEN_FONT = "008000"
GRAY_FONT = "666666"
ORANGE = "FCE4D6"


@dataclass(frozen=True)
class Subject:
    subject: str
    ccid: str
    precise_age: float
    workbook_age: float
    sex: str


@dataclass(frozen=True)
class Band:
    index: int
    lower: float
    upper: float
    center: float

    @property
    def label(self) -> str:
        return f"{self.lower:.2f}-{self.upper:.2f}"

    def contains(self, age: float) -> bool:
        if self.index == TARGET_AGE_BANDS:
            return self.lower <= age <= self.upper
        return self.lower <= age < self.upper


@dataclass(frozen=True)
class Candidate:
    subject: Subject
    band: Band
    rank: int

    @property
    def distance(self) -> float:
        return abs(self.subject.precise_age - self.band.center)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_subject_ids(path: Path, *, expected_count: int | None = None) -> list[str]:
    subject_ids = [line.strip() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if expected_count is not None and len(subject_ids) != expected_count:
        raise ValueError(f"{path} contains {len(subject_ids)} subjects; expected {expected_count}.")
    if len(subject_ids) != len(set(subject_ids)):
        raise ValueError(f"{path} contains duplicate subject IDs.")
    invalid = [subject for subject in subject_ids if not subject.startswith("sub-CC")]
    if invalid:
        raise ValueError(f"{path} contains invalid subject IDs: {invalid[:5]}")
    return subject_ids


def read_workbook_demographics(path: Path) -> dict[str, tuple[float, str]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    if WORKBOOK_SHEET not in workbook.sheetnames:
        raise ValueError(f"Workbook does not contain required sheet {WORKBOOK_SHEET!r}.")
    demographics: dict[str, tuple[float, str]] = {}
    for row in workbook[WORKBOOK_SHEET].iter_rows(values_only=True):
        if not row or not isinstance(row[0], str) or not row[0].startswith("sub-CC"):
            continue
        subject = row[0].strip()
        if subject in demographics:
            raise ValueError(f"Duplicate workbook row for {subject}.")
        if row[1] is None or row[3] is None:
            raise ValueError(f"Workbook row for {subject} is missing age or sex.")
        demographics[subject] = (float(row[1]), str(row[3]).strip().upper())
    return demographics


def read_precise_demographics(path: Path) -> dict[str, tuple[float, str]]:
    demographics: dict[str, tuple[float, str]] = {}
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        required = {"CCID", "Age", "Sex"}
        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise ValueError(f"{path} is missing columns: {sorted(required)}")
        for row in reader:
            subject = f"sub-{row['CCID'].strip()}"
            if subject in demographics:
                raise ValueError(f"Duplicate standard-data row for {subject}.")
            demographics[subject] = (float(row["Age"]), row["Sex"].strip().upper())
    return demographics


def merge_demographics(
    approved_ids: list[str],
    workbook_data: dict[str, tuple[float, str]],
    precise_data: dict[str, tuple[float, str]],
) -> tuple[list[Subject], float]:
    missing_workbook = [subject for subject in approved_ids if subject not in workbook_data]
    missing_precise = [subject for subject in approved_ids if subject not in precise_data]
    if missing_workbook or missing_precise:
        raise ValueError(
            "Demographic join is incomplete: "
            f"workbook_missing={missing_workbook}, precise_missing={missing_precise}"
        )

    subjects: list[Subject] = []
    max_age_delta = 0.0
    for subject in approved_ids:
        workbook_age, workbook_sex = workbook_data[subject]
        precise_age, precise_sex = precise_data[subject]
        if workbook_sex != precise_sex:
            raise ValueError(
                f"Sex mismatch for {subject}: workbook={workbook_sex}, standard={precise_sex}"
            )
        age_delta = abs(workbook_age - precise_age)
        max_age_delta = max(max_age_delta, age_delta)
        if age_delta > 1.01:
            raise ValueError(
                f"Age mismatch for {subject} exceeds one year: "
                f"workbook={workbook_age}, standard={precise_age}"
            )
        subjects.append(
            Subject(
                subject=subject,
                ccid=subject.removeprefix("sub-"),
                precise_age=precise_age,
                workbook_age=workbook_age,
                sex=precise_sex,
            )
        )
    return subjects, max_age_delta


def build_bands(subjects: Iterable[Subject]) -> list[Band]:
    ages = [subject.precise_age for subject in subjects]
    lower = min(ages)
    upper = max(ages)
    width = (upper - lower) / TARGET_AGE_BANDS
    edges = [lower + index * width for index in range(TARGET_AGE_BANDS + 1)]
    return [
        Band(
            index=index + 1,
            lower=edges[index],
            upper=edges[index + 1],
            center=(edges[index] + edges[index + 1]) / 2,
        )
        for index in range(TARGET_AGE_BANDS)
    ]


def rank_candidates(subjects: list[Subject], bands: list[Band]) -> list[Candidate]:
    candidates: list[Candidate] = []
    for band in bands:
        for sex in SEX_ORDER:
            eligible = [
                subject for subject in subjects if subject.sex == sex and band.contains(subject.precise_age)
            ]
            if not eligible:
                raise ValueError(f"No {sex} candidates in age band {band.label}.")
            ranked = sorted(
                eligible,
                key=lambda subject: (
                    abs(subject.precise_age - band.center),
                    subject.precise_age,
                    subject.subject,
                ),
            )
            candidates.extend(
                Candidate(subject=subject, band=band, rank=rank)
                for rank, subject in enumerate(ranked, start=1)
            )
    return candidates


def selected_candidates(candidates: Iterable[Candidate]) -> list[Candidate]:
    selected = [candidate for candidate in candidates if candidate.rank == 1]
    selected.sort(key=lambda candidate: (candidate.band.index, SEX_ORDER.index(candidate.subject.sex)))
    if len(selected) != TARGET_SUBJECTS:
        raise ValueError(f"Selection produced {len(selected)} subjects; expected {TARGET_SUBJECTS}.")
    if len({candidate.subject.subject for candidate in selected}) != TARGET_SUBJECTS:
        raise ValueError("Selection produced duplicate subjects.")
    return selected


def age_band_for(subject: Subject, bands: list[Band]) -> Band:
    matching = [band for band in bands if band.contains(subject.precise_age)]
    if len(matching) != 1:
        raise ValueError(f"Expected one age band for {subject.subject}; found {len(matching)}.")
    return matching[0]


def write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def add_table(worksheet, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def set_column_widths(worksheet, widths: dict[str, float]) -> None:
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def style_header(row) -> None:
    for cell in row:
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def write_audit_workbook(
    path: Path,
    *,
    subjects: list[Subject],
    selected: list[Candidate],
    candidates: list[Candidate],
    bands: list[Band],
    source_paths: dict[str, Path],
    source_hashes: dict[str, str],
    max_age_delta: float,
    old_selected_ids: set[str],
) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    selected_sheet = workbook.create_sheet("Selected 10")
    eligible_sheet = workbook.create_sheet("Eligible 132")
    alternatives_sheet = workbook.create_sheet("Ranked Candidates")
    sources_sheet = workbook.create_sheet("Sources")

    for worksheet in workbook.worksheets:
        worksheet.sheet_view.showGridLines = False

    summary.merge_cells("A1:D1")
    summary["A1"] = "Final-132 Repeatability Cohort: Balanced 10"
    summary["A1"].fill = PatternFill("solid", fgColor=NAVY)
    summary["A1"].font = Font(color=WHITE, bold=True, size=16)
    summary["A1"].alignment = Alignment(horizontal="center")
    summary["A3"] = "Metric"
    summary["B3"] = "Value"
    summary["C3"] = "Validation"
    summary["D3"] = "Notes"
    style_header(summary[3])
    metrics = [
        ("Approved pool size", EXPECTED_APPROVED_SUBJECTS, "PASS", "Authoritative final_132 list"),
        ("Selected subjects", "=COUNTA('Selected 10'!A2:A11)", "PASS", "Requested full cohort"),
        ("Female subjects", '=COUNTIF(\'Selected 10\'!C2:C11,"FEMALE")', "PASS", "Target = 5"),
        ("Male subjects", '=COUNTIF(\'Selected 10\'!C2:C11,"MALE")', "PASS", "Target = 5"),
        ("Age bands represented", "=COUNTA(A16:A20)", "PASS", "Target = 5"),
        ("Selected mean age", "=AVERAGE('Selected 10'!B2:B11)", "PASS", "Years"),
        ("Selected median age", "=MEDIAN('Selected 10'!B2:B11)", "PASS", "Years"),
        ("Selected minimum age", "=MIN('Selected 10'!B2:B11)", "PASS", "Years"),
        ("Selected maximum age", "=MAX('Selected 10'!B2:B11)", "PASS", "Years"),
        ("Maximum source age delta", max_age_delta, "PASS", "Workbook is integer-age; CSV is precise"),
        ("Old/new selected overlap", len({item.subject.subject for item in selected} & old_selected_ids), "INFO", "No overlap was required"),
    ]
    for row_index, (metric, value, validation, notes) in enumerate(metrics, start=4):
        summary.cell(row=row_index, column=1, value=metric)
        summary.cell(row=row_index, column=2, value=value)
        summary.cell(row=row_index, column=3, value=validation)
        summary.cell(row=row_index, column=4, value=notes)
        summary.cell(row=row_index, column=1).font = Font(color=GRAY_FONT)
        if isinstance(value, str) and value.startswith("="):
            summary.cell(row=row_index, column=2).font = Font(color="000000")
        else:
            summary.cell(row=row_index, column=2).font = Font(color=GRAY_FONT)
        summary.cell(row=row_index, column=3).fill = PatternFill("solid", fgColor=LIGHT_GREEN)
    summary["A15"] = "Age Band"
    summary["B15"] = "Female"
    summary["C15"] = "Male"
    summary["D15"] = "Total"
    style_header(summary[15])
    for offset, band in enumerate(bands, start=16):
        summary.cell(offset, 1, band.label)
        summary.cell(
            offset,
            2,
            f'=COUNTIFS(\'Selected 10\'!$D$2:$D$11,A{offset},\'Selected 10\'!$C$2:$C$11,"FEMALE")',
        )
        summary.cell(
            offset,
            3,
            f'=COUNTIFS(\'Selected 10\'!$D$2:$D$11,A{offset},\'Selected 10\'!$C$2:$C$11,"MALE")',
        )
        summary.cell(offset, 4, f"=SUM(B{offset}:C{offset})")
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Selected Subjects by Age Band and Sex"
    chart.y_axis.title = "Subjects"
    chart.x_axis.title = "Age band (years)"
    chart.height = 7
    chart.width = 14
    chart.add_data(Reference(summary, min_col=2, max_col=3, min_row=15, max_row=20), titles_from_data=True)
    chart.set_categories(Reference(summary, min_col=1, min_row=16, max_row=20))
    summary.add_chart(chart, "F3")
    set_column_widths(summary, {"A": 30, "B": 18, "C": 14, "D": 42})
    summary.freeze_panes = "A4"

    selected_headers = [
        "Subject ID",
        "Precise Age",
        "Sex",
        "Age Band",
        "Band Lower",
        "Band Upper",
        "Band Center",
        "Distance to Center",
        "Workbook Age",
        "Rank",
        "Approved",
        "Old 10",
    ]
    selected_sheet.append(selected_headers)
    style_header(selected_sheet[1])
    for row_index, candidate in enumerate(selected, start=2):
        subject = candidate.subject
        selected_sheet.append(
            [
                subject.subject,
                subject.precise_age,
                subject.sex,
                candidate.band.label,
                candidate.band.lower,
                candidate.band.upper,
                f"=(E{row_index}+F{row_index})/2",
                f"=ABS(B{row_index}-G{row_index})",
                subject.workbook_age,
                candidate.rank,
                f'=IF(COUNTIF(\'Eligible 132\'!$A$2:$A$133,A{row_index})=1,"YES","ERROR")',
                "YES" if subject.subject in old_selected_ids else "NO",
            ]
        )
    for row in selected_sheet.iter_rows(min_row=2, max_row=11):
        row[0].alignment = Alignment(vertical="center", horizontal="left")
        for cell in row[1:]:
            cell.alignment = Alignment(vertical="center", horizontal="center")
        for column in (1, 2, 3, 4, 5, 6, 9, 10, 12):
            row[column - 1].font = Font(color=GREEN_FONT, size=9)
        for column in (7, 8, 11):
            row[column - 1].font = Font(color="000000", size=9)
    for column in (2, 5, 6, 7, 8, 9):
        for cell in selected_sheet[get_column_letter(column)][1:]:
            cell.number_format = "0.00"
    selected_sheet.freeze_panes = "A2"
    selected_sheet.auto_filter.ref = "A1:L11"
    add_table(selected_sheet, "Selected10", "A1:L11")
    set_column_widths(
        selected_sheet,
        {
            "A": 19,
            "B": 13,
            "C": 12,
            "D": 17,
            "E": 13,
            "F": 13,
            "G": 14,
            "H": 19,
            "I": 15,
            "J": 9,
            "K": 12,
            "L": 10,
        },
    )

    eligible_headers = [
        "Subject ID",
        "CCID",
        "Precise Age",
        "Workbook Age",
        "Sex",
        "Age Band",
        "Band Index",
        "Selected",
    ]
    eligible_sheet.append(eligible_headers)
    style_header(eligible_sheet[1])
    selected_ids = {candidate.subject.subject for candidate in selected}
    for row_index, subject in enumerate(sorted(subjects, key=lambda item: item.subject), start=2):
        band = age_band_for(subject, bands)
        eligible_sheet.append(
            [
                subject.subject,
                subject.ccid,
                subject.precise_age,
                subject.workbook_age,
                subject.sex,
                band.label,
                band.index,
                f'=IF(COUNTIF(\'Selected 10\'!$A$2:$A$11,A{row_index})>0,"YES","")',
            ]
        )
        for cell in eligible_sheet[row_index]:
            cell.font = Font(color=GREEN_FONT, size=9)
        eligible_sheet.cell(row_index, 1).alignment = Alignment(horizontal="left")
        for column in range(2, 9):
            eligible_sheet.cell(row_index, column).alignment = Alignment(horizontal="center")
        eligible_sheet.cell(row_index, 8).font = Font(color="000000", size=9)
    for column in (3, 4):
        for cell in eligible_sheet[get_column_letter(column)][1:]:
            cell.number_format = "0.00"
    eligible_sheet.conditional_formatting.add(
        f"A2:H{len(subjects) + 1}",
        FormulaRule(
            formula=["$H2=\"YES\""],
            fill=PatternFill("solid", fgColor=LIGHT_GREEN),
        ),
    )
    eligible_sheet.freeze_panes = "A2"
    eligible_sheet.auto_filter.ref = f"A1:H{len(subjects) + 1}"
    add_table(eligible_sheet, "Eligible132", f"A1:H{len(subjects) + 1}")
    set_column_widths(
        eligible_sheet,
        {"A": 19, "B": 15, "C": 13, "D": 15, "E": 12, "F": 17, "G": 12, "H": 12},
    )

    alternatives_headers = [
        "Age Band",
        "Band Index",
        "Sex",
        "Rank",
        "Subject ID",
        "Precise Age",
        "Band Center",
        "Distance",
        "Selected",
    ]
    alternatives_sheet.append(alternatives_headers)
    style_header(alternatives_sheet[1])
    top_candidates = sorted(
        (candidate for candidate in candidates if candidate.rank <= 3),
        key=lambda candidate: (
            candidate.band.index,
            SEX_ORDER.index(candidate.subject.sex),
            candidate.rank,
        ),
    )
    for candidate in top_candidates:
        alternatives_sheet.append(
            [
                candidate.band.label,
                candidate.band.index,
                candidate.subject.sex,
                candidate.rank,
                candidate.subject.subject,
                candidate.subject.precise_age,
                candidate.band.center,
                candidate.distance,
                "YES" if candidate.subject.subject in selected_ids else "NO",
            ]
        )
    for row in alternatives_sheet.iter_rows(min_row=2, max_row=len(top_candidates) + 1):
        row[0].alignment = Alignment(horizontal="left")
        for cell in row[1:]:
            cell.alignment = Alignment(horizontal="center")
        for cell in row:
            cell.font = Font(size=9)
    for column in (6, 7, 8):
        for cell in alternatives_sheet[get_column_letter(column)][1:]:
            cell.number_format = "0.00"
    alternatives_sheet.conditional_formatting.add(
        f"A2:I{len(top_candidates) + 1}",
        FormulaRule(
            formula=["$I2=\"YES\""],
            fill=PatternFill("solid", fgColor=LIGHT_GREEN),
        ),
    )
    alternatives_sheet.freeze_panes = "A2"
    alternatives_sheet.auto_filter.ref = f"A1:I{len(top_candidates) + 1}"
    add_table(alternatives_sheet, "RankedCandidates", f"A1:I{len(top_candidates) + 1}")
    set_column_widths(
        alternatives_sheet,
        {"A": 17, "B": 12, "C": 12, "D": 9, "E": 19, "F": 13, "G": 14, "H": 12, "I": 12},
    )

    sources_sheet.append(["Source", "Path", "SHA-256", "Use"])
    style_header(sources_sheet[1])
    source_uses = {
        "Approved final-132 list": "Eligibility gate",
        "CamCAN workbook": "Age/sex validation; integer age",
        "Precise demographics CSV": "Primary age ranking and sex",
        "Old repeatability config": "Old/new overlap audit",
    }
    for name, source_path in source_paths.items():
        sources_sheet.append([name, str(source_path), source_hashes[name], source_uses[name]])
    for row in sources_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(color=GREEN_FONT, size=8)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sources_sheet["A8"] = "Selection logic"
    sources_sheet["B8"] = (
        "Five equal-width bands across the final-132 precise age range; "
        "one female and one male nearest each band center; ties by lower age then subject ID."
    )
    sources_sheet["A8"].fill = PatternFill("solid", fgColor=LIGHT_PURPLE)
    sources_sheet["A8"].font = Font(bold=True)
    sources_sheet["B8"].fill = PatternFill("solid", fgColor=LIGHT_PURPLE)
    sources_sheet["B8"].alignment = Alignment(wrap_text=True, vertical="top")
    set_column_widths(sources_sheet, {"A": 29, "B": 86, "C": 68, "D": 42})
    sources_sheet.freeze_panes = "A2"

    for worksheet in workbook.worksheets:
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_LETTER
        worksheet.page_setup.orientation = (
            worksheet.ORIENTATION_PORTRAIT
            if worksheet.title == "Summary"
            else worksheet.ORIENTATION_LANDSCAPE
        )
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_view.zoomScale = 90

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)


def markdown_table(selected: list[Candidate]) -> str:
    lines = [
        "| Band | Female | Female age | Male | Male age | Pair gap |",
        "| --- | --- | ---: | --- | ---: | ---: |",
    ]
    by_band: dict[int, dict[str, Candidate]] = {}
    for candidate in selected:
        by_band.setdefault(candidate.band.index, {})[candidate.subject.sex] = candidate
    for band_index in sorted(by_band):
        female = by_band[band_index]["FEMALE"]
        male = by_band[band_index]["MALE"]
        lines.append(
            f"| {female.band.label} | `{female.subject.subject}` | {female.subject.precise_age:.2f} "
            f"| `{male.subject.subject}` | {male.subject.precise_age:.2f} "
            f"| {abs(female.subject.precise_age - male.subject.precise_age):.2f} |"
        )
    return "\n".join(lines)


def write_readme(
    path: Path,
    *,
    subjects: list[Subject],
    selected: list[Candidate],
    bands: list[Band],
    old_pool_ids: list[str],
    old_selected_ids: set[str],
    source_paths: dict[str, Path],
    source_hashes: dict[str, str],
) -> None:
    selected_ages = [candidate.subject.precise_age for candidate in selected]
    pool_ages = [subject.precise_age for subject in subjects]
    pool_sexes = Counter(subject.sex for subject in subjects)
    old_overlap = {candidate.subject.subject for candidate in selected} & old_selected_ids
    pair_gaps = []
    for band in bands:
        pair = [candidate for candidate in selected if candidate.band.index == band.index]
        pair_gaps.append(abs(pair[0].subject.precise_age - pair[1].subject.precise_age))
    source_root = "/mnt/parscratch/users/cop23bi/ti_dataset_final_132_balanced_10_corrected"
    readme = f"""# Final-132 Repeatability Balanced-10 Cohort

This package defines the new 10-subject repeatability cohort selected only from
the supervisor-approved `final_132` CamCAN cohort.

## Outcome

- Subjects: `{TARGET_SUBJECTS}`
- Sex balance: `5 FEMALE / 5 MALE`
- Age balance: `2 subjects in each of 5 equal-width age bands`
- Approved-pool age range: `{min(pool_ages):.2f}-{max(pool_ages):.2f}` years
- Selected age range: `{min(selected_ages):.2f}-{max(selected_ages):.2f}` years
- Selected mean age: `{mean(selected_ages):.2f}` years
- Selected median age: `{median(selected_ages):.2f}` years
- Maximum female/male within-band age gap: `{max(pair_gaps):.2f}` years
- Overlap with the previous 10-subject repeatability set: `{len(old_overlap)}/10`

{markdown_table(selected)}

## Selection Rule

1. Restrict eligibility to the `{EXPECTED_APPROVED_SUBJECTS}` IDs in the authoritative
   `final_132/subjects.txt` file.
2. Verify every approved ID has age and sex in the supplied workbook.
3. Use the more precise decimal age in `standard_data.csv` for ranking; the Excel
   workbook stores age at whole-year precision. Sex agrees for all 132 subjects,
   and the maximum age-source difference is one year.
4. Divide the final-132 age range into five equal-width bands:
{chr(10).join(f"   - `{band.label}` (center `{band.center:.3f}`)" for band in bands)}
5. In each band, select one female and one male by:
   - smallest absolute distance to the band center
   - lower age
   - lexical subject ID

This is the same age/sex balancing principle used for the old 175-subject pool,
recomputed from the new approved cohort rather than carrying forward the old
subjects or old band boundaries.

## Cohort Delta

- Eligible subjects: `{len(old_pool_ids)} -> {len(subjects)}`
- Old eligible sex counts: computed in the prior selection package
- New eligible sex counts: `{pool_sexes['FEMALE']} FEMALE / {pool_sexes['MALE']} MALE`
- Old selected subjects retained by the new deterministic rule: `{len(old_overlap)}`

## Files

- `subjects.txt`: canonical 10-subject list for staging and runner configs
- `selection_manifest.csv`: selected demographics and ranking evidence
- `eligible_132_demographics.csv`: full approved pool with selected flags
- `selection_audit.xlsx`: formatted audit workbook with formulas and top-three alternates
- `cohort.json`: machine-readable provenance and validation record
- `paired_repeatability_experiment.candidate.json`: candidate 40-remesh + 40-fixed config
- `build_balanced_repeatability_cohort.py`: exact reproducible cohort builder
- `stage_repeatability_dataset.py`: safe audit/staging utility for the 30 required NIfTI inputs
- `SHA256SUMS`: package checksums

## Imaging Dataset Staging

The supplied demographics directory does not contain the T1, T2, and corrected
segmentation NIfTI payload required by the repeatability runner. On HPC, use the
included staging utility against the canonical corrected-v4 scaffold tree. This
mode requires a completed per-subject provenance record, verifies all three
recorded source hashes, and decompresses `.nii.gz` inputs to the exact uncompressed
filenames required by the repeatability runner. It never modifies the scaffold.

Audit first:

```bash
python stage_repeatability_dataset.py \\
  --source-layout final132-scaffold \\
  --source-root /mnt/parscratch/users/cop23bi/CamCan_Corrected_v4_Scaffolds \\
  --output-root {source_root} \\
  --subjects-file subjects.txt
```

Create the physical 10-subject dataset only after the audit reports
`subjects_ready=10` and `required_files_ready=30`:

```bash
python stage_repeatability_dataset.py \\
  --source-layout final132-scaffold \\
  --source-root /mnt/parscratch/users/cop23bi/CamCan_Corrected_v4_Scaffolds \\
  --output-root {source_root} \\
  --subjects-file subjects.txt \\
  --apply
```

The candidate config assumes the staged directory will be uploaded to:

`{source_root}`

That HPC path is a candidate, not a known-good live path. Confirm it after upload
before launching the runner.

## Requested Full Execution Scope

- Dataset: final-132 balanced repeatability cohort
- ROI: left hippocampus
- Subjects: `10`
- Conditions: `2` (`remesh`, `fixed_mesh`)
- Repeats per condition: `40`
- Total task count: `10 x 2 x 40 = 800`
- Expected repeat outputs: `800`
- Scope: full requested 10-subject repeatability study

No Slurm job is submitted by this package.

## Automated Two-Condition Execution

After the dataset is staged and the updated repository is available on HPC,
initialize the current-repair pipeline:

```bash
export CURRENT_REPAIR_DIR=/users/cop23bi/Repos/TI_Pipeline/SimNIBS/Scripts/ti_current_repair
export EXPERIMENT_ROOT=/mnt/parscratch/users/cop23bi/current-repair/final_132_balanced_10

python "$CURRENT_REPAIR_DIR/pipeline/staged_median_fixed_experiment.py" init \\
  --source-root {source_root} \\
  --experiment-root "$EXPERIMENT_ROOT" \\
  --subjects {','.join(candidate.subject.subject for candidate in selected)} \\
  --repeat-count 40 \\
  --atlas-dir /mnt/parscratch/users/cop23bi/ZIPs/atlases \\
  --roi-preset left-hippocampus
```

Run the read-only full-scope submission preflight:

```bash
python "$CURRENT_REPAIR_DIR/pipeline/staged_median_fixed_experiment.py" submit-all \\
  --experiment-root "$EXPERIMENT_ROOT" \\
  --max-concurrent 50 \\
  --analysis-max-concurrent 10 \\
  --dry-run
```

It must report 400 remesh tasks (`0-399%50`), 400 fixed-mesh tasks
(`0-399%50`), and 800 expected `TI.msh` outputs. The production command is the
same without `--dry-run`.

The command initially submits only the remesh array and an `afterok` controller.
The controller validates all remesh outputs, analyzes them, selects each
subject's median representative mesh, physically seeds and validates the fixed
workspaces, then releases the fixed array, paired analysis, and final figures.
Any failed job or validation gate stops the chain before downstream submission.

## Provenance

- Approved-list SHA-256: `{source_hashes['Approved final-132 list']}`
- Workbook SHA-256: `{source_hashes['CamCAN workbook']}`
- Precise-demographics SHA-256: `{source_hashes['Precise demographics CSV']}`
- Old-config SHA-256: `{source_hashes['Old repeatability config']}`
- Generated: `{datetime.now().astimezone().isoformat(timespec='seconds')}`

Source paths are also recorded in `cohort.json` and the workbook `Sources` sheet.
"""
    path.write_text(readme, encoding="utf-8")


def copy_staging_utility(destination: Path) -> None:
    source = Path(__file__).with_name("stage_repeatability_dataset.py")
    if not source.is_file():
        raise FileNotFoundError(f"Staging utility is missing: {source}")
    shutil.copy2(source, destination)


def build(args: argparse.Namespace) -> Path:
    demographics_dir = args.demographics_dir.resolve()
    approved_path = args.approved_subjects.resolve()
    output_dir = args.output_dir.resolve()
    workbook_path = demographics_dir / WORKBOOK_NAME
    precise_path = demographics_dir / STANDARD_DATA_NAME
    old_pool_path = demographics_dir / OLD_POOL_NAME
    old_config_path = args.old_config.resolve()

    for required_path in (
        workbook_path,
        precise_path,
        old_pool_path,
        approved_path,
        old_config_path,
    ):
        if not required_path.is_file():
            raise FileNotFoundError(required_path)

    approved_ids = read_subject_ids(approved_path, expected_count=EXPECTED_APPROVED_SUBJECTS)
    old_pool_ids = read_subject_ids(old_pool_path)
    workbook_data = read_workbook_demographics(workbook_path)
    precise_data = read_precise_demographics(precise_path)
    subjects, max_age_delta = merge_demographics(approved_ids, workbook_data, precise_data)
    bands = build_bands(subjects)
    candidates = rank_candidates(subjects, bands)
    selected = selected_candidates(candidates)

    old_config = json.loads(old_config_path.read_text(encoding="utf-8"))
    old_selected_ids = set(old_config["subjects"])
    if len(old_selected_ids) != TARGET_SUBJECTS:
        raise ValueError(f"Old config contains {len(old_selected_ids)} subjects; expected 10.")

    output_dir.mkdir(parents=True, exist_ok=True)
    source_paths = {
        "Approved final-132 list": approved_path,
        "CamCAN workbook": workbook_path,
        "Precise demographics CSV": precise_path,
        "Old repeatability config": old_config_path,
    }
    source_hashes = {name: sha256_file(path) for name, path in source_paths.items()}

    (output_dir / "subjects.txt").write_text(
        "".join(f"{candidate.subject.subject}\n" for candidate in selected),
        encoding="utf-8",
    )

    selected_rows = []
    selected_ids = {candidate.subject.subject for candidate in selected}
    for candidate in selected:
        subject = candidate.subject
        selected_rows.append(
            {
                "subject": subject.subject,
                "CCID": subject.ccid,
                "age_years_precise": f"{subject.precise_age:.2f}",
                "age_years_workbook": f"{subject.workbook_age:.2f}",
                "sex": subject.sex,
                "age_band_index": candidate.band.index,
                "age_band": candidate.band.label,
                "band_lower": f"{candidate.band.lower:.3f}",
                "band_upper": f"{candidate.band.upper:.3f}",
                "band_center": f"{candidate.band.center:.3f}",
                "distance_to_band_center": f"{candidate.distance:.3f}",
                "selection_rank": candidate.rank,
                "approved_final_132": "YES",
                "old_175_repeatability_subject": (
                    "YES" if subject.subject in old_selected_ids else "NO"
                ),
            }
        )
    write_csv(
        output_dir / "selection_manifest.csv",
        list(selected_rows[0].keys()),
        selected_rows,
    )

    eligible_rows = []
    for subject in sorted(subjects, key=lambda item: item.subject):
        band = age_band_for(subject, bands)
        eligible_rows.append(
            {
                "subject": subject.subject,
                "CCID": subject.ccid,
                "age_years_precise": f"{subject.precise_age:.2f}",
                "age_years_workbook": f"{subject.workbook_age:.2f}",
                "sex": subject.sex,
                "age_band_index": band.index,
                "age_band": band.label,
                "selected": "YES" if subject.subject in selected_ids else "NO",
            }
        )
    write_csv(
        output_dir / "eligible_132_demographics.csv",
        list(eligible_rows[0].keys()),
        eligible_rows,
    )

    pair_gaps = []
    for band in bands:
        pair = [candidate for candidate in selected if candidate.band.index == band.index]
        pair_gaps.append(abs(pair[0].subject.precise_age - pair[1].subject.precise_age))
    cohort_payload = {
        "schema_version": 1,
        "cohort_id": COHORT_ID,
        "status": "selection_complete_imaging_staging_pending",
        "created_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "source_cohort_id": "final_132",
        "selection_scope": {
            "eligible_subjects": len(subjects),
            "selected_subjects": len(selected),
            "female_selected": sum(candidate.subject.sex == "FEMALE" for candidate in selected),
            "male_selected": sum(candidate.subject.sex == "MALE" for candidate in selected),
            "age_bands": len(bands),
            "subjects_per_age_band": 2,
        },
        "selection_rule": {
            "age_source": "standard_data.csv:Age (decimal precision)",
            "sex_source": "CamCAN workbook Raw data plus standard_data.csv agreement",
            "band_definition": "five equal-width bands over the final-132 precise age range",
            "per_band": "one FEMALE and one MALE",
            "ranking": [
                "absolute distance to band center",
                "lower age",
                "lexical subject ID",
            ],
        },
        "age_bands": [
            {
                "index": band.index,
                "lower": band.lower,
                "upper": band.upper,
                "center": band.center,
                "label": band.label,
            }
            for band in bands
        ],
        "selected_subjects": [candidate.subject.subject for candidate in selected],
        "validation": {
            "approved_subjects_expected": EXPECTED_APPROVED_SUBJECTS,
            "approved_subjects_matched_to_workbook": len(subjects),
            "missing_age": 0,
            "missing_sex": 0,
            "sex_source_mismatches": 0,
            "max_age_source_delta_years": max_age_delta,
            "exact_sex_balance": True,
            "exact_two_per_age_band": True,
            "maximum_within_band_pair_age_gap_years": max(pair_gaps),
            "old_selected_overlap": len(selected_ids & old_selected_ids),
        },
        "source_files": {
            name: {"path": str(path), "sha256": source_hashes[name]}
            for name, path in source_paths.items()
        },
        "hpc_profile": {
            "status": "candidate",
            "closest_known_profile": "camcan-corrected-v4-final-132-four-roi-qos-sequential",
            "note": "The final-132 source cohort is known-good; the new 10-subject repeatability paths are unconfirmed.",
        },
    }
    (output_dir / "cohort.json").write_text(
        json.dumps(cohort_payload, indent=2) + "\n",
        encoding="utf-8",
    )

    candidate_config = {
        "source_root": "/mnt/parscratch/users/cop23bi/ti_dataset_final_132_balanced_10_corrected",
        "experiment_root": "/mnt/parscratch/users/cop23bi/current-repair/final_132_balanced_10",
        "subjects": [candidate.subject.subject for candidate in selected],
        "conditions": [
            {
                "name": "remesh",
                "mesh_mode": "remesh",
                "repeat_count": 40,
                "description": "Generate a fresh mesh for every repeat to capture remeshing plus FEM variation.",
            },
            {
                "name": "fixed_mesh",
                "mesh_mode": "fixed_mesh",
                "repeat_count": 40,
                "description": "Generate one mesh once per subject, then reuse it across repeats to isolate FEM variation.",
            },
        ],
        "analysis": {
            "roi_preset": "left-hippocampus",
            "atlas_dir": "/mnt/parscratch/users/cop23bi/ZIPs/atlases",
            "compare_metric": "median_roi",
        },
    }
    (output_dir / "paired_repeatability_experiment.candidate.json").write_text(
        json.dumps(candidate_config, indent=2) + "\n",
        encoding="utf-8",
    )

    shutil.copy2(Path(__file__), output_dir / "build_balanced_repeatability_cohort.py")
    copy_staging_utility(output_dir / "stage_repeatability_dataset.py")
    write_audit_workbook(
        output_dir / "selection_audit.xlsx",
        subjects=subjects,
        selected=selected,
        candidates=candidates,
        bands=bands,
        source_paths=source_paths,
        source_hashes=source_hashes,
        max_age_delta=max_age_delta,
        old_selected_ids=old_selected_ids,
    )
    write_readme(
        output_dir / "README.md",
        subjects=subjects,
        selected=selected,
        bands=bands,
        old_pool_ids=old_pool_ids,
        old_selected_ids=old_selected_ids,
        source_paths=source_paths,
        source_hashes=source_hashes,
    )

    checksum_targets = sorted(
        path
        for path in output_dir.iterdir()
        if path.is_file() and path.name not in {"SHA256SUMS"}
    )
    checksum_lines = [f"{sha256_file(path)}  {path.name}\n" for path in checksum_targets]
    (output_dir / "SHA256SUMS").write_text("".join(checksum_lines), encoding="utf-8")
    return output_dir


def parse_args() -> argparse.Namespace:
    scripts_root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--demographics-dir",
        type=Path,
        required=True,
        help=f"Directory containing {WORKBOOK_NAME!r} and {STANDARD_DATA_NAME!r}.",
    )
    parser.add_argument(
        "--approved-subjects",
        type=Path,
        default=scripts_root
        / "CamCan_Experiment"
        / "cohort_pipeline"
        / "cohorts"
        / "final_132"
        / "subjects.txt",
    )
    parser.add_argument(
        "--old-config",
        type=Path,
        default=scripts_root / "mesh_repeat_analysis" / "paired_repeatability_experiment.example.json",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=scripts_root / "output" / "spreadsheet" / COHORT_ID,
    )
    return parser.parse_args()


if __name__ == "__main__":
    destination = build(parse_args())
    print(f"[OK] Wrote cohort package: {destination}")
