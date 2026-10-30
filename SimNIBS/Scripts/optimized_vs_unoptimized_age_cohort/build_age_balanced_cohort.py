#!/usr/bin/env python3
"""Build the final-132 age/recorded-sex-balanced TI comparison cohort.

The seven previously optimized best/worst cases are retained as a separate
descriptive cohort and excluded from the new selection because their original
selection used the simulated outcome. The new cohort is selected without using
any TI field outcome.
"""

from __future__ import annotations

import csv
import hashlib
import json
import math
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from statistics import mean, median
from typing import Iterable

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


REPO_ROOT = Path(__file__).resolve().parents[1]
TASK_DIR = REPO_ROOT / "optimized_vs_unoptimized_age_cohort"
PACKAGE_DIR = (
    REPO_ROOT
    / "output"
    / "spreadsheet"
    / "final_132_optimized_vs_unoptimized_age_balanced_14"
)
CANONICAL_DIR = (
    REPO_ROOT
    / "CamCan_Experiment"
    / "cohort_pipeline"
    / "cohorts"
    / "optimized_vs_unoptimized_age_balanced_14"
)

FINAL_132_FILE = (
    REPO_ROOT
    / "CamCan_Experiment"
    / "cohort_pipeline"
    / "cohorts"
    / "final_132"
    / "subjects.txt"
)
DEMOGRAPHICS_FILE = (
    REPO_ROOT
    / "output"
    / "spreadsheet"
    / "final_132_repeatability_balanced_10"
    / "eligible_132_demographics.csv"
)
OPTIMIZED_7_DIR = (
    REPO_ROOT
    / "CamCan_Experiment"
    / "cohort_pipeline"
    / "cohorts"
    / "optimized_best_worst_7"
)
OPTIMIZED_7_FILE = OPTIMIZED_7_DIR / "subjects.txt"
OPTIMIZED_7_COHORT_FILE = OPTIMIZED_7_DIR / "cohort.json"

EXPECTED_FINAL_132 = 132
EXPECTED_OPTIMIZED = 7
EXPECTED_FD_BINS = 7
SEX_ORDER = ("FEMALE", "MALE")
FD_REFERENCE_URL = "https://doi.org/10.1007/BF01025868"

NAVY = "17324D"
TEAL = "177E89"
LIGHT_TEAL = "DDEFF1"
LIGHT_GREEN = "E2F0D9"
LIGHT_ORANGE = "FCE4D6"
LIGHT_GRAY = "E7E6E6"
WHITE = "FFFFFF"


@dataclass(frozen=True)
class Subject:
    subject: str
    age_precise: float
    age_workbook: float
    sex: str


@dataclass(frozen=True)
class Band:
    index: int
    lower: float
    upper: float
    center: float

    @property
    def label(self) -> str:
        return f"{self.lower:.2f}-{self.upper:.2f}"

    def contains(self, age: float, total_bands: int) -> bool:
        if self.index == total_bands:
            return self.lower <= age <= self.upper
        return self.lower <= age < self.upper


@dataclass(frozen=True)
class RankedCandidate:
    subject: Subject
    band: Band
    rank: int
    candidates_in_stratum: int

    @property
    def distance_to_center(self) -> float:
        return abs(self.subject.age_precise - self.band.center)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_subject_ids(path: Path, expected: int) -> list[str]:
    values = [line.strip() for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if len(values) != expected:
        raise ValueError(f"{path} has {len(values)} subjects; expected {expected}.")
    if len(values) != len(set(values)):
        raise ValueError(f"{path} contains duplicate subject IDs.")
    return values


def read_demographics(path: Path, final_ids: set[str]) -> list[Subject]:
    subjects: list[Subject] = []
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        required = {"subject", "age_years_precise", "age_years_workbook", "sex"}
        if not reader.fieldnames or not required.issubset(reader.fieldnames):
            raise ValueError(f"{path} is missing required columns {sorted(required)}.")
        for row in reader:
            subject_id = row["subject"].strip()
            subjects.append(
                Subject(
                    subject=subject_id,
                    age_precise=float(row["age_years_precise"]),
                    age_workbook=float(row["age_years_workbook"]),
                    sex=row["sex"].strip().upper(),
                )
            )

    ids = [subject.subject for subject in subjects]
    if len(ids) != EXPECTED_FINAL_132 or set(ids) != final_ids:
        raise ValueError(
            "Demographic join does not exactly match final_132: "
            f"rows={len(ids)}, missing={sorted(final_ids - set(ids))}, "
            f"extra={sorted(set(ids) - final_ids)}"
        )
    if len(ids) != len(set(ids)):
        raise ValueError("Demographics contain duplicate subject IDs.")
    invalid_sex = sorted({subject.sex for subject in subjects} - set(SEX_ORDER))
    if invalid_sex:
        raise ValueError(f"Unexpected recorded-sex values: {invalid_sex}")
    return subjects


def linear_quantile(values: list[float], probability: float) -> float:
    ordered = sorted(values)
    position = (len(ordered) - 1) * probability
    lower_index = math.floor(position)
    upper_index = math.ceil(position)
    if lower_index == upper_index:
        return ordered[lower_index]
    fraction = position - lower_index
    return ordered[lower_index] * (1 - fraction) + ordered[upper_index] * fraction


def build_bands(minimum: float, maximum: float, count: int) -> list[Band]:
    width = (maximum - minimum) / count
    edges = [minimum + index * width for index in range(count + 1)]
    return [
        Band(
            index=index + 1,
            lower=edges[index],
            upper=edges[index + 1],
            center=(edges[index] + edges[index + 1]) / 2,
        )
        for index in range(count)
    ]


def age_band(age: float, bands: list[Band]) -> Band:
    matches = [band for band in bands if band.contains(age, len(bands))]
    if len(matches) != 1:
        raise ValueError(f"Age {age} maps to {len(matches)} bands, expected exactly one.")
    return matches[0]


def rank_candidates(subjects: Iterable[Subject], bands: list[Band]) -> list[RankedCandidate]:
    output: list[RankedCandidate] = []
    materialized = list(subjects)
    for band in bands:
        for sex in SEX_ORDER:
            eligible = [
                subject
                for subject in materialized
                if subject.sex == sex and band.contains(subject.age_precise, len(bands))
            ]
            if not eligible:
                raise ValueError(f"No {sex} candidates in band {band.label}.")
            ordered = sorted(
                eligible,
                key=lambda item: (
                    abs(item.age_precise - band.center),
                    item.age_precise,
                    item.subject,
                ),
            )
            output.extend(
                RankedCandidate(
                    subject=subject,
                    band=band,
                    rank=index,
                    candidates_in_stratum=len(ordered),
                )
                for index, subject in enumerate(ordered, start=1)
            )
    return output


def write_csv(path: Path, fieldnames: list[str], rows: Iterable[dict[str, object]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def add_table(worksheet, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)


def style_header(cells) -> None:
    for cell in cells:
        cell.fill = PatternFill("solid", fgColor=TEAL)
        cell.font = Font(color=WHITE, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autosize(worksheet, maximum: int = 46) -> None:
    for column_cells in worksheet.columns:
        values = [str(cell.value) for cell in column_cells if cell.value is not None]
        width = min(maximum, max([len(value) for value in values] + [8]) + 2)
        worksheet.column_dimensions[column_cells[0].column_letter].width = width


def write_rows_sheet(workbook: Workbook, title: str, rows: list[dict[str, object]], table_name: str):
    worksheet = workbook.create_sheet(title)
    if not rows:
        return worksheet
    headers = list(rows[0])
    worksheet.append(headers)
    for row in rows:
        worksheet.append([row[header] for header in headers])
    style_header(worksheet[1])
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    add_table(worksheet, table_name, worksheet.dimensions)
    autosize(worksheet)
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
    worksheet.page_setup.fitToWidth = 1
    worksheet.page_setup.fitToHeight = 0
    worksheet.print_title_rows = "1:1"
    worksheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    worksheet.page_margins.left = 0.25
    worksheet.page_margins.right = 0.25
    worksheet.page_margins.top = 0.4
    worksheet.page_margins.bottom = 0.4
    worksheet.print_area = worksheet.dimensions
    return worksheet


def optimized_roles(cohort_path: Path) -> tuple[dict[str, list[str]], dict]:
    cohort = json.loads(cohort_path.read_text(encoding="utf-8"))
    roles: dict[str, list[str]] = defaultdict(list)
    for roi, mapping in cohort["selection"].items():
        for direction in ("best", "worst"):
            roles[mapping[direction]].append(f"{roi}:{direction}")
    return dict(roles), cohort


def feasibility_rows(
    all_subjects: list[Subject],
    eligible_subjects: list[Subject],
    minimum: float,
    maximum: float,
) -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for count in range(4, 13):
        bands = build_bands(minimum, maximum, count)
        full_counts = []
        eligible_counts = []
        for band in bands:
            full_counts.append(
                Counter(
                    subject.sex
                    for subject in all_subjects
                    if band.contains(subject.age_precise, count)
                )
            )
            eligible_counts.append(
                Counter(
                    subject.sex
                    for subject in eligible_subjects
                    if band.contains(subject.age_precise, count)
                )
            )
        rows.append(
            {
                "bin_count": count,
                "equal_width_years": round((maximum - minimum) / count, 3),
                "full_pool_all_sexes_each_bin": "YES"
                if all(c["FEMALE"] > 0 and c["MALE"] > 0 for c in full_counts)
                else "NO",
                "eligible_pool_all_sexes_each_bin": "YES"
                if all(c["FEMALE"] > 0 and c["MALE"] > 0 for c in eligible_counts)
                else "NO",
                "eligible_min_female_candidates": min(c["FEMALE"] for c in eligible_counts),
                "eligible_min_male_candidates": min(c["MALE"] for c in eligible_counts),
                "eligible_min_total_candidates": min(sum(c.values()) for c in eligible_counts),
                "eligible_max_total_candidates": max(sum(c.values()) for c in eligible_counts),
            }
        )
    return rows


def build_workbook(
    path: Path,
    selected_rows: list[dict[str, object]],
    optimized_rows: list[dict[str, object]],
    bin_rows: list[dict[str, object]],
    feasibility: list[dict[str, object]],
    audit_rows: list[dict[str, object]],
    source_rows: list[dict[str, object]],
    statistics: dict[str, float | int | str],
) -> None:
    workbook = Workbook()
    readme = workbook.active
    readme.title = "Read me"
    readme.sheet_view.showGridLines = False
    readme.merge_cells("A1:F1")
    readme["A1"] = "Optimized vs unoptimized TI: age-balanced cohort audit"
    readme["A1"].fill = PatternFill("solid", fgColor=NAVY)
    readme["A1"].font = Font(color=WHITE, bold=True, size=15)
    readme["A1"].alignment = Alignment(horizontal="left")
    readme.sheet_properties.pageSetUpPr.fitToPage = True
    readme.page_setup.orientation = "landscape"
    readme.page_setup.paperSize = readme.PAPERSIZE_A4
    readme.page_setup.fitToWidth = 1
    readme.page_setup.fitToHeight = 1
    readme.page_margins.left = 0.35
    readme.page_margins.right = 0.35
    readme.page_margins.top = 0.45
    readme.page_margins.bottom = 0.45

    notes = [
        ("Decision", "Seven equal-width age bands; one recorded-female and one recorded-male subject per band."),
        ("Scientific basis", "Freedman-Diaconis age-bin width on all 132 precise ages gives 7 bins."),
        ("Selection", "Within each recorded-sex x band stratum: nearest to midpoint, then lower age, then subject ID."),
        ("Independence", "The existing optimized best/worst seven are held out because they were selected using simulated TI outcomes."),
        ("Analysis", "Treat age as continuous; bins are a sampling device, not seven powered inferential groups."),
        ("Reference", FD_REFERENCE_URL),
    ]
    for row_index, (label, value) in enumerate(notes, start=3):
        readme.cell(row=row_index, column=1, value=label).font = Font(bold=True, color=NAVY)
        readme.cell(row=row_index, column=2, value=value)
        readme.merge_cells(start_row=row_index, start_column=2, end_row=row_index, end_column=6)
        readme.cell(row=row_index, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    metric_start = 11
    readme.cell(metric_start, 1, "Validation metric")
    readme.cell(metric_start, 2, "Value")
    style_header(readme[metric_start][0:2])
    metrics = [
        ("Final-132 participants", statistics["n"]),
        ("Precise-age range (years)", statistics["range_label"]),
        ("IQR (years)", statistics["iqr"]),
        ("Freedman-Diaconis width (years)", statistics["fd_width"]),
        ("Freedman-Diaconis bin count", statistics["fd_bins"]),
        ("Selected participants", "=COUNTA('Selected 14'!A2:A15)"),
        ("Selected FEMALE", '=COUNTIF(\'Selected 14\'!D2:D15,"FEMALE")'),
        ("Selected MALE", '=COUNTIF(\'Selected 14\'!D2:D15,"MALE")'),
        ("Overlap with current optimized 7", statistics["optimized_overlap"]),
        ("Largest selected F/M within-band age gap", statistics["max_pair_gap"]),
    ]
    for offset, (label, value) in enumerate(metrics, start=1):
        row_index = metric_start + offset
        readme.cell(row_index, 1, label)
        readme.cell(row_index, 2, value)
        if offset in (6, 7, 8):
            readme.cell(row_index, 2).font = Font(color="000000")
    readme.column_dimensions["A"].width = 42
    readme.column_dimensions["B"].width = 28
    for column in ("C", "D", "E", "F"):
        readme.column_dimensions[column].width = 16
    readme.print_area = "A1:F21"

    selected_sheet = write_rows_sheet(workbook, "Selected 14", selected_rows, "Selected14")
    optimized_sheet = write_rows_sheet(workbook, "Current optimized 7", optimized_rows, "CurrentOptimized7")
    bin_sheet = write_rows_sheet(workbook, "Bin counts", bin_rows, "BinCounts")
    feasibility_sheet = write_rows_sheet(workbook, "Rule checks", feasibility, "RuleChecks")
    audit_sheet = write_rows_sheet(workbook, "Eligible 132", audit_rows, "Eligible132")
    source_sheet = write_rows_sheet(workbook, "Sources", source_rows, "SourcesTable")

    for sheet in (selected_sheet, optimized_sheet, bin_sheet, feasibility_sheet, audit_sheet, source_sheet):
        sheet.sheet_view.showGridLines = False

    for cell in selected_sheet[1]:
        cell.fill = PatternFill("solid", fgColor=TEAL)
    for row in selected_sheet.iter_rows(min_row=2):
        row[2].number_format = "0.00"
        row[7].number_format = "0.000"
        row[8].number_format = "0.000"
        row[9].number_format = "0.000"

    for row in optimized_sheet.iter_rows(min_row=2):
        row[1].number_format = "0.00"

    for row in audit_sheet.iter_rows(min_row=2):
        if row[4].value == "YES":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=LIGHT_ORANGE)
        elif row[5].value == "YES":
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=LIGHT_GREEN)

    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = "Eligible candidates after holding out optimized seven"
    chart.y_axis.title = "Candidates"
    chart.x_axis.title = "Age band"
    data = Reference(bin_sheet, min_col=8, max_col=9, min_row=1, max_row=1 + len(bin_rows))
    categories = Reference(bin_sheet, min_col=2, min_row=2, max_row=1 + len(bin_rows))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 14
    bin_sheet.add_chart(chart, "L2")
    bin_sheet.print_area = f"A1:S{max(9, 1 + len(bin_rows))}"

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.save(path)


def main() -> None:
    final_ids = read_subject_ids(FINAL_132_FILE, EXPECTED_FINAL_132)
    optimized_ids = read_subject_ids(OPTIMIZED_7_FILE, EXPECTED_OPTIMIZED)
    subjects = read_demographics(DEMOGRAPHICS_FILE, set(final_ids))
    by_id = {subject.subject: subject for subject in subjects}
    if not set(optimized_ids).issubset(by_id):
        raise ValueError("At least one optimized subject is absent from final_132 demographics.")

    roles, optimized_cohort = optimized_roles(OPTIMIZED_7_COHORT_FILE)
    if set(roles) != set(optimized_ids):
        raise ValueError("Optimized cohort roles do not match optimized subjects.txt.")

    ages = [subject.age_precise for subject in subjects]
    minimum = min(ages)
    maximum = max(ages)
    q1 = linear_quantile(ages, 0.25)
    q3 = linear_quantile(ages, 0.75)
    iqr = q3 - q1
    fd_width = 2 * iqr * (len(ages) ** (-1 / 3))
    fd_bins = math.ceil((maximum - minimum) / fd_width)
    sturges_bins = math.ceil(math.log2(len(ages)) + 1)
    sqrt_bins = math.ceil(math.sqrt(len(ages)))
    if fd_bins != EXPECTED_FD_BINS:
        raise ValueError(f"Freedman-Diaconis result changed: {fd_bins} bins, expected 7.")

    bands = build_bands(minimum, maximum, fd_bins)
    optimized_set = set(optimized_ids)
    eligible = [subject for subject in subjects if subject.subject not in optimized_set]
    ranked = rank_candidates(eligible, bands)
    selected = sorted(
        [candidate for candidate in ranked if candidate.rank == 1],
        key=lambda candidate: (candidate.band.index, SEX_ORDER.index(candidate.subject.sex)),
    )
    expected_selected = fd_bins * len(SEX_ORDER)
    if len(selected) != expected_selected or len({item.subject.subject for item in selected}) != expected_selected:
        raise ValueError("Selection did not produce 14 unique subjects.")
    selected_ids = {candidate.subject.subject for candidate in selected}
    if selected_ids & optimized_set:
        raise ValueError("New selection overlaps the held-out optimized seven.")

    rank_by_id = {candidate.subject.subject: candidate for candidate in ranked}
    selected_rows: list[dict[str, object]] = []
    for candidate in selected:
        subject = candidate.subject
        selected_rows.append(
            {
                "subject": subject.subject,
                "CCID": subject.subject.removeprefix("sub-"),
                "age_years_precise": round(subject.age_precise, 2),
                "recorded_sex": subject.sex,
                "age_band_index": candidate.band.index,
                "age_band_years": candidate.band.label,
                "band_center_years": round(candidate.band.center, 3),
                "distance_to_center_years": round(candidate.distance_to_center, 3),
                "band_lower_years": round(candidate.band.lower, 3),
                "band_upper_years": round(candidate.band.upper, 3),
                "within_sex_band_rank": candidate.rank,
                "candidates_in_sex_band_after_holdout": candidate.candidates_in_stratum,
                "current_optimized_7_overlap": "NO",
            }
        )

    optimized_rows: list[dict[str, object]] = []
    for subject_id in sorted(optimized_ids, key=lambda value: by_id[value].age_precise):
        subject = by_id[subject_id]
        band = age_band(subject.age_precise, bands)
        optimized_rows.append(
            {
                "subject": subject.subject,
                "age_years_precise": round(subject.age_precise, 2),
                "age_years_workbook": round(subject.age_workbook, 2),
                "recorded_sex": subject.sex,
                "new_7_bin_index": band.index,
                "new_7_bin_age_range": band.label,
                "original_best_worst_selection_roles": "; ".join(sorted(roles[subject_id])),
                "original_extreme_role_count": len(roles[subject_id]),
                "individualized_optimization_scope": "4 manuscript ROIs",
            }
        )

    bin_rows: list[dict[str, object]] = []
    pair_gaps: list[float] = []
    selected_by_band: dict[int, list[RankedCandidate]] = defaultdict(list)
    for candidate in selected:
        selected_by_band[candidate.band.index].append(candidate)
    for band in bands:
        full_in_band = [subject for subject in subjects if band.contains(subject.age_precise, len(bands))]
        eligible_in_band = [subject for subject in eligible if band.contains(subject.age_precise, len(bands))]
        selected_in_band = selected_by_band[band.index]
        female = next(item.subject for item in selected_in_band if item.subject.sex == "FEMALE")
        male = next(item.subject for item in selected_in_band if item.subject.sex == "MALE")
        pair_gap = abs(female.age_precise - male.age_precise)
        pair_gaps.append(pair_gap)
        bin_rows.append(
            {
                "age_band_index": band.index,
                "age_band_years": band.label,
                "lower_years": round(band.lower, 3),
                "upper_years": round(band.upper, 3),
                "center_years": round(band.center, 3),
                "full_pool_female": sum(subject.sex == "FEMALE" for subject in full_in_band),
                "full_pool_male": sum(subject.sex == "MALE" for subject in full_in_band),
                "eligible_female_after_holdout": sum(subject.sex == "FEMALE" for subject in eligible_in_band),
                "eligible_male_after_holdout": sum(subject.sex == "MALE" for subject in eligible_in_band),
                "selected_female": female.subject,
                "selected_male": male.subject,
                "selected_pair_age_gap_years": round(pair_gap, 2),
            }
        )

    audit_rows: list[dict[str, object]] = []
    for subject in sorted(subjects, key=lambda item: (item.age_precise, item.sex, item.subject)):
        band = age_band(subject.age_precise, bands)
        candidate = rank_by_id.get(subject.subject)
        audit_rows.append(
            {
                "subject": subject.subject,
                "age_years_precise": round(subject.age_precise, 2),
                "age_years_workbook": round(subject.age_workbook, 2),
                "recorded_sex": subject.sex,
                "held_out_current_optimized_7": "YES" if subject.subject in optimized_set else "NO",
                "selected_new_14": "YES" if subject.subject in selected_ids else "NO",
                "age_band_index": band.index,
                "age_band_years": band.label,
                "distance_to_center_years": ""
                if candidate is None
                else round(candidate.distance_to_center, 3),
                "within_sex_band_rank_after_holdout": "" if candidate is None else candidate.rank,
            }
        )

    feasibility = feasibility_rows(subjects, eligible, minimum, maximum)
    fd_feasibility = next(row for row in feasibility if row["bin_count"] == fd_bins)
    sturges_feasibility = next(row for row in feasibility if row["bin_count"] == sturges_bins)
    sqrt_feasibility = next(row for row in feasibility if row["bin_count"] == sqrt_bins)
    if fd_feasibility["eligible_pool_all_sexes_each_bin"] != "YES":
        raise ValueError("Freedman-Diaconis bins do not support both recorded sexes after holdout.")

    PACKAGE_DIR.mkdir(parents=True, exist_ok=True)
    CANONICAL_DIR.mkdir(parents=True, exist_ok=True)

    subjects_path = PACKAGE_DIR / "subjects.txt"
    subjects_path.write_text(
        "".join(f"{candidate.subject.subject}\n" for candidate in selected),
        encoding="utf-8",
    )
    selection_path = PACKAGE_DIR / "selection_manifest.csv"
    write_csv(selection_path, list(selected_rows[0]), selected_rows)
    optimized_summary_path = PACKAGE_DIR / "current_optimized_7_summary.csv"
    write_csv(optimized_summary_path, list(optimized_rows[0]), optimized_rows)
    bin_path = PACKAGE_DIR / "bin_diagnostics.csv"
    write_csv(bin_path, list(bin_rows[0]), bin_rows)
    feasibility_path = PACKAGE_DIR / "bin_rule_feasibility.csv"
    write_csv(feasibility_path, list(feasibility[0]), feasibility)
    audit_path = PACKAGE_DIR / "eligible_132_selection_audit.csv"
    write_csv(audit_path, list(audit_rows[0]), audit_rows)

    source_files = {
        "final_132_subjects": FINAL_132_FILE,
        "demographics": DEMOGRAPHICS_FILE,
        "current_optimized_7_subjects": OPTIMIZED_7_FILE,
        "current_optimized_7_cohort": OPTIMIZED_7_COHORT_FILE,
        "builder": Path(__file__).resolve(),
    }
    source_rows = [
        {"source": label, "path": str(path), "sha256": sha256_file(path)}
        for label, path in source_files.items()
    ]
    source_rows.append(
        {
            "source": "Freedman-Diaconis reference",
            "path": FD_REFERENCE_URL,
            "sha256": "not_applicable_web_reference",
        }
    )

    optimized_ages = [by_id[subject_id].age_precise for subject_id in optimized_ids]
    selected_ages = [candidate.subject.age_precise for candidate in selected]
    selected_sexes = Counter(candidate.subject.sex for candidate in selected)
    generated_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    cohort = {
        "schema_version": 1,
        "cohort_id": "optimized_vs_unoptimized_age_balanced_14",
        "status": "selection_complete_optimization_and_simulation_pending",
        "created_at_utc": generated_at,
        "source_cohort": "final_132",
        "subjects_file": "subjects.txt",
        "selection_manifest": "selection_manifest.csv",
        "expected_subjects": expected_selected,
        "recorded_female": selected_sexes["FEMALE"],
        "recorded_male": selected_sexes["MALE"],
        "selection_independence": {
            "outcome_metrics_used": False,
            "held_out_cohort": "optimized_best_worst_7",
            "held_out_subjects": optimized_ids,
            "reason": (
                "The existing seven were originally selected as best/worst simulated TI cases. "
                "Holding them out avoids outcome-informed sampling in the new age-balanced cohort."
            ),
        },
        "binning": {
            "method": "Freedman-Diaconis rule applied to precise ages of all final_132 participants",
            "reference": FD_REFERENCE_URL,
            "n": len(ages),
            "age_min_years": minimum,
            "age_max_years": maximum,
            "q1_years": round(q1, 4),
            "q3_years": round(q3, 4),
            "iqr_years": round(iqr, 4),
            "fd_target_width_years": round(fd_width, 4),
            "fd_bin_count": fd_bins,
            "implemented_equal_width_years": round((maximum - minimum) / fd_bins, 4),
            "sturges_bin_count": sturges_bins,
            "sturges_feasible_after_holdout": sturges_feasibility[
                "eligible_pool_all_sexes_each_bin"
            ]
            == "YES",
            "sqrt_bin_count": sqrt_bins,
            "sqrt_feasible_after_holdout": sqrt_feasibility[
                "eligible_pool_all_sexes_each_bin"
            ]
            == "YES",
            "bands": [
                {
                    "index": band.index,
                    "lower": round(band.lower, 6),
                    "upper": round(band.upper, 6),
                    "center": round(band.center, 6),
                    "label": band.label,
                }
                for band in bands
            ],
        },
        "selection_rule": {
            "per_band": "one recorded-FEMALE and one recorded-MALE participant",
            "ranking": [
                "absolute distance to age-band midpoint",
                "lower precise age",
                "lexical subject ID",
            ],
            "age_role": (
                "Age bands enforce sampling coverage. Analyze age continuously in the "
                "within-participant optimized-minus-unoptimized comparison."
            ),
        },
        "selected_subjects": [candidate.subject.subject for candidate in selected],
        "validation": {
            "source_subjects": len(subjects),
            "missing_demographics": 0,
            "eligible_after_holdout": len(eligible),
            "selected_subjects": len(selected),
            "exact_recorded_sex_balance": selected_sexes == Counter({"FEMALE": 7, "MALE": 7}),
            "two_subjects_per_band": all(len(selected_by_band[index]) == 2 for index in range(1, 8)),
            "optimized_7_overlap": len(selected_ids & optimized_set),
            "selected_age_min_years": min(selected_ages),
            "selected_age_max_years": max(selected_ages),
            "selected_age_mean_years": round(mean(selected_ages), 4),
            "selected_age_median_years": round(median(selected_ages), 4),
            "maximum_within_band_pair_age_gap_years": round(max(pair_gaps), 4),
            "minimum_female_candidates_in_any_band_after_holdout": fd_feasibility[
                "eligible_min_female_candidates"
            ],
            "minimum_male_candidates_in_any_band_after_holdout": fd_feasibility[
                "eligible_min_male_candidates"
            ],
        },
        "current_optimized_7_summary": {
            "subjects": len(optimized_ids),
            "recorded_female": sum(by_id[value].sex == "FEMALE" for value in optimized_ids),
            "recorded_male": sum(by_id[value].sex == "MALE" for value in optimized_ids),
            "age_min_years": min(optimized_ages),
            "age_max_years": max(optimized_ages),
            "age_mean_years": round(mean(optimized_ages), 4),
            "age_median_years": round(median(optimized_ages), 4),
            "original_selection_metric": optimized_cohort["selection_metric"],
        },
        "source_files": {
            label: {"path": str(path), "sha256": sha256_file(path)}
            for label, path in source_files.items()
        },
        "artifact_hashes": {
            "subjects.txt": sha256_file(subjects_path),
            "selection_manifest.csv": sha256_file(selection_path),
            "current_optimized_7_summary.csv": sha256_file(optimized_summary_path),
        },
    }
    cohort_path = PACKAGE_DIR / "cohort.json"
    cohort_path.write_text(json.dumps(cohort, indent=2) + "\n", encoding="utf-8")

    selected_table = "\n".join(
        "| {age_band_years} | `{subject}` | {female_age:.2f} | `{male}` | {male_age:.2f} | {gap:.2f} |".format(
            age_band_years=row["age_band_years"],
            subject=row["selected_female"],
            female_age=by_id[str(row["selected_female"])].age_precise,
            male=row["selected_male"],
            male_age=by_id[str(row["selected_male"])].age_precise,
            gap=float(row["selected_pair_age_gap_years"]),
        )
        for row in bin_rows
    )
    optimized_table = "\n".join(
        f"| `{row['subject']}` | {float(row['age_years_precise']):.2f} | {row['recorded_sex']} | {row['original_best_worst_selection_roles']} |"
        for row in optimized_rows
    )
    readme_text = f"""# Optimized vs unoptimized TI age-balanced cohort

This package defines a new 14-participant selection from the final-132 CamCan
cohort for a future within-participant comparison of optimized and unoptimized
TI stimulation. Optimization and simulations have **not** been run by this
selection task.

## Decision

- Seven equal-width age bands across {minimum:.2f}-{maximum:.2f} years.
- One recorded-female and one recorded-male participant per band: 14 total.
- The seven current optimized best/worst cases are excluded from the new cohort.
- Selection uses age and recorded sex only; it does not use a TI outcome.

The bin count is determined by the Freedman-Diaconis rule, `h = 2 x IQR x
n^(-1/3)`. For the 132 precise ages, IQR = {iqr:.2f} years and the target width
is {fd_width:.2f} years, giving `ceil(({maximum:.2f}-{minimum:.2f})/{fd_width:.2f}) = {fd_bins}`
bins. Sturges gives {sturges_bins} and the square-root rule gives {sqrt_bins}, but those
equal-width designs are infeasible here because at least one oldest-age band has
no eligible female participant. Seven bands are feasible after the optimized
seven are held out and retain at least two female candidates and six male
candidates in every band. Reference: Freedman & Diaconis (1981),
{FD_REFERENCE_URL}.

Equal-width rather than equal-count bands are used because the design goal is
coverage of the age range. Within each recorded-sex x band stratum, the selected
participant is nearest the midpoint; ties are resolved by lower age and then
lexical subject ID.

## Selected new cohort

| Age band (years) | Female | Age | Male | Age | Pair gap |
| --- | --- | ---: | --- | ---: | ---: |
{selected_table}

Selected age range: {min(selected_ages):.2f}-{max(selected_ages):.2f} years;
mean {mean(selected_ages):.2f}; median {median(selected_ages):.2f}. Maximum
within-band female/male age gap: {max(pair_gaps):.2f} years. Overlap with the
current optimized seven: 0/14.

## Current optimized seven (kept separate)

| Subject | Age (years) | Recorded sex | Original best/worst role(s) |
| --- | ---: | --- | --- |
{optimized_table}

Summary: 7 participants (3 recorded-female, 4 recorded-male), age range
{min(optimized_ages):.2f}-{max(optimized_ages):.2f} years, mean
{mean(optimized_ages):.2f}, median {median(optimized_ages):.2f}. These cases were
originally selected using the historical full-anatomical-parcel ROI median TI
field; they remain useful as an external extreme-case/sensitivity cohort, but
mixing them into the new primary sample would make sampling outcome-informed.

## Analysis boundary

The age bands are a sampling device, not seven powered inferential groups. For
the optimized-vs-unoptimized study, use the within-participant difference as the
primary contrast and model age continuously. Recorded sex can be included as a
prespecified covariate; interaction estimates should be described as exploratory
unless a separate power analysis supports them.

## Files

- `subjects.txt`: canonical ordered 14-subject list.
- `selection_manifest.csv`: selected demographics and deterministic ranking.
- `current_optimized_7_summary.csv`: ages, recorded sex, and original roles.
- `bin_diagnostics.csv`: full and eligible counts by age band.
- `bin_rule_feasibility.csv`: equal-width feasibility for 4-12 bins.
- `eligible_132_selection_audit.csv`: full selection audit.
- `selection_audit.xlsx`: formatted review workbook.
- `cohort.json`: machine-readable method, validation, and provenance.
- `SHA256SUMS`: artifact checksums.

Builder: `optimized_vs_unoptimized_age_cohort/build_age_balanced_cohort.py`.
Generated: {generated_at}.
"""
    readme_path = PACKAGE_DIR / "README.md"
    readme_path.write_text(readme_text, encoding="utf-8")

    statistics = {
        "n": len(ages),
        "range_label": f"{minimum:.2f}-{maximum:.2f}",
        "iqr": round(iqr, 2),
        "fd_width": round(fd_width, 3),
        "fd_bins": fd_bins,
        "optimized_overlap": len(selected_ids & optimized_set),
        "max_pair_gap": round(max(pair_gaps), 2),
    }
    workbook_path = PACKAGE_DIR / "selection_audit.xlsx"
    build_workbook(
        workbook_path,
        selected_rows,
        optimized_rows,
        bin_rows,
        feasibility,
        audit_rows,
        source_rows,
        statistics,
    )

    package_artifacts = [
        subjects_path,
        selection_path,
        optimized_summary_path,
        bin_path,
        feasibility_path,
        audit_path,
        cohort_path,
        readme_path,
        workbook_path,
    ]
    checksum_path = PACKAGE_DIR / "SHA256SUMS"
    checksum_path.write_text(
        "".join(f"{sha256_file(path)}  {path.name}\n" for path in sorted(package_artifacts)),
        encoding="utf-8",
    )

    canonical_names = [
        "subjects.txt",
        "selection_manifest.csv",
        "current_optimized_7_summary.csv",
        "bin_diagnostics.csv",
        "cohort.json",
        "README.md",
    ]
    for name in canonical_names:
        shutil.copy2(PACKAGE_DIR / name, CANONICAL_DIR / name)
    (CANONICAL_DIR / "SOURCE_PACKAGE.txt").write_text(
        f"{PACKAGE_DIR.relative_to(REPO_ROOT)}\n",
        encoding="utf-8",
    )

    for name in canonical_names:
        if sha256_file(PACKAGE_DIR / name) != sha256_file(CANONICAL_DIR / name):
            raise ValueError(f"Canonical cohort copy differs for {name}.")

    print(f"Built {len(selected)}-subject cohort in {PACKAGE_DIR}")
    print(f"Canonical cohort: {CANONICAL_DIR}")
    print(f"Bins: {fd_bins}; selected sex counts: {dict(selected_sexes)}")
    print(f"Optimized-seven overlap: {len(selected_ids & optimized_set)}")


if __name__ == "__main__":
    main()
